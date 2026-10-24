"""Genera un proyecto de ejemplo en formato XML de MS Project (MSPDI).

Sirve para probar el exportador sin MS Project. El archivo imita lo que escribe
MS Project al usar Archivo > Guardar como > XML: línea base 0, campos
personalizados con alias ("Método EVM" en Texto10 y "Peso hito" en Número1) y,
opcionalmente, avance a una fecha de estado. El XML incluye también costos
(multiplicados por 100, como los escribe MS Project), que el exportador no usa:
son el presupuesto y el costo real ficticios que --libro escribe en el libro como
BAC y AC.

El proyecto tiene dos planificaciones, con códigos EDT de alto nivel "INTERNO" y
"CLIENTE" (la del cliente con costos un 15 % mayores). Debajo de cada una, los
elementos tienen códigos EDT del tipo 1.0.0, 1.1.0, 1.1.1 y 1.1.1.1; los hitos de
los paquetes y la reunión de arranque no tienen código EDT. Los códigos EDT están
en el campo EDT (WBS) y también en el campo personalizado "Item EDT".

Para multi-empresa, el campo "Empresa" indica EmpresaA, EmpresaB o Ambas, y el
campo "% Reparto" el porcentaje de EmpresaA en las tareas de ambas empresas.

Con --proyecto 2 se genera un segundo proyecto, "Accesos y servicios", con el
código EDT 9.0.0 y sin planificaciones, para combinarlo con el primero en un
mismo libro (varios archivos de MS Project, secciones [[proyectos]]).

Uso:
    python ejemplos/generar_ejemplo.py                       # solo línea base
    python ejemplos/generar_ejemplo.py --fecha-estado 2027-06-30 --salida ejemplos/avance_2027-06.xml
    python ejemplos/generar_ejemplo.py --proyecto 2 --salida ejemplos/proyecto2_ejemplo.xml
"""
from __future__ import annotations

import argparse
import calendar
import sys
from dataclasses import dataclass, field, replace
from datetime import date, datetime, time, timedelta
from pathlib import Path
from xml.sax.saxutils import escape

TEXT10 = 188743750      # Texto10 (tarea): "Método EVM"
NUMBER1 = 188743767     # Número1 (tarea): "Peso hito"
TEXT2 = 188743734       # Texto2 (tarea): "Item EDT"
TEXT3 = 188743737       # Texto3 (tarea): "Empresa"
NUMBER2 = 188743768     # Número2 (tarea): "% Reparto"


@dataclass
class T:
    name: str
    level: int
    start: date | None = None
    finish: date | None = None
    cost: float = 0.0
    method: str = ""
    milestone: bool = False
    weight: float | None = None
    children: list["T"] = field(default_factory=list)
    code: str | None = None     # código EDT fijo ("" = sin código); None = se numera
    company: str = ""           # empresa responsable (EmpresaA, EmpresaB, Ambas)
    split: float | None = None  # % de EmpresaA en tareas "Ambas"
    # calculados
    uid: int = 0
    wbs: str = ""
    outline: str = ""
    pct: float = 0.0
    actual_cost: float = 0.0
    actual_start: date | None = None
    actual_finish: date | None = None


def wdays(a: date, b: date) -> list[date]:
    out, d = [], a
    while d <= b:
        if d.weekday() < 5:
            out.append(d)
        d += timedelta(days=1)
    return out


def add_wd(d: date, n: int) -> date:
    while d.weekday() >= 5:
        d += timedelta(days=1)
    k = 1
    while k < n:
        d += timedelta(days=1)
        if d.weekday() < 5:
            k += 1
    return d


def work_package(start: date, factor: float = 1.0):
    """Constructor de paquetes de trabajo (nivel 3) a partir del inicio del proyecto."""
    def wp(name, offset_m, dur_d, cost, method="", acts=None):
        s = add_wd(start + timedelta(days=int(offset_m * 30.4)), 1)
        t = T(name, 3, method=method)
        if acts:
            cur = s
            for (aname, adur, acost, *rest) in acts:
                if adur == 0:  # hito (sin código EDT)
                    t.children.append(T(aname, 4, cur, cur, acost * factor, milestone=True,
                                        weight=rest[0] if rest else None, code=""))
                else:
                    f = add_wd(cur, adur)
                    t.children.append(T(aname, 4, cur, f, acost * factor))
                    cur = add_wd(f + timedelta(days=1), 1)
        else:
            t.start, t.finish, t.cost = s, add_wd(s, dur_d), cost * factor
        return t
    return wp


def build_program(factor: float = 1.0) -> list[T]:
    """Programa de 4 elementos de nivel EDT 1, hasta nivel 4, ~5 años."""
    wp = work_package(date(2026, 1, 5), factor)
    prog = [
        T("Gestión del programa", 1, children=[
            T("Dirección y control", 2, children=[
                wp("Oficina de gestión del programa (PMO)", 0, 1250, 1_800_000, "Avance lineal (LOE)"),
                wp("Control de proyectos y EVM", 0, 1250, 650_000, "Avance lineal (LOE)"),
                wp("Gestión de riesgos", 1, 1200, 300_000, "Avance lineal (LOE)"),
            ]),
            T("Gestión de contratos", 2, children=[
                wp("Licitación de ingeniería", 0, 80, 120_000, "0/100"),
                wp("Licitación de obras", 8, 100, 180_000, "50/50"),
                wp("Administración de contratos", 10, 1000, 420_000, "Avance lineal (LOE)"),
            ]),
        ]),
        T("Ingeniería", 1, children=[
            T("Ingeniería básica", 2, children=[
                wp("Estudios de terreno", 2, 60, 350_000, "50/50"),
                wp("Diseño conceptual", 3, 0, 0, "Hitos ponderados", acts=[
                    ("Bases de diseño", 25, 180_000), ("Hito: bases aprobadas", 0, 0, 30),
                    ("Alternativas y selección", 40, 260_000), ("Hito: alternativa seleccionada", 0, 0, 30),
                    ("Informe conceptual", 30, 160_000), ("Hito: informe aprobado", 0, 0, 40),
                ]),
                wp("Ingeniería básica civil", 6, 140, 900_000, "% Completado"),
                wp("Ingeniería básica eléctrica", 6, 140, 750_000, "% Completado"),
            ]),
            T("Ingeniería de detalle", 2, children=[
                wp("Detalle civil", 12, 220, 1_600_000, "% Físico completado"),
                wp("Detalle mecánico", 13, 220, 1_450_000, "% Físico completado"),
                wp("Detalle eléctrico y control", 14, 240, 1_300_000, "% Completado"),
                wp("Revisión y aprobación", 22, 0, 0, "Hitos ponderados", acts=[
                    ("Revisión 60 %", 30, 120_000), ("Hito: revisión 60 % aprobada", 0, 0, 40),
                    ("Revisión 90 %", 30, 120_000), ("Hito: emitido para construcción", 0, 0, 60),
                ]),
            ]),
        ]),
        T("Construcción", 1, children=[
            T("Obras civiles", 2, children=[
                wp("Movimiento de tierras", 18, 160, 3_200_000, "Unidades completadas"),
                wp("Fundaciones", 22, 220, 5_400_000, "Unidades completadas"),
                wp("Estructuras", 28, 300, 7_800_000, "% Físico completado"),
            ]),
            T("Montaje electromecánico", 2, children=[
                wp("Suministro de equipos", 20, 0, 0, "Hitos ponderados", acts=[
                    ("Fabricación de equipos", 200, 6_000_000), ("Hito: equipos en fábrica (FAT)", 0, 0, 40),
                    ("Transporte", 40, 900_000), ("Hito: equipos en obra", 0, 0, 60),
                ]),
                wp("Montaje mecánico", 32, 280, 4_300_000, "% Completado"),
                wp("Montaje eléctrico", 34, 300, 3_900_000, "% Completado"),
            ]),
        ]),
        T("Puesta en marcha", 1, children=[
            T("Pruebas", 2, children=[
                wp("Pruebas en frío", 45, 60, 650_000, "0/100"),
                wp("Pruebas en caliente", 48, 60, 850_000, "20/80"),
            ]),
            T("Cierre", 2, children=[
                wp("Capacitación y entrega", 50, 60, 400_000, "50/50"),
                wp("Cierre administrativo", 52, 60, 150_000, ""),
            ]),
        ]),
    ]
    # Multi-empresa: gestión compartida 60/40, ingeniería de A (salvo los estudios de
    # terreno, de B), construcción de B y puesta en marcha compartida (50 % por defecto).
    prog[0].company, prog[0].split = "Ambas", 60
    prog[1].company = "EmpresaA"
    prog[1].children[0].children[0].company = "EmpresaB"
    prog[2].company = "EmpresaB"
    prog[3].company = "Ambas"
    return prog


def build_plans() -> list[T]:
    """Dos planificaciones en el mismo proyecto: INTERNO y CLIENTE."""
    plans = []
    for name, code, factor in (("Planificación interna", "INTERNO", 1.0),
                               ("Planificación del cliente", "CLIENTE", 1.15)):
        kickoff = T("Reunión de arranque", 2, date(2026, 1, 5), date(2026, 1, 5), milestone=True, code="")
        plans.append(T(name, 1, code=code, children=[kickoff] + build_program(factor)))
    return plans


def build_second() -> list[T]:
    """Segundo proyecto (código EDT 9.0.0, sin planificaciones): accesos y servicios de la obra, ~2 años."""
    wp = work_package(date(2026, 2, 2))
    prog = [
        T("Accesos y servicios", 1, children=[
            T("Camino de acceso", 2, children=[
                wp("Diseño del camino", 0, 60, 220_000, "50/50"),
                wp("Construcción del camino", 3, 200, 2_400_000, "Unidades completadas"),
            ]),
            T("Servicios de obra", 2, children=[
                wp("Suministro eléctrico provisional", 2, 120, 900_000, "% Completado"),
                wp("Campamento", 1, 0, 0, "Hitos ponderados", acts=[
                    ("Montaje de módulos", 60, 650_000), ("Hito: módulos habilitados", 0, 0, 60),
                    ("Servicios del campamento", 40, 250_000), ("Hito: campamento en servicio", 0, 0, 40),
                ]),
                wp("Gestión del proyecto de accesos", 0, 500, 380_000, "Avance lineal (LOE)"),
            ]),
        ]),
    ]
    # Multi-empresa: el camino es de B; los servicios, compartidos 30/70
    prog[0].children[0].company = "EmpresaB"
    prog[0].children[1].company, prog[0].children[1].split = "Ambas", 30
    return prog


PROJECTS = {
    1: ("Programa de ejemplo EVM", "programa_ejemplo.xml", build_plans, 1),
    2: ("Accesos y servicios (proyecto 2 de ejemplo)", "proyecto2_ejemplo.xml", build_second, 9),
}


def walk(nodes, depth=1, outline="", numbers=(), first=1):
    """Asigna nivel de esquema y código EDT.

    Las planificaciones conservan su código (INTERNO / CLIENTE); debajo, los
    elementos se numeran 1.0.0, 1.1.0, 1.1.1 y 1.1.1.1 (solo los que tienen código).
    first: número del primer elemento de nivel 1 (9 -> 9.0.0).
    """
    k = first - 1
    for i, n in enumerate(nodes, start=1):
        n.level = depth
        n.outline = f"{outline}{i}"
        if n.code is not None:
            n.wbs = n.code
            child_numbers = () if n.code else numbers
        else:
            k += 1
            segs = list(numbers) + [k]
            n.wbs = ".".join(str(x) for x in (segs + [0] * (3 - len(segs)) if len(segs) < 3 else segs))
            child_numbers = tuple(segs)
        yield n
        yield from walk(n.children, depth + 1, n.outline + ".", child_numbers)


def rollup(n: T) -> None:
    if not n.children:
        return
    for c in n.children:
        rollup(c)
    n.start = min(c.start for c in n.children)
    n.finish = max(c.finish for c in n.children)
    n.cost = sum(c.cost for c in n.children)


def monthly(n: T) -> dict[tuple[int, int], float]:
    """Costo de línea base por mes (lineal por días hábiles), como lo hace MS Project."""
    if n.children:
        out: dict[tuple[int, int], float] = {}
        for c in n.children:
            for k, v in monthly(c).items():
                out[k] = out.get(k, 0.0) + v
        return out
    if not n.cost:
        return {}
    days = wdays(n.start, n.finish) or [n.start]
    out = {}
    for d in days:
        out[(d.year, d.month)] = out.get((d.year, d.month), 0.0) + n.cost / len(days)
    return out


def progress(n: T, status: date) -> None:
    """Avance simulado: atraso del 15 % y sobrecosto del 8 %."""
    if n.children:
        for c in n.children:
            progress(c, status)
        done_cost = sum(c.cost * c.pct for c in n.children)
        tot = sum(c.cost for c in n.children)
        dur = sum(max(1, len(wdays(c.start, c.finish))) * c.pct for c in n.children)
        dtot = sum(max(1, len(wdays(c.start, c.finish))) for c in n.children)
        n.pct = (dur / dtot) if dtot else 0.0
        n.actual_cost = sum(c.actual_cost for c in n.children)
        starts = [c.actual_start for c in n.children if c.actual_start]
        n.actual_start = min(starts) if starts else None
        n.actual_finish = max(c.actual_finish for c in n.children) if all(c.actual_finish for c in n.children) else None
        _ = done_cost, tot
        return
    if status < n.start:
        return
    total = max(1, len(wdays(n.start, n.finish)))
    elapsed = len(wdays(n.start, min(status, n.finish)))
    pct = min(1.0, elapsed / total * 0.85)
    if n.milestone:
        pct = 1.0 if status >= n.finish + timedelta(days=20) else 0.0
    n.pct = round(pct, 2)
    n.actual_start = n.start if n.pct > 0 or n.milestone and n.pct else None
    n.actual_finish = n.finish + timedelta(days=20) if n.pct >= 1.0 else None
    n.actual_cost = round(n.cost * n.pct * 1.08, 2)


def dt(d: date, end: bool = False) -> str:
    return datetime.combine(d, time(17, 0) if end else time(8, 0)).isoformat()


def money(v: float) -> str:
    """Montos en MSPDI: multiplicados por 100."""
    text = f"{v * 100:.2f}"
    return text[:-3] if text.endswith(".00") else text


def write(nodes: list[T], path: Path, status: date | None, project: int = 1) -> None:
    title, file_name, _build, first = PROJECTS[project]
    tasks = list(walk(nodes, first=first))
    for r in nodes:
        rollup(r)
    for uid, t in enumerate(tasks, start=1):
        t.uid = uid
    if status:
        for r in nodes:
            progress(r, status)
    start = min(t.start for t in tasks)
    finish = max(t.finish for t in tasks)
    total = sum(r.cost for r in nodes)
    total_ac = sum(r.actual_cost for r in nodes)
    total_ev = sum(t.cost * t.pct for t in tasks if not t.children)

    out = ['<?xml version="1.0" encoding="UTF-8" standalone="yes"?>',
           '<Project xmlns="http://schemas.microsoft.com/project">',
           "\t<SaveVersion>14</SaveVersion>",
           f"\t<Name>{file_name}</Name>",
           f"\t<Title>{escape(title)}</Title>",
           f"\t<StartDate>{dt(start)}</StartDate>",
           f"\t<FinishDate>{dt(finish, True)}</FinishDate>",
           "\t<CurrencyDigits>2</CurrencyDigits>",
           "\t<CurrencySymbol>$</CurrencySymbol>",
           "\t<CurrencyCode>USD</CurrencyCode>",
           "\t<DefaultFixedCostAccrual>3</DefaultFixedCostAccrual>",
           "\t<BaselineForEarnedValue>0</BaselineForEarnedValue>"]
    if status:
        out.append(f"\t<StatusDate>{dt(status, True)}</StatusDate>")
    out += ["\t<DefaultTaskEVMethod>0</DefaultTaskEVMethod>",
            "\t<ExtendedAttributes>",
            f"\t\t<ExtendedAttribute><FieldID>{TEXT10}</FieldID><FieldName>Text10</FieldName>"
            "<Alias>Método EVM</Alias></ExtendedAttribute>",
            f"\t\t<ExtendedAttribute><FieldID>{NUMBER1}</FieldID><FieldName>Number1</FieldName>"
            "<Alias>Peso hito</Alias></ExtendedAttribute>",
            f"\t\t<ExtendedAttribute><FieldID>{TEXT2}</FieldID><FieldName>Text2</FieldName>"
            "<Alias>Item EDT</Alias></ExtendedAttribute>",
            f"\t\t<ExtendedAttribute><FieldID>{TEXT3}</FieldID><FieldName>Text3</FieldName>"
            "<Alias>Empresa</Alias></ExtendedAttribute>",
            f"\t\t<ExtendedAttribute><FieldID>{NUMBER2}</FieldID><FieldName>Number2</FieldName>"
            "<Alias>% Reparto</Alias></ExtendedAttribute>",
            "\t</ExtendedAttributes>",
            "\t<Tasks>"]

    def task_xml(uid, tid, name, level, wbs, t: T | None):
        summary = bool(t and t.children)
        s = t.start if t else start
        f = t.finish if t else finish
        cost = t.cost if t else total
        ac = t.actual_cost if t else total_ac
        pct = (t.pct if t else (total_ev / total if total else 0))
        ev = (sum(x.cost * x.pct for x in walk_leaves(t)) if t else total_ev)
        lines = ["\t\t<Task>", f"\t\t\t<UID>{uid}</UID>", f"\t\t\t<ID>{tid}</ID>",
                 f"\t\t\t<Name>{escape(name)}</Name>", "\t\t\t<Active>1</Active>", "\t\t\t<IsNull>0</IsNull>",
                 f"\t\t\t<WBS>{wbs}</WBS>",
                 f"\t\t\t<OutlineNumber>{t.outline if t else '0'}</OutlineNumber>",
                 f"\t\t\t<OutlineLevel>{level}</OutlineLevel>",
                 f"\t\t\t<Start>{dt(s)}</Start>", f"\t\t\t<Finish>{dt(f, True)}</Finish>",
                 f"\t\t\t<Milestone>{1 if t and t.milestone else 0}</Milestone>",
                 f"\t\t\t<Summary>{1 if summary or t is None else 0}</Summary>",
                 "\t\t\t<ExternalTask>0</ExternalTask>",
                 "\t\t\t<FixedCost>0</FixedCost>", "\t\t\t<FixedCostAccrual>3</FixedCostAccrual>",
                 f"\t\t\t<PercentComplete>{int(round(pct * 100))}</PercentComplete>",
                 f"\t\t\t<Cost>{money(cost)}</Cost>"]
        if t and t.actual_start:
            lines.append(f"\t\t\t<ActualStart>{dt(t.actual_start)}</ActualStart>")
        if t and t.actual_finish:
            lines.append(f"\t\t\t<ActualFinish>{dt(t.actual_finish, True)}</ActualFinish>")
        lines += [f"\t\t\t<ActualCost>{money(ac)}</ActualCost>",
                  f"\t\t\t<ACWP>{round(ac * 100, 2):.2f}</ACWP>",
                  f"\t\t\t<BCWP>{round(ev * 100, 2):.2f}</BCWP>",
                  f"\t\t\t<PhysicalPercentComplete>{int(round(pct * 100))}</PhysicalPercentComplete>",
                  "\t\t\t<EarnedValueMethod>0</EarnedValueMethod>"]
        if t and t.method:
            lines.append(f"\t\t\t<ExtendedAttribute><FieldID>{TEXT10}</FieldID><Value>{escape(t.method)}</Value>"
                         "</ExtendedAttribute>")
        if t and t.weight is not None:
            lines.append(f"\t\t\t<ExtendedAttribute><FieldID>{NUMBER1}</FieldID><Value>{t.weight:g}</Value>"
                         "</ExtendedAttribute>")
        if t and wbs:
            lines.append(f"\t\t\t<ExtendedAttribute><FieldID>{TEXT2}</FieldID><Value>{escape(wbs)}</Value>"
                         "</ExtendedAttribute>")
        if t and t.company:
            lines.append(f"\t\t\t<ExtendedAttribute><FieldID>{TEXT3}</FieldID><Value>{escape(t.company)}</Value>"
                         "</ExtendedAttribute>")
        if t and t.split is not None:
            lines.append(f"\t\t\t<ExtendedAttribute><FieldID>{NUMBER2}</FieldID><Value>{t.split:g}</Value>"
                         "</ExtendedAttribute>")
        lines += ["\t\t\t<Baseline>", "\t\t\t\t<Number>0</Number>",
                  f"\t\t\t\t<Start>{dt(s)}</Start>", f"\t\t\t\t<Finish>{dt(f, True)}</Finish>",
                  f"\t\t\t\t<Cost>{money(cost)}</Cost>", "\t\t\t</Baseline>"]
        months = monthly(t) if t else {}
        if t is None:
            for r in nodes:
                for k, v in monthly(r).items():
                    months[k] = months.get(k, 0.0) + v
        for (y, m), v in sorted(months.items()):
            a = max(s, date(y, m, 1))
            b = min(f, date(y, m, calendar.monthrange(y, m)[1]))
            lines += ["\t\t\t<TimephasedData>", "\t\t\t\t<Type>10</Type>", f"\t\t\t\t<UID>{uid}</UID>",
                      f"\t\t\t\t<Start>{dt(a)}</Start>", f"\t\t\t\t<Finish>{dt(b, True)}</Finish>",
                      "\t\t\t\t<Unit>1</Unit>", f"\t\t\t\t<Value>{money(v)}</Value>",
                      "\t\t\t</TimephasedData>"]
        lines.append("\t\t</Task>")
        return lines

    out += task_xml(0, 0, title, 0, "0", None)
    for i, t in enumerate(tasks, start=1):
        out += task_xml(t.uid, i, t.name, t.level, t.wbs, t)
    out += ["\t</Tasks>", "</Project>"]
    path.write_text("\n".join(out) + "\n", encoding="utf-8")


def walk_leaves(t: T):
    if not t.children:
        yield t
    for c in t.children:
        yield from walk_leaves(c)


# ----------------------------------------------------------------------------
# Datos ficticios para el libro (BAC, AC y previsión de la depuración de costes)
# ----------------------------------------------------------------------------
def _import_evm() -> None:
    """Permite usar el paquete evm al ejecutar este archivo desde la carpeta ejemplos."""
    try:
        import evm  # noqa: F401
    except ModuleNotFoundError:
        sys.path.insert(0, str(Path(__file__).resolve().parents[1]))


def _every(nodes):
    """Todas las tareas del ejemplo (resúmenes y tareas de detalle)."""
    for n in nodes:
        yield n
        yield from _every(n.children)


def fill_book(book_path: Path, status: date | None = None, projects: dict[str, int] | None = None,
              single: int = 1) -> tuple[int, int]:
    """Escribe en el libro los datos que en un proyecto real se ingresan a mano.

    - BAC de cada fila (solo en las celdas vacías): el presupuesto ficticio del proyecto de ejemplo (la suma del
      costo de sus tareas en el ejemplo; el programa no lee costos de MS Project).
    - Con `status`: AC del mes de esa fecha, el costo real simulado del ejemplo (como si viniera del ERP de cada
      empresa). El costo real simulado tiene desvíos: algunos paquetes quedan por encima de su presupuesto.

    projects: {nombre del archivo en el libro: proyecto de ejemplo (1 o 2)}. Por defecto, el del libro: un
    solo archivo es el proyecto `single` y, con varios archivos ([[proyectos]]), el primero es el proyecto 1 y el
    segundo el proyecto 2.
    Devuelve la cantidad de celdas de BAC y de AC escritas.
    """
    _import_evm()
    import tempfile

    from openpyxl import load_workbook

    from evm.companies import resolve_shares
    from evm.config import Config
    from evm.reader_xml import read_xml
    from evm.timeline import Horizon
    from evm.update import _refresh_charts, _save, apply_book_settings, check_version, read_book
    from evm.wbs import build_part, build_tree, select_scope
    from evm.workbook import (book_sources, company_sheet, read_meta, read_structure, row_key, source_meta,
                              split_key)

    wb = load_workbook(book_path)
    meta = read_meta(wb)
    check_version(meta)
    if not projects:
        names = book_sources(meta)
        projects = {"": single} if names == [""] else {n: i + 1 for i, n in enumerate(names[:2])}
    cfg = apply_book_settings(Config(), meta)
    cfg.max_level = int(meta.get("nivel_maximo") or cfg.max_level)
    multi = bool(meta.get("multiempresa"))
    roles = [("A", company_sheet("PV", cfg.company_a), company_sheet("AC", cfg.company_a)),
             ("B", company_sheet("PV", cfg.company_b), company_sheet("AC", cfg.company_b))] if multi \
        else [("", "PV", "AC")]
    structure = [s for s in read_structure(wb) if s.is_input]
    horizon = Horizon(meta["horizonte_inicio"], meta["meses"])
    bac_cells = ac_cells = 0
    with tempfile.TemporaryDirectory() as tmp:
        for source, number in projects.items():
            pcfg = replace(cfg, plan=str(source_meta(meta, "planificacion", source) or ""))
            plans = PROJECTS[number][2]()
            path = Path(tmp) / f"ejemplo{number}.xml"
            write(plans, path, status, number)          # numera las tareas y, con fecha, simula el avance
            model = {t.uid: t for t in _every(plans)}   # costo y costo real simulado de cada tarea del ejemplo
            project = read_xml(path, pcfg)
            rows = {row_key(r): r for r in build_part(project, pcfg).rows if r.is_input}
            scope = select_scope(build_tree([t for t in project.tasks if t.outline_level >= 1]), pcfg.plan)
            shares = resolve_shares(scope.roots, pcfg)[0] if multi else {}

            def split(value) -> dict[str, float]:
                """Monto de la fila por empresa, según la parte de cada empresa en sus tareas de detalle."""
                total = sum(value(model[t.uid]) for t in leaves)
                a = sum(value(model[t.uid]) * shares.get(t.uid, 0.5) for t in leaves)
                return {"": total, "A": a, "B": total - a}

            idx = horizon.index(status) if status else -1
            for s in (x for x in structure if split_key(x.key)[0] == source):
                _src, uid, residual = split_key(s.key)
                row = rows.get(uid + ("-D" if residual else ""))
                if row is None:
                    continue
                leaves = row.leaves
                budget = split(lambda t: t.cost)
                for role, pv_sheet, _ac in roles:
                    cell = wb[pv_sheet].cell(row=s.row, column=4)
                    if cell.value in (None, "") and budget[role] > 0.005:
                        cell.value = round(budget[role], 2)
                        bac_cells += 1
                if not 0 <= idx < horizon.months:
                    continue
                actual = split(lambda t: t.actual_cost)          # acumulado a la fecha de estado
                col = 7 + idx
                for role, _pv, ac_sheet in roles:
                    ws = wb[ac_sheet]
                    prev = sum(float(ws.cell(row=s.row, column=7 + i).value or 0) for i in range(idx))
                    value = round(actual[role] - prev, 2)
                    if actual[role] or ws.cell(row=s.row, column=col).value is not None:
                        ws.cell(row=s.row, column=col, value=value)
                        ac_cells += 1
    if status is None:
        fill_review(wb, meta)
    # openpyxl no conserva los gráficos al guardar: se vuelven a crear
    book = read_book(book_path, wb)
    when = status or book.status_date
    _refresh_charts(wb, book, book.horizon.index(when) if when is not None else None, meta)
    _save(wb, Path(book_path), False)
    return bac_cells, ac_cells


def fill_review(wb, meta: dict) -> int:
    """Libro con depuración de costes: escribe en la hoja PV-depuración (si está vacía) una previsión ficticia de
    la empresa, como la pegaría desde su plan de costes: el PV planificado de cada paquete con algunos meses
    desplazados y algunos totales distintos del BAC, un paquete con hitos de pago (HP-01...) y otro (y los paquetes
    «Otras tareas») con subtareas sin código numérico (XRF01...). Devuelve la cantidad de líneas escritas."""
    from evm import cost_review
    from evm.timeline import round_preserving_total
    from evm.workbook import FIRST_MONTH_COL, PCT_PREFIX, read_structure

    if cost_review.SHEET_PASTE not in wb.sheetnames:
        return 0
    paste = wb[cost_review.SHEET_PASTE]
    first = cost_review.P_FIRST
    if any(paste.cell(row=first + i, column=1).value not in (None, "") for i in range(5)):
        return 0
    sheet = str(meta.get("depuracion_hoja") or "PV")
    pv, pct = wb[sheet], wb[PCT_PREFIX + sheet]
    months = int(meta["meses"])
    shifts = [0, 1, 0, -1, 2, 0, 1, 0]
    scales = [1.0, 1.0, 0.95, 1.0, 1.0, 1.08, 1.0, 1.0]
    lines: list[list] = []
    k = 0
    for s in read_structure(wb):
        bac = pv.cell(row=s.row, column=4).value
        if not s.is_input or not isinstance(bac, (int, float)) or bac <= 0:
            continue
        shares = [pct.cell(row=s.row, column=FIRST_MONTH_COL + i).value or 0.0 for i in range(months)]
        moved = [0.0] * months
        for i, v in enumerate(shares):
            moved[min(months - 1, max(0, i + shifts[k % len(shifts)]))] += v * bac * scales[k % len(scales)]
        values = round_preserving_total(moved, 0)
        code = s.code[:-2] if s.kind == "residual" else s.code
        if s.kind == "residual" or k == 4:  # subtareas sin código numérico bajo el código padre
            parts = [values[:], [0.0] * months]
            for i, v in enumerate(values):
                parts[0][i], parts[1][i] = round(v * 0.6), v - round(v * 0.6)
            lines.append([code, s.name.replace("Otras tareas de ", "Subtareas de "), None])
            lines += [[f"XRF{j + 1:02d}", f"Subtarea {j + 1}", sum(p)] + p for j, p in enumerate(parts)]
        elif k == 1:                        # hitos de pago: 30 %, 40 % y 30 % en tres meses del paquete
            active = [i for i, v in enumerate(values) if v]
            lines.append([code, s.name, None])
            if active:
                picks = [active[len(active) // 3], active[2 * len(active) // 3], active[-1]]
                for j, (i, part) in enumerate(zip(picks, (0.3, 0.4, 0.3))):
                    row = [0.0] * months
                    row[i] = round(sum(values) * part)
                    lines.append([f"HP-{j + 1:02d}", f"Hito de pago {j + 1}", row[i]] + row)
        else:
            lines.append([code, s.name, sum(values)] + values)
        k += 1
    for r, line in enumerate(lines, start=first):
        for c, v in enumerate(line, start=1):
            if v not in (None, 0, 0.0) or c <= 2:
                cell = paste.cell(row=r, column=c, value=v)
                if c >= 3:
                    cell.number_format = "#,##0"
    return len(lines)


def main() -> None:
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--proyecto", type=int, choices=sorted(PROJECTS), default=1,
                   help="1 = programa de ejemplo (por defecto), 2 = segundo proyecto (código EDT 9.0.0)")
    p.add_argument("--salida", help="archivo XML a crear (por defecto el archivo de ejemplo del proyecto)")
    p.add_argument("--fecha-estado", help="AAAA-MM-DD: agrega avance simulado a esa fecha")
    p.add_argument("--libro", help="libro del exportador: escribe en él el BAC de cada fila (si está vacío) y, "
                                   "con --fecha-estado, el AC del mes, con datos ficticios del ejemplo; en un libro "
                                   "con depuración de costes, también la previsión de la hoja PV-depuración")
    a = p.parse_args()
    status = date.fromisoformat(a.fecha_estado) if a.fecha_estado else None
    _title, file_name, build, _first = PROJECTS[a.proyecto]
    if a.libro and not (a.salida or a.fecha_estado):
        bac, ac = fill_book(Path(a.libro), single=a.proyecto)
        print(f"Libro {a.libro}: BAC escrito en {bac} celdas.")
        return
    salida = a.salida or str(Path(__file__).with_name(file_name))
    write(build(), Path(salida), status, a.proyecto)
    print(f"Archivo generado: {salida}")
    if a.libro:
        bac, ac = fill_book(Path(a.libro), status, single=a.proyecto)
        print(f"Libro {a.libro}: BAC escrito en {bac} celdas" + (f"; AC de {status:%m/%Y} en {ac} celdas." if status
                                                                  else "."))


if __name__ == "__main__":
    main()
