"""Escribe la versión del programa en la portada de la presentación y vuelve a exportar el PDF.

La versión se toma de evm/__init__.py y se escribe en la forma "Versión" de la primera diapositiva
("Versión 1.5.0  ·  Exportador EVM"). El PDF se exporta con PowerPoint (automatización COM).

Requisitos: python-pptx (python -m pip install python-pptx); para el PDF, Windows con PowerPoint y pywin32.

Uso (desde la carpeta raíz del repositorio):
    python docs\\herramientas\\version_presentacion.py
"""
from __future__ import annotations

import copy
import re
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

ROOT = Path(__file__).resolve().parents[2]
DECK = ROOT / "docs" / "presentacion-jefes-de-proyecto.pptx"
NAME = "Versión"


def program_version() -> str:
    text = (ROOT / "evm" / "__init__.py").read_text(encoding="utf-8")
    return re.search(r'__version__ = "([^"]+)"', text).group(1)


def set_version(deck: Path, version: str) -> None:
    prs = Presentation(deck)
    cover = prs.slides[0]
    shape = next((s for s in cover.shapes if s.name == NAME), None)
    if shape is None:
        # Primera vez: copia de la línea "30 minutos · PV, EV y AC..." (mismo estilo), debajo del presentador
        model = next(s for s in cover.shapes if s.has_text_frame and s.text_frame.text.startswith("30 minutos"))
        el = copy.deepcopy(model._element)
        model._element.addnext(el)
        el.nvSpPr.cNvPr.id = max(s.shape_id for s in cover.shapes) + 1
        el.nvSpPr.cNvPr.name = NAME
        shape = next(s for s in cover.shapes if s._element is el)
        shape.top = Inches(6.80)
    runs = shape.text_frame.paragraphs[0].runs
    runs[0].text = f"Versión {version}"
    runs[1].text = "  ·  Exportador EVM"
    for r in runs[2:]:
        r._r.getparent().remove(r._r)
    prs.save(deck)


def export_pdf(deck: Path) -> Path:
    import pythoncom
    import win32com.client
    pdf = deck.with_suffix(".pdf")
    pythoncom.CoInitialize()
    app = win32com.client.DispatchEx("PowerPoint.Application")
    try:
        pres = app.Presentations.Open(str(deck), ReadOnly=True, Untitled=False, WithWindow=False)
        pres.SaveAs(str(pdf), 32)          # ppSaveAsPDF
        pres.Close()
    finally:
        app.Quit()
    return pdf


def main() -> None:
    version = program_version()
    set_version(DECK, version)
    print(f"Portada: Versión {version}")
    try:
        print("PDF:", export_pdf(DECK))
    except ImportError:
        print("Sin pywin32: el PDF se debe exportar desde PowerPoint (Archivo > Exportar > PDF).")


if __name__ == "__main__":
    main()
