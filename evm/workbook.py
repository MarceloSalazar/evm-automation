"""Generación y lectura del libro Excel de Valor Ganado."""
from __future__ import annotations

import io
import os
import shutil
import tempfile
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.cell.cell import ERROR_CODES, ILLEGAL_CHARACTERS_RE
from openpyxl.chart.axis import DateAxis
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import from_excel
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.properties import Outline

from evm import __version__
from evm.config import MODE_PERCENT, Config
from evm.wbs import KIND_LEAF, KIND_RESIDUAL, KIND_SUMMARY, Report, ReportRow

# ----------------------------------------------------------------------------
# Disposición del libro (no modificar: el comando "actualizar" depende de ella)
# ----------------------------------------------------------------------------
SHEET_SUMMARY = "Resumen"
SHEET_PV = "PV"
SHEET_EV = "EV"
SHEET_AC = "AC"
SHEET_CURVE = "Evolución mensual"
LEGACY_SHEET_CURVE = "Curva S"      # libros generados con versiones anteriores a la 1.3.0
SHEET_IND = "Indicadores EDT"
SHEET_VALID = "Validación"
SHEET_LISTS = "Listas"
SHEET_STRUCT = "_estructura"
SHEET_META = "_meta"
META_SOURCES = "proyectos"          # _meta: nombres de los archivos de un libro con varios archivos ("P1;P2")
PCT_PREFIX = "% "                   # hojas '% PV' y '% EV' (o '% PV-EmpresaA'...)
KIND_COST = "coste"                 # _estructura: fila «(coste)» de la depuración de costes (sin clave)
COST_SUFFIX = " (coste)"

TITLE_ROW, INFO_ROW, MONTHNUM_ROW, HEADER_ROW, FIRST_ROW = 1, 2, 3, 4, 5
COL_CODE, COL_NAME, COL_METHOD, COL_BAC, COL_START, COL_FINISH = 1, 2, 3, 4, 5, 6
FIRST_MONTH_COL = 7  # columna G

# Evolución mensual (tabla vertical; datos de la curva S)
CURVE_HEADER_ROW = 4
CURVE_FIRST_ROW = 5
(C_NUM, C_DATE, C_PV, C_PVC, C_EV, C_EVC, C_AC, C_ACC,
 C_PVPCT, C_EVPCT, C_SPI, C_CPI, C_BAC, C_EVN, C_ACN) = range(1, 16)
CHART_ANCHOR = "Q4"

# Resumen
STATUS_CELL = "B5"
CUT_MONTH_CELL = "B6"
NAME_METHODS = "ListaMetodos"
START_CELL = "B7"
NAME_STATUS = "FechaCorte"
NAME_START = "InicioHorizonte"

# ----------------------------------------------------------------------------
# Estilos
# ----------------------------------------------------------------------------
DARK = "1F3864"
FILL_HEADER = PatternFill("solid", fgColor=DARK)
FILL_L1 = PatternFill("solid", fgColor="BDD7EE")
FILL_L2 = PatternFill("solid", fgColor="DDEBF7")
FILL_INPUT = PatternFill("solid", fgColor="FFF2CC")
FILL_TOTAL = PatternFill("solid", fgColor="D9D9D9")
# Colores de semáforo (fgColor y bgColor: Excel usa bgColor en el formato condicional)
FILL_RED = PatternFill("solid", fgColor="F8CBAD", bgColor="F8CBAD")
FILL_YELLOW = PatternFill("solid", fgColor="FFE699", bgColor="FFE699")
FILL_GREEN = PatternFill("solid", fgColor="C6E0B4", bgColor="C6E0B4")
FILL_SECTION = PatternFill("solid", fgColor="D6DCE4")
FILL_COST = PatternFill("solid", fgColor="FCE4D6")      # filas «(coste)» de la depuración de costes
COLOR_COST = "833C0B"
FONT_HEADER = Font(bold=True, color="FFFFFF")
FONT_TITLE = Font(bold=True, size=14, color=DARK)
FONT_BOLD = Font(bold=True)
FONT_NOTE = Font(italic=True, size=9, color="595959")
FONT_ITALIC = Font(italic=True)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
FMT_DATE = "dd/mm/yyyy"
FMT_MONTH = "mmm-yyyy"
FMT_RATIO = "0.00"
FMT_PCT = "0.0%"
FMT_PCT_MONTH = '0.0%;-0.0%;""'      # % del mes sin mostrar los meses en 0

# Regla de cálculo de cada método (el avance se pondera por la duración de las tareas y el peso de los hitos)
METHOD_HELP = {
    "Por definir": "Aún no se asignó método. Para calcular el % ganado se usa % completado.",
    "0/100": "0 % al comenzar y 100 % al terminar. Para tareas que empiezan y terminan en el mismo mes.",
    "50/50": "50 % al comenzar y 50 % al terminar. Para tareas de hasta 2 meses.",
    "25/75": "25 % al comenzar y 75 % al terminar. Tareas de hasta 2 meses con el trabajo concentrado al final.",
    "20/80": "20 % al comenzar y 80 % al terminar. Tareas de hasta 2 meses con el trabajo concentrado al final.",
    "Hitos ponderados": "% ganado = suma de los 'Peso hito' de los hitos cumplidos / suma de todos los pesos. Para "
                        "paquetes de 3 o más meses, con al menos un hito por mes; los pesos deben sumar 100.",
    "% Completado": "% ganado = % completado de MS Project de las subtareas (mide % de duración), ponderado por su "
                    "duración de línea base; un hito cuenta su 'Peso hito' al cumplirse.",
    "% Físico completado": "% ganado = '% físico completado' de MS Project de las subtareas (avance medido "
                           "objetivamente), ponderado por su duración de línea base y el peso de los hitos.",
    "Unidades completadas": "% ganado = unidades terminadas / unidades totales. El cociente se registra en "
                            "'% físico completado' de MS Project.",
    "Avance lineal (LOE)": "Nivel de esfuerzo: % ganado = % planificado (EV = PV). No genera variación de plazo; "
                           "conviene limitarlo.",
}


def row_key(row: ReportRow) -> str:
    """Clave estable de una fila: UniqueID de la tarea en MS Project ('123', o '123-D' en una fila de otros
    costos). Con varios archivos, el UniqueID se repite entre archivos: la clave lleva el nombre del archivo
    ('P2:123')."""
    key = f"{row.task.uid}-D" if row.kind == KIND_RESIDUAL else str(row.task.uid)
    return f"{row.source}:{key}" if row.source else key


def split_key(key: str) -> tuple[str, str, bool]:
    """Partes de una clave de fila: (archivo, UniqueID, fila de otros costos). 'P2:123-D' -> ('P2', '123', True)."""
    source, _, rest = str(key).rpartition(":")
    residual = rest.endswith("-D")
    return source, rest[:-2] if residual else rest, residual


def money_format(decimals: int) -> str:
    return "#,##0" if decimals <= 0 else "#,##0." + "0" * decimals


def money_format_hide_zero(decimals: int) -> str:
    """Formato de importe que no muestra los ceros (meses calculados como BAC × % sin valor)."""
    m = money_format(decimals)
    return f'{m};-{m};""'


def col(c: int) -> str:
    return get_column_letter(c)


def clean_text(value):
    """Quita caracteres de control que Excel no admite."""
    return ILLEGAL_CHARACTERS_RE.sub("", value) if isinstance(value, str) else value


def put_text(ws, row: int, column: int, value):
    """Escribe un texto libre (nombre de tarea, código...) siempre como texto, nunca como fórmula."""
    cell = ws.cell(row=row, column=column, value=clean_text(value))
    if isinstance(cell.value, str) and (cell.value.startswith("=") or cell.value in ERROR_CODES):
        cell.data_type = "s"
    return cell


def _texts_in_last_row(ws) -> None:
    for cell in ws[ws.max_row]:
        if isinstance(cell.value, str):
            cell.value = clean_text(cell.value)
            if cell.value.startswith("=") or cell.value in ERROR_CODES:
                cell.data_type = "s"


MAX_FORMULA = 8000      # Excel admite fórmulas de hasta 8.192 caracteres


def _sum_formula(letter: str, rows: list[int], block: tuple[int, int] | None = None, review: bool = False,
                 cost: bool = False) -> str:
    """=SUM de las filas indicadas (las consecutivas se agrupan en rangos).

    Si la fórmula resultara demasiado larga, suma las filas de ingreso del bloque
    `block` con SUMIFS sobre el tipo de fila de la hoja _estructura.
    review: libro con depuración de costes (las filas «(coste)» del bloque no se suman).
    cost: suma de filas «(coste)»: se suman las filas «(coste)» de las filas de ingreso del bloque.
    """
    if not rows:
        return "=0"
    rows = sorted(rows)
    refs: list[str] = []
    start = prev = rows[0]
    for r in rows[1:] + [None]:
        if r is not None and r == prev + 1:
            prev = r
            continue
        refs.append(f"{letter}{start}" if start == prev else f"{letter}{start}:{letter}{prev}")
        if r is not None:
            start = prev = r
    chunks = ["SUM(" + ",".join(refs[i:i + 250]) + ")" for i in range(0, len(refs), 250)]
    formula = "=" + (chunks[0] if len(chunks) == 1 else "SUM(" + ",".join(chunks) + ")")
    if len(formula) > MAX_FORMULA and block is not None:
        a, b = block
        off = FIRST_ROW - 2   # fila de datos r <-> fila r - off de _estructura
        kinds = f"'{SHEET_STRUCT}'!$E${a - off}:$E${b - off}"
        if cost:              # filas «(coste)» cuya fila de la EDT (la de arriba) es de ingreso
            above = f"'{SHEET_STRUCT}'!$E${a - off - 1}:$E${b - off - 1}"
            return (f"=SUMIFS({letter}{a}:{letter}{b},{kinds},\"{KIND_COST}\",{above},"
                    f"\"<>{KIND_SUMMARY}\")")
        formula = f"=SUMIFS({letter}{a}:{letter}{b},{kinds},\"<>{KIND_SUMMARY}\""
        formula += f",{kinds},\"<>{KIND_COST}\")" if review else ")"
    return formula


def _header(ws, row: int, labels: list[tuple[int, str]], height: float = 30) -> None:
    for c, text in labels:
        cell = ws.cell(row=row, column=c, value=text)
        cell.fill = FILL_HEADER
        cell.font = FONT_HEADER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[row].height = height


def _naive(d):
    if isinstance(d, datetime):
        return d.replace(tzinfo=None)
    return d


# ----------------------------------------------------------------------------
# Escritura
# ----------------------------------------------------------------------------
class _Layout:
    """Filas y columnas de las hojas mensuales (las mismas en todas las hojas).

    Con depuración de costes, debajo de cada fila de la EDT (y de la fila TOTAL) va su fila «(coste)»: la hoja que
    se depura la completa y en las demás hojas queda vacía y oculta. Así cada fila conserva la misma posición en
    todas las hojas.
    """

    def __init__(self, report: Report):
        self.months = report.horizon.months
        self.first_col = FIRST_MONTH_COL
        self.last_col = FIRST_MONTH_COL + self.months - 1
        self.total_col = self.last_col + 1
        self.review = report.cost_review is not None
        self.stride = 2 if self.review else 1
        self.row_of: dict[int, int] = {}
        for i, r in enumerate(report.rows):
            self.row_of[id(r)] = FIRST_ROW + self.stride * i
        self.last_row = FIRST_ROW + self.stride * len(report.rows) - 1
        self.total_row = self.last_row + 1
        self.block_end: dict[int, int] = {}

        def size(row) -> int:
            n = 1 + sum(size(c) for c in row.children)
            self.block_end[id(row)] = self.row_of[id(row)] + self.stride * n - 1
            return n

        for root in report.roots:
            size(root)
        self.month_dates = [report.horizon.month_start(i) for i in range(self.months)]

    def r(self, row: ReportRow) -> int:
        return self.row_of[id(row)]

    def cost_rows(self, report: Report) -> list[int]:
        """Depuración de costes: filas «(coste)» (debajo de cada fila de la EDT y de la fila TOTAL)."""
        return [self.r(x) + 1 for x in report.rows] + [self.total_row + 1] if self.review else []

    def struct_row(self, r: int) -> int:
        """Fila de la hoja _estructura que corresponde a la fila r de las hojas mensuales."""
        return r - (FIRST_ROW - 2)

    def block(self, row: ReportRow) -> tuple[int, int]:
        """Filas de los descendientes de una fila resumen."""
        return self.row_of[id(row)] + 1, self.block_end[id(row)]

    def month_range(self, sheet: str, row: int) -> str:
        return f"'{sheet}'!${col(self.first_col)}{row}:${col(self.last_col)}{row}"

    def header_range(self, sheet: str) -> str:
        return f"'{sheet}'!${col(self.first_col)}${HEADER_ROW}:${col(self.last_col)}${HEADER_ROW}"


def _info_line(report: Report, cfg: Config, generated: datetime) -> str:
    p = report.project
    cur = cfg.currency or report.currency or "-"
    companies = f"Empresas: {report.company_a} y {report.company_b}  |  " if report.multi else ""
    if report.multi_file:
        files = ", ".join(f"{s.name} = {Path(s.project.source).name}" for s in report.sources)
        return (f"Informe: {report.title}  |  {companies}Archivos: {files}  |  Moneda: {cur}  |  "
                f"Generado: {generated:%d/%m/%Y %H:%M}")
    plan = f"Planificación: {report.plan}  |  " if report.plan else ""
    return (f"Proyecto: {report.title}  |  {plan}{companies}Origen: {Path(p.source).name}  |  {p.baseline_label}  |  "
            f"Moneda: {cur}  |  Generado: {generated:%d/%m/%Y %H:%M}")


def per_file(report: Report, value) -> str:
    """Dato de cada archivo de un informe con varios archivos: 'P1: valor  |  P2: valor'."""
    return "  |  ".join(f"{s.name}: {value(s)}" for s in report.sources)


def company_sheet(kind: str, company: str) -> str:
    """Nombre de la hoja de una empresa: 'PV-EmpresaA'."""
    return f"{kind}-{company}"


def pct_sheet(kind: str, company: str = "") -> str:
    """Hoja con el % planificado o ganado de cada mes: '% PV' o '% PV-EmpresaA'."""
    return PCT_PREFIX + (company_sheet(kind, company) if company else kind)


def role_key(kind: str, role: str | None) -> str:
    """Serie de una hoja para --migrar-desde: 'EV', 'EV-A'..."""
    return kind if not role else f"{kind}-{role}"


def _write_data_sheet(ws, kind: str, report: Report, cfg: Config, lay: _Layout,
                      generated: datetime, prior=None, role: str | None = None) -> None:
    """Escribe una hoja mensual (PV, EV o AC).

    role: None = hoja total; "A" / "B" = hoja de cada empresa (solo multi-empresa).
    En multi-empresa la hoja total suma, celda por celda, las hojas de A y de B.
    """
    multi_total = report.multi and role is None
    company = {"A": report.company_a, "B": report.company_b}.get(role or "", "")
    is_pv = kind == SHEET_PV
    # PV y EV = BAC × % de la hoja '% PV' o '% EV'; el BAC se ingresa en la columna D de las hojas PV
    by_pct = kind in (SHEET_PV, SHEET_EV) and not multi_total
    pct_name = pct_sheet(kind, company) if by_pct else ""
    bac_input = is_pv and not multi_total
    input_sheet = kind == SHEET_AC and not multi_total    # celdas amarillas de ingreso: el AC del mes
    month_key = kind if role is None else f"{kind}-{role}"   # serie para --migrar-desde
    # Depuración de costes: la hoja PV de la empresa que se depura lleva las filas «(coste)»
    review_sheet = lay.review and is_pv and not multi_total and (role or "") == report.cost_review

    names = {SHEET_PV: "Valor Planificado (PV)", SHEET_EV: "Valor Ganado (EV)", SHEET_AC: "Costo Real (AC)"}
    if multi_total:
        title = (f"{names[kind]} por mes: suma de las hojas {company_sheet(kind, report.company_a)} y "
                 f"{company_sheet(kind, report.company_b)}")
    elif bac_input:
        title = (f"{names[kind]} por mes" + (f"  -  {company}" if company else "") +
                 f"  -  BAC × % planificado (hoja '{pct_name}'); el BAC se debe ingresar en las celdas amarillas")
        if review_sheet:
            from evm.cost_review import SHEET_PASTE
            title += f"  -  filas «(coste)»: previsión pegada en la hoja '{SHEET_PASTE}'"
    elif by_pct:
        title = (f"{names[kind]} por mes" + (f"  -  {company}" if company else "") +
                 f"  -  BAC × % ganado (hoja '{pct_name}'), que 'actualizar' calcula desde MS Project")
    else:
        title = (f"{names[kind]} por mes" + (f"  -  {company}" if company else "") +
                 "  -  en las celdas amarillas se debe ingresar el valor DEL MES (no acumulado)")
    money = money_format(cfg.decimals)
    month_money = money_format_hide_zero(cfg.decimals) if kind != SHEET_AC else money
    ws.cell(row=TITLE_ROW, column=1, value=title).font = FONT_TITLE
    put_text(ws, INFO_ROW, 1, _info_line(report, cfg, generated)).font = FONT_NOTE

    # Fila 3: número de mes; fila 4: encabezados y fecha de cada mes
    ws.cell(row=MONTHNUM_ROW, column=COL_FINISH, value="Mes N°").font = FONT_BOLD
    ws.cell(row=MONTHNUM_ROW, column=COL_FINISH).alignment = Alignment(horizontal="right")
    for i in range(lay.months):
        c = ws.cell(row=MONTHNUM_ROW, column=lay.first_col + i, value=i + 1)
        c.alignment = Alignment(horizontal="center")
        c.font = FONT_BOLD
    _header(ws, HEADER_ROW, [
        (COL_CODE, "Código EDT"), (COL_NAME, "Actividad"), (COL_METHOD, "Método de medición"),
        (COL_BAC, "BAC (Presupuesto)" + (f" {company}" if company else "")),
        (COL_START, "Fecha inicio Plan (PV)"), (COL_FINISH, "Fecha fin Plan (PV)"),
        (lay.total_col, "Total horizonte"),
    ])
    for i, d in enumerate(lay.month_dates):
        c = ws.cell(row=HEADER_ROW, column=lay.first_col + i, value=d)
        c.number_format = FMT_MONTH
        c.fill = FILL_HEADER
        c.font = FONT_HEADER
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = BORDER

    def sum_of_companies(letter: str, r: int) -> str:
        return (f"='{company_sheet(kind, report.company_a)}'!{letter}{r}"
                f"+'{company_sheet(kind, report.company_b)}'!{letter}{r}")

    for row in report.rows:
        r = lay.r(row)
        put_text(ws, r, COL_CODE, row.code)
        name = put_text(ws, r, COL_NAME, row.name)
        name.alignment = Alignment(indent=max(0, row.level - 1))

        # Columnas C a F
        if is_pv and role is None:
            method = (prior.method_for(row_key(row)) if prior is not None and row.is_input else None) or row.method
            put_text(ws, r, COL_METHOD, method if row.is_input else None)
            ws.cell(row=r, column=COL_START, value=_naive(row.start))
            ws.cell(row=r, column=COL_FINISH, value=_naive(row.finish))
        else:
            for c_ in (COL_METHOD, COL_START, COL_FINISH):
                ref = f"'{SHEET_PV}'!{col(c_)}{r}"
                ws.cell(row=r, column=c_, value=f'=IF({ref}="","",{ref})')
        if is_pv and multi_total:
            ws.cell(row=r, column=COL_BAC, value=sum_of_companies(col(COL_BAC), r))
        elif bac_input and row.is_input:
            bac_cell = ws.cell(row=r, column=COL_BAC)
            value = prior.bac_for(row_key(row), role or "") if prior is not None else None
            if value is not None:
                bac_cell.value = round(value, 2)
            bac_cell.fill = FILL_INPUT
        elif is_pv:                     # subtotales: suma de sus hijos
            ws.cell(row=r, column=COL_BAC,
                    value=_sum_formula(col(COL_BAC), [lay.r(c) for c in row.children], lay.block(row), lay.review))
        else:
            pv_sheet = SHEET_PV if role is None else company_sheet(SHEET_PV, company)
            ref = f"'{pv_sheet}'!{col(COL_BAC)}{r}"
            ws.cell(row=r, column=COL_BAC, value=f'=IF({ref}="","",{ref})')
        ws.cell(row=r, column=COL_BAC).number_format = money
        ws.cell(row=r, column=COL_START).number_format = FMT_DATE
        ws.cell(row=r, column=COL_FINISH).number_format = FMT_DATE

        # Columnas mensuales
        if multi_total:
            for i in range(lay.months):
                ws.cell(row=r, column=lay.first_col + i, value=sum_of_companies(col(lay.first_col + i), r))
        elif row.is_input and by_pct:
            for i in range(lay.months):
                L = col(lay.first_col + i)
                ws.cell(row=r, column=lay.first_col + i, value=f"=N($D{r})*'{pct_name}'!{L}{r}")
        elif row.is_input:                  # AC del mes: se ingresa a mano (o se traslada con --migrar-desde)
            if prior is not None:
                for k, v in prior.values_for(row_key(row), month_key).items():
                    i = report.horizon.index(k)
                    if 0 <= i < lay.months:
                        ws.cell(row=r, column=lay.first_col + i, value=round(v, 2))
            if input_sheet:
                for i in range(lay.months):
                    ws.cell(row=r, column=lay.first_col + i).fill = FILL_INPUT
        else:
            refs_rows = [lay.r(c) for c in row.children]
            for i in range(lay.months):
                L = col(lay.first_col + i)
                ws.cell(row=r, column=lay.first_col + i,
                        value=_sum_formula(L, refs_rows, lay.block(row), lay.review))
        for i in range(lay.months):
            ws.cell(row=r, column=lay.first_col + i).number_format = month_money
        tot = ws.cell(row=r, column=lay.total_col,
                      value=f"=SUM({col(lay.first_col)}{r}:{col(lay.last_col)}{r})")
        tot.number_format = money

        # Estilo por nivel
        font = FONT_BOLD if row.kind == KIND_SUMMARY else (FONT_ITALIC if row.kind == KIND_RESIDUAL else None)
        fill = FILL_L1 if (row.kind == KIND_SUMMARY and row.level == 1) else (
            FILL_L2 if row.kind == KIND_SUMMARY else None)
        for c_ in range(1, lay.total_col + 1):
            cell = ws.cell(row=r, column=c_)
            if font:
                cell.font = font
            if fill and not (row.is_input and input_sheet and c_ >= lay.first_col):
                cell.fill = fill
        ws.row_dimensions[r].outline_level = min(7, max(0, row.level - 1))
        if review_sheet:
            _write_cost_row(ws, report, cfg, lay, row, r + 1)
        elif lay.review:
            ws.row_dimensions[r + 1].hidden = True

    # Fila de total
    tr = lay.total_row
    ws.cell(row=tr, column=COL_CODE, value="TOTAL")
    total = "Total programa" if report.multi_file else "Total proyecto"
    ws.cell(row=tr, column=COL_NAME, value=total + (f" - {company}" if company else ""))
    root_rows = [lay.r(x) for x in report.roots]
    for c_ in [COL_BAC] + list(range(lay.first_col, lay.last_col + 1)):
        L = col(c_)
        if multi_total:
            value = sum_of_companies(L, tr)
        else:
            value = _sum_formula(L, root_rows, (FIRST_ROW, lay.last_row), lay.review)
        ws.cell(row=tr, column=c_, value=value).number_format = money if c_ == COL_BAC else month_money
    ws.cell(row=tr, column=lay.total_col,
            value=f"=SUM({col(lay.first_col)}{tr}:{col(lay.last_col)}{tr})").number_format = money
    for c_ in range(1, lay.total_col + 1):
        cell = ws.cell(row=tr, column=c_)
        cell.fill = FILL_TOTAL
        cell.font = FONT_BOLD
        cell.border = Border(top=Side(style="medium", color=DARK))
    if review_sheet:
        _write_cost_row(ws, report, cfg, lay, None, tr + 1, total + (f" - {company}" if company else ""))

    if is_pv:
        # Resaltar filas cuyo total en el horizonte no coincide con el BAC
        rng = f"{col(lay.total_col)}{FIRST_ROW}:{col(lay.total_col)}{lay.total_row}"
        ws.conditional_formatting.add(rng, FormulaRule(
            formula=[f"ABS({col(lay.total_col)}{FIRST_ROW}-$D{FIRST_ROW})>0.5"], fill=FILL_RED))
    if is_pv and role is None and report.rows:
        dv = DataValidation(type="list", formula1=NAME_METHODS,
                            allow_blank=True, showErrorMessage=True, errorStyle="warning",
                            errorTitle="Método de medición",
                            error="El valor no está en la lista de la hoja 'Listas'. ¿Mantener el valor?")
        ws.add_data_validation(dv)
        for row in report.rows:
            if row.is_input:
                dv.add(f"{col(COL_METHOD)}{lay.r(row)}")

    widths = {COL_CODE: 12, COL_NAME: 48, COL_METHOD: 20, COL_BAC: 16, COL_START: 12, COL_FINISH: 12}
    for c_, w in widths.items():
        ws.column_dimensions[col(c_)].width = w
    for i in range(lay.months):
        ws.column_dimensions[col(lay.first_col + i)].width = 12
    ws.column_dimensions[col(lay.total_col)].width = 15
    ws.freeze_panes = ws.cell(row=FIRST_ROW, column=lay.first_col)
    ws.sheet_properties.outlinePr = Outline(summaryBelow=False, summaryRight=False)
    ws.sheet_format.outlineLevelRow = min(7, max([r.level - 1 for r in report.rows] + [0]))
    if review_sheet:
        from evm import cost_review
        ws.sheet_format.outlineLevelRow = cost_outline(report)
        if cfg.cost_review_alignment:
            cost_review.alignment_format(ws, report, cfg, lay)


def cost_outline(report: Report) -> int:
    """Nivel de esquema de las filas «(coste)»: uno más que la fila más profunda de la EDT, para mostrarlas u
    ocultarlas todas con el botón de ese nivel."""
    return min(7, max([r.level - 1 for r in report.rows] + [0]) + 1)


def _write_cost_row(ws, report: Report, cfg: Config, lay: _Layout, row: ReportRow | None, c: int,
                    total_name: str = "") -> None:
    """Depuración de costes: fila «(coste)» c, debajo de la fila de la EDT (row; None = fila TOTAL).

    Filas de ingreso: cada mes suma la previsión pegada en la hoja PV-depuración para la fila. Subtotales y TOTAL:
    suma de las filas «(coste)» de sus hijos. Columna D: suma de los meses (se compara con el BAC de la fila de
    arriba).
    """
    from evm import cost_review
    money = money_format(cfg.decimals)
    month_money = money_format_hide_zero(cfg.decimals)
    r = c - 1
    if row is None:
        put_text(ws, c, COL_CODE, "TOTAL")
        put_text(ws, c, COL_NAME, total_name + COST_SUFFIX)
        children = [lay.r(x) + 1 for x in report.roots]
        block = (FIRST_ROW + 1, lay.last_row)
    else:
        put_text(ws, c, COL_CODE, row.code)
        put_text(ws, c, COL_NAME, row.name + COST_SUFFIX).alignment = Alignment(indent=max(0, row.level - 1))
        children = [lay.r(x) + 1 for x in row.children]
        a, b = lay.block(row)
        block = (a + 1, b)
    for i in range(lay.months):
        if row is not None and row.is_input:
            value = cost_review.paste_sum(i, r)
        else:
            value = _sum_formula(col(lay.first_col + i), children, block, cost=True)
        ws.cell(row=c, column=lay.first_col + i, value=value).number_format = month_money
    span = f"{col(lay.first_col)}{c}:{col(lay.last_col)}{c}"
    ws.cell(row=c, column=COL_BAC, value=f"=SUM({span})").number_format = money
    ws.cell(row=c, column=lay.total_col, value=f"=SUM({span})").number_format = money
    font = Font(italic=True, bold=row is None or row.kind == KIND_SUMMARY, color=COLOR_COST)
    for c_ in range(1, lay.total_col + 1):
        cell = ws.cell(row=c, column=c_)
        cell.fill = FILL_COST
        cell.font = font
    ws.row_dimensions[c].outline_level = cost_outline(report)
    if cfg.cost_review_alignment:
        cost_review.alignment_cells(ws, cfg, lay, r, c)


def _write_pct_sheet(ws, kind: str, report: Report, cfg: Config, lay: _Layout,
                     generated: datetime, prior=None, role: str | None = None) -> None:
    """Hoja '% PV' o '% EV' (de cada empresa en multi-empresa).

    Filas de ingreso: el % del mes (no acumulado) que calcula 'generar' (% PV) o 'actualizar' (% EV) desde MS
    Project; no se editan. Filas de nivel 1 y 2 y TOTAL: el % ponderado por el BAC (importe / BAC).
    """
    company = {"A": report.company_a, "B": report.company_b}.get(role or "", "")
    amount = company_sheet(kind, company) if company else kind
    is_pv = kind == SHEET_PV
    if is_pv:
        title = ("% planificado por mes" + (f"  -  {company}" if company else "") +
                 "  -  calculado por 'generar' desde MS Project (duración de la línea base y peso de los hitos)")
    else:
        title = ("% ganado por mes" + (f"  -  {company}" if company else "") +
                 "  -  valor DEL MES (no acumulado) que 'actualizar' calcula desde MS Project")
    ws.cell(row=TITLE_ROW, column=1, value=title).font = FONT_TITLE
    put_text(ws, INFO_ROW, 1, _info_line(report, cfg, generated)).font = FONT_NOTE
    ws.cell(row=MONTHNUM_ROW, column=COL_FINISH, value="Mes N°").font = FONT_BOLD
    ws.cell(row=MONTHNUM_ROW, column=COL_FINISH).alignment = Alignment(horizontal="right")
    for i in range(lay.months):
        c = ws.cell(row=MONTHNUM_ROW, column=lay.first_col + i, value=i + 1)
        c.alignment = Alignment(horizontal="center")
        c.font = FONT_BOLD
    _header(ws, HEADER_ROW, [
        (COL_CODE, "Código EDT"), (COL_NAME, "Actividad"), (COL_METHOD, "Método de medición"),
        (COL_BAC, "BAC (Presupuesto)" + (f" {company}" if company else "")),
        (COL_START, "Fecha inicio Plan (PV)"), (COL_FINISH, "Fecha fin Plan (PV)"),
        (lay.total_col, "Total horizonte" if is_pv else "% ganado acumulado"),
    ])
    for i, d in enumerate(lay.month_dates):
        c = ws.cell(row=HEADER_ROW, column=lay.first_col + i, value=d)
        c.number_format = FMT_MONTH
        c.fill = FILL_HEADER
        c.font = FONT_HEADER
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = BORDER
    money = money_format(cfg.decimals)
    series = role_key(pct_sheet(kind), role)            # serie para --migrar-desde ('% EV-A')
    for row in report.rows:
        r = lay.r(row)
        put_text(ws, r, COL_CODE, row.code)
        put_text(ws, r, COL_NAME, row.name).alignment = Alignment(indent=max(0, row.level - 1))
        for c_ in (COL_METHOD, COL_START, COL_FINISH):
            ref = f"'{SHEET_PV}'!{col(c_)}{r}"
            ws.cell(row=r, column=c_, value=f'=IF({ref}="","",{ref})')
        ref = f"'{amount}'!{col(COL_BAC)}{r}"
        ws.cell(row=r, column=COL_BAC, value=f'=IF({ref}="","",{ref})').number_format = money
        ws.cell(row=r, column=COL_START).number_format = FMT_DATE
        ws.cell(row=r, column=COL_FINISH).number_format = FMT_DATE
        if row.is_input:
            if is_pv:
                for i, v in enumerate(row.pct_vec.get(role or "", [])):
                    if abs(v) > 1e-12:
                        ws.cell(row=r, column=lay.first_col + i, value=v)
            elif prior is not None:
                for k, v in prior.values_for(row_key(row), series).items():
                    i = report.horizon.index(k)
                    if 0 <= i < lay.months:
                        ws.cell(row=r, column=lay.first_col + i, value=v)
        else:
            for i in range(lay.months):
                L = col(lay.first_col + i)
                ws.cell(row=r, column=lay.first_col + i,
                        value=f"=IFERROR('{amount}'!{L}{r}/'{amount}'!$D{r},\"\")")
        for i in range(lay.months):
            ws.cell(row=r, column=lay.first_col + i).number_format = FMT_PCT_MONTH
        ws.cell(row=r, column=lay.total_col,
                value=f"=SUM({col(lay.first_col)}{r}:{col(lay.last_col)}{r})").number_format = FMT_PCT
        font = FONT_BOLD if row.kind == KIND_SUMMARY else (FONT_ITALIC if row.kind == KIND_RESIDUAL else None)
        fill = FILL_L1 if (row.kind == KIND_SUMMARY and row.level == 1) else (
            FILL_L2 if row.kind == KIND_SUMMARY else None)
        for c_ in range(1, lay.total_col + 1):
            cell = ws.cell(row=r, column=c_)
            if font:
                cell.font = font
            if fill:
                cell.fill = fill
        ws.row_dimensions[r].outline_level = min(7, max(0, row.level - 1))
        if lay.review:
            ws.row_dimensions[r + 1].hidden = True

    tr = lay.total_row
    ws.cell(row=tr, column=COL_CODE, value="TOTAL")
    total = "Total programa" if report.multi_file else "Total proyecto"
    ws.cell(row=tr, column=COL_NAME, value=total + (f" - {company}" if company else ""))
    ref = f"'{amount}'!{col(COL_BAC)}{tr}"
    ws.cell(row=tr, column=COL_BAC, value=f'=IF({ref}="","",{ref})').number_format = money
    for i in range(lay.months):
        L = col(lay.first_col + i)
        ws.cell(row=tr, column=lay.first_col + i,
                value=f"=IFERROR('{amount}'!{L}{tr}/'{amount}'!$D{tr},\"\")").number_format = FMT_PCT_MONTH
    ws.cell(row=tr, column=lay.total_col,
            value=f"=SUM({col(lay.first_col)}{tr}:{col(lay.last_col)}{tr})").number_format = FMT_PCT
    for c_ in range(1, lay.total_col + 1):
        cell = ws.cell(row=tr, column=c_)
        cell.fill = FILL_TOTAL
        cell.font = FONT_BOLD
        cell.border = Border(top=Side(style="medium", color=DARK))
    if is_pv:
        # Filas de ingreso cuyo % planificado no suma 100 % (p. ej. tareas sin fechas o fuera del horizonte),
        # por bloques de filas de ingreso consecutivas
        # (con depuración de costes, entre las filas quedan las filas «(coste)», vacías en esta hoja)
        total_col = col(lay.total_col)
        inputs = sorted(lay.r(x) for x in report.rows if x.is_input)
        blocks: list[list[int]] = []
        for r in inputs:
            if blocks and r == blocks[-1][-1] + lay.stride:
                blocks[-1].append(r)
            else:
                blocks.append([r])
        for b in blocks:
            rule = f"ABS({total_col}{b[0]}-1)>0.0001"
            if lay.review:
                rule = f"AND(ISNUMBER({total_col}{b[0]}),{rule})"
            ws.conditional_formatting.add(f"{total_col}{b[0]}:{total_col}{b[-1]}", FormulaRule(
                formula=[rule], fill=FILL_RED))

    widths = {COL_CODE: 12, COL_NAME: 48, COL_METHOD: 20, COL_BAC: 16, COL_START: 12, COL_FINISH: 12}
    for c_, w in widths.items():
        ws.column_dimensions[col(c_)].width = w
    for i in range(lay.months):
        ws.column_dimensions[col(lay.first_col + i)].width = 11
    ws.column_dimensions[col(lay.total_col)].width = 15
    ws.freeze_panes = ws.cell(row=FIRST_ROW, column=lay.first_col)
    ws.sheet_properties.outlinePr = Outline(summaryBelow=False, summaryRight=False)
    ws.sheet_format.outlineLevelRow = min(7, max([r.level - 1 for r in report.rows] + [0]))


def _method_list(cfg: Config) -> list[str]:
    return cfg.methods or ["Por definir"]


def _write_lists(ws, methods: list[str]) -> None:
    _header(ws, 1, [(1, "Métodos de medición"), (2, "Descripción / regla de cálculo")])
    for i, m in enumerate(methods, start=2):
        put_text(ws, i, 1, m)
        ws.cell(row=i, column=2, value=METHOD_HELP.get(m, ""))
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 110
    ws.column_dimensions["D"].width = 70
    ws.cell(row=1, column=4, value="Se pueden agregar métodos al final de la columna A, sin dejar filas vacías: "
            "aparecen en el desplegable de la columna C de la hoja PV. Para calcular el EV se reconocen: "
            "fórmulas fijas X/Y (0/100, 50/50...), hitos ponderados, % completado, % físico completado, "
            "unidades completadas y avance lineal (LOE). Otros textos se calculan como % completado.").font = FONT_NOTE
    ws.cell(row=1, column=4).alignment = Alignment(wrap_text=True, vertical="top")


def input_sheets(report: Report, kind: str) -> list[str]:
    """Hojas donde se cargan los valores de EV o AC (las de cada empresa en multi-empresa): el EV se carga como
    % ganado en las hojas '% EV'."""
    if kind == SHEET_EV:
        if report.multi:
            return [pct_sheet(kind, report.company_a), pct_sheet(kind, report.company_b)]
        return [pct_sheet(kind)]
    if report.multi:
        return [company_sheet(kind, report.company_a), company_sheet(kind, report.company_b)]
    return [kind]


def _write_curve(ws, report: Report, cfg: Config, lay: _Layout) -> None:
    money = money_format(cfg.decimals)
    ws.cell(row=1, column=1, value="Evolución mensual  -  valores del proyecto").font = FONT_TITLE
    ws.cell(row=2, column=1, value="Tabla calculada con fórmulas a partir de la fila TOTAL de las hojas PV, EV y AC. "
            "EV y AC acumulados se muestran desde el primer mes con datos cargados hasta la fecha de corte de la "
            "hoja 'Resumen'. Los gráficos están en la hoja 'Cuadro de Mando'.").font = FONT_NOTE
    _header(ws, CURVE_HEADER_ROW, [
        (C_NUM, "Mes N°"), (C_DATE, "Mes"), (C_PV, "PV del mes"), (C_PVC, "PV acumulado"),
        (C_EV, "EV del mes"), (C_EVC, "EV acumulado"), (C_AC, "AC del mes"), (C_ACC, "AC acumulado"),
        (C_PVPCT, "% PV acum."), (C_EVPCT, "% EV acum."), (C_SPI, "SPI acum."), (C_CPI, "CPI acum."),
        (C_BAC, "BAC"), (C_EVN, "Celdas EV cargadas (acum.)"), (C_ACN, "Celdas AC cargadas (acum.)"),
    ])
    tr = lay.total_row
    f0 = CURVE_FIRST_ROW
    bac = f"'{SHEET_PV}'!$D${tr}"
    kinds = f"'{SHEET_STRUCT}'!$E${lay.struct_row(FIRST_ROW)}:$E${lay.struct_row(lay.last_row)}"
    for i in range(lay.months):
        r = f0 + i
        mc = col(lay.first_col + i)
        month_cells = f"{mc}${FIRST_ROW}:{mc}${lay.last_row}"
        prev_evn = f"+{col(C_EVN)}{r - 1}" if i else ""
        prev_acn = f"+{col(C_ACN)}{r - 1}" if i else ""
        # EV y AC acumulados: desde el primer mes con datos cargados hasta la fecha de corte
        cut_ev = f'OR({NAME_STATUS}="",{col(C_DATE)}{r}>{NAME_STATUS},{col(C_EVN)}{r}=0)'
        cut_ac = f'OR({NAME_STATUS}="",{col(C_DATE)}{r}>{NAME_STATUS},{col(C_ACN)}{r}=0)'
        cells = {
            C_NUM: i + 1,
            C_DATE: lay.month_dates[i],
            C_PV: f"='{SHEET_PV}'!{mc}{tr}",
            C_PVC: f"=SUM(${col(C_PV)}${f0}:{col(C_PV)}{r})",
            C_EV: f"='{SHEET_EV}'!{mc}{tr}",
            C_EVC: f"=IF({cut_ev},NA(),SUM(${col(C_EV)}${f0}:{col(C_EV)}{r}))",
            C_AC: f"='{SHEET_AC}'!{mc}{tr}",
            C_ACC: f"=IF({cut_ac},NA(),SUM(${col(C_AC)}${f0}:{col(C_AC)}{r}))",
            C_PVPCT: f"=IFERROR({col(C_PVC)}{r}/{bac},0)",
            C_EVPCT: f"=IFERROR({col(C_EVC)}{r}/{bac},NA())",
            C_SPI: f"=IFERROR({col(C_EVC)}{r}/{col(C_PVC)}{r},NA())",
            C_CPI: f"=IFERROR({col(C_EVC)}{r}/{col(C_ACC)}{r},NA())",
            C_BAC: f"={bac}",
            C_EVN: "=" + "+".join(f"COUNTIFS({kinds},\"<>{KIND_SUMMARY}\",'{sh}'!{month_cells},\"<>\")"
                                  for sh in input_sheets(report, SHEET_EV)) + prev_evn,
            C_ACN: "=" + "+".join(f"COUNTIFS({kinds},\"<>{KIND_SUMMARY}\",'{sh}'!{month_cells},\"<>\")"
                                  for sh in input_sheets(report, SHEET_AC)) + prev_acn,
        }
        for c_, v in cells.items():
            cell = ws.cell(row=r, column=c_, value=v)
            if c_ == C_DATE:
                cell.number_format = FMT_MONTH
            elif c_ in (C_PVPCT, C_EVPCT):
                cell.number_format = FMT_PCT
            elif c_ in (C_SPI, C_CPI):
                cell.number_format = FMT_RATIO
            elif c_ in (C_EVN, C_ACN):
                cell.number_format = "0"
            elif c_ != C_NUM:
                cell.number_format = money
    for c_ in (C_EVN, C_ACN):
        ws.column_dimensions[col(c_)].hidden = True
    for c_, w in {C_NUM: 7, C_DATE: 10}.items():
        ws.column_dimensions[col(c_)].width = w
    for c_ in range(C_PV, C_BAC + 1):
        ws.column_dimensions[col(c_)].width = 14
    ws.freeze_panes = ws.cell(row=CURVE_FIRST_ROW, column=C_PV)


def add_s_curve_chart(ws, months: int, project_name: str, currency: str = "") -> None:
    """Agrega (o vuelve a agregar) el gráfico de curva S en la hoja 'Curva S' de los libros
    generados con versiones anteriores a la 1.3.0 (sin hoja 'Cuadro de Mando').

    El gráfico toma sus datos de la tabla de la hoja, calculada con fórmulas de
    Excel, y abarca todo el horizonte.
    """
    ws._charts = []  # openpyxl no conserva gráficos al reabrir: se recrea siempre
    last = CURVE_FIRST_ROW + max(1, months) - 1
    chart = LineChart()
    chart.title = f"Curva S - {project_name}" if project_name else "Curva S"
    chart.height = 15
    chart.width = 30
    chart.y_axis.title = f"Monto acumulado ({currency})" if currency else "Monto acumulado"
    chart.y_axis.number_format = "#,##0"
    chart.y_axis.crossAx = 500
    chart.x_axis = DateAxis(crossAx=100)
    chart.x_axis.number_format = "mmm-yy"
    chart.x_axis.majorTimeUnit = "months"
    chart.x_axis.title = "Mes"
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    chart.legend.position = "b"
    chart.display_blanks = "gap"
    series_spec = [(C_PVC, "1F4E79", None), (C_EVC, "70AD47", None), (C_ACC, "C00000", None),
                   (C_BAC, "7F7F7F", "dash")]
    for c_, color, dash in series_spec:
        ref = Reference(ws, min_col=c_, min_row=CURVE_HEADER_ROW, max_row=last)
        chart.add_data(ref, titles_from_data=True)
        s = chart.series[-1]
        s.smooth = False
        s.marker.symbol = "none"
        s.graphicalProperties.line.solidFill = color
        s.graphicalProperties.line.width = 28575 if c_ != C_BAC else 15875
        if dash:
            s.graphicalProperties.line.dashStyle = dash
    chart.set_categories(Reference(ws, min_col=C_DATE, min_row=CURVE_FIRST_ROW, max_row=last))
    ws.add_chart(chart, CHART_ANCHOR)


def _write_summary(ws, report: Report, cfg: Config, lay: _Layout, generated: datetime,
                   status: date | None) -> None:
    money = money_format(cfg.decimals)
    p = report.project
    tr = lay.total_row
    ws.cell(row=1, column=1, value="Informe de Valor Ganado (EVM)").font = FONT_TITLE
    put_text(ws, 2, 1, report.title).font = Font(bold=True, size=12)
    several = report.multi_file

    def section(r: int, text: str) -> None:
        for c_ in (1, 2, 3):
            ws.cell(row=r, column=c_).fill = FILL_SECTION
        ws.cell(row=r, column=1, value=text).font = FONT_BOLD

    section(4, "Parámetros")
    st = ws[STATUS_CELL]
    ws["A5"] = "Fecha de corte (fecha de estado)"
    st.value = status
    st.number_format = FMT_DATE
    st.fill = FILL_INPUT
    st.border = BORDER
    ws["C5"] = "Se debe ingresar el último día del mes informado (p. ej. 31/10/2026). Celda con nombre 'FechaCorte'."
    ws["A6"] = "Mes de corte (N°)"
    ws["B6"] = (f'=IF({NAME_STATUS}="","",(YEAR({NAME_STATUS})-YEAR({NAME_START}))*12'
                f'+MONTH({NAME_STATUS})-MONTH({NAME_START})+1)')
    ws["C6"] = "Mes del horizonte que contiene la fecha de corte."
    ws["A7"] = "Inicio del horizonte (mes 1)"
    ws[START_CELL] = report.horizon.start
    ws[START_CELL].number_format = FMT_MONTH
    ws["A8"] = "Meses del horizonte"
    ws["B8"] = report.horizon.months
    ws["A9"] = "Moneda"
    put_text(ws, 9, 2, cfg.currency or report.currency or "-")
    ws["A10"] = "Línea base"
    put_text(ws, 10, 2, per_file(report, lambda s: s.project.baseline_label) if several else p.baseline_label)
    ws["A11"] = "Origen" if not several else "Archivos de MS Project"
    put_text(ws, 11, 2, per_file(report, lambda s: s.project.source) if several else str(p.source))
    ws["A12"] = "Fuente del PV mensual"
    put_text(ws, 12, 2, report.pv_source)
    ws["A13"] = "Generado"
    ws["B13"] = generated.replace(microsecond=0)
    ws["B13"].number_format = "dd/mm/yyyy hh:mm"
    ws["A14"] = "Planificación"
    put_text(ws, 14, 2, per_file(report, lambda s: s.plan or "todas") if several
             else report.plan or "Todas (sin filtro)")

    pv_tot = lay.month_range(SHEET_PV, tr)
    ev_tot = lay.month_range(SHEET_EV, tr)
    ac_tot = lay.month_range(SHEET_AC, tr)
    hdr = lay.header_range(SHEET_PV)

    def cum(rng: str) -> str:
        return f'=IF({NAME_STATUS}="","",SUMIFS({rng},{hdr},"<="&{NAME_STATUS}))'

    section(15, "Indicadores a la fecha de corte")
    rows = [
        (16, "BAC - Presupuesto a la conclusión", f"='{SHEET_PV}'!D{tr}", money,
         "Suma del BAC ingresado en la columna D de la hoja PV."),
        (17, "PV - Valor planificado acumulado", cum(pv_tot), money, "Trabajo planificado hasta la fecha de corte."),
        (18, "EV - Valor ganado acumulado", cum(ev_tot), money, "Valor del trabajo realizado."),
        (19, "AC - Costo real acumulado", cum(ac_tot), money, "Costo incurrido."),
        (20, "SV = EV - PV", '=IFERROR(B18-B17,"")', money, "< 0: atraso."),
        (21, "CV = EV - AC", '=IFERROR(B18-B19,"")', money, "< 0: sobrecosto."),
        (22, "SPI = EV / PV", '=IFERROR(B18/B17,"")', FMT_RATIO, "< 1: atraso."),
        (23, "CPI = EV / AC", '=IFERROR(B18/B19,"")', FMT_RATIO, "< 1: sobrecosto."),
        (24, "EAC = BAC / CPI", '=IFERROR(B16/B23,"")', money, "Estimación a la conclusión si el desempeño de costo continúa."),
        (25, "EAC = AC + (BAC - EV)", '=IFERROR(B19+B16-B18,"")', money, "Si el trabajo restante se ejecuta según presupuesto."),
        (26, "EAC = AC + (BAC - EV) / (CPI × SPI)", '=IFERROR(B19+(B16-B18)/(B23*B22),"")', money,
         "Si costo y plazo siguen influyendo."),
        (27, "ETC = EAC - AC", '=IFERROR(B24-B19,"")', money, "Costo estimado para terminar (EAC = BAC / CPI)."),
        (28, "VAC = BAC - EAC", '=IFERROR(B16-B24,"")', money, "< 0: se espera terminar por encima del presupuesto."),
        (29, "TCPI = (BAC - EV) / (BAC - AC)", '=IFERROR((B16-B18)/(B16-B19),"")', FMT_RATIO,
         "Eficiencia de costo necesaria para terminar dentro del BAC (> 1: exigente)."),
        (30, "% Planificado = PV / BAC", '=IFERROR(B17/B16,"")', FMT_PCT, ""),
        (31, "% Completado = EV / BAC", '=IFERROR(B18/B16,"")', FMT_PCT, ""),
        (32, "% Gastado = AC / BAC", '=IFERROR(B19/B16,"")', FMT_PCT, ""),
    ]
    curve_pvc = (f"'{SHEET_CURVE}'!${col(C_PVC)}${CURVE_FIRST_ROW}:"
                 f"${col(C_PVC)}${CURVE_FIRST_ROW + lay.months - 1}")
    rows += [
        (34, "Duración planificada (meses)",
         f'=IFERROR(IF(INDEX({curve_pvc},{lay.months})<B16-0.005,"",COUNTIF({curve_pvc},"<"&(B16-0.005))+1),"")',
         "0.0", "Meses del horizonte hasta completar el BAC."),
        (35, "ES - Cronograma ganado (meses)",
         '=IFERROR(IF(B18="","",IF(B18>=B16-0.005,B34,IF(B41>B44,MIN(MAX(B36,IF(B18>0,B44+1,0)),B41),'
         'B41+(B18-B42)/(B43-B42)))),"")', "0.00",
         "Momento en que el plan preveía el EV actual."),
        (36, "AT - Tiempo transcurrido (meses)", '=B6', "0", "Meses desde el inicio del horizonte hasta el corte."),
        (37, "SV(t) = ES - AT (meses)", '=IFERROR(B35-B36,"")', "0.00", "< 0: meses de atraso."),
        (38, "SPI(t) = ES / AT", '=IFERROR(B35/B36,"")', FMT_RATIO, "No converge a 1 al final, a diferencia del SPI."),
        (39, "Duración estimada = duración plan / SPI(t)", '=IFERROR(B34/B38,"")', "0.0", "Meses."),
        (40, "Término estimado (aprox.)", f'=IFERROR(EOMONTH({NAME_START},ROUNDUP(B39,0)-1),"")', FMT_MONTH, ""),
        (41, "Auxiliar ES: meses con PV acumulado <= EV", f'=IFERROR(COUNTIF({curve_pvc},"<="&(B18+0.005)),"")',
         "0", ""),
        (42, "Auxiliar ES: PV acumulado en C", f'=IFERROR(IF(B41=0,0,INDEX({curve_pvc},B41)),"")', money, ""),
        (43, "Auxiliar ES: PV acumulado en C+1", f'=IFERROR(INDEX({curve_pvc},B41+1),"")', money, ""),
        (44, "Auxiliar ES: meses con PV acumulado < EV", f'=IFERROR(COUNTIF({curve_pvc},"<"&(B18-0.005)),"")',
         "0", ""),
    ]
    section(33, "Cronograma ganado (Earned Schedule)")
    for r, label, formula, fmt, note in rows:
        ws.cell(row=r, column=1, value=label)
        c = ws.cell(row=r, column=2, value=formula)
        c.number_format = fmt
        c.border = BORDER
        if note:
            ws.cell(row=r, column=3, value=note).font = FONT_NOTE
    for r in (41, 42, 43, 44):
        ws.cell(row=r, column=1).font = FONT_NOTE

    for ref in ("B22", "B23", "B38"):
        _traffic_light(ws, ref, Thresholds.from_config(cfg))
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 70
    ws.cell(row=46, column=1, value="Valores mensuales: hojas PV, EV y AC (celdas amarillas = ingreso). "
            "Gráficos: hoja 'Cuadro de Mando'. Tabla mensual: hoja 'Evolución mensual'. Indicadores por elemento "
            "EDT: hoja 'Indicadores EDT'.").font = FONT_NOTE


@dataclass(frozen=True)
class Thresholds:
    """Umbrales comunes del semáforo de SPI y CPI. Rojo: menor que red_low o mayor que red_high; amarillo:
    menor que yellow_low o mayor que yellow_high; verde: entre yellow_low y yellow_high (inclusive)."""

    red_low: float
    yellow_low: float
    yellow_high: float
    red_high: float

    META_KEYS = ("semaforo_rojo_inferior", "semaforo_amarillo_inferior", "semaforo_amarillo_superior",
                 "semaforo_rojo_superior")

    @classmethod
    def from_config(cls, cfg: Config) -> "Thresholds":
        return cls(cfg.threshold_red, cfg.threshold_yellow, cfg.threshold_yellow_upper, cfg.threshold_red_upper)

    @classmethod
    def from_meta(cls, meta: dict, cfg: Config) -> "Thresholds":
        """Los umbrales con que se generó el libro (libros de versiones anteriores: los de la configuración)."""
        try:
            return cls(*(float(meta[k]) for k in cls.META_KEYS))
        except (KeyError, TypeError, ValueError):
            return cls.from_config(cfg)

    def meta(self) -> dict:
        return dict(zip(self.META_KEYS, (self.red_low, self.yellow_low, self.yellow_high, self.red_high)))

    def light(self, value: float | None) -> str:
        """'Rojo', 'Amarillo', 'Verde' ('' sin valor)."""
        if value is None:
            return ""
        if value < self.red_low or value > self.red_high:
            return "Rojo"
        if value < self.yellow_low or value > self.yellow_high:
            return "Amarillo"
        return "Verde"

    def bands(self) -> list[tuple[str, str]]:
        """Escala del semáforo, de menor a mayor: (texto, color)."""
        f = fmt_decimal
        return [(f"< {f(self.red_low)}", "Rojo"), (f"{f(self.red_low)} a {f(self.yellow_low)}", "Amarillo"),
                (f"{f(self.yellow_low)} a {f(self.yellow_high)}", "Verde"),
                (f"{f(self.yellow_high)} a {f(self.red_high)}", "Amarillo"), (f"> {f(self.red_high)}", "Rojo")]

    def describe(self) -> str:
        """'verde de 0,95 a 1,05; amarillo de 0,90 a 0,95 y de 1,05 a 1,10; rojo fuera de 0,90 a 1,10'."""
        f = fmt_decimal
        return (f"verde de {f(self.yellow_low)} a {f(self.yellow_high)}; amarillo de {f(self.red_low)} a "
                f"{f(self.yellow_low)} y de {f(self.yellow_high)} a {f(self.red_high)}; rojo por debajo de "
                f"{f(self.red_low)} o por encima de {f(self.red_high)}")


def fmt_decimal(v: float) -> str:
    """Número con dos decimales y coma decimal: 0.9 -> '0,90'."""
    return f"{v:.2f}".replace(".", ",")


def _traffic_light(ws, rng: str, th: Thresholds) -> None:
    """Semáforo (formato condicional) de un rango de índices SPI o CPI."""
    first = rng.split(":")[0].replace("$", "")
    for low, high, fill in ((th.red_low, th.red_high, FILL_RED), (th.yellow_low, th.yellow_high, FILL_YELLOW)):
        ws.conditional_formatting.add(rng, FormulaRule(
            formula=[f'AND(ISNUMBER({first}),OR({first}<{low!r},{first}>{high!r}))'], fill=fill, stopIfTrue=True))
    ws.conditional_formatting.add(rng, FormulaRule(
        formula=[f'ISNUMBER({first})'], fill=FILL_GREEN, stopIfTrue=True))


def _write_indicators(ws, report: Report, cfg: Config, lay: _Layout) -> None:
    money = money_format(cfg.decimals)
    ws.cell(row=TITLE_ROW, column=1, value="Indicadores por elemento EDT a la fecha de corte").font = FONT_TITLE
    ws.cell(row=INFO_ROW, column=1, value="Fecha de corte: celda 'FechaCorte' de la hoja Resumen. "
            "Valores acumulados hasta el mes de la fecha de corte inclusive.").font = FONT_NOTE
    heads = ["Código EDT", "Actividad", "Método de medición", "BAC", "PV acum.", "EV acum.", "AC acum.",
             "SV", "CV", "SPI", "CPI", "EAC", "VAC", "% Planificado", "% Completado", "% Gastado"]
    _header(ws, HEADER_ROW, list(enumerate(heads, start=1)))
    hdr = lay.header_range(SHEET_PV)
    all_rows = [(lay.r(x), x) for x in report.rows] + [(lay.total_row, None)]
    for r, row in all_rows:
        def cum(sheet: str) -> str:
            return f'=IF({NAME_STATUS}="","",SUMIFS({lay.month_range(sheet, r)},{hdr},"<="&{NAME_STATUS}))'
        f = {
            1: f"='{SHEET_PV}'!A{r}", 2: f"='{SHEET_PV}'!B{r}",
            3: f"=IF('{SHEET_PV}'!C{r}=\"\",\"\",'{SHEET_PV}'!C{r})", 4: f"='{SHEET_PV}'!D{r}",
            5: cum(SHEET_PV), 6: cum(SHEET_EV), 7: cum(SHEET_AC),
            8: f'=IFERROR(F{r}-E{r},"")', 9: f'=IFERROR(F{r}-G{r},"")',
            10: f'=IFERROR(F{r}/E{r},"")', 11: f'=IFERROR(F{r}/G{r},"")',
            12: f'=IFERROR(IF(OR(K{r}="",K{r}=0),G{r}+D{r}-F{r},D{r}/K{r}),"")', 13: f'=IFERROR(D{r}-L{r},"")',
            14: f'=IFERROR(E{r}/D{r},"")', 15: f'=IFERROR(F{r}/D{r},"")', 16: f'=IFERROR(G{r}/D{r},"")',
        }
        for c_, v in f.items():
            cell = ws.cell(row=r, column=c_, value=v)
            if c_ in (4, 5, 6, 7, 8, 9, 12, 13):
                cell.number_format = money
            elif c_ in (10, 11):
                cell.number_format = FMT_RATIO
            elif c_ >= 14:
                cell.number_format = FMT_PCT
        if row is None:
            for c_ in range(1, 17):
                ws.cell(row=r, column=c_).font = FONT_BOLD
                ws.cell(row=r, column=c_).fill = FILL_TOTAL
        else:
            ws.cell(row=r, column=2).alignment = Alignment(indent=max(0, row.level - 1))
            if row.kind == KIND_SUMMARY:
                for c_ in range(1, 17):
                    ws.cell(row=r, column=c_).font = FONT_BOLD
                    ws.cell(row=r, column=c_).fill = FILL_L1 if row.level == 1 else FILL_L2
            ws.row_dimensions[r].outline_level = min(7, max(0, row.level - 1))
            if lay.review:
                ws.row_dimensions[r + 1].hidden = True
    last = lay.total_row
    _traffic_light(ws, f"J{FIRST_ROW}:K{last}", Thresholds.from_config(cfg))
    for c_, w in {1: 12, 2: 46, 3: 20}.items():
        ws.column_dimensions[col(c_)].width = w
    for c_ in range(4, 17):
        ws.column_dimensions[col(c_)].width = 14
    ws.freeze_panes = ws.cell(row=FIRST_ROW, column=3)
    ws.sheet_properties.outlinePr = Outline(summaryBelow=False, summaryRight=False)
    ws.sheet_format.outlineLevelRow = min(7, max([r.level - 1 for r in report.rows] + [0]))


def _write_validation(ws, report: Report) -> None:
    ws.cell(row=1, column=1, value="Validación de la exportación").font = FONT_TITLE
    _header(ws, 3, [(1, "Severidad"), (2, "Código EDT"), (3, "Tarea"), (4, "Detalle")])
    if not report.issues:
        ws.cell(row=4, column=1, value="OK")
        ws.cell(row=4, column=4, value="No se encontraron observaciones.")
    for i, issue in enumerate(report.issues, start=4):
        ws.cell(row=i, column=1, value=issue.severity)
        put_text(ws, i, 2, issue.code)
        put_text(ws, i, 3, issue.name)
        put_text(ws, i, 4, issue.message).alignment = Alignment(wrap_text=True, vertical="top")
        if issue.severity != "Info":
            ws.cell(row=i, column=1).fill = FILL_YELLOW if issue.severity == "Advertencia" else FILL_RED
    for c_, w in {1: 13, 2: 14, 3: 40, 4: 110}.items():
        ws.column_dimensions[col(c_)].width = w
    ws.auto_filter.ref = f"A3:D{max(4, 3 + len(report.issues))}"


def _append_texts(ws, values: list) -> None:
    """Agrega una fila con los textos limpios (sin caracteres de control) y siempre como texto."""
    ws.append([clean_text(v) for v in values])
    _texts_in_last_row(ws)


def _write_structure(ws, report: Report, lay: _Layout) -> None:
    """Hoja _estructura: una fila por fila de la EDT, alineada con las hojas mensuales (la fila r de las hojas va en
    la fila r - 3), para las fórmulas que filtran por el tipo de fila. Con depuración de costes, las filas «(coste)»
    solo llevan el tipo 'coste' (sin fila ni clave: no son filas del libro)."""
    ws.append(["fila", "clave", "edt", "nivel", "tipo", "fila_padre", "actividad", "metodo_origen"])
    for row in report.rows:
        r = lay.r(row)
        values = [r, row_key(row), row.code, row.level, row.kind, lay.r(row.parent) if row.parent else None,
                  row.name, row.method if row.is_input else None]
        for c_, v in enumerate(values, start=1):
            put_text(ws, lay.struct_row(r), c_, v)
        if lay.review:
            ws.cell(row=lay.struct_row(r + 1), column=5, value=KIND_COST)


def write_meta(ws, meta: dict) -> None:
    ws.delete_rows(1, ws.max_row)
    ws.append(["clave", "valor"])
    for k, v in meta.items():
        _append_texts(ws, [k, v])


def write_workbook(report: Report, cfg: Config, path: str | Path, prior=None,
                   generated: datetime | None = None) -> dict:
    """Crea el libro Excel completo y lo guarda en `path`. Devuelve los metadatos."""
    from evm import analysis, dashboard
    generated = generated or datetime.now()
    lay = _Layout(report)
    wb = Workbook()
    ws_dash = wb.active
    ws_dash.title = dashboard.SHEET_DASHBOARD
    ws_sum = wb.create_sheet(SHEET_SUMMARY)
    data_sheets = []                     # (hoja, tipo, empresa)
    pct_sheets = []                      # hojas '% PV' y '% EV'
    review = None                        # depuración de costes: hojas del gráfico y de la previsión pegada
    for kind in (SHEET_PV, SHEET_EV, SHEET_AC):
        data_sheets.append((wb.create_sheet(kind), kind, None))
        roles = (("A", report.company_a), ("B", report.company_b)) if report.multi else ()
        for role, company in roles:
            data_sheets.append((wb.create_sheet(company_sheet(kind, company)), kind, role))
        if kind != SHEET_AC:
            for role, company in roles or ((None, ""),):
                pct_sheets.append((wb.create_sheet(pct_sheet(kind, company)), kind, role))
        if kind == SHEET_PV and lay.review:
            from evm import cost_review
            review = (wb.create_sheet(cost_review.SHEET_CHART), wb.create_sheet(cost_review.SHEET_PASTE))
    ws_curve = wb.create_sheet(SHEET_CURVE)
    ws_ind = wb.create_sheet(SHEET_IND)
    ws_val = wb.create_sheet(SHEET_VALID)
    ws_lists = wb.create_sheet(SHEET_LISTS)
    ws_charts = wb.create_sheet(dashboard.SHEET_CHARTDATA)
    ws_struct = wb.create_sheet(SHEET_STRUCT)
    ws_meta = wb.create_sheet(SHEET_META)
    hidden = [ws_charts, ws_struct, ws_meta]
    if review is not None:
        hidden.append(wb.create_sheet(cost_review.SHEET_HELPER))
    for ws in hidden:
        ws.sheet_state = "hidden"

    status = prior.status_date if prior else None
    methods = list(_method_list(cfg))
    if prior is not None:
        methods += [m for m in prior.methods_list if m not in methods]
    for r in report.rows:
        m = (prior.method_for(row_key(r)) if prior is not None and r.is_input else None) or r.method
        if r.is_input and m and m not in methods:
            methods.append(m)

    _write_lists(ws_lists, methods)
    for ws, kind, role in data_sheets:
        _write_data_sheet(ws, kind, report, cfg, lay, generated, prior, role)
    for ws, kind, role in pct_sheets:
        _write_pct_sheet(ws, kind, report, cfg, lay, generated, prior, role)
    _write_curve(ws_curve, report, cfg, lay)
    _write_summary(ws_sum, report, cfg, lay, generated, status)
    _write_indicators(ws_ind, report, cfg, lay)
    _write_validation(ws_val, report)
    _write_structure(ws_struct, report, lay)
    packages = dashboard.chart_packages((lay.r(r), r.code, r.level, r.kind) for r in report.rows)
    dashboard.write_chart_data(ws_charts, cfg, lay, packages)
    dashboard.write_dashboard(ws_dash, report, cfg, lay, packages, status)
    analysis.write_sheet(wb, prior.analysis if prior is not None else [],
                         {row_key(r): lay.r(r) for r in report.rows if r.is_input}, lay.months,
                         Thresholds.from_config(cfg), cfg.decimals)

    wb.defined_names[NAME_STATUS] = DefinedName(NAME_STATUS, attr_text=f"'{SHEET_SUMMARY}'!$B$5")
    wb.defined_names[NAME_START] = DefinedName(NAME_START, attr_text=f"'{SHEET_SUMMARY}'!$B$7")
    # Lista de métodos con tamaño dinámico (crece si se agregan métodos en la hoja Listas)
    wb.defined_names[NAME_METHODS] = DefinedName(
        NAME_METHODS, attr_text=f"OFFSET('{SHEET_LISTS}'!$A$2,0,0,MAX(1,COUNTA('{SHEET_LISTS}'!$A:$A)-1),1)")
    dashboard.define_names(wb, lay)
    cutoff = report.horizon.index(status) + 1 if status else None
    # % planificado (fracciones): último mes con algún %; el BAC se ingresa después, en el libro
    plan_end = max((i + 1 for i, v in enumerate(report.total_vector) if abs(v) > 1e-9), default=0)
    dashboard.add_dashboard_charts(wb, packages, lay.months, plan_end, cutoff, 0.0,
                                   cfg.currency or report.currency or "")
    if review is not None:
        ws_chart, ws_paste = review
        pasted = prior.review_rows if prior is not None and prior.review_keep else None
        cost_review.write_helper(wb[cost_review.SHEET_HELPER], report, lay, cfg)
        cost_review.write_paste(ws_paste, report, cfg, lay, pasted)
        cost_review.write_chart_sheet(ws_chart, report, cfg, lay)
        cost_review.add_chart(wb, lay.months, max(plan_end, cost_review.last_paste_month(ws_paste, lay.months)),
                              cfg.currency or report.currency or "")
    wb.active = 0

    p = report.project
    several = report.multi_file
    meta = {
        "version": __version__,
        "proyecto": report.title if several else p.name,
        "titulo": report.title,
        "origen": per_file(report, lambda s: s.project.source) if several else str(p.source),
        "motor": per_file(report, lambda s: s.project.backend) if several else p.backend,
        "linea_base": per_file(report, lambda s: s.project.baseline_label) if several else p.baseline_label,
        "fuente_pv": report.pv_source,
        "planificacion": report.plan,
        "moneda": cfg.currency or report.currency or "",
        "horizonte_inicio": report.horizon.start,
        "meses": report.horizon.months,
        "primera_fila": FIRST_ROW,
        "fila_total": lay.total_row,
        "nivel_maximo": cfg.max_level,
        "decimales": cfg.decimals,
        **Thresholds.from_config(cfg).meta(),
        "generado": generated.replace(microsecond=0),
        "modo": MODE_PERCENT,
        "campo_visible": cfg.visible_field,
        "multiempresa": bool(report.multi),
        "depuracion_costes": review is not None,
    }
    if review is not None:
        meta.update({
            "depuracion_empresa": cost_review.review_company(report),
            "depuracion_hoja": cost_review.review_sheet(report),
            "depuracion_alineamiento": bool(cfg.cost_review_alignment),
            "depuracion_umbral": cfg.cost_review_threshold,
        })
    if report.multi:
        meta.update({
            "empresa_a": report.company_a,
            "empresa_b": report.company_b,
            "ambas": cfg.company_both,
            "campo_empresa": cfg.company_field,
            "campo_reparto": cfg.split_field,
            "reparto_por_defecto": cfg.split_default,
        })
    if several:
        # Datos de cada archivo: 'actualizar' y --migrar-desde asocian las filas de cada uno por su nombre
        meta[META_SOURCES] = ";".join(s.name for s in report.sources)
        for s in report.sources:
            meta.update({f"proyecto.{s.name}": s.project.name, f"origen.{s.name}": str(s.project.source),
                         f"motor.{s.name}": s.project.backend, f"linea_base.{s.name}": s.project.baseline_label,
                         f"planificacion.{s.name}": s.plan})
    write_meta(ws_meta, meta)

    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    wb.properties.creator = "evm_export"
    wb.properties.title = clean_text(f"EVM - {report.title}")
    save_workbook(wb, path)
    return meta


def _save_error(path: Path) -> str:
    return f"No se pudo guardar '{path}'. Si el libro está abierto en Excel, se debe cerrar y volver a intentar."


def save_workbook(wb, path: str | Path) -> None:
    """Guarda el libro sin riesgo de dañarlo: escribe un archivo temporal en la misma carpeta y luego
    reemplaza el libro. Si el guardado falla (libro abierto en Excel, disco lleno, interrupción), el
    libro anterior queda intacto y el archivo temporal se elimina."""
    path = Path(path)
    tmp = None
    try:
        buffer = io.BytesIO()
        wb.save(buffer)                       # se arma el libro en memoria: un error no toca el disco
        fd, name = tempfile.mkstemp(prefix=f"~{path.stem}.", suffix=".tmp", dir=path.parent)
        tmp = Path(name)
        with os.fdopen(fd, "wb") as fh:
            fh.write(buffer.getbuffer())
        if path.exists():
            try:
                shutil.copymode(path, tmp)
            except OSError:
                pass
        os.replace(tmp, path)
        tmp = None
    except PermissionError as exc:
        raise PermissionError(_save_error(path)) from exc
    finally:
        if tmp is not None:
            try:
                tmp.unlink(missing_ok=True)
            except OSError:
                pass


def check_writable(path: str | Path) -> None:
    """Verifica que un libro existente se pueda reemplazar (no está abierto en Excel), antes de leer
    MS Project, para no esperar la lectura y fallar al guardar."""
    path = Path(path)
    try:
        with open(path, "r+b"):
            pass
    except PermissionError as exc:
        raise PermissionError(_save_error(path)) from exc


# ----------------------------------------------------------------------------
# Lectura (para los comandos actualizar / curva / migración)
# ----------------------------------------------------------------------------
class WorkbookFormatError(ValueError):
    pass


@dataclass
class StructRow:
    row: int
    key: str
    code: str
    level: int
    kind: str
    parent_row: int | None
    name: str
    origin_method: str | None = None    # método que traía la fila al generarse (None en libros antiguos)

    @property
    def is_input(self) -> bool:
        return self.kind in (KIND_LEAF, KIND_RESIDUAL)


def _as_bool(value) -> bool:
    """Valor lógico de _meta (Excel o LibreOffice pueden guardarlo como texto o fórmula)."""
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return value != 0
    return str(value or "").strip().upper().lstrip("=").rstrip("()") in ("TRUE", "VERDADERO", "SI", "SÍ", "1")


def read_meta(wb) -> dict:
    if SHEET_META not in wb.sheetnames or SHEET_STRUCT not in wb.sheetnames:
        raise WorkbookFormatError(
            "El libro no fue generado por este programa (faltan las hojas ocultas _meta/_estructura).")
    meta = {}
    for k, v in wb[SHEET_META].iter_rows(min_row=2, values_only=True):
        if k:
            meta[str(k)] = v
    for k in ("horizonte_inicio", "meses", "primera_fila", "fila_total"):
        if k not in meta:
            raise WorkbookFormatError(f"Falta el dato '{k}' en la hoja _meta.")

    def invalid(k: str) -> WorkbookFormatError:
        return WorkbookFormatError(f"El dato '{k}' de la hoja oculta _meta no es válido ({meta[k]!r}). Esa hoja no "
                                   "se debe modificar: se debe recuperar el libro desde un respaldo.")

    hs = meta["horizonte_inicio"]
    if isinstance(hs, (int, float)) and not isinstance(hs, bool):
        try:
            hs = from_excel(hs)                 # fecha guardada como número de serie
        except (ValueError, OverflowError):
            raise invalid("horizonte_inicio") from None
    if isinstance(hs, datetime):
        hs = hs.date()
    if not isinstance(hs, date):
        raise invalid("horizonte_inicio")
    meta["horizonte_inicio"] = hs
    for k in ("meses", "primera_fila", "fila_total") + (("nivel_maximo",) if meta.get("nivel_maximo") else ()):
        try:
            meta[k] = int(meta[k])
        except (TypeError, ValueError):
            raise invalid(k) from None
    if not 1 <= meta["meses"] <= 16384 - FIRST_MONTH_COL:
        raise invalid("meses")
    meta["multiempresa"] = _as_bool(meta.get("multiempresa"))
    return meta


def book_sources(meta: dict) -> list[str]:
    """Nombres de los archivos de MS Project del libro (['P1', 'P2']); [''] en un libro de un solo archivo."""
    names = [n for n in str(meta.get(META_SOURCES) or "").split(";") if n]
    return names or [""]


def source_meta(meta: dict, key: str, source: str):
    """Dato de _meta de un archivo del libro ('proyecto', 'planificacion'...). Libro de un solo archivo: el dato
    general."""
    return meta.get(f"{key}.{source}") if source else meta.get(key)


def read_structure(wb) -> list[StructRow]:
    out = []
    for vals in wb[SHEET_STRUCT].iter_rows(min_row=2, values_only=True):
        if not vals or vals[0] is None:
            continue
        fila, key, code, level, kind, parent, name, origin = (list(vals) + [None] * 8)[:8]
        try:
            out.append(StructRow(int(fila), str(key), str(code or ""), int(level or 0), str(kind),
                                 int(parent) if parent else None, str(name or ""),
                                 None if len(vals) < 8 else str(origin or "")))
        except (TypeError, ValueError):
            raise WorkbookFormatError(f"La hoja oculta _estructura no es válida (fila {fila!r}). Esa hoja no se "
                                      "debe modificar: se debe recuperar el libro desde un respaldo.") from None
    return out
