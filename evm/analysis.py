"""Hoja "Análisis de desviaciones": paquetes de trabajo fuera de umbral en el mes de la fecha de corte,
con las columnas que se completan a mano (causa, impacto, acción correctiva, responsable y fecha).

'actualizar' (y 'analizar', sin leer MS Project) agrega los paquetes del mes de corte. Las filas se
identifican con la clave de su fila del libro (UniqueID de la tarea, o 'P2:UniqueID' con varios archivos),
de modo que el texto escrito se conserva al actualizar y al regenerar el libro con --migrar-desde. Los
valores (PV, EV, AC e índices, también los del mes anterior y la tendencia) son fórmulas sobre las hojas PV, EV y
AC. El desvío de fin, la holgura y la marca de crítica son los del cronograma de MS Project en el mes del análisis
(valores, que se conservan en los meses siguientes).
"""
from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, datetime

from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.properties import PageSetupProperties

from evm.workbook import (BORDER, COL_BAC, FILL_GREEN, FILL_INPUT, FILL_RED, FILL_YELLOW, FIRST_MONTH_COL,
                          FMT_DATE, FMT_DAYS, FMT_MONTH, FMT_RATIO, FMT_SLACK, FONT_NOTE, FONT_TITLE, HEADER_ROW,
                          NAME_START, SHEET_AC, SHEET_EV, SHEET_IND, SHEET_PV, Thresholds, _header, _traffic_light,
                          col, fmt_decimal, loaded_month, money_format, put_text)

SHEET_ANALYSIS = "Análisis de desviaciones"
STATUS_ROW = 3              # resumen del último análisis
HEAD_ROW = 4                # encabezados
FIRST = 5                   # primera fila de datos
(A_MONTH, A_CODE, A_NAME, A_BAC, A_PV, A_EV, A_AC, A_SV, A_CV, A_SPI, A_CPI, A_LIGHT, A_SPI_PREV, A_CPI_PREV,
 A_TREND, A_VARIANCE, A_SLACK, A_CRITICAL, A_ORIGIN, A_CAUSE, A_IMPACT, A_ACTION, A_OWNER, A_DATE,
 A_KEY) = range(1, 26)
TEXT_COLS = (A_CAUSE, A_IMPACT, A_ACTION, A_OWNER, A_DATE)
PLAN_COLS = (A_VARIANCE, A_SLACK, A_CRITICAL)      # plazo según MS Project en el mes del análisis (valores)
HEADERS = {
    A_MONTH: "Mes", A_CODE: "Código EDT", A_NAME: "Actividad", A_BAC: "BAC", A_PV: "PV acum.", A_EV: "EV acum.",
    A_AC: "AC acum.", A_SV: "SV", A_CV: "CV", A_SPI: "SPI", A_CPI: "CPI", A_LIGHT: "Semáforo",
    A_SPI_PREV: "SPI mes ant.", A_CPI_PREV: "CPI mes ant.", A_TREND: "Tendencia",
    A_VARIANCE: "Desvío fin (días lab.)", A_SLACK: "Holgura (días lab.)", A_CRITICAL: "Crítica",
    A_ORIGIN: "Antecedente", A_CAUSE: "Causa", A_IMPACT: "Impacto en plazo y costo",
    A_ACTION: "Acción correctiva", A_OWNER: "Responsable", A_DATE: "Fecha compromiso", A_KEY: "clave",
}
# Columnas de las hojas de versiones anteriores a la 2.3.0 (sin encabezados reconocibles)
LEGACY_COLS = {c_: i for i, c_ in enumerate((A_MONTH, A_CODE, A_NAME, A_BAC, A_PV, A_EV, A_AC, A_SV, A_CV, A_SPI,
                                             A_CPI, A_LIGHT, A_ORIGIN) + TEXT_COLS + (A_KEY,), start=1)}
WIDTHS = {A_MONTH: 10, A_CODE: 11, A_NAME: 34, A_BAC: 13, A_PV: 13, A_EV: 13, A_AC: 13, A_SV: 12, A_CV: 12,
          A_SPI: 7, A_CPI: 7, A_LIGHT: 10, A_SPI_PREV: 8, A_CPI_PREV: 8, A_TREND: 10, A_VARIANCE: 9, A_SLACK: 9,
          A_CRITICAL: 9, A_ORIGIN: 24, A_CAUSE: 38, A_IMPACT: 30, A_ACTION: 38, A_OWNER: 16, A_DATE: 12}
NEW = "Nuevo"
TREND_TOL = 0.01            # variación de la distancia de un índice a 1 que cuenta como mejora o empeoramiento


@dataclass
class Entry:
    """Fila del análisis: un paquete en un mes."""

    month: date | None          # primer día del mes analizado
    key: str                    # clave de la fila del libro ("" = fila agregada a mano)
    code: str
    name: str
    texts: list = field(default_factory=lambda: [None] * len(TEXT_COLS))   # causa ... fecha compromiso
    origin: str = ""            # antecedente: "Nuevo", "También en 06/2027", "Texto copiado de 06/2027..."
    raw: list | None = None     # fila agregada a mano: sus valores (sin fórmulas), tal como estaban
    plan: tuple = (None, None, None)    # desvío de fin, holgura y crítica en el mes del análisis

    @property
    def has_text(self) -> bool:
        return any(v is not None and str(v).strip() != "" for v in self.texts)


@dataclass
class Deviation:
    """Paquete fuera de umbral a la fecha de corte (valores acumulados)."""

    key: str
    code: str
    name: str
    spi: float | None
    cpi: float | None


def thresholds(meta: dict, cfg) -> Thresholds:
    """Umbrales del semáforo con que se generó el libro (libros antiguos: los de la configuración)."""
    return Thresholds.from_meta(meta, cfg)


def _month(value) -> date | None:
    if isinstance(value, datetime):
        value = value.date()
    return date(value.year, value.month, 1) if isinstance(value, date) else None


def _columns(ws) -> dict[int, int]:
    """Columna de la hoja de cada dato (A_MONTH...), según los encabezados: las hojas de versiones anteriores
    tienen menos columnas."""
    found = {str(c.value).strip(): c.column for c in ws[HEAD_ROW] if c.value is not None}
    if found.get(HEADERS[A_KEY]) is None:
        return dict(LEGACY_COLS)
    return {c_: found[text] for c_, text in HEADERS.items() if text in found}


def read_entries(ws) -> list[Entry]:
    """Filas del análisis de una hoja existente. Las filas sin clave (agregadas a mano) se conservan tal cual."""
    out = []
    cols = _columns(ws)
    width = max(cols.values())
    for row in ws.iter_rows(min_row=FIRST, max_col=width, values_only=True):
        row = list(row) + [None] * (width - len(row))
        vals = [row[cols[c_] - 1] if c_ in cols else None for c_ in range(1, A_KEY + 1)]   # en el orden actual
        if all(v is None or str(v).strip() == "" for v in vals):
            continue
        key = str(vals[A_KEY - 1] or "").strip()
        month = _month(vals[A_MONTH - 1])
        code, name = str(vals[A_CODE - 1] or ""), str(vals[A_NAME - 1] or "")
        texts = [vals[c - 1] for c in TEXT_COLS]
        if not key or month is None:
            raw = [None if isinstance(v, str) and v.startswith("=") else v for v in vals[:A_KEY - 1]]
            out.append(Entry(month, "", code, name, texts, raw=raw))
            continue
        out.append(Entry(month, key, code, name, texts, str(vals[A_ORIGIN - 1] or ""),
                         plan=tuple(vals[c - 1] for c in PLAN_COLS)))
    return out


def deviations(book, idx: int, th: Thresholds) -> list[Deviation]:
    """Filas de ingreso con SPI o CPI acumulado fuera del rango verde del semáforo (menor que el amarillo
    inferior o mayor que el amarillo superior) en el mes idx del horizonte (mismos cálculos que la hoja
    Indicadores EDT)."""
    out = []
    for s in book.input_rows:
        pv = sum(book.pv[s.row][: idx + 1])
        ev = sum(v or 0.0 for v in book.ev[s.row][: idx + 1])
        ac = sum(v or 0.0 for v in book.ac[s.row][: idx + 1])
        spi = ev / pv if abs(pv) > 1e-9 else None
        cpi = ev / ac if abs(ac) > 1e-9 else None
        if any(th.light(x) in ("Rojo", "Amarillo") for x in (spi, cpi)):
            out.append(Deviation(s.key, s.code, s.name, spi, cpi))
    return out


def refresh(entries: list[Entry], found: list[Deviation], month: date,
            plan: dict[str, tuple] | None = None) -> list[Entry]:
    """Análisis del mes `month`: los paquetes fuera de umbral (`found`) más las filas de ese mes con texto.

    - Un paquete que ya tenía fila en el mes conserva su texto.
    - Un paquete nuevo en el mes copia el texto de su último análisis anterior (para revisarlo).
    - Las filas de otros meses se conservan solo si tienen texto (o si se agregaron a mano).
    - plan: desvío de fin, holgura y crítica de cada fila (clave -> valores) para las filas del mes; las de meses
      anteriores conservan los suyos.
    """
    current = {e.key: e for e in entries if e.key and e.month == month}
    earlier: dict[str, Entry] = {}
    for e in entries:
        if e.key and e.month is not None and e.month < month:
            if e.key not in earlier or e.month > earlier[e.key].month:
                earlier[e.key] = e
    out: list[Entry] = []
    flagged = set()
    for d in found:
        flagged.add(d.key)
        e = current.get(d.key)
        if e is None:
            prev = earlier.get(d.key)
            if prev is None:
                e = Entry(month, d.key, d.code, d.name, origin=NEW)
            elif prev.has_text:
                e = Entry(month, d.key, d.code, d.name, list(prev.texts),
                          f"Texto copiado de {prev.month:%m/%Y}: se debe revisar")
            else:
                e = Entry(month, d.key, d.code, d.name, origin=f"También en {prev.month:%m/%Y}")
        else:
            e.code, e.name = d.code, d.name
        out.append(e)
    out += [e for k, e in current.items() if k not in flagged and e.has_text]
    if plan:
        for e in out:
            if e.key in plan:
                e.plan = tuple(plan[e.key])
    out += [e for e in entries if not e.key or (e.month != month and e.has_text)]
    return out


def _light_formula(rng: str, th: Thresholds) -> str:
    """Semáforo del peor de los dos índices de la fila (SPI y CPI)."""
    low, high = f"MIN({rng})", f"MAX({rng})"
    return (f'=IF(COUNT({rng})=0,"",IF(OR({low}<{th.red_low!r},{high}>{th.red_high!r}),"Rojo",'
            f'IF(OR({low}<{th.yellow_low!r},{high}>{th.yellow_high!r}),"Amarillo","Verde")))')


def _trend_formula(i: int) -> str:
    """Tendencia del paquete: 'Empeora' si el SPI o el CPI se alejan de 1 más que TREND_TOL respecto del mes anterior,
    'Mejora' si alguno se acerca y ninguno se aleja, 'Estable' en otro caso ('' sin índices de los dos meses)."""
    def pair(cur: int, prev: int) -> tuple[str, str]:
        a, b = f"{col(cur)}{i}", f"{col(prev)}{i}"
        return f"AND(ISNUMBER({a}),ISNUMBER({b}))", f"ABS(N({a})-1)-ABS(N({b})-1)"

    (ps, ds), (pc, dc) = pair(A_SPI, A_SPI_PREV), pair(A_CPI, A_CPI_PREV)
    tol = repr(TREND_TOL)
    return (f'=IF(NOT(OR({ps},{pc})),"",IF(OR(AND({ps},{ds}>{tol}),AND({pc},{dc}>{tol})),"Empeora",'
            f'IF(OR(AND({ps},{ds}<-{tol}),AND({pc},{dc}<-{tol})),"Mejora","Estable")))')


def _put(ws, row: int, column: int, value) -> None:
    """Texto siempre como texto (nunca como fórmula); números y fechas, tal cual."""
    if isinstance(value, str):
        put_text(ws, row, column, value)
    else:
        ws.cell(row=row, column=column, value=value)


def write_sheet(wb, entries: list[Entry], row_of: dict[str, int], months: int, th: Thresholds,
                decimals: int, month: date | None = None, found: int | None = None):
    """Escribe (o vuelve a escribir) la hoja. row_of: clave -> fila de las hojas PV, EV y AC."""
    if SHEET_ANALYSIS in wb.sheetnames:
        ws = wb[SHEET_ANALYSIS]
        ws.delete_rows(1, ws.max_row)
        ws.conditional_formatting = type(ws.conditional_formatting)()
        ws.auto_filter.ref = None
    else:
        index = wb.sheetnames.index(SHEET_IND) + 1 if SHEET_IND in wb.sheetnames else len(wb.sheetnames)
        ws = wb.create_sheet(SHEET_ANALYSIS, index)
    money = money_format(decimals)
    first, last = col(FIRST_MONTH_COL), col(FIRST_MONTH_COL + months - 1)
    header = f"'{SHEET_PV}'!${first}${HEADER_ROW}:${last}${HEADER_ROW}"

    ws.cell(row=1, column=1, value="Análisis de desviaciones").font = FONT_TITLE
    ws.cell(row=2, column=1, value=(
        f"Paquetes con SPI o CPI acumulado fuera del rango verde ({fmt_decimal(th.yellow_low)} a "
        f"{fmt_decimal(th.yellow_high)}) en el mes de corte. Semáforo: {th.describe()}. Tendencia: SPI y CPI "
        f"frente a los del mes anterior (Empeora si alguno se aleja de 1 más de {fmt_decimal(TREND_TOL)}). Desvío "
        "de fin, holgura y crítica: cronograma de MS Project en el mes del análisis. Las columnas amarillas se "
        "completan a mano y se conservan al actualizar; las filas de meses anteriores sin texto se "
        "eliminan.")).font = FONT_NOTE
    if month is not None and found is not None:
        ws.cell(row=STATUS_ROW, column=1, value=(
            f"Último análisis: {month:%m/%Y}, {found} "
            f"{'paquete fuera de umbral' if found == 1 else 'paquetes fuera de umbral'}.")).font = Font(bold=True)
    _header(ws, HEAD_ROW, list(HEADERS.items()))

    # Último mes primero; dentro del mes, en el orden de la EDT
    entries = sorted(entries, key=lambda e: (-(e.month.toordinal() if e.month else 0),
                                             row_of.get(e.key, 10 ** 7), e.code))
    for i, e in enumerate(entries, start=FIRST):
        if not e.key:
            for c_, v in enumerate(e.raw or [], start=1):
                _put(ws, i, c_, v)
        else:
            ws.cell(row=i, column=A_MONTH, value=e.month)
            put_text(ws, i, A_CODE, e.code)
            put_text(ws, i, A_NAME, e.name)
            put_text(ws, i, A_ORIGIN, e.origin)
            for c_, v in zip(TEXT_COLS, e.texts):
                _put(ws, i, c_, v)
            for c_, v in zip(PLAN_COLS, e.plan):
                _put(ws, i, c_, v)
            put_text(ws, i, A_KEY, e.key)
            r = row_of.get(e.key)
            if r is not None:
                def cum(sheet: str, until: str = f"${col(A_MONTH)}{i}") -> str:
                    return (f"SUMIFS('{sheet}'!${first}${r}:${last}${r},{header},"
                            f"\"<=\"&{until})")
                # Mes anterior con EV cargado (con informes bimestrales, el del informe anterior)
                a = f"${col(A_MONTH)}{i}"
                month_n = f"((YEAR({a})-YEAR({NAME_START}))*12+MONTH({a})-MONTH({NAME_START})+1)"
                prev_n = loaded_month(month_n, 1, months)
                before = f"EDATE({NAME_START},{prev_n}-1)"
                pv1, ev1, ac1 = cum(SHEET_PV, before), cum(SHEET_EV, before), cum(SHEET_AC, before)
                L = {c_: col(c_) for c_ in HEADERS}
                formulas = {
                    A_BAC: f"='{SHEET_PV}'!${col(COL_BAC)}${r}",
                    A_PV: "=" + cum(SHEET_PV), A_EV: "=" + cum(SHEET_EV), A_AC: "=" + cum(SHEET_AC),
                    A_SV: f'=IFERROR({L[A_EV]}{i}-{L[A_PV]}{i},"")',
                    A_CV: f'=IFERROR({L[A_EV]}{i}-{L[A_AC]}{i},"")',
                    A_SPI: f'=IFERROR({L[A_EV]}{i}/{L[A_PV]}{i},"")',
                    A_CPI: f'=IFERROR({L[A_EV]}{i}/{L[A_AC]}{i},"")',
                    A_LIGHT: _light_formula(f"{L[A_SPI]}{i}:{L[A_CPI]}{i}", th),
                    A_SPI_PREV: f'=IFERROR(IF({prev_n}="","",{ev1}/{pv1}),"")',
                    A_CPI_PREV: f'=IFERROR(IF({prev_n}="","",{ev1}/{ac1}),"")',
                    A_TREND: _trend_formula(i),
                }
                for c_, f in formulas.items():
                    ws.cell(row=i, column=c_, value=f)
        for c_ in range(1, A_KEY):
            cell = ws.cell(row=i, column=c_)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=c_ in (A_NAME, A_ORIGIN) + TEXT_COLS)
            if c_ in TEXT_COLS:
                cell.fill = FILL_INPUT
        ws.cell(row=i, column=A_MONTH).number_format = FMT_MONTH
        for c_ in (A_BAC, A_PV, A_EV, A_AC, A_SV, A_CV):
            ws.cell(row=i, column=c_).number_format = money
        for c_ in (A_SPI, A_CPI, A_SPI_PREV, A_CPI_PREV):
            ws.cell(row=i, column=c_).number_format = FMT_RATIO
        ws.cell(row=i, column=A_VARIANCE).number_format = FMT_DAYS
        ws.cell(row=i, column=A_SLACK).number_format = FMT_SLACK
        for c_ in (A_TREND, A_VARIANCE, A_SLACK, A_CRITICAL):
            ws.cell(row=i, column=c_).alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(row=i, column=A_DATE).number_format = FMT_DATE

    end = max(FIRST, FIRST + len(entries) - 1)
    _traffic_light(ws, f"{col(A_SPI)}{FIRST}:{col(A_CPI)}{end}", th)
    _traffic_light(ws, f"{col(A_SPI_PREV)}{FIRST}:{col(A_CPI_PREV)}{end}", th)
    light = f"{col(A_LIGHT)}{FIRST}:{col(A_LIGHT)}{end}"
    for text, fill in (("Rojo", FILL_RED), ("Amarillo", FILL_YELLOW), ("Verde", FILL_GREEN)):
        ws.conditional_formatting.add(light, FormulaRule(formula=[f'${col(A_LIGHT)}{FIRST}="{text}"'], fill=fill))
    trend = f"{col(A_TREND)}{FIRST}:{col(A_TREND)}{end}"
    for text, fill in (("Empeora", FILL_RED), ("Mejora", FILL_GREEN)):
        ws.conditional_formatting.add(trend, FormulaRule(formula=[f'${col(A_TREND)}{FIRST}="{text}"'], fill=fill))
    v, s, k = (f"${col(c_)}{FIRST}" for c_ in (A_VARIANCE, A_SLACK, A_CRITICAL))
    variance = f"{col(A_VARIANCE)}{FIRST}:{col(A_VARIANCE)}{end}"
    ws.conditional_formatting.add(variance, FormulaRule(
        formula=[f'AND(ISNUMBER({v}),{v}>0,OR({k}="Sí",AND(ISNUMBER({s}),{s}<0)))'], fill=FILL_RED, stopIfTrue=True))
    ws.conditional_formatting.add(variance, FormulaRule(formula=[f"AND(ISNUMBER({v}),{v}>0)"], fill=FILL_YELLOW,
                                                        stopIfTrue=True))
    ws.conditional_formatting.add(f"{col(A_CRITICAL)}{FIRST}:{col(A_CRITICAL)}{end}",
                                  FormulaRule(formula=[f'{k}="Sí"'], font=Font(bold=True, color="C00000")))
    for c_, w in WIDTHS.items():
        ws.column_dimensions[col(c_)].width = w
        ws.column_dimensions[col(c_)].hidden = False      # hojas de versiones anteriores: la clave estaba en la S
    ws.column_dimensions[col(A_KEY)].hidden = True
    # El filtro incluye la columna oculta de la clave: al ordenar desde el filtro, la clave acompaña a su fila
    ws.auto_filter.ref = f"A{HEAD_ROW}:{col(A_KEY)}{end}"
    ws.freeze_panes = ws.cell(row=FIRST, column=A_BAC)
    ws.print_title_rows = f"{HEAD_ROW}:{HEAD_ROW}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    return ws
