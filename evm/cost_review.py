"""Depuración de costes: PV planificado de una empresa frente a su previsión de costes.

Mientras dura el ajuste del cronograma, la empresa pega en la hoja PV-depuración su previsión mensual (EDT cliente,
actividad, EDT interna, actividad interna, BAC y un importe por mes, tal como sale de su plan de costes). Con fórmulas
de Excel, cada línea pegada se asigna a una fila del libro y la hoja PV de la empresa muestra, debajo de cada fila, su
fila «(coste)» con la previsión. La hoja 'Depuración costes' compara las dos curvas acumuladas de un elemento EDT. Al
terminar, el libro de entrega se genera sin depuración (con --migrar-desde, que conserva el BAC).

Reglas de asignación de cada línea pegada (columna A, EDT cliente):
- código numérico (2.3.4.1): su fila del libro; un código más detallado que el libro (2.3.4.1.5) va a la fila de la
  que depende. Los ceros finales no cuentan (1.2 = 1.2.0) y se admite el código convertido en número por Excel;
- código no numérico o vacío (líneas de EDT interna como ABC01 en la columna C, o códigos como XRF01 en la columna
  A): pertenece al último código numérico de las líneas anteriores;
- una línea de EDT interna que repite la EDT cliente (columnas A y C) pertenece a esa EDT cliente, como si la columna
  A estuviera vacía;
- una línea de EDT cliente (sin EDT interna) con líneas que dependen de ella (EDT interna o códigos más detallados)
  es un subtotal: no se suma. Una línea de EDT interna nunca es un subtotal.
"""
from __future__ import annotations

from datetime import date

from openpyxl.chart import LineChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from evm.config import Config
from evm.fields import fmt_number
from evm.messages import Issue
from evm.wbs import KIND_RESIDUAL, Report, plural
from evm.workbook import (BORDER, COLOR_COST, FILL_COST, FILL_HEADER, FILL_INPUT, FILL_RED, FILL_TOTAL, FMT_MONTH,
                          FMT_PCT, FONT_BOLD, FONT_HEADER, FONT_NOTE, FONT_TITLE, HEADER_ROW, SHEET_PV, _header,
                          _Layout, col, company_sheet, money_format, put_text)

SHEET_PASTE = "PV-depuración"           # previsión pegada por la empresa
SHEET_CHART = "Depuración costes"       # gráfico comparativo
SHEET_HELPER = "_depuracion"            # hoja oculta con los cálculos de asignación

# Hoja PV-depuración
P_TOTAL_ROW, P_ASSIGNED_ROW, P_MISSING_ROW = 3, 4, 5
P_MONTHNUM_ROW, P_HEADER_ROW, P_FIRST = 6, 7, 8
P_LINES = 2000
P_LAST = P_FIRST + P_LINES - 1
P_CODE, P_DESC, P_INT_CODE, P_INT_DESC, P_BAC, P_FIRST_MONTH = 1, 2, 3, 4, 5, 6
PASTE_MONTHS = 120                      # meses del plan de costes de la empresa (al menos los del horizonte)
# Libros generados antes de las columnas de EDT interna: código, descripción, BAC (con este encabezado) y meses desde D
LEGACY_BAC_HEADER, LEGACY_FIRST_MONTH = "BAC (total)", 4

# Hoja _depuracion: columnas por línea pegada (filas P_FIRST a P_LAST, alineadas con PV-depuración)
(H_TEXT, H_NUM, H_NORM, H_TOTAL, H_CHILD, H_GROUP, H_EFF, H_PARENT, H_COUNT, H_TRUNC, H_TARGET, H_HORIZON,
 H_SUBDIFF, H_BACDIFF, H_INT) = range(1, 16)
# Tabla de códigos del libro (desde la fila 2): código sin ceros finales, fila si el código coincide y fila si la
# línea depende del código. Un resumen del libro sin fila propia lleva -1 (la línea queda sin asignar).
L_KEY, L_EXACT, L_PREFIX = 17, 18, 19
NO_ROW = -1
# Lista del selector de la hoja 'Depuración costes' (desde la fila 2) y fila elegida
S_CODE, S_ROW = 21, 22
SELECTED = "$X$1"
# Indicador de alineamiento: diferencia acumulada (previsión - PV planificado) por mes, en la fila de cada elemento
CUM_FIRST = 26

# Hoja 'Depuración costes'
G_SELECT = "B4"
G_HEADER_ROW, G_FIRST = 11, 12
(G_NUM, G_DATE, G_PV, G_PVC, G_COST, G_COSTC, G_DIFF) = range(1, 8)
STATUS_SUBTOTAL = "Subtotal (no se suma)"
STATUS_SUMMARY = "Resumen sin fila propia"
STATUS_NONE = "Sin fila del libro"
FILL_OUTSIDE = PatternFill("solid", fgColor="7F7F7F")     # meses posteriores al horizonte del libro
FILL_ALERT = PatternFill("solid", fgColor="F28B82", bgColor="F28B82")   # se distingue del naranja de «(coste)»


def paste_months(months: int) -> int:
    """Columnas de meses de la hoja PV-depuración: 120 o los meses del horizonte, si son más."""
    return max(PASTE_MONTHS, months)


def review_sheet(report: Report) -> str:
    """Hoja que se depura: 'PV' o 'PV-EmpresaA'."""
    if report.cost_review == "A":
        return company_sheet(SHEET_PV, report.company_a)
    if report.cost_review == "B":
        return company_sheet(SHEET_PV, report.company_b)
    return SHEET_PV


def review_company(report: Report) -> str:
    """Empresa que se depura ('' en un libro de una sola empresa)."""
    return {"A": report.company_a, "B": report.company_b}.get(report.cost_review or "", "")


def norm_code(code) -> str | None:
    """Código EDT numérico sin los ceros finales ('1.2.0' -> '1.2'); None si no es numérico (las mismas reglas que
    las fórmulas de la hoja _depuracion)."""
    text = str(code if code is not None else "").strip().replace(",", ".")
    parts = text.split(".")
    if not text or not all(p and p.isascii() and p.isdigit() for p in parts):
        return None
    while len(parts) > 1 and parts[-1] == "0":
        parts.pop()
    return ".".join(parts)


def lookup_rows(report: Report) -> tuple[dict[str, tuple[object, object]], list, list[str]]:
    """Códigos del libro que pueden recibir la previsión: {código sin ceros finales: (fila exacta, fila de lo que
    depende del código)}, con las filas de la EDT (ReportRow) o NO_ROW. Devuelve también las filas con código no
    numérico y los códigos repetidos."""
    entries: dict[str, tuple[object, object]] = {}
    nonnumeric = []
    repeated: list[str] = []
    for row in report.rows:
        if row.kind == KIND_RESIDUAL:
            continue                    # recibe lo que depende de su resumen sin fila propia
        key = norm_code(row.code)
        if key is None:
            nonnumeric.append(row)
            continue
        if row.is_input:
            value = (row, row)
        else:
            residual = next((c for c in row.children if c.kind == KIND_RESIDUAL), None)
            value = (NO_ROW, residual or NO_ROW)
        if key in entries:
            repeated.append(row.code)
            continue
        entries[key] = value
    return entries, nonnumeric, repeated


def review_issues(report: Report) -> list[Issue]:
    """Observaciones de la hoja Validación de un libro con depuración de costes."""
    if report.cost_review is None:
        return []
    company = review_company(report)
    sheet = review_sheet(report)
    issues = [Issue("DEP-001", "", "",
                    f"Depuración de costes{' de ' + company if company else ''}: se debe pegar la previsión mensual "
                    f"en la hoja {SHEET_PASTE} (desde la celda A{P_FIRST}); las filas «(coste)» de la hoja {sheet} y "
                    f"la hoja {SHEET_CHART} la comparan con el PV planificado. El libro de entrega se debe generar sin "
                    "depuracion_costes, con --migrar-desde para conservar el BAC.")]
    _, nonnumeric, repeated = lookup_rows(report)
    if nonnumeric:
        codes = ", ".join(r.code or "(sin código)" for r in nonnumeric[:10])
        more = f" y {len(nonnumeric) - 10} más" if len(nonnumeric) > 10 else ""
        issues.append(Issue("DEP-002", "", "",
                            f"Depuración de costes: {plural(len(nonnumeric), 'fila tiene', 'filas tienen')} código EDT "
                            f"no numérico ({codes}{more}) y no pueden recibir la previsión de la hoja {SHEET_PASTE}."))
    if repeated:
        issues.append(Issue("DEP-003", "", "",
                            "Depuración de costes: códigos EDT repetidos en el libro (" + ", ".join(repeated[:10]) +
                            f"); la previsión de la hoja {SHEET_PASTE} se asigna a la primera fila de cada código."))
    return issues


def paste_sum(i: int, r: int) -> str:
    """Fórmula del mes i (base 0) de la fila «(coste)» de la fila de ingreso r: suma de las líneas pegadas asignadas
    a la fila."""
    L = col(P_FIRST_MONTH + i)
    return (f"=SUMIFS('{SHEET_PASTE}'!${L}${P_FIRST}:${L}${P_LAST},"
            f"'{SHEET_HELPER}'!${col(H_TARGET)}${P_FIRST}:${col(H_TARGET)}${P_LAST},{r})")


# ----------------------------------------------------------------------------
# Indicador de alineamiento (depuracion_alineamiento)
# ----------------------------------------------------------------------------
def _threshold(cfg: Config) -> float:
    return cfg.cost_review_threshold / 100


def alignment_cells(ws, cfg: Config, lay: _Layout, r: int, c: int) -> None:
    """Columnas del indicador en la fila «(coste)» c: diferencia de BAC y diferencia máxima acumulada entre las dos
    curvas, en % del BAC de la fila r."""
    first, last = col(CUM_FIRST), col(CUM_FIRST + lay.months - 1)
    cum = f"'{SHEET_HELPER}'!${first}{r}:${last}{r}"
    values = {lay.total_col + 1: f'=IF(N($D{r})=0,"",($D{c}-$D{r})/$D{r})',
              lay.total_col + 2: f'=IF(N($D{r})=0,"",MAX(MAX({cum}),-MIN({cum}))/$D{r})'}
    for c_, v in values.items():
        cell = ws.cell(row=c, column=c_, value=v)
        cell.number_format = FMT_PCT
        cell.fill = FILL_COST
        cell.font = Font(italic=True, color=COLOR_COST)


def alignment_format(ws, report: Report, cfg: Config, lay: _Layout) -> None:
    """Encabezados del indicador y resaltado: meses cuya diferencia acumulada supera el umbral y diferencias de BAC o
    acumuladas que lo superan."""
    th = cfg.cost_review_threshold
    _header(ws, HEADER_ROW, [(lay.total_col + 1, "Dif. BAC (% del BAC)"),
                             (lay.total_col + 2, f"Dif. máx. acumulada (umbral {fmt_number(th, trim=True)} %)")])
    ws.row_dimensions[HEADER_ROW].height = 30
    for c_ in (lay.total_col + 1, lay.total_col + 2):
        ws.column_dimensions[col(c_)].width = 14
    rows = lay.cost_rows(report)
    if not rows:
        return
    t = _threshold(cfg)
    c0 = rows[0]
    first, last = col(lay.first_col), col(lay.last_col)
    months = " ".join(f"{first}{c}:{last}{c}" for c in rows)
    ws.conditional_formatting.add(months, FormulaRule(
        formula=[f"AND(N($D{c0 - 1})>0,ABS(SUM(${first}{c0}:{first}{c0})-SUM(${first}{c0 - 1}:{first}{c0 - 1}))"
                 f">{t!r}*$D{c0 - 1})"], fill=FILL_ALERT))
    x = col(lay.total_col + 1)
    y = col(lay.total_col + 2)
    ws.conditional_formatting.add(" ".join(f"{x}{c}:{y}{c}" for c in rows), FormulaRule(
        formula=[f"AND(ISNUMBER({x}{c0}),ABS({x}{c0})>{t!r})"], fill=FILL_ALERT))


# ----------------------------------------------------------------------------
# Hojas
# ----------------------------------------------------------------------------
def _digits_removed(ref: str) -> str:
    out = ref
    for ch in "0123456789.":
        out = f'SUBSTITUTE({out},"{ch}","")'
    return out


def _trunc_chain(ref: str, depth: int, keys: str, prefix: str) -> str:
    """Fila de la que depende el código ref (sin contar el código completo): el código más largo de la tabla del
    libro que es un prefijo de ref, por segmentos ('2.3.4.1.5' -> '2.3.4.1' -> '2.3.4' -> ...); 0 si no hay."""
    out = "0"
    for k in range(1, depth + 1):
        part = f'LEFT({ref},FIND("|",SUBSTITUTE({ref},".","|",{k}))-1)'
        out = f"IFERROR(INDEX({prefix},MATCH({part},{keys},0)),{out})"
    return out


def write_helper(ws, report: Report, lay: _Layout, cfg: Config) -> None:
    """Hoja oculta _depuracion: asignación de cada línea pegada, tabla de códigos del libro, lista del selector y
    diferencias acumuladas del indicador."""
    rs = review_sheet(report)
    n_paste = paste_months(lay.months)
    last_paste = col(P_FIRST_MONTH + n_paste - 1)
    last_horizon = col(P_FIRST_MONTH + lay.months - 1)
    paste = f"'{SHEET_PASTE}'"
    heads = {H_TEXT: "codigo", H_NUM: "numerico", H_NORM: "codigo_normalizado", H_TOTAL: "total_linea",
             H_CHILD: "subtarea", H_GROUP: "fila_grupo", H_EFF: "codigo_efectivo", H_PARENT: "subtotal",
             H_COUNT: "se_suma", H_TRUNC: "fila_por_dependencia", H_TARGET: "fila_libro", H_HORIZON: "total_horizonte",
             H_SUBDIFF: "dif_subtotal", H_BACDIFF: "dif_bac", H_INT: "edt_interna"}
    for c_, text in heads.items():
        ws.cell(row=P_FIRST - 1, column=c_, value=text)
    tables = {L_KEY: "edt", L_EXACT: "fila_exacta", L_PREFIX: "fila_dependencia", S_CODE: "selector", S_ROW: "fila",
              CUM_FIRST - 1: "fila_libro"}
    for c_, text in tables.items():
        ws.cell(row=1, column=c_, value=text)

    entries, _nonnumeric, _repeated = lookup_rows(report)

    def row_number(value) -> int:
        return value if value == NO_ROW else lay.r(value)

    for i, (key, (exact, prefix)) in enumerate(entries.items(), start=2):
        put_text(ws, i, L_KEY, key)
        ws.cell(row=i, column=L_EXACT, value=row_number(exact))
        ws.cell(row=i, column=L_PREFIX, value=row_number(prefix))
    n_keys = max(2, len(entries) + 1)
    keys = f"${col(L_KEY)}$2:${col(L_KEY)}${n_keys}"
    exact_rng = f"${col(L_EXACT)}$2:${col(L_EXACT)}${n_keys}"
    prefix_rng = f"${col(L_PREFIX)}$2:${col(L_PREFIX)}${n_keys}"
    depth = max([key.count(".") + 1 for key in entries] + [1])

    def rng(c_: int) -> str:
        return f"${col(c_)}${P_FIRST}:${col(c_)}${P_LAST}"

    norm = f"${col(H_NORM)}$1:${col(H_NORM)}${P_LAST}"
    for r in range(P_FIRST, P_LAST + 1):
        a, b, cn, d, e, f, g, h, i_, j, k, n = (f"{col(x)}{r}" for x in (H_TEXT, H_NUM, H_NORM, H_TOTAL, H_CHILD,
                                                                        H_GROUP, H_EFF, H_PARENT, H_COUNT, H_TRUNC,
                                                                        H_TARGET, H_INT))
        prev = f"{col(H_GROUP)}{r - 1}"
        int_code = f'LEN(TRIM({paste}!{col(P_INT_CODE)}{r}&""))>0'
        cells = {
            H_TEXT: f'=TRIM(SUBSTITUTE({paste}!A{r}&"",",","."))',
            H_NUM: (f'=IF({a}="",0,IF(AND(LEN({_digits_removed(a)})=0,LEFT({a},1)<>".",RIGHT({a},1)<>".",'
                    f'ISERROR(FIND("..",{a}))),1,0))'),
            H_NORM: (f'=IF({b}=1,LEFT({a},LEN({a})-2*SUMPRODUCT(--(RIGHT({a},{{2,4,6,8,10,12,14,16,18}})'
                     f'=REPT(".0",{{1,2,3,4,5,6,7,8,9}})))),"")'),
            H_TOTAL: f"=SUM({paste}!{col(P_FIRST_MONTH)}{r}:{last_paste}{r})",
            H_CHILD: (f'=IF(AND({b}=0,OR({a}<>"",{int_code},'
                      f'COUNT({paste}!{col(P_BAC)}{r}:{last_paste}{r})>0)),1,0)'),
            # Grupo: la línea numérica de la que dependen las líneas siguientes. Una línea de EDT interna que repite
            # la EDT cliente de su grupo sigue en ese grupo.
            H_GROUP: (f"=IF({b}=1,ROW(),0)" if r == P_FIRST else
                      f"=IF({b}=1,IF(AND({n}=1,{prev}>0),IF(INDEX({norm},{prev})={cn},{prev},ROW()),ROW()),{prev})"),
            H_EFF: f'=IF({b}=1,{cn},IF({f}>0,INDEX({norm},{f}),""))',
            H_PARENT: (f'=IF(AND({b}=1,{n}=0),IF(OR(COUNTIFS({rng(H_GROUP)},ROW(),{rng(H_CHILD)},1)>0,'
                       f'COUNTIFS({rng(H_GROUP)},ROW(),{rng(H_NUM)},1,{rng(H_INT)},1)>0,'
                       f'COUNTIF({rng(H_NORM)},{cn}&".*")>0),1,0),0)'),
            H_COUNT: f"=IF(OR(AND({b}=1,{h}=0),{e}=1),1,0)",
            H_TRUNC: f'=IF({g}="",0,{_trunc_chain(g, depth, keys, prefix_rng)})',
            H_TARGET: (f'=IF({i_}=0,0,IF({g}="",0,IF(AND({b}=1,{n}=0),'
                       f'IFERROR(INDEX({exact_rng},MATCH({g},{keys},0)),{j}),'
                       f'IFERROR(INDEX({prefix_rng},MATCH({g},{keys},0)),{j}))))'),
            H_HORIZON: f"=SUM({paste}!{col(P_FIRST_MONTH)}{r}:{last_horizon}{r})",
            H_SUBDIFF: (f'=IF({h}=1,{d}-SUMIFS({rng(H_TOTAL)},{rng(H_COUNT)},1,{rng(H_GROUP)},ROW())'
                        f'-SUMIFS({rng(H_TOTAL)},{rng(H_COUNT)},1,{rng(H_EFF)},{cn}&".*"),0)'),
            H_BACDIFF: (f"=IF(AND(ISNUMBER({paste}!{col(P_BAC)}{r}),ABS(N({paste}!{col(P_BAC)}{r})-{d})>0.5),1,0)"),
            H_INT: f'=IF(OR({int_code},LEN(TRIM({paste}!{col(P_INT_DESC)}{r}&""))>0),1,0)',
        }
        for c_, v in cells.items():
            ws.cell(row=r, column=c_, value=v)

    # Lista del selector: TOTAL y los elementos EDT del libro, con su fila en la hoja que se depura
    put_text(ws, 2, S_CODE, "TOTAL")
    ws.cell(row=2, column=S_ROW, value=lay.total_row)
    for i, row in enumerate(report.rows, start=3):
        put_text(ws, i, S_CODE, row.code)
        ws.cell(row=i, column=S_ROW, value=lay.r(row))
    n_list = len(report.rows) + 2
    ws[SELECTED] = (f"=IFERROR(INDEX(${col(S_ROW)}$2:${col(S_ROW)}${n_list},MATCH('{SHEET_CHART}'!${G_SELECT[0]}$"
                    f"{G_SELECT[1:]},${col(S_CODE)}$2:${col(S_CODE)}${n_list},0)),{lay.total_row})")

    if cfg.cost_review_alignment:
        rows = [lay.r(x) for x in report.rows] + [lay.total_row]
        for r in rows:
            ws.cell(row=r, column=CUM_FIRST - 1, value=r)
            for i in range(lay.months):
                L = col(lay.first_col + i)
                prev = f"{col(CUM_FIRST + i - 1)}{r}+" if i else ""
                ws.cell(row=r, column=CUM_FIRST + i, value=f"={prev}('{rs}'!{L}{r + 1}-'{rs}'!{L}{r})")


def selector_range(report: Report) -> str:
    return f"'{SHEET_HELPER}'!${col(S_CODE)}$2:${col(S_CODE)}${len(report.rows) + 2}"


def write_paste(ws, report: Report, cfg: Config, lay: _Layout, rows: list[list] | None = None) -> None:
    """Hoja PV-depuración: el bloque de la previsión de la empresa (desde A8), los encabezados de los meses y los
    controles de la asignación. rows: líneas trasladadas del libro anterior (--migrar-desde)."""
    company = review_company(report)
    rs = review_sheet(report)
    money = money_format(cfg.decimals)
    n_paste = paste_months(lay.months)
    status_col = P_FIRST_MONTH + n_paste + 1
    obs_col = status_col + 1
    helper = f"'{SHEET_HELPER}'"
    amounts, last_paste = col(P_BAC), col(P_FIRST_MONTH + n_paste - 1)      # BAC y meses de cada línea
    ws.cell(row=1, column=1, value="Previsión de costes para la depuración del PV" +
            (f"  -  {company}" if company else "")).font = FONT_TITLE
    ws.cell(row=2, column=1, value=(
        f"Pegar desde la celda A{P_FIRST} (pegado especial > valores): EDT cliente, actividad, EDT interna, actividad "
        f"interna, BAC y un importe por mes, con el mes 1 en la columna {col(P_FIRST_MONTH)}. Las líneas de EDT "
        "interna (p. ej. ABC01) se suman en la última EDT cliente de arriba; las líneas de EDT cliente con EDT interna "
        "o códigos más detallados debajo son subtotales y no se suman.")).font = FONT_NOTE

    labels = [(P_TOTAL_ROW, "Total pegado (sin subtotales)",
               f"=SUMIFS({helper}!${col(H_TOTAL)}${P_FIRST}:${col(H_TOTAL)}${P_LAST},"
               f"{helper}!${col(H_COUNT)}${P_FIRST}:${col(H_COUNT)}${P_LAST},1)"),
              (P_ASSIGNED_ROW, "Asignado a filas del libro (meses del horizonte)",
               f"=SUMIFS({helper}!${col(H_HORIZON)}${P_FIRST}:${col(H_HORIZON)}${P_LAST},"
               f"{helper}!${col(H_TARGET)}${P_FIRST}:${col(H_TARGET)}${P_LAST},\">0\")"),
              (P_MISSING_ROW, "Sin asignar (sin fila del libro o fuera del horizonte)",
               f"={amounts}{P_TOTAL_ROW}-{amounts}{P_ASSIGNED_ROW}")]
    for r, label, formula in labels:
        ws.cell(row=r, column=P_DESC, value=label).font = FONT_BOLD
        cell = ws.cell(row=r, column=P_BAC, value=formula)
        cell.number_format = money
        cell.border = BORDER
    ws.conditional_formatting.add(f"{amounts}{P_MISSING_ROW}", FormulaRule(
        formula=[f"ABS({amounts}{P_MISSING_ROW})>0.5"], fill=FILL_RED))
    ws.cell(row=P_TOTAL_ROW, column=P_FIRST_MONTH,
            value=f"Resultado: filas «(coste)» de la hoja {rs} y hoja '{SHEET_CHART}'.").font = FONT_NOTE
    ws.cell(row=P_MISSING_ROW, column=P_FIRST_MONTH,
            value=f"Fila del libro de cada línea: columnas {col(status_col)} ('Asignado a') y {col(obs_col)} "
                  "('Observación'), a la derecha de los meses.").font = FONT_NOTE

    ws.cell(row=P_MONTHNUM_ROW, column=P_BAC, value="Mes N°").font = FONT_BOLD
    ws.cell(row=P_MONTHNUM_ROW, column=P_BAC).alignment = Alignment(horizontal="right")
    _header(ws, P_HEADER_ROW, [(P_CODE, "EDT cliente"), (P_DESC, "Actividad"), (P_INT_CODE, "EDT interna"),
                               (P_INT_DESC, "Actividad interna"), (P_BAC, "BAC"),
                               (status_col, "Asignado a"), (obs_col, "Observación")])
    for i in range(n_paste):
        c_ = P_FIRST_MONTH + i
        num = ws.cell(row=P_MONTHNUM_ROW, column=c_, value=i + 1)
        num.alignment = Alignment(horizontal="center")
        num.font = FONT_BOLD
        head = ws.cell(row=P_HEADER_ROW, column=c_, value=report.horizon.month_start(i))
        head.number_format = FMT_MONTH
        head.font = FONT_HEADER
        head.alignment = Alignment(horizontal="center", vertical="center")
        head.border = BORDER
        head.fill = FILL_HEADER if i < lay.months else FILL_OUTSIDE
        ws.column_dimensions[col(c_)].width = 11
        ws.column_dimensions[col(c_)].number_format = money
    if n_paste > lay.months:
        ws.cell(row=P_HEADER_ROW, column=P_FIRST_MONTH + lay.months).comment = Comment(
            f"Meses posteriores al horizonte del libro ({lay.months} meses): no se asignan a las filas del libro.",
            "evm_export")
    ws.column_dimensions[col(P_CODE)].number_format = "@"
    ws.column_dimensions[col(P_INT_CODE)].number_format = "@"
    ws.column_dimensions[col(P_BAC)].number_format = money

    for k, values in enumerate((rows or [])[:P_LINES]):
        for c_, v in enumerate(values, start=1):
            if v is not None and c_ < status_col:
                cell = put_text(ws, P_FIRST + k, c_, v)
                if c_ >= P_BAC and isinstance(v, (int, float)) and not isinstance(v, bool):
                    cell.number_format = money

    for r in range(P_FIRST, P_LAST + 1):
        h = {x: f"{helper}!{col(x)}{r}" for x in (H_NUM, H_CHILD, H_GROUP, H_PARENT, H_TARGET, H_BACDIFF, H_TOTAL,
                                                  H_SUBDIFF, H_HORIZON)}
        ws.cell(row=r, column=status_col, value=(
            f'=IF(AND({h[H_NUM]}=0,{h[H_CHILD]}=0),"",IF({h[H_PARENT]}=1,"{STATUS_SUBTOTAL}",IF({h[H_TARGET]}>0,'
            f'INDEX(\'{rs}\'!$A$1:$A${lay.total_row + 1},{h[H_TARGET]}),IF({h[H_TARGET]}={NO_ROW},'
            f'"{STATUS_SUMMARY}","{STATUS_NONE}"))))'))
        s = f"{col(status_col)}{r}"
        ws.cell(row=r, column=obs_col, value=(
            f'=IF({s}="","",TRIM(IF({h[H_BACDIFF]}=1,"BAC distinto de la suma de los meses. ","")'
            f'&IF(AND({h[H_NUM]}=0,{h[H_GROUP]}=0),"Sin EDT cliente en las líneas anteriores. ","")'
            f'&IF(AND({h[H_PARENT]}=1,{h[H_TOTAL]}<>0,ABS({h[H_SUBDIFF]})>0.5),'
            f'"Subtotal distinto de la suma de sus líneas. ","")'
            f'&IF(AND({h[H_TARGET]}>0,ABS({h[H_TOTAL]}-{h[H_HORIZON]})>0.5),'
            f'"Importes fuera del horizonte del libro. ","")'
            f'&IF(SUMPRODUCT(ISTEXT({amounts}{r}:{last_paste}{r})*(LEN({amounts}{r}:{last_paste}{r})>0))>0,'
            f'"Importes guardados como texto: no se suman. ","")))'))
    status = f"{col(status_col)}{P_FIRST}"
    obs = f"{col(obs_col)}{P_FIRST}"
    area = f"{col(status_col)}{P_FIRST}:{col(obs_col)}{P_LAST}"
    ws.conditional_formatting.add(area, FormulaRule(
        formula=[f'OR(${status}="{STATUS_NONE}",${status}="{STATUS_SUMMARY}",${obs}<>"")'], fill=FILL_RED))
    ws.conditional_formatting.add(f"A{P_FIRST}:{col(P_BAC)}{P_LAST}", FormulaRule(
        formula=[f'${status}="{STATUS_SUBTOTAL}"'], fill=FILL_TOTAL, font=Font(italic=True, color="595959")))

    for c_, w in {P_CODE: 12, P_DESC: 34, P_INT_CODE: 12, P_INT_DESC: 30, P_BAC: 14, status_col - 1: 3,
                  status_col: 20, obs_col: 60}.items():
        ws.column_dimensions[col(c_)].width = w
    ws.row_dimensions[P_HEADER_ROW].height = 30
    ws.freeze_panes = ws.cell(row=P_FIRST, column=P_FIRST_MONTH)


def write_chart_sheet(ws, report: Report, cfg: Config, lay: _Layout) -> None:
    """Hoja 'Depuración costes': curvas acumuladas del PV planificado (BAC × %) y de la previsión del elemento EDT
    elegido en la celda B4 (por defecto, TOTAL)."""
    company = review_company(report)
    rs = review_sheet(report)
    money = money_format(cfg.decimals)
    sel = f"'{SHEET_HELPER}'!{SELECTED}"
    ws.cell(row=1, column=1, value="Depuración de costes  -  PV planificado frente a la previsión" +
            (f" de {company}" if company else "")).font = FONT_TITLE
    ws.cell(row=2, column=1, value=(
        f"Se elige el elemento EDT en la celda {G_SELECT} (TOTAL o un código de la hoja {rs}). PV planificado: BAC × "
        f"% planificado de MS Project. Previsión: fila «(coste)», con lo pegado en la hoja {SHEET_PASTE}. El objetivo "
        "es que los dos BAC coincidan y que las curvas se parezcan; se ajustan en MS Project las tareas, sus fechas "
        "y el peso de los hitos, y se regenera el libro con --migrar-desde.")).font = FONT_NOTE
    ws.cell(row=2, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=8)
    ws.row_dimensions[2].height = 45

    ws.cell(row=4, column=1, value="Elemento EDT").font = FONT_BOLD
    pick = ws[G_SELECT]
    pick.value = "TOTAL"
    pick.fill = FILL_INPUT
    pick.border = BORDER
    pick.alignment = Alignment(horizontal="center")
    ws.cell(row=4, column=3, value=f"=IFERROR(INDEX('{rs}'!$B$1:$B${lay.total_row + 1},{sel}),\"\")")
    dv = DataValidation(type="list", formula1=f"={selector_range(report)}", allow_blank=False,
                        showErrorMessage=True, errorTitle="Elemento EDT",
                        error="Se debe elegir TOTAL o un código EDT de la lista.")
    ws.add_data_validation(dv)
    dv.add(G_SELECT)

    end = G_FIRST + lay.months - 1
    diff = f"${col(G_DIFF)}${G_FIRST}:${col(G_DIFF)}${end}"
    rows = [
        (5, "BAC (fila de porcentaje)", f"=N(INDEX('{rs}'!$D$1:$D${lay.total_row + 1},{sel}))", money),
        (6, "Previsión (fila «(coste)»)", f"=N(INDEX('{rs}'!$D$1:$D${lay.total_row + 1},{sel}+1))", money),
        (7, "Diferencia: previsión - BAC", "=B6-B5", money),
        (8, "Diferencia máxima acumulada", f"=MAX(MAX({diff}),-MIN({diff}))", money),
    ]
    for r, label, formula, fmt in rows:
        ws.cell(row=r, column=1, value=label)
        cell = ws.cell(row=r, column=2, value=formula)
        cell.number_format = fmt
        cell.border = BORDER
    for r in (7, 8):
        pct = ws.cell(row=r, column=3, value=f'=IF(B5=0,"",B{r}/B5)')
        pct.number_format = FMT_PCT
    ws.cell(row=7, column=4, value="(% del BAC)").font = FONT_NOTE
    t = _threshold(cfg)
    for r in (7, 8):
        ws.conditional_formatting.add(f"C{r}", FormulaRule(formula=[f"AND(ISNUMBER(C{r}),ABS(C{r})>{t!r})"],
                                                           fill=FILL_RED))

    _header(ws, G_HEADER_ROW, [(G_NUM, "Mes N°"), (G_DATE, "Mes"), (G_PV, "PV planificado del mes"),
                               (G_PVC, "PV planificado acumulado"), (G_COST, "Previsión del mes"),
                               (G_COSTC, "Previsión acumulada"), (G_DIFF, "Diferencia acumulada")])
    block = f"'{rs}'!${col(lay.first_col)}$1:${col(lay.last_col)}${lay.total_row + 1}"
    for i in range(lay.months):
        r = G_FIRST + i
        cells = {
            G_NUM: i + 1,
            G_DATE: lay.month_dates[i],
            G_PV: f"=N(INDEX({block},{sel},{i + 1}))",
            G_PVC: f"=SUM(${col(G_PV)}${G_FIRST}:{col(G_PV)}{r})",
            G_COST: f"=N(INDEX({block},{sel}+1,{i + 1}))",
            G_COSTC: f"=SUM(${col(G_COST)}${G_FIRST}:{col(G_COST)}{r})",
            G_DIFF: f"={col(G_COSTC)}{r}-{col(G_PVC)}{r}",
        }
        for c_, v in cells.items():
            cell = ws.cell(row=r, column=c_, value=v)
            cell.number_format = FMT_MONTH if c_ == G_DATE else ("0" if c_ == G_NUM else money)
            if c_ in (G_NUM, G_DATE):
                cell.alignment = Alignment(horizontal="center")
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 14
    for c_ in range(G_PV, G_DIFF + 1):
        ws.column_dimensions[col(c_)].width = 15
    ws.freeze_panes = ws.cell(row=G_FIRST, column=G_PV)


# Eje de montos del gráfico: la escala depende del elemento EDT elegido (millones, miles o unidades)
AXIS_FORMAT = '[>=1000000]#,##0.0,,"M";[>=1000]#,##0,"k";#,##0'


def add_chart(wb, months: int, chart_months: int, currency: str) -> None:
    """Agrega (o vuelve a agregar: openpyxl no conserva los gráficos) el gráfico de la hoja 'Depuración costes'."""
    from evm.dashboard import _date_axis, _frame, _legend, _line, _style_axis
    from openpyxl.chart.shapes import GraphicalProperties
    from openpyxl.chart.marker import Marker
    ws = wb[SHEET_CHART]
    ws._charts = []
    n = max(1, min(months, chart_months))
    last = G_FIRST + n - 1
    line = LineChart()
    for c_, color, dash in ((G_PVC, "1F4E79", None), (G_COSTC, "C55A11", "dash")):
        line.add_data(Reference(ws, min_col=c_, min_row=G_HEADER_ROW, max_row=last), titles_from_data=True)
        s = line.series[-1]
        s.smooth = False
        s.marker = Marker(symbol="none")
        s.graphicalProperties = GraphicalProperties(ln=_line(color, 2.25, dash))
    x = _date_axis(n)
    line.x_axis = x
    line.y_axis.crossAx = x.axId
    _style_axis(line.y_axis, title=currency or "Monto", fmt=AXIS_FORMAT, grid=True)
    line.y_axis.scaling.min = 0
    line.set_categories(Reference(ws, min_col=G_DATE, min_row=G_FIRST, max_row=last))
    _legend(line)
    _frame(line, "PV planificado y previsión de costes acumulados", 760, 400)
    ws.add_chart(line, "I4")


def is_legacy_paste(ws) -> bool:
    """True si la hoja PV-depuración es de un libro anterior a las columnas de EDT interna."""
    return ws.cell(row=P_HEADER_ROW, column=P_INT_CODE).value == LEGACY_BAC_HEADER


def last_paste_month(ws, months: int) -> int:
    """Último mes (N°, dentro del horizonte) con algún importe pegado en la hoja PV-depuración."""
    first = LEGACY_FIRST_MONTH if is_legacy_paste(ws) else P_FIRST_MONTH
    last = 0
    for values in ws.iter_rows(min_row=P_FIRST, max_row=min(ws.max_row, P_LAST), min_col=first,
                               max_col=first + months - 1, values_only=True):
        for i, v in enumerate(values):
            if isinstance(v, (int, float)) and not isinstance(v, bool) and v and i + 1 > last:
                last = i + 1
    return last


def read_paste(ws, months: int) -> tuple[list[list], bool]:
    """Líneas pegadas en la hoja PV-depuración (para --migrar-desde), sin las filas vacías del final, y si la hoja
    tenía el formato anterior a las columnas de EDT interna. Las líneas de ese formato se devuelven en el actual: un
    código numérico y su descripción quedan como EDT cliente y actividad; un código no numérico (XRF01) y su
    descripción, como EDT interna y actividad interna."""
    legacy = is_legacy_paste(ws)
    n = (LEGACY_FIRST_MONTH if legacy else P_FIRST_MONTH) + paste_months(months) - 1
    rows = [list(v) for v in ws.iter_rows(min_row=P_FIRST, max_row=min(ws.max_row, P_LAST), max_col=n,
                                          values_only=True)]
    while rows and all(v is None or v == "" for v in rows[-1]):
        rows.pop()
    if legacy:
        rows = [([code, desc, None, None] if norm_code(code) is not None else [None, None, code, desc]) + rest
                for code, desc, *rest in rows]
    return rows, legacy


def shift_rows(rows: list[list], offset: int, n_paste: int) -> tuple[list[list], float]:
    """Desplaza los meses de las líneas pegadas cuando cambia el mes 1 del horizonte (offset: meses que avanza el
    inicio). Devuelve las líneas y el importe que queda fuera de las columnas de meses."""
    lost = 0.0
    out = []
    for values in rows:
        head = (list(values[:P_FIRST_MONTH - 1]) + [None] * (P_FIRST_MONTH - 1))[:P_FIRST_MONTH - 1]
        months = list(values[P_FIRST_MONTH - 1:])
        new = [None] * n_paste
        for i, v in enumerate(months):
            j = i - offset
            if 0 <= j < n_paste:
                new[j] = v
            elif isinstance(v, (int, float)) and not isinstance(v, bool):
                lost += abs(v)
        out.append(head + new)
    return out, lost


def month_offset(old: date, new: date) -> int:
    return (new.year - old.year) * 12 + (new.month - old.month)
