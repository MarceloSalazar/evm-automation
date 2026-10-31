"""Conservación del BAC, el % ganado, el AC y los métodos de un libro anterior al regenerar (--migrar-desde)."""
from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date

from evm import changes
from evm.analysis import A_CODE, A_KEY, A_MONTH, A_NAME, A_ORIGIN, SHEET_ANALYSIS, TEXT_COLS, read_entries
from evm.config import ConfigError
from evm.cost_review import SHEET_PASTE, month_offset, paste_months, read_paste, review_company, shift_rows
from evm.fields import same_name
from evm.messages import SEV_INFO, SEV_WARN, Issue
from evm.methods import is_undefined
from evm.model import MonthKey, Task
from evm.wbs import (KIND_LEAF, KIND_RESIDUAL, Report, ReportRow, code_key, fmt_money, has_interior_zero,
                     iter_subtree, plural)
from evm.workbook import (PCT_PREFIX, SHEET_EV, SHEET_LISTS, book_sources, clean_text, pct_sheet, role_key, row_key,
                          source_meta, split_key)


@dataclass
class PriorRow:
    key: str                     # UniqueID (o UniqueID-D); con varios archivos, 'P2:UniqueID'
    code: str
    name: str
    uid: int | None
    residual: bool
    # Valores mensuales por serie: el % ganado de cada mes ("% EV", o en multi-empresa "% EV-A" y "% EV-B") y el
    # AC ("AC", o "AC-A" y "AC-B")
    series: dict[str, dict[MonthKey, float]]
    method: str
    origin_method: str | None = None     # método al generarse el libro (None: libro antiguo)
    source: str = ""                     # varios archivos: nombre del archivo de la fila
    bac: dict[str, float] = field(default_factory=dict)     # BAC ingresado ("" o "A" y "B")

    @property
    def method_edited(self) -> bool:
        """True si el método de la columna C fue modificado en Excel."""
        if not self.method:
            return False
        if self.origin_method is not None:
            return self.method != self.origin_method
        return self.method.strip().lower() not in ("por definir", "a definir", "")

    @property
    def has_values(self) -> bool:
        """True si tiene algún valor distinto de 0 (los meses cargados con 0 también se trasladan) o un BAC
        ingresado."""
        return any(v for data in self.series.values() for v in data.values()) or any(self.bac.values())

    def total(self, kind: str) -> float:
        """Suma de las series de EV o de AC (el EV es el BAC × el % ganado)."""
        out = 0.0
        for k, v in self.series.items():
            if k.startswith(PCT_PREFIX) and kind == SHEET_EV:
                out += self.bac.get(_series_role(k), 0.0) * sum(v.values())
            elif k.split("-")[0] == kind:
                out += sum(v.values())
        return out


def _series_role(kind: str) -> str:
    """Empresa de una serie: '% EV-A' -> 'A'; 'EV' o '% EV' -> ''."""
    head, sep, tail = kind.rpartition("-")
    return tail if sep and tail in ("A", "B") else ""


@dataclass
class PriorData:
    rows: list[PriorRow]
    status_date: date | None
    project_name: str
    max_level: int | None
    methods_list: list[str]
    multi: bool = False
    ignore_interior_zeros: bool = False     # ignorar_ceros_intermedios con que se generó el libro anterior
    # Proyecto de MS Project de cada archivo del libro anterior ({'': nombre} en un libro de un solo archivo)
    project_names: dict[str, str] = field(default_factory=dict)
    analysis: list = field(default_factory=list)     # filas de la hoja Análisis de desviaciones (analysis.Entry)
    # Resultado de plan_migration: clave de la fila nueva -> datos a escribir
    values: dict[str, dict[str, dict[MonthKey, float]]] = field(default_factory=dict)
    methods: dict[str, str] = field(default_factory=dict)
    matched: int = 0
    unmatched: list[PriorRow] = field(default_factory=list)
    bac: dict[str, dict[str, float]] = field(default_factory=dict)    # clave nueva -> BAC
    # Depuración de costes: líneas de la hoja PV-depuración (en el formato actual), empresa que se depuraba ('' =
    # libro de una sola empresa), mes 1 del horizonte, si la hoja tenía el formato anterior a las columnas de EDT
    # interna y si se trasladan al libro nuevo (lo decide plan_migration)
    review_rows: list[list] = field(default_factory=list)
    review_company: str = ""
    review_start: date | None = None
    review_legacy: bool = False
    review_keep: bool = False
    # Registro de cambios: filas de la hoja y referencia de la última ejecución (con las claves del libro nuevo
    # después de plan_migration)
    change_log: list[list] = field(default_factory=list)
    reference: changes.Reference | None = None
    # Celdas del libro anterior con un valor escrito sobre la fórmula (subtotales, TOTAL, suma de empresas)
    overwritten: list[str] = field(default_factory=list)

    def values_for(self, key: str, kind: str) -> dict[MonthKey, float]:
        return self.values.get(key, {}).get(kind, {})

    def bac_for(self, key: str, role: str) -> float | None:
        """BAC ingresado en el libro anterior para la fila nueva (None si no tenía)."""
        value = self.bac.get(key, {}).get(role)
        return value if value else None

    def method_for(self, key: str) -> str | None:
        return self.methods.get(key)


def prior_from_book(path, project_name: str) -> PriorData:
    """Lee EV, AC, métodos y fecha de corte de un libro generado anteriormente."""
    from evm.update import book_ignores_interior_zeros, read_book
    book = read_book(path)
    sources = {role_key(pct_sheet(SHEET_EV), role or None): book.pct_ev[role] for role, _c in book.roles()}
    sources.update({role_key("AC", role or None): (book.ac_a if role == "A" else book.ac_b if role == "B"
                                                   else book.ac) for role, _c in book.roles()})
    rows = []
    for s in book.input_rows:
        source, uid_text, residual = split_key(s.key)
        # Los meses cargados con 0 se conservan: indican que el mes ya se informó (curvas desde ese mes)
        series = {k: {book.horizon.key(i): v for i, v in enumerate(src[s.row]) if v is not None}
                  for k, src in sources.items()}
        bac = {role: book.bac_role[role][s.row] for role, _c in book.roles()}
        rows.append(PriorRow(
            key=s.key, code=s.code, name=s.name, uid=int(uid_text) if uid_text.isdigit() else None,
            residual=residual, series={k: v for k, v in series.items() if v},
            method=book.method[s.row], origin_method=s.origin_method, source=source,
            bac={k: v for k, v in bac.items() if v}))
    methods_list = []
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True)
    if SHEET_LISTS in wb.sheetnames:
        for (value,) in wb[SHEET_LISTS].iter_rows(min_row=2, max_col=1, values_only=True):
            if not isinstance(value, str) or not value.strip():
                break                   # la lista termina en la primera fila vacía
            if len(value.strip()) <= 60:
                methods_list.append(value.strip())
    analysis = read_entries(wb[SHEET_ANALYSIS]) if SHEET_ANALYSIS in wb.sheetnames else []
    review_rows, review_legacy = (read_paste(wb[SHEET_PASTE], book.horizon.months) if SHEET_PASTE in wb.sheetnames
                                  else ([], False))
    change_log = changes.read_rows(wb[changes.SHEET_LOG]) if changes.SHEET_LOG in wb.sheetnames else []
    reference = changes.read_reference(wb)
    wb.close()
    level = book.meta.get("nivel_maximo")
    names = {n: str(source_meta(book.meta, "proyecto", n) or "") for n in book_sources(book.meta)}
    return PriorData(rows=rows, status_date=book.status_date,
                     project_name=str(book.meta.get("proyecto") or ""),
                     max_level=int(level) if level else None, methods_list=methods_list, multi=book.multi,
                     ignore_interior_zeros=book_ignores_interior_zeros(book.meta),
                     project_names=names, analysis=analysis, review_rows=review_rows,
                     review_company=str(book.meta.get("depuracion_empresa") or ""), review_start=book.horizon.start,
                     review_legacy=review_legacy, change_log=change_log, reference=reference,
                     overwritten=list(book.formula_values))


def _migrate_analysis(prior: PriorData, key_map: dict[str, str]) -> list[Issue]:
    """Las filas del análisis de desviaciones pasan a la fila nueva de su paquete. Si el paquete ya no tiene
    fila, el texto se conserva como fila sin valores (sin clave, para no asociarla a otra fila)."""
    issues: list[Issue] = []
    moved = 0
    for e in prior.analysis:
        if not e.key:
            continue
        if e.key in key_map:
            e.key = key_map[e.key]
            moved += 1
            continue
        raw = [None] * (A_KEY - 1)
        raw[A_MONTH - 1], raw[A_CODE - 1], raw[A_NAME - 1] = e.month, e.code, e.name
        raw[A_ORIGIN - 1] = "Fila inexistente en el libro nuevo"
        for c_, v in zip(TEXT_COLS, e.texts):
            raw[c_ - 1] = v
        e.key, e.raw = "", raw
        if e.has_text:
            issues.append(Issue("MIG-015", e.code, e.name,
                                f"Análisis de desviaciones de {e.month:%m/%Y}: el paquete ya no tiene fila en el libro "
                                "nuevo; el texto se conservó, sin valores."))
    if moved:
        issues.append(Issue("MIG-016", "", "", f"Se {'trasladó' if moved == 1 else 'trasladaron'} "
                                              f"{plural(moved, 'fila', 'filas')} del análisis de desviaciones."))
    return issues


def _migrate_changes(prior: PriorData, key_map: dict[str, str], sources: dict[str, list]) -> list[Issue]:
    """El registro de cambios pasa al libro nuevo con las claves nuevas (las filas cuyo paquete ya no tiene fila
    quedan sin clave). La referencia solo se traslada en las filas asociadas una a una: una fila nueva que reúne
    varias anteriores guarda su referencia en la primera ejecución, sin registrar cambios."""
    for vals in prior.change_log:
        vals[changes.L_KEY - 1] = key_map.get(str(vals[changes.L_KEY - 1] or ""))
    moved = 0
    if prior.reference is not None:
        rows = {}
        for (key, role), ref in prior.reference.rows.items():
            new = key_map.get(key)
            if new is not None and len(sources.get(new, [])) == 1:
                rows[(new, role)] = ref
        moved = len({key for key, _role in rows})
        prior.reference = changes.Reference(prior.reference.start, prior.reference.cutoff, rows) if rows else None
    if not prior.change_log and not moved:
        return []
    parts = []
    if prior.change_log:
        parts.append(plural(len(prior.change_log), "fila de la hoja", "filas de la hoja"))
    if moved:
        parts.append(f"la referencia de {plural(moved, 'fila', 'filas')} del libro")
    return [Issue("MIG-018", "", "", f"Registro de cambios trasladado: {' y '.join(parts)}.")]


def _migrate_review(report: Report, prior: PriorData) -> list[Issue]:
    """Depuración de costes: la previsión pegada en la hoja PV-depuración pasa al libro nuevo si también se depura
    y es la misma empresa. Si cambia el mes 1 del horizonte, los meses se desplazan para conservar su fecha."""
    if not prior.review_rows:
        return []
    lines = plural(len(prior.review_rows), "línea", "líneas")
    if report.cost_review is None:
        return [Issue("DEP-004", "", "", f"La hoja {SHEET_PASTE} del libro anterior ({lines}) no se trasladó: el libro "
                                        "nuevo no tiene depuración de costes.")]
    company = review_company(report)
    if not same_name(company, prior.review_company):
        old = prior.review_company or "el libro de una sola empresa"
        return [Issue("DEP-005", "", "",
                      f"La hoja {SHEET_PASTE} del libro anterior no se trasladó: tiene la previsión de {old} y el "
                      f"libro nuevo depura {company or 'un libro de una sola empresa'}. Se debe pegar la previsión "
                      "de nuevo.")]
    prior.review_keep = True
    issues = [Issue("DEP-006", "", "", f"Se conservó el contenido de la hoja {SHEET_PASTE} ({lines}).")]
    if prior.review_legacy:
        issues.append(Issue("DEP-007", "", "", f"La hoja {SHEET_PASTE} del libro anterior tenía el formato sin EDT "
                                              "interna: los códigos numéricos quedaron como EDT cliente y los demás "
                                              "(p. ej. XRF01), como EDT interna."))
    offset = month_offset(prior.review_start, report.horizon.start) if prior.review_start else 0
    if offset:
        prior.review_rows, lost = shift_rows(prior.review_rows, offset, paste_months(report.horizon.months))
        issues.append(Issue("DEP-008", "", "",
                            f"El mes 1 del horizonte cambió ({prior.review_start:%m/%Y} a "
                            f"{report.horizon.start:%m/%Y}): los meses de la hoja {SHEET_PASTE} se desplazaron "
                            f"{abs(offset)} columnas para conservar su fecha"
                            + (f"; quedaron fuera {fmt_money(lost)}." if lost > 0.005 else "."),
                            SEV_WARN if lost > 0.005 else SEV_INFO))
    return issues


def _whole_branches(row: ReportRow) -> list[Task]:
    """Tareas de una fila -D cuya rama completa forma parte de la fila: todas sus tareas de detalle están en ella."""
    leaves = {t.uid for t in row.leaves}
    out: list[Task] = []

    def visit(t: Task) -> None:
        for c in t.children:
            if all(n.uid in leaves for n in iter_subtree(c) if not n.children):
                out.extend(iter_subtree(c))
            else:
                visit(c)

    visit(row.task)
    return out


def plan_migration(report: Report, prior: PriorData, max_level: int,
                   ignore_interior_zeros: bool = False) -> list[Issue]:
    """Asigna cada fila del libro anterior a una fila del libro nuevo (sin duplicar historiales).

    Mismo proyecto: por UniqueID; si la fila anterior quedó dentro de un paquete o de una fila -D
    nueva (p. ej. al bajar el nivel máximo), su historial se suma a esa fila.
    Proyecto distinto: por código EDT, una fila anterior por fila nueva.
    """
    issues: list[Issue] = []
    if prior.multi != report.multi:
        raise ConfigError("MIG-001",
                          "--migrar-desde no se admite entre un libro de una sola empresa y uno multi-empresa: "
                          "el historial de EV y AC no se puede repartir entre empresas. Se debe generar el libro "
                          "nuevo sin --migrar-desde.")
    # Mismo proyecto: se compara cada archivo del libro anterior con el archivo del mismo nombre del nuevo
    new_names = {s.name: s.project.name for s in report.sources} or {"": report.project.name}
    old_names = prior.project_names or {"": prior.project_name}
    same = {n: n in new_names and old_names[n] == clean_text(new_names[n]) for n in old_names}
    if "" in old_names and "" not in new_names:
        issues.append(Issue("MIG-002", "", "",
                            f"El libro anterior es de un solo archivo de MS Project ('{old_names['']}') y el nuevo "
                            "combina varios archivos: los EV/AC se asocian por código EDT."))
    elif "" in new_names and "" not in old_names:
        issues.append(Issue("MIG-003", "", "",
                            f"El libro anterior combina varios archivos de MS Project ({', '.join(old_names)}) y el "
                            "nuevo es de un solo archivo: los EV/AC se asocian por código EDT."))
    else:
        for n, ok in same.items():
            if ok:
                continue
            if n not in new_names:
                code_id = "MIG-004"
                text = (f"El archivo '{n}' del libro anterior no forma parte del libro nuevo: los EV/AC de sus "
                        "filas se asocian por código EDT.")
            else:
                code_id = "MIG-005"
                text = (f"El libro anterior corresponde al proyecto '{old_names[n]}' y el nuevo a "
                        f"'{new_names[n]}': los EV/AC se asocian por código EDT.")
            issues.append(Issue(code_id, "", "", f"[{n}] {text}" if n else text))
    if prior.max_level and prior.max_level != max_level:
        issues.append(Issue("MIG-006", "", "",
                            f"El libro anterior usaba nivel máximo {prior.max_level} y el nuevo {max_level}. Las "
                            "filas que ahora se dividen en varias no pueden repartir su historial (ver errores)."))
    if prior.ignore_interior_zeros != ignore_interior_zeros and any(
            has_interior_zero(c) for c in [o.code for o in prior.rows] + [r.code for r in report.rows]):
        old, new = ("true" if v else "false" for v in (prior.ignore_interior_zeros, ignore_interior_zeros))
        issues.append(Issue("MIG-020", "", "",
                            f"El libro anterior se generó con ignorar_ceros_intermedios = {old} y el nuevo con "
                            f"{new}: cambia el nivel de los códigos con un 0 intermedio (p. ej. 1.2.0.1). Las filas "
                            "que ahora se dividen en varias no pueden repartir su historial (ver errores)."))

    inputs = [r for r in report.rows if r.is_input]
    by_key = {row_key(r): r for r in inputs}
    by_code: dict[str, ReportRow] = {}      # código sin los ceros finales (1.2.0 = 1.2.0.0) -> fila
    for r in inputs:
        by_code.setdefault(code_key(r.code), r)
    contains: dict[tuple[str, int], ReportRow] = {}      # (archivo, UniqueID) -> paquete que contiene la tarea
    for r in inputs:
        if r.kind == KIND_LEAF:
            for n in iter_subtree(r.task):
                contains.setdefault((r.source, n.uid), r)
        elif r.kind == KIND_RESIDUAL:          # p. ej. un paquete que pasó a tener un nivel mayor al máximo
            for n in _whole_branches(r):
                contains.setdefault((r.source, n.uid), r)

    used_codes: set[str] = set()
    sources: dict[str, list[tuple[PriorRow, str]]] = {}
    key_map: dict[str, str] = {}            # clave de la fila anterior -> clave de la fila nueva
    for old in prior.rows:
        target, how = None, ""
        if same.get(old.source):
            if old.key in by_key:
                target, how = by_key[old.key], "uid"
            elif old.uid is not None and (old.source, old.uid) in contains:
                target, how = contains[(old.source, old.uid)], "rama"
        else:
            cand = by_code.get(code_key(old.code))
            if cand is not None and row_key(cand) not in used_codes:
                target, how = cand, "codigo"
                used_codes.add(row_key(cand))
        if target is None:
            if old.has_values:
                prior.unmatched.append(old)
            continue
        sources.setdefault(row_key(target), []).append((old, how))
        key_map[old.key] = row_key(target)
    issues += _migrate_analysis(prior, key_map)

    horizon = report.horizon
    clamped = early = 0.0
    for key, lst in sources.items():
        target = by_key[key]
        merged: dict[str, dict[MonthKey, float]] = {}
        merged_bac: dict[str, float] = {}
        several = len(lst) > 1
        for old, _how in lst:
            for role, v in old.bac.items():
                merged_bac[role] = merged_bac.get(role, 0.0) + v
        for old, how in lst:
            for kind, data in old.series.items():
                merged.setdefault(kind, {})
                # Varias filas anteriores en una: el % ganado se combina ponderado por su BAC
                bac = old.bac.get(_series_role(kind), 0.0)
                factor = bac if (several and kind.startswith(PCT_PREFIX)) else 1.0
                # Importe para los avisos: el EV es el BAC × el % ganado
                worth = bac if kind.startswith(PCT_PREFIX) else 1.0
                for k, v in data.items():
                    i = horizon.index(k)
                    if i >= horizon.months:
                        clamped += abs(v) * worth
                    elif i < 0:
                        early += abs(v) * worth
                    i = min(max(i, 0), horizon.months - 1)
                    nk = horizon.key(i)
                    merged[kind][nk] = merged[kind].get(nk, 0.0) + v * factor
            if how == "rama" and old.has_values:
                issues.append(Issue("MIG-007", target.code, target.name,
                                    f"Se sumó el historial de EV/AC de la fila anterior {old.code} '{old.name}'."))
            if how == "codigo" and old.has_values and old.name != target.name:
                issues.append(Issue("MIG-008", target.code, target.name,
                                    f"Recibió el historial de EV/AC de la fila anterior {old.code} '{old.name}' "
                                    "(asociada por código EDT)."))
        if several:
            for kind, data in merged.items():
                if not kind.startswith(PCT_PREFIX):
                    continue
                total_bac = merged_bac.get(_series_role(kind), 0.0)
                if total_bac > 0:
                    merged[kind] = {k: v / total_bac for k, v in data.items()}
                elif data:
                    merged[kind] = {}
                    issues.append(Issue("MIG-009", target.code, target.name,
                                        "Recibió el % ganado de varias filas del libro anterior sin BAC: no se pudo "
                                        "combinar. Se debe revisar la hoja % EV de esta fila."))
        if any(merged_bac.values()):
            prior.bac[key] = merged_bac
        if any(merged.values()):
            prior.values[key] = merged
            if any(v for data in merged.values() for v in data.values()):
                prior.matched += 1
        # El método editado en Excel en el libro anterior se conserva si la asociación es uno a uno.
        # Si el método cambió en MS Project (y no se editó en Excel), se usa el nuevo.
        if len(lst) == 1 and lst[0][1] in ("uid", "codigo"):
            old = lst[0][0]
            if old.method_edited and old.method != target.method:
                prior.methods[key] = old.method

    if prior.methods:
        legacy = any(o.origin_method is None for o in prior.rows)
        issues.append(Issue("MIG-010", "", "",
                            f"Se conservaron {len(prior.methods)} métodos de medición de la columna C del libro "
                            "anterior"
                            + (" (el libro anterior no indica si se editaron en Excel o cambiaron en MS Project; "
                               "se recomienda revisarlos)." if legacy else ", editados en Excel.")))
        # Recalcular el aviso de filas sin método
        undefined = sum(1 for r in inputs if is_undefined(prior.methods.get(row_key(r)) or r.method))
        report.issues[:] = [i for i in report.issues if i.id != "MET-005"]
        if undefined:
            issues.append(Issue("MET-005", "", "", f"{plural(undefined, 'fila', 'filas')} sin método de medición "
                                                   "definido (columna C)."))
    if early > 0.005:
        issues.append(Issue("MIG-011", "", "",
                            f"Hay EV/AC por {fmt_money(early)} cargados antes del primer mes del nuevo horizonte "
                            f"({horizon.start:%m/%Y}): se acumularon en el mes 1."))
    if clamped > 0.005:
        issues.append(Issue("MIG-012", "", "",
                            f"Hay EV/AC por {fmt_money(clamped)} cargados después del último mes del nuevo horizonte: "
                            "se acumularon en el último mes."))
    if prior.bac:
        issues.append(Issue("MIG-013", "", "", f"Se conservó el BAC de {plural(len(prior.bac), 'fila', 'filas')} del "
                                              "libro anterior."))
    issues += _migrate_changes(prior, key_map, sources)
    issues += _migrate_review(report, prior)
    if prior.overwritten:
        shown = ", ".join(prior.overwritten[:8])
        more = f" y {len(prior.overwritten) - 8} más" if len(prior.overwritten) > 8 else ""
        issues.append(Issue("MIG-019", "", "",
                            f"El libro anterior tiene valores escritos sobre celdas con fórmula (subtotales, fila "
                            f"TOTAL o suma de las empresas) que no se trasladan: {shown}{more}. Si son el AC o el BAC "
                            "de un paquete, se deben cargar en sus filas de ingreso del libro nuevo."))
    for old in prior.unmatched:
        bac = f"BAC {fmt_money(sum(old.bac.values()))}, "
        issues.append(Issue("MIG-014", old.code, old.name,
                            f"No se pudo trasladar su historial ({bac}EV {fmt_money(old.total('EV'))}, AC "
                            f"{fmt_money(old.total('AC'))}): la tarea no existe en el proyecto nuevo o se dividió "
                            "en varias filas. Se debe cargar a mano en las filas que correspondan."))
    return issues
