"""Comando 'actualizar' (incorpora el % ganado del mes) y lectura de libros existentes."""
from __future__ import annotations

import shutil
import zipfile
from dataclasses import dataclass, field, replace
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.workbook.properties import CalcProperties

from evm import analysis, changes, consistency, progress
from evm.config import MODE_PERCENT, Config, ConfigError
from evm.fields import fmt_number
from evm.messages import InputError, Issue, MissingFileError
from evm.methods import parse_method
from evm.model import ProjectData, Task
from evm.timeline import Horizon
from evm.wbs import (KIND_RESIDUAL, KIND_SUMMARY, Scope, Visibility, build_tree, code_key, has_interior_zero,
                     iter_subtree, plural, prune_inactive, row_visibility, select_scope)
from evm.workbook import (COL_BAC, COL_CODE, COL_METHOD, FIRST_MONTH_COL, FMT_PCT_MONTH, LEGACY_SHEET_CURVE,
                          PCT_PREFIX, SHEET_AC, SHEET_EV, SHEET_META, SHEET_PV, SHEET_SUMMARY, STATUS_CELL, StructRow,
                          WorkbookFormatError, add_s_curve_chart, book_sources, clean_text, col, company_sheet,
                          pct_sheet, read_meta, read_schedule, read_structure, read_wbs_list, save_workbook,
                          source_meta, split_key, write_meta, write_schedule, write_summary_trend)

PLAN_TOL = 1e-6         # diferencia del % planificado acumulado (fracción del BAC) que se considera un cambio


# ----------------------------------------------------------------------------
# Lectura de libros existentes
# ----------------------------------------------------------------------------
@dataclass
class BookData:
    path: Path
    meta: dict
    structure: list[StructRow]
    horizon: Horizon
    pv: dict[int, list[float]]           # fila -> valores mensuales (solo filas de ingreso)
    ev: dict[int, list[float | None]]
    ac: dict[int, list[float | None]]
    bac: dict[int, float]
    method: dict[int, str]
    status_date: date | None
    # Celdas de ingreso que no son números: (hoja, fila, columna, valor numérico, motivo)
    fixable: list[tuple[str, int, int, float, str]] = field(default_factory=list)
    # Multi-empresa: EV y AC de la empresa A y B por fila (ev y ac totales son la suma)
    multi: bool = False
    company_a: str = ""
    company_b: str = ""
    ev_a: dict[int, list[float | None]] = field(default_factory=dict)
    ev_b: dict[int, list[float | None]] = field(default_factory=dict)
    ac_a: dict[int, list[float | None]] = field(default_factory=dict)
    ac_b: dict[int, list[float | None]] = field(default_factory=dict)
    # Celdas que deben tener fórmula (subtotales, fila TOTAL, suma de empresas) y no la tienen: vacías o con un valor
    formula_gaps: list[str] = field(default_factory=list)
    formula_values: list[str] = field(default_factory=list)     # las que tienen un valor escrito
    # BAC ingresado, % planificado y % ganado de cada mes por empresa
    # (clave "" en un libro de una sola empresa; "A" y "B" en multi-empresa)
    bac_role: dict[str, dict[int, float]] = field(default_factory=dict)
    pct_pv: dict[str, dict[int, list[float]]] = field(default_factory=dict)
    pct_ev: dict[str, dict[int, list[float | None]]] = field(default_factory=dict)
    wbs_list: list[str] | None = None    # lista EDT con que se generó el libro (None = sin lista_edt)

    @property
    def input_rows(self) -> list[StructRow]:
        return [s for s in self.structure if s.is_input]

    def roles(self) -> list[tuple[str, str]]:
        """(clave, empresa) de cada serie: [("", "")] o [("A", empresa A), ("B", empresa B)]."""
        if self.multi:
            return [("A", self.company_a), ("B", self.company_b)]
        return [("", "")]

    def ac_of(self, role: str) -> dict[int, list[float | None]]:
        """AC ingresado de cada mes de una empresa ("" en un libro de una sola empresa)."""
        return self.ac_a if role == "A" else self.ac_b if role == "B" else self.ac


def _num(value, cached, where: str, fixable: list | None = None, key=None) -> float | None:
    """Valor numérico de una celda.

    Fórmula: usa el valor calculado que guardó Excel. Texto con un número: Excel no
    lo suma, así que se rechaza (o, si se indica `fixable`, se anota para convertirlo).
    """
    if isinstance(value, str) and value.startswith("="):
        if cached is None or isinstance(cached, str):
            raise WorkbookFormatError(
                "LIB-012",
                f"La celda {where} tiene una fórmula sin valor numérico calculado. Se debe abrir y guardar el "
                "libro en Excel, o reemplazar la fórmula por un número.")
        if fixable is not None and key is not None and not isinstance(cached, bool):
            fixable.append((*key, float(cached), "fórmula reemplazada por su valor"))
        value = cached
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        raise WorkbookFormatError("LIB-013", f"Valor no numérico en {where}: {value!r}")
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    if not text:
        return None
    try:
        number = float(text.replace(" ", "").replace(",", "."))
    except ValueError as exc:
        raise WorkbookFormatError("LIB-013", f"Valor no numérico en {where}: {value!r}") from exc
    if fixable is None or key is None:
        raise WorkbookFormatError(
            "LIB-014", f"La celda {where} contiene el número {value!r} guardado como texto: Excel no lo suma. "
            "Se debe ingresar como número.")
    fixable.append((*key, number, "número guardado como texto convertido a número"))
    return number


def _write_fixable(wb, book: "BookData") -> None:
    """Reemplaza por su valor numérico las celdas de ingreso con texto o fórmulas (en las hojas de %, con el formato
    de %)."""
    for sheet, r, c, v, _why in book.fixable:
        cell = wb[sheet].cell(row=r, column=c, value=v)
        if sheet.startswith(PCT_PREFIX):
            cell.number_format = FMT_PCT_MONTH


def _as_date(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return None


def _company_names(meta: dict) -> tuple[str, str]:
    return str(meta.get("empresa_a") or ""), str(meta.get("empresa_b") or "")


def check_version(meta: dict) -> None:
    """Los libros de versiones anteriores calculaban el PV y el EV con los costos de MS Project: no se admiten."""
    if str(meta.get("modo") or "").strip().lower() != MODE_PERCENT:
        raise WorkbookFormatError(
            "LIB-009",
            "El libro se generó con una versión anterior del programa, que calculaba el PV y el EV con los costos "
            "de MS Project: no se puede usar con esta versión (para ese libro se debe usar la versión 1.x). Se debe "
            "generar un libro nuevo.")


def company_sheets(meta: dict) -> list[str]:
    """Hojas de cada empresa en multi-empresa (PV-EmpresaA, % PV-EmpresaA...)."""
    if not meta.get("multiempresa"):
        return []
    a, b = _company_names(meta)
    names = [company_sheet(k, c) for k in (SHEET_PV, SHEET_EV, SHEET_AC) for c in (a, b)]
    return names + [pct_sheet(k, c) for k in (SHEET_PV, SHEET_EV) for c in (a, b)]


def data_sheets(meta: dict) -> list[str]:
    """Hojas mensuales del libro: PV, EV y AC, las hojas '% PV' y '% EV' y, en multi-empresa, las de cada
    empresa."""
    names = [SHEET_PV, SHEET_EV, SHEET_AC]
    if meta.get("multiempresa"):
        return names + company_sheets(meta)
    return names + [pct_sheet(SHEET_PV), pct_sheet(SHEET_EV)]


def check_sheets(sheetnames: list[str], meta: dict, charts: bool = False) -> None:
    """Verifica que estén las hojas que el programa lee o escribe (no se deben eliminar ni renombrar).

    charts: también la hoja de datos de los gráficos del Cuadro de Mando (la usa 'actualizar').
    """
    from evm.dashboard import SHEET_CHARTDATA, SHEET_DASHBOARD
    names = [SHEET_SUMMARY] + data_sheets(meta)
    if charts and SHEET_DASHBOARD in sheetnames:
        names.append(SHEET_CHARTDATA)
    companies = set(company_sheets(meta))
    for name in names:
        if name in sheetnames:
            continue
        if name in companies:
            raise WorkbookFormatError("LIB-010", f"Falta la hoja '{name}' del libro multi-empresa.")
        raise WorkbookFormatError(
            "LIB-010",
            f"Falta la hoja '{name}' del libro. Las hojas del libro no se deben eliminar ni cambiar de nombre: se "
            "debe restaurar su nombre original o recuperar el libro desde un respaldo.")


def _check_layout(wb, structure: list[StructRow], meta: dict) -> None:
    """Las filas de las hojas mensuales deben estar donde las ubicó 'generar' ('actualizar' escribe por
    posición): se compara el código EDT de la columna A con la hoja _estructura."""
    expected = [(s.row, s.code) for s in structure] + [(meta["fila_total"], "TOTAL")]
    for name in data_sheets(meta):
        ws = wb[name]
        for r, code in expected:
            value = ws.cell(row=r, column=COL_CODE).value
            text = "" if value is None else str(value).strip()
            if text != code.strip():
                raise WorkbookFormatError(
                    "LIB-011",
                    f"La hoja '{name}' no tiene las filas con las que se generó el libro: la fila {r} debe tener el "
                    f"código EDT '{code}' y tiene '{text}'. Como 'actualizar' ubica cada fila por su posición, se "
                    "deben deshacer las filas insertadas, eliminadas u ordenadas en esa hoja (o recuperar el libro "
                    "desde un respaldo). Para cambiar la EDT se debe modificar MS Project y regenerar el libro con "
                    "--migrar-desde.")


def _formula_gaps(wb, structure: list[StructRow], meta: dict) -> tuple[list[str], list[str]]:
    """Celdas que deben tener fórmula y tienen un valor o están vacías: en las hojas PV y EV los meses de todas las
    filas (BAC × %) y el BAC de los subtotales y de la fila TOTAL; en AC, los subtotales y la fila TOTAL; en
    multi-empresa, todas las filas de las hojas PV, EV y AC (suma de las dos empresas). Devuelve todas y las que
    tienen un valor escrito."""
    total = meta["fila_total"]
    months = list(range(FIRST_MONTH_COL, FIRST_MONTH_COL + meta["meses"]))
    subtotals = [s.row for s in structure if s.kind == KIND_SUMMARY] + [total]
    every = [s.row for s in structure] + [total]
    multi = bool(meta.get("multiempresa"))
    out, typed = [], []
    for name in data_sheets(meta):
        is_total_sheet = multi and name in (SHEET_PV, SHEET_EV, SHEET_AC)
        is_pv = name == SHEET_PV or name.startswith(SHEET_PV + "-")
        is_ev = name == SHEET_EV or name.startswith(SHEET_EV + "-")
        if is_total_sheet:
            cells = [(r, c) for r in every for c in months + ([COL_BAC] if is_pv else [])]
        elif is_pv or is_ev:
            cells = [(r, c) for r in every for c in months] + ([(r, COL_BAC) for r in subtotals] if is_pv else [])
        elif name.startswith(SHEET_AC):
            cells = [(r, c) for r in subtotals for c in months]
        else:
            continue                    # hojas '% PV' y '% EV'
        ws = wb[name]
        for r, c in cells:
            value = ws.cell(row=r, column=c).value
            if not (isinstance(value, str) and value.startswith("=")):
                out.append(f"{name}!{col(c)}{r}")
                if value is not None and str(value).strip():
                    typed.append(out[-1])
    return out, typed


def read_book(path: str | Path, wb=None, allow_fix: bool = False) -> BookData:
    """Lee un libro generado por este programa.

    allow_fix: en lugar de rechazar celdas de ingreso con texto o fórmulas, las
    anota en BookData.fixable para que 'actualizar' las convierta en números.
    """
    path = Path(path)
    if not path.is_file():
        raise MissingFileError("LIB-002", f"No existe el libro: {path}")
    if path.suffix.lower() != ".xlsx":
        raise WorkbookFormatError("LIB-003", f"'{path.name}' no es un libro .xlsx generado por este programa.")
    try:
        wb = wb or load_workbook(path)
    except (zipfile.BadZipFile, InvalidFileException, OSError) as exc:
        raise WorkbookFormatError("LIB-004", f"No se pudo abrir '{path.name}' como libro de Excel: {exc}") from exc
    cached_values = []

    def cached(sheet: str, r: int, c: int):
        """Valor que calculó Excel en una celda con fórmula. El libro con los valores calculados se lee solo si hace
        falta (las celdas de ingreso con fórmula son raras): leerlo duplica el tiempo de lectura."""
        if not cached_values:
            cached_values.append(load_workbook(path, data_only=True))
        return cached_values[0][sheet].cell(row=r, column=c).value
    meta = read_meta(wb)
    check_version(meta)
    structure = read_structure(wb)
    check_sheets(wb.sheetnames, meta)
    _check_layout(wb, structure, meta)
    gaps, typed = _formula_gaps(wb, structure, meta)
    horizon = Horizon(meta["horizonte_inicio"], meta["meses"])
    months = horizon.months
    inputs = [s for s in structure if s.is_input]
    fixable: list = []

    def value(name: str, ws, r: int, c: int, fix: bool) -> float | None:
        v = ws.cell(row=r, column=c).value
        formula = isinstance(v, str) and v.startswith("=")
        return _num(v, cached(name, r, c) if formula else None, f"{name}!{col(c)}{r}", fixable if fix else None,
                    (name, r, c))

    def sheet_values(name: str, fix: bool) -> dict[int, list[float | None]]:
        ws = wb[name]
        return {s.row: [value(name, ws, s.row, FIRST_MONTH_COL + i, fix) for i in range(months)] for s in inputs}

    def bac_values(name: str) -> dict[int, float]:
        ws = wb[name]
        return {s.row: value(name, ws, s.row, COL_BAC, allow_fix) or 0.0 for s in inputs}

    def zeros(series: dict[int, list[float | None]]) -> dict[int, list[float]]:
        return {r: [v or 0.0 for v in vals] for r, vals in series.items()}

    def combine(x: dict, y: dict) -> dict:
        return {r: [None if a is None and b is None else (a or 0.0) + (b or 0.0) for a, b in zip(x[r], y[r])]
                for r in x}

    ws_pv = wb[SHEET_PV]
    method = {s.row: str(ws_pv.cell(row=s.row, column=COL_METHOD).value or "").strip() for s in inputs}
    status = _as_date(wb[SHEET_SUMMARY][STATUS_CELL].value)
    # PV y EV = BAC ingresado × % de las hojas '% PV' y '% EV'
    multi = bool(meta.get("multiempresa"))
    a, b = _company_names(meta) if multi else ("", "")
    roles = [("A", a), ("B", b)] if multi else [("", "")]
    bac_role, pct_pv, pct_ev, pv_r, ev_r, ac_r = {}, {}, {}, {}, {}, {}
    for role, company in roles:
        bac_role[role] = bac_values(company_sheet(SHEET_PV, company) if company else SHEET_PV)
        pct_pv[role] = zeros(sheet_values(pct_sheet(SHEET_PV, company), allow_fix))
        pct_ev[role] = sheet_values(pct_sheet(SHEET_EV, company), allow_fix)
        bacs = bac_role[role]
        pv_r[role] = {r: [bacs[r] * v for v in vals] for r, vals in pct_pv[role].items()}
        ev_r[role] = {r: [None if v is None else bacs[r] * v for v in vals] for r, vals in pct_ev[role].items()}
        ac_r[role] = sheet_values(company_sheet(SHEET_AC, company) if company else SHEET_AC, allow_fix)
    extra = dict(bac_role=bac_role, pct_pv=pct_pv, pct_ev=pct_ev, formula_gaps=gaps, formula_values=typed,
                 fixable=fixable, wbs_list=read_wbs_list(wb))
    if multi:
        pv = {r: [x + y for x, y in zip(pv_r["A"][r], pv_r["B"][r])] for r in pv_r["A"]}
        bac = {r: bac_role["A"][r] + bac_role["B"][r] for r in bac_role["A"]}
        return BookData(path, meta, structure, horizon, pv, combine(ev_r["A"], ev_r["B"]),
                        combine(ac_r["A"], ac_r["B"]), bac, method, status, multi=True, company_a=a, company_b=b,
                        ev_a=ev_r["A"], ev_b=ev_r["B"], ac_a=ac_r["A"], ac_b=ac_r["B"], **extra)
    return BookData(path, meta, structure, horizon, pv_r[""], ev_r[""], ac_r[""], bac_role[""], method, status,
                    **extra)


def book_meta(path: str | Path) -> dict:
    """Metadatos de un libro generado por este programa (sin leer sus valores)."""
    path = Path(path)
    if not path.is_file():
        raise MissingFileError("LIB-002", f"No existe el libro: {path}")
    try:
        wb = load_workbook(path, read_only=True)
    except (zipfile.BadZipFile, InvalidFileException, OSError) as exc:
        raise WorkbookFormatError("LIB-004", f"No se pudo abrir '{path.name}' como libro de Excel: {exc}") from exc
    try:
        meta = read_meta(wb)
        check_version(meta)
        check_sheets(wb.sheetnames, meta, charts=True)
        return meta
    finally:
        wb.close()


def apply_book_settings(cfg: Config, meta: dict) -> Config:
    """Configuración con la que se generó el libro: campo Visible y, en multi-empresa, nombres y campos de
    MS Project."""
    cfg = replace(cfg)
    # Campo Visible con que se generó el libro (libros anteriores: no se lee)
    cfg.visible_field = str(meta.get("campo_visible") or "")
    cfg.visible_field_explicit = False
    cfg.multi_enabled = bool(meta.get("multiempresa"))
    if cfg.multi_enabled:
        cfg.company_a = str(meta.get("empresa_a") or cfg.company_a)
        cfg.company_b = str(meta.get("empresa_b") or cfg.company_b)
        cfg.company_both = str(meta.get("ambas") or cfg.company_both)
        cfg.company_field = str(meta.get("campo_empresa") or cfg.company_field)
        cfg.split_field = str(meta.get("campo_reparto") or cfg.split_field)
        if meta.get("reparto_por_defecto") is not None:
            cfg.split_default = float(meta["reparto_por_defecto"])
    return cfg


def check_sources(meta: dict, names: list[str]) -> None:
    """Los archivos de MS Project indicados deben ser los del libro: todos y con los mismos nombres.

    names: nombres de [[proyectos]]; [''] para un solo archivo (--origen).
    """
    book = book_sources(meta)
    if book == names or sorted(book) == sorted(names):
        return
    if book == [""]:
        raise ConfigError("ACT-001",
                          "El libro se generó con un solo archivo de MS Project: se debe indicar ese archivo con "
                          "--origen, con un archivo de configuración sin secciones [[proyectos]].")
    if names == [""]:
        raise ConfigError("ACT-002",
                          f"El libro se generó con varios archivos de MS Project ({', '.join(book)}): se deben "
                          "indicar en secciones [[proyectos]] del archivo de configuración, con los mismos nombres.")
    raise ConfigError("ACT-003", f"Los archivos de [[proyectos]] ({', '.join(names)}) no coinciden con los del libro "
                      f"({', '.join(book)}). Se deben indicar todos los archivos del libro, con los mismos nombres; "
                      "para agregar o quitar un archivo se debe generar un libro nuevo con --migrar-desde.")


# ----------------------------------------------------------------------------
# Actualización mensual
# ----------------------------------------------------------------------------
@dataclass
class RowResult:
    code: str
    name: str
    method: str
    bac: float
    ev_cum: float | None
    ev_month: float | None
    notes: list[Issue] = field(default_factory=list)            # observaciones de la fila
    by_company: dict[str, float] = field(default_factory=dict)   # EV del mes por hoja de empresa

    @property
    def note(self) -> str:
        return "; ".join(n.message for n in self.notes)


@dataclass
class UpdateResult:
    status_date: date
    month_index: int
    rows: list[RowResult]
    warnings: list[Issue]
    backup: Path | None
    book: BookData | None = None
    dry_run: bool = False
    deviations: list = field(default_factory=list)   # paquetes fuera de umbral (analysis.Deviation)
    thresholds: object | None = None                 # umbrales del semáforo (workbook.Thresholds)
    schedule: object | None = None                   # plazo según MS Project (schedule.ScheduleSummary)


def book_level(book: BookData, cfg: Config) -> int:
    """Nivel máximo del código EDT con que se generó el libro (libros antiguos: el de la configuración)."""
    return int(book.meta.get("nivel_maximo") or cfg.max_level)


def book_ignores_interior_zeros(meta: dict) -> bool:
    """ignorar_ceros_intermedios con que se generó el libro (los libros que no lo guardan: false)."""
    return meta.get("ignorar_ceros_intermedios") is True


def _interior_zero_warnings(scope: Scope, cfg: Config, visibility: Visibility, configured: bool) -> list[Issue]:
    """Tareas con un 0 intermedio en el código que tendrían otra fila con ignorar_ceros_intermedios de la
    configuración (configured): el libro conserva el criterio con que se generó."""
    other = row_visibility(scope, replace(cfg, ignore_interior_zeros=configured))
    changed = [t for t in scope.tasks() if has_interior_zero(scope.raw_code(t))
               and other.has_row(t) != visibility.has_row(t)]
    if not changed:
        return []
    codes = ", ".join(scope.raw_code(t) for t in changed[:10])
    more = f" y {len(changed) - 10} más" if len(changed) > 10 else ""
    used = "true" if cfg.ignore_interior_zeros else "false"
    return [Issue("ACT-050", message=f"El libro se generó con ignorar_ceros_intermedios = {used} y la configuración "
                                     f"indica {'true' if configured else 'false'}: "
                                     f"{plural(len(changed), 'tarea', 'tareas')} con un 0 intermedio en el código "
                                     f"({codes}{more}) {'tendría' if len(changed) == 1 else 'tendrían'} otra fila. "
                                     "El libro conserva sus filas. Para aplicar la configuración se debe regenerar el "
                                     "libro con --migrar-desde.")]


def _match_tasks(project: ProjectData, book: BookData, cfg: Config,
                 source: str = "") -> tuple[dict[int, Task], list[Issue], Scope]:
    """Asocia cada fila del libro con la tarea de MS Project.

    Mismo proyecto: por UniqueID (si la tarea ya no existe, la fila se omite).
    Proyecto distinto: por código EDT, sin asignar una tarea a dos filas.
    Se usa la planificación con la que se generó el libro.
    source: con varios archivos, nombre del archivo; solo se asocian las filas de ese archivo.
    """
    warnings: list[Issue] = []
    tasks = [t for t in project.tasks if t.outline_level >= 1]
    roots = build_tree(tasks)
    active: set[int] | None = None
    if not cfg.include_inactive:
        roots, _, _ = prune_inactive(roots)     # el avance de tareas inactivas no cuenta
        active = {t.uid for r in roots for t in iter_subtree(r)}
    plan_code = str(source_meta(book.meta, "planificacion", source) or "")
    if cfg.plan and cfg.plan.strip().upper() != plan_code.upper():
        warnings.append(Issue("ACT-009", message=f"El libro se generó para la planificación '{plan_code or 'todas'}': "
                                                 f"se usa esa planificación y no '{cfg.plan}'."))
    scope = select_scope(roots, plan_code)
    by_uid = {t.uid: t for t in tasks}
    by_code: dict[str, Task] = {}       # código sin los ceros finales (1.2.0 = 1.2.0.0) -> tarea
    for t in scope.tasks():
        raw = scope.raw_code(t)
        if raw:
            by_code.setdefault(code_key(clean_text(raw)), t)
    book_name = source_meta(book.meta, "proyecto", source)
    same = str(book_name or "") == clean_text(project.name)
    if not same:
        warnings.append(Issue("ACT-010", message=f"El libro fue generado para el proyecto '{book_name}' y el origen "
                                                 f"es '{project.name}': las filas se asocian por código EDT. Se "
                                                 "recomienda verificar los resultados."))
    out: dict[int, Task] = {}
    used: set[tuple[int, bool]] = set()
    for s in book.structure:
        row_source, uid_text, residual = split_key(s.key)
        if row_source != source:
            continue
        if same and uid_text.isdigit():
            task = by_uid.get(int(uid_text))
            if task is None:
                if s.is_input:
                    warnings.append(Issue("ACT-011", s.code, s.name,
                                          f"La tarea (UID {uid_text}) ya no existe en MS Project; se omite. Si cambió "
                                          "la EDT, se debe regenerar el libro con --migrar-desde."))
                continue
            if active is not None and task.uid not in active:
                if s.is_input:
                    warnings.append(Issue("ACT-012", s.code, s.name,
                                          f"La tarea (UID {uid_text}) está inactiva en MS Project; se omite y se "
                                          "conservan los valores ya cargados. Si quedó fuera del alcance, se "
                                          "recomienda regenerar el libro con --migrar-desde."))
                continue
        else:
            code = s.code[:-2] if residual else s.code
            task = by_code.get(code_key(code))
            if task is None or (task.uid, residual) in used:
                if s.is_input:
                    warnings.append(Issue("ACT-013", s.code, s.name,
                                          "No se encontró una tarea con ese código EDT; se omite."))
                continue
        used.add((task.uid, residual))
        out[s.row] = task
    return out, warnings, scope


def _contributors(book: BookData, row: StructRow, task: Task, matches: dict[int, Task],
                  unlisted: list[Task] = ()) -> list[Task]:
    """Tareas cuyo costo forma la fila: la rama del paquete, o (fila -D) lo que cuelga de la
    tarea resumen sin pertenecer a sus filas hijas ni a paquetes agregados después de generar el libro."""
    if row.kind != KIND_RESIDUAL:
        return list(iter_subtree(task))
    taken = {n.uid for c in _row_children(book, row, matches) + list(unlisted) for n in iter_subtree(c)}
    return [n for n in iter_subtree(task) if n.uid not in taken]


def _row_children(book: BookData, row: StructRow, matches: dict[int, Task]) -> list[Task]:
    """Tareas de las filas hijas de la tarea resumen de una fila de otros costos (-D)."""
    return [matches[s.row] for s in book.structure
            if s.parent_row == row.parent_row and s.kind != "residual" and s.row in matches]


def _unlisted_packages(book: BookData, matches: dict[int, Task], visibility: Visibility) -> dict[int, list[Task]]:
    """Tareas que tendrían fila propia (código EDT de nivel no mayor al máximo del libro, o Visible = Sí) y que
    cuelgan de la tarea resumen de una fila -D sin tener fila en el libro: se agregaron después de generarlo. Al
    generar el libro tendrían fila propia, por eso no se suman en la fila -D (fila -> tareas)."""
    with_row = {t.uid for t in matches.values()}
    out: dict[int, list[Task]] = {}

    out_ids = visibility.excluded_uids()

    def visit(node: Task, found: list[Task]) -> None:
        for c in node.children:
            if c.uid in with_row or c.uid in out_ids:     # ramas con Visible = No: fuera del cálculo
                continue
            if visibility.has_row(c):
                found.append(c)
            else:
                visit(c, found)         # sin código, de nivel mayor u oculta: forma parte de la fila -D

    for s in book.input_rows:
        task = matches.get(s.row)
        if s.kind == KIND_RESIDUAL and task is not None:
            found: list[Task] = []
            visit(task, found)
            if found:
                out[s.row] = found
    return out


def _visibility_warnings(book: BookData, matches: dict[int, Task], scope: Scope, visibility: Visibility,
                         source: str = "") -> list[Issue]:
    """Tareas con valor en el campo Visible que no coincide con las filas del libro: el campo cambió después de
    generarlo. El libro conserva sus filas hasta regenerarlo."""
    from evm.fields import parse_visible
    if visibility.listed:
        return []                       # con la lista EDT, el campo Visible no se usa
    rows = {matches[s.row].uid for s in book.structure
            if s.row in matches and s.kind != KIND_RESIDUAL and split_key(s.key)[0] == source}
    changed = [t for t in scope.tasks()
               if parse_visible(t.visible)[0] is not None and visibility.has_row(t) != (t.uid in rows)]
    if not changed:
        return []
    codes = ", ".join(scope.raw_code(t) or t.name for t in changed[:10])
    more = f" y {len(changed) - 10} más" if len(changed) > 10 else ""
    return [Issue("ACT-014", message=f"El campo Visible de {plural(len(changed), 'tarea', 'tareas')} ({codes}{more}) "
                                     "no coincide con las filas del libro (cambió después de generarlo): el libro "
                                     "conserva sus filas. Para aplicar el cambio se debe regenerar el libro con "
                                     "--migrar-desde.")]


def _new_task_warnings(book: BookData, matches: dict[int, Task], scope: Scope, visibility: Visibility,
                       source: str = "") -> list[Issue]:
    """Tareas que tendrían fila propia y no la tienen en el libro: se agregaron en MS Project después de generarlo
    (las de Visible con valor las informa _visibility_warnings). Su avance no se incluye en el % ganado."""
    from evm.fields import parse_visible
    rows = {matches[s.row].uid for s in book.structure
            if s.row in matches and s.kind != KIND_RESIDUAL and split_key(s.key)[0] == source}
    out_ids = visibility.excluded_uids()
    new = [t for t in scope.tasks() if visibility.has_row(t) and t.uid not in rows and t.uid not in out_ids
           and (visibility.listed or parse_visible(t.visible)[0] is None)]
    uids = {t.uid for t in new}
    new = [t for t in new if t.parent is None or t.parent.uid not in uids]      # solo la tarea superior de la rama
    if not new:
        return []
    tasks = ", ".join(f"{scope.raw_code(t)} '{t.name}'" for t in new[:10])
    more = f" y {len(new) - 10} más" if len(new) > 10 else ""
    if visibility.listed:
        text = (f"{plural(len(new), 'tarea de la lista EDT no tiene', 'tareas de la lista EDT no tienen')} fila en "
                f"el libro (se {'agregó' if len(new) == 1 else 'agregaron'} en MS Project después de generarlo): "
                f"{tasks}{more}. Su avance no se incluye en el % ganado. Para incluirlo, se debe regenerar el libro "
                "con --migrar-desde.")
    else:
        text = (f"{plural(len(new), 'tarea con código EDT no tiene', 'tareas con código EDT no tienen')} fila en el "
                f"libro (se {'agregó' if len(new) == 1 else 'agregaron'} después de generarlo): {tasks}{more}. Su "
                "avance no se incluye en el % ganado. Para incluirlo, se debe regenerar el libro con --migrar-desde.")
    return [Issue("ACT-015", message=text)]


def _month_end(d: date) -> date:
    return date(d.year + (d.month == 12), d.month % 12 + 1, 1) - timedelta(days=1)


def _backup_path(book_path: Path) -> Path:
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    path = book_path.with_name(f"{book_path.stem}.respaldo-{stamp}{book_path.suffix}")
    n = 2
    while path.exists():
        path = book_path.with_name(f"{book_path.stem}.respaldo-{stamp}-{n}{book_path.suffix}")
        n += 1
    return path


def _tag(source: str, issue: Issue) -> Issue:
    """Con varios archivos, el mensaje indica el archivo: '[P2] ...'."""
    return replace(issue, message=f"[{source}] {issue.message}") if source else issue


def _percent_row(book: BookData, s: StructRow, contributors: list[Task], shares: dict[int, float], calendar,
                 cfg: Config, idx: int) -> tuple[RowResult, list[tuple[str, str, float, float]],
                                                 dict[str, list[float]] | None]:
    """% ganado acumulado de cada empresa en una fila a la fecha de corte.

    Devuelve la fila del resultado, por empresa (clave, empresa, % ganado acumulado, BAC) y el % planificado de cada
    mes que resulta hoy de MS Project, por clave de empresa (como en 'generar'; None si alguna tarea no tiene fechas
    de línea base: su % planificado sale de las fechas actuales).
    """
    from evm import percent as pct
    method = parse_method(book.method[s.row])
    times = pct.leaf_times([t for t in contributors if not t.children], calendar, cfg.distribution)
    full = pct.full_weights(times)
    out = []
    notes: list[Issue] = []
    planned: dict[str, list[float]] | None = {}
    if any(tt.dist and not tt.from_baseline for tt in times):
        planned = None
    for role, company in book.roles():
        share = None
        if role == "A":
            share = lambda t: shares.get(t.uid, 0.5)            # noqa: E731
        elif role == "B":
            share = lambda t: 1.0 - shares.get(t.uid, 0.5)      # noqa: E731
        rw = pct.company_weights(times, share, company, full)
        if planned is not None:
            vec, before, _after = book.horizon.to_vector(pct.curve(times, rw))
            vec[0] += before                # como en 'generar': lo anterior al horizonte, en el mes 1
            planned[role] = vec
        pv_cum = sum(book.pct_pv[role][s.row][: idx + 1])
        frac, note = pct.earned_fraction(method, rw, pv_cum)
        bac = book.bac_role[role][s.row]
        if note is not None and all(n.message != note.message for n in notes):
            notes.append(replace(note, code=s.code, name=s.name))
        if rw.fallback and bac > 0:
            notes.append(Issue("ACT-028", s.code, s.name,
                               f"{company} no participa en las tareas de la fila: se usó el avance de la fila "
                               "completa"))
        out.append((role, company, frac, bac))
    res = RowResult(s.code, s.name, book.method[s.row], book.bac[s.row], None, None, notes)
    return res, out, planned


def _baseline_drift(book: BookData, s: StructRow, planned: dict[str, list[float]],
                    idx: int) -> tuple[float, float] | None:
    """PV acumulado de la fila a la fecha de corte con el % planificado de MS Project y con el del libro, si el %
    planificado acumulado de algún mes difiere (None si coinciden)."""
    now = book_pv = 0.0
    differs = False
    for role, _company in book.roles():
        new, old = planned.get(role, []), book.pct_pv[role][s.row]
        cum_new = cum_old = 0.0
        for i in range(len(old)):
            cum_new += new[i] if i < len(new) else 0.0
            cum_old += old[i]
            differs = differs or abs(cum_new - cum_old) > PLAN_TOL
        bac = book.bac_role[role][s.row]
        now += bac * sum(new[: idx + 1])
        book_pv += bac * sum(old[: idx + 1])
    return (now, book_pv) if differs else None


def _drift_warning(drift: list[tuple[StructRow, float, float]], book: BookData, idx: int) -> Issue:
    """ACT-047: filas cuyo % planificado no coincide con la línea base actual de MS Project."""
    decimals = int(book.meta.get("decimales") or 0)
    codes = ", ".join(s.code for s, _n, _o in drift[:10]) + (f" y {len(drift) - 10} más" if len(drift) > 10 else "")
    now, old = sum(n for _s, n, _o in drift), sum(o for _s, _n, o in drift)
    month = book.horizon.month_start(idx)
    return Issue("ACT-047", message=(
        f"El % planificado de {plural(len(drift), 'fila', 'filas')} no coincide con el que resulta hoy de MS "
        f"Project ({codes}): el libro sigue con la línea base con la que se generó. Pasa cuando se guarda de nuevo la "
        "línea base, o cambian las tareas o el calendario de la fila. Con la línea base actual, el PV acumulado de "
        f"{'esa fila' if len(drift) == 1 else 'esas filas'} a {month:%m/%Y} sería {fmt_number(now, decimals)} "
        f"({fmt_number(old, decimals)} en el libro). Si el cambio está aprobado, se debe generar un libro nuevo con "
        "--migrar-desde; si no, restaurar la línea base en MS Project."))


def run_update(projects: ProjectData | dict[str, ProjectData], cfg: Config, book_path: str | Path,
               status: date | None = None, backup: bool = True, force: bool = False,
               dry_run: bool = False) -> UpdateResult:
    """Carga el % ganado del mes de la fecha de corte (el AC se carga a mano en el libro).

    projects: el proyecto leído de MS Project o, en un libro con varios archivos, {nombre: proyecto} con
    todos los archivos de [[proyectos]].
    """
    from evm.percent import calendar_for
    if isinstance(projects, ProjectData):
        projects = {"": projects}
    book_path = Path(book_path)
    progress.stage("libro", "Lectura del libro")
    wb = load_workbook(book_path)
    book = read_book(book_path, wb, allow_fix=True)
    check_sheets(wb.sheetnames, book.meta, charts=True)
    check_sources(book.meta, list(projects))
    warnings: list[Issue] = [_tag(n, w) for n, p in projects.items() for w in p.warnings]
    warnings += _list_warnings(cfg, book)
    configured_zeros = cfg.ignore_interior_zeros
    cfg = apply_book_settings(cfg, book.meta)
    cfg.wbs_list = book.wbs_list            # las filas del libro se decidieron con su lista (o sin lista)
    cfg.ignore_interior_zeros = book_ignores_interior_zeros(book.meta)     # y con su nivel de los códigos
    multi = book.multi

    if status is None:
        missing = [n for n, p in projects.items() if p.status_date is None]
        if missing == [""]:
            raise InputError("ACT-004", "El proyecto no tiene fecha de estado: se debe indicar la fecha de corte con "
                             "--fecha-corte AAAA-MM-DD o definir Proyecto > Fecha de estado en MS Project.")
        if missing:
            raise InputError("ACT-004",
                             f"{'El archivo' if len(missing) == 1 else 'Los archivos'} {', '.join(missing)} no "
                             f"{'tiene' if len(missing) == 1 else 'tienen'} fecha de estado: se debe indicar la fecha "
                             "de corte con --fecha-corte AAAA-MM-DD o definir Proyecto > Fecha de estado en MS "
                             "Project.")
        dates = {n: p.status_date.date() for n, p in projects.items()}
        if len(set(dates.values())) > 1:
            raise InputError("ACT-005", "Los archivos tienen fechas de estado distintas ("
                             + ", ".join(f"{n}: {d:%d/%m/%Y}" for n, d in dates.items()) + "): se debe indicar la "
                             "fecha de corte con --fecha-corte AAAA-MM-DD o definir la misma fecha de estado en MS "
                             "Project.")
        status = next(iter(dates.values()))
    idx = book.horizon.index(status)
    if not 0 <= idx < book.horizon.months:
        raise InputError("ACT-006", f"La fecha de corte {status:%d/%m/%Y} está fuera del horizonte del libro "
                         f"({book.horizon.start:%m/%Y} a {book.horizon.end:%m/%Y}).")
    if status != _month_end(status):
        warnings.append(Issue("ACT-018", message="La fecha de corte no es el último día del mes: el PV acumulado "
                                                 "incluye el mes completo."))
    for n, project in projects.items():
        if project.status_date and project.status_date.date() != status:
            warnings.append(_tag(n, Issue("ACT-019", message=f"La fecha de estado de MS Project "
                                                             f"({project.status_date:%d/%m/%Y}) difiere de la fecha de "
                                                             f"corte ({status:%d/%m/%Y}): el avance registrado en MS "
                                                             "Project debe corresponder a la fecha de corte.")))
    if book.status_date and status < book.status_date:
        warnings.append(Issue("ACT-020", message=f"La fecha de corte ({status:%d/%m/%Y}) es anterior a la última "
                                                 f"cargada en el libro ({book.status_date:%d/%m/%Y})."))
    # Cambios del BAC y de los meses ya informados desde la ejecución anterior (antes de cargar el mes)
    decimals = int(book.meta.get("decimales") or 0)
    reference = changes.read_reference(wb)
    changed = changes.detect(book, reference, idx)
    warnings += changes.issues(changed, decimals, registered=not dry_run)
    if reference is None and not dry_run:
        warnings.append(changes.first_reference())

    # Cada archivo asocia sus filas (el UniqueID se repite entre archivos)
    matches: dict[int, Task] = {}
    unlisted: dict[int, list[Task]] = {}
    shares_of: dict[str, dict[int, float]] = {}
    calendars: dict[str, object] = {}                # calendario de cada archivo
    hidden: dict[str, tuple[set[int], set[int]]] = {}  # Visible = No: tareas y ramas excluidas
    level = book_level(book, cfg)
    leaves_of: dict[int, list[Task]] = {}         # fila -> tareas de detalle (para el plazo)
    if multi:
        from evm.companies import resolve_shares
    for n, project in projects.items():
        pcfg = cfg.project_config(n)
        found, w, scope = _match_tasks(project, book, pcfg, n)
        warnings += [_tag(n, x) for x in w]
        visibility = row_visibility(scope, replace(pcfg, max_level=level))
        if configured_zeros != cfg.ignore_interior_zeros and not visibility.listed:
            warnings += [_tag(n, x) for x in _interior_zero_warnings(scope, replace(pcfg, max_level=level),
                                                                     visibility, configured_zeros)]
        warnings += [_tag(n, x) for x in _visibility_warnings(book, found, scope, visibility, n)]
        warnings += [_tag(n, x) for x in _new_task_warnings(book, found, scope, visibility, n)]
        hidden[n] = (visibility.excluded_uids(), {t.uid for t in visibility.excluded})
        matches.update(found)
        unlisted.update(_unlisted_packages(book, found, visibility))
        if multi:
            shares_of[n], _ = resolve_shares(scope.roots, pcfg)
        calendars[n] = calendar_for(project, pcfg)

    rows: list[RowResult] = []
    later_filled: list[str] = []
    later_months: set[int] = set()
    writes: list[tuple[str, int, int, float]] = []   # (hoja, fila, columna, valor)
    no_bac: list[str] = []                           # filas con avance y sin BAC
    drift: list[tuple[StructRow, float, float]] = []  # filas con otro % planificado en MS Project (PV hoy y libro)
    progress.stage("calcular", "Cálculo del % ganado")
    inputs = book.input_rows
    for k, s in enumerate(inputs, start=1):
        progress.step(k, len(inputs), "filas")
        task = matches.get(s.row)
        if task is None:
            continue
        src = split_key(s.key)[0]
        # Tareas que forman la fila: en una fila -D, solo las tareas sin fila propia; las ramas con Visible = No no
        # forman parte de la fila
        contributors = _contributors(book, s, task, matches, unlisted.get(s.row, []))
        out_ids, out_roots = hidden.get(src, (set(), set()))
        if any(t.uid in out_roots for t in contributors):
            contributors = [t for t in contributors if t.uid not in out_ids]
        leaves_of[s.row] = [t for t in contributors if not t.children]
        res, parts, planned = _percent_row(book, s, contributors, shares_of.get(src, {}), calendars[src], cfg, idx)
        if planned is not None:
            moved = _baseline_drift(book, s, planned, idx)
            if moved is not None:
                drift.append((s, *moved))
        for role, company, frac, bac_c in parts:
            sheet = pct_sheet(SHEET_EV, company)
            series = book.pct_ev[role][s.row]
            later = [i for i in range(idx + 1, len(series)) if series[i]]
            if later:
                later_filled.append(f"{sheet} {s.code}")
                later_months.update(later)
            prev = sum(v or 0.0 for v in series[:idx])
            month_value = round(frac - prev, 12)
            res.ev_cum = (res.ev_cum or 0.0) + round(bac_c * frac, 2)
            res.ev_month = (res.ev_month or 0.0) + round(bac_c * month_value, 2)
            if multi:
                res.by_company[company_sheet(SHEET_EV, company)] = round(bac_c * month_value, 2)
            if month_value < -1e-9:
                res.notes.append(Issue("ACT-027", s.code, s.name, f"{sheet} del mes negativo (corrección)"))
            if frac > 1e-9 and bac_c <= 0:
                no_bac.append(f"{s.code} ({company})" if company else s.code)
            writes.append((sheet, s.row, FIRST_MONTH_COL + idx, month_value))
        rows.append(res)

    if later_filled and not force:
        raise InputError(
            "ACT-007",
            "Hay valores cargados en meses posteriores a la fecha de corte (" + ", ".join(later_filled[:10]) +
            ("..." if len(later_filled) > 10 else "") + "). Se deben borrar esos valores o usar --forzar.")
    if later_filled:
        months = [f"{book.horizon.month_start(i):%m/%Y}" for i in sorted(later_months)]
        warnings.append(Issue("ACT-021", message=(
            f"Se cargó la fecha de corte {status:%d/%m/%Y} con --forzar, pero hay valores en meses posteriores ("
            + ", ".join(months[:6]) + ("..." if len(months) > 6 else "") + "; filas " + ", ".join(later_filled[:10])
            + ("..." if len(later_filled) > 10 else "") + "). Esos valores se calcularon con el acumulado anterior: "
            "los acumulados de esos meses ya no coinciden con MS Project. Se deben revisar esos meses o volver a "
            "cargarlos con 'actualizar' (fecha de corte de cada mes, en orden).")))

    if drift:
        warnings.append(_drift_warning(drift, book, idx))

    where = (f"las hojas {company_sheet(SHEET_PV, book.company_a)} y {company_sheet(SHEET_PV, book.company_b)}"
             if multi else f"la hoja {SHEET_PV}")
    if not any(book.bac[s.row] for s in book.input_rows):
        warnings.append(Issue("ACT-022", message=f"Ninguna fila tiene BAC: el PV, el EV y los indicadores quedan en 0 "
                                                 f"hasta ingresar el BAC de cada fila en la columna D de {where}."))
    elif no_bac:
        warnings.append(Issue("ACT-023", message=(
            f"{plural(len(no_bac), 'fila tiene', 'filas tienen')} % ganado pero no "
            f"{'tiene' if len(no_bac) == 1 else 'tienen'} BAC: el EV no se puede valorar hasta "
            f"ingresar el BAC en la columna D de {where} (" + ", ".join(no_bac[:15]) +
            (f" y {len(no_bac) - 15} más" if len(no_bac) > 15 else "") + ").")))

    if book.formula_gaps:
        cells = ", ".join(book.formula_gaps[:8])
        more = f" y {len(book.formula_gaps) - 8} más" if len(book.formula_gaps) > 8 else ""
        warnings.append(Issue("ACT-024", message=(
            f"Hay celdas sin la fórmula del libro (subtotales, fila TOTAL o suma de las empresas): "
            f"{cells}{more}. Los totales e indicadores del libro no coinciden con los que se muestran "
            "aquí: se deben restaurar esas fórmulas (por ejemplo, copiándolas de una celda vecina de la "
            "misma fila) o recuperar el libro desde un respaldo.")))

    if book.fixable:
        cells = ", ".join(f"{sh}!{col(c)}{r}" for sh, r, c, _v, _why in book.fixable[:8])
        more = f" y {len(book.fixable) - 8} más" if len(book.fixable) > 8 else ""
        verb = "reemplazarán" if dry_run else "reemplazaron"
        warnings.append(Issue("ACT-025", message=f"Se {verb} por su valor numérico celdas de ingreso con texto o "
                                                 f"fórmulas: {cells}{more}."))

    # Valores en memoria (también en simulación, para mostrar los indicadores resultantes)
    pct_roles = {pct_sheet(SHEET_EV, company): role for role, company in book.roles()}
    for sheet, r, c, v in writes:
        i = c - FIRST_MONTH_COL
        role = pct_roles[sheet]
        book.pct_ev[role][r][i] = v
        amount = book.bac_role[role][r] * v
        if not multi:
            book.ev[r][i] = amount
            continue
        (book.ev_a if role == "A" else book.ev_b)[r][i] = amount
        x, y = book.ev_a[r][i], book.ev_b[r][i]
        book.ev[r][i] = None if x is None and y is None else (x or 0.0) + (y or 0.0)

    # Coherencia del EV y el AC y meses sin informe (con el % ganado del mes ya calculado)
    warnings += consistency.ev_ac_checks(book, idx)
    warnings += consistency.month_gaps(book, idx, latest_only=True)
    warnings += consistency.eac_realism(book, idx)
    warnings += consistency.loe_share(book, cfg.max_loe_share)

    # Paquetes fuera de umbral en el mes de corte (con el EV y el AC que quedan en el libro)
    th = analysis.thresholds(book.meta, cfg)
    found = analysis.deviations(book, idx, th)
    plan = _schedule(book, projects, calendars, leaves_of, status)
    entries, notes = _analysis_entries(wb, book, idx, found, status, _schedule_by_key(book, plan))
    warnings += notes

    backup_path = None
    if not dry_run:
        progress.stage("guardar", "Guardado del libro")
        _write_fixable(wb, book)
        for sheet, r, c, v in writes:
            # Las celdas vacías de las hojas '% EV' no tienen formato: se da al cargar el valor
            wb[sheet].cell(row=r, column=c, value=v).number_format = FMT_PCT_MONTH
        wb[SHEET_SUMMARY][STATUS_CELL].value = status
        meta = dict(book.meta)
        meta.pop("meses_grafico", None)
        meta["ultima_actualizacion"] = datetime.now().replace(microsecond=0)
        write_meta(wb[SHEET_META], meta)
        write_schedule(wb, plan, [x.row for x in book.structure], book.meta["fila_total"], "actualizar")
        write_summary_trend(wb[SHEET_SUMMARY], book.meta["fila_total"], book.horizon.months, th)
        _write_analysis(wb, book, idx, entries, found, th)
        _write_changes(wb, book, idx, status, changed, decimals)
        _refresh_charts(wb, book, idx, meta)
        backup_path = _save(wb, book_path, backup, cfg.fast_save)
    book.status_date = status
    return UpdateResult(status, idx, rows, warnings, backup_path, book, dry_run, deviations=found, thresholds=th,
                        schedule=plan)


def _list_warnings(cfg: Config, book: BookData) -> list[Issue]:
    """La lista EDT de la configuración frente a la del libro: las filas se decidieron al generarlo."""
    if cfg.wbs_list is None or book.wbs_list == cfg.wbs_list:
        return []
    if book.wbs_list is None:
        return [Issue("ACT-016", message="La configuración indica lista_edt, pero el libro se generó sin lista EDT: "
                                         "se mantienen sus filas (nivel máximo y campo Visible). Para aplicar la "
                                         "lista, se debe regenerar el libro con --migrar-desde.")]
    if {code_key(c) for c in book.wbs_list} == {code_key(c) for c in cfg.wbs_list}:     # 1.2.0 = 1.2.0.0
        return []
    return [Issue("ACT-017", message="La lista EDT de la configuración cambió después de generar el libro: se usa la "
                                     "lista con la que se generó. Para aplicar los cambios, se debe regenerar el libro "
                                     "con --migrar-desde.")]


def _schedule(book: BookData, projects: dict[str, ProjectData], calendars: dict[str, object],
              leaves_of: dict[int, list[Task]], status: date):
    """Plazo de cada fila con las tareas de detalle leídas de MS Project (schedule.ScheduleSummary)."""
    from evm.schedule import ScheduleSummary, row_schedule
    default = next(iter(calendars.values()))
    rows, cal_of, children = {}, {}, {}
    for s in book.structure:
        src = split_key(s.key)[0]
        cal_of[s.row] = calendars.get(src, default)
        if s.parent_row is not None:
            children.setdefault(s.parent_row, []).append(s.row)
        if s.is_input and leaves_of.get(s.row):
            rows[s.row] = row_schedule(leaves_of[s.row], cal_of[s.row], projects[src].minutes_per_day)
    inputs = [s.row for s in book.input_rows]
    return ScheduleSummary.build(status, rows, inputs, children, cal_of, default)


def _schedule_by_key(book: BookData, plan) -> dict[str, tuple]:
    """Plazo de cada fila de ingreso para el análisis de desviaciones: {clave: (desvío, holgura, crítica)}."""
    out = {}
    for s in book.input_rows:
        x = plan.rows.get(s.row)
        if x is not None:
            out[s.key] = (x.variance, x.slack, x.status or None)
    return out


def _save(wb, book_path: Path, backup: bool, fast: bool = True) -> Path | None:
    """Guarda el libro (fórmulas recalculadas al abrirlo) con una copia de respaldo previa, si se pide. fast:
    guardado rápido de las celdas ([exportacion] guardado_rapido)."""
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    backup_path = None
    if backup:
        backup_path = _backup_path(book_path)
        shutil.copy2(book_path, backup_path)
    try:
        save_workbook(wb, book_path, fast)
    except BaseException:
        if backup_path is not None:     # el libro no cambió: el respaldo no hace falta
            backup_path.unlink(missing_ok=True)
        raise
    return backup_path


def _analysis_entries(wb, book: BookData, idx: int, found: list, status: date,
                      plan: dict[str, tuple] | None = None) -> tuple[list, list[Issue]]:
    """Filas de la hoja Análisis de desviaciones con los paquetes fuera de umbral del mes idx, conservando el texto
    ya escrito, y las observaciones del análisis: el del mes anterior incompleto y las acciones vencidas. plan: plazo
    de cada fila (clave -> desvío, holgura, crítica); sin él, el de la hoja Indicadores EDT (última lectura de MS
    Project)."""
    entries = analysis.read_entries(wb[analysis.SHEET_ANALYSIS]) if analysis.SHEET_ANALYSIS in wb.sheetnames else []
    month = book.horizon.month_start(idx)
    notes = analysis.incomplete(entries, month)
    if plan is None:
        plan = read_schedule(wb, {s.row: s.key for s in book.input_rows})
    entries = analysis.refresh(entries, found, month, plan)
    return entries, notes + analysis.overdue(entries, found, month, status)


def _write_analysis(wb, book: BookData, idx: int, entries: list, found: list, th) -> None:
    """Escribe la hoja Análisis de desviaciones (la crea en los libros de versiones anteriores)."""
    analysis.write_sheet(wb, entries, {s.key: s.row for s in book.input_rows}, book.horizon.months, th,
                         int(book.meta.get("decimales") or 0), book.horizon.month_start(idx), len(found))


def _write_changes(wb, book: BookData, idx: int, status: date, changed: list, decimals: int) -> None:
    """Agrega los cambios a la hoja Registro de cambios (la crea en los libros de versiones anteriores) y guarda la
    referencia para la ejecución siguiente."""
    if changed:
        changes.append(wb, changed, datetime.now().replace(microsecond=0), status, decimals)
    else:
        changes.write_sheet(wb, decimals=decimals)
    changes.write_reference(wb, changes.reference_from_book(book, idx, status))


def run_analysis(cfg: Config, book_path: str | Path, backup: bool = True) -> UpdateResult:
    """Comando 'analizar': actualiza la hoja Análisis de desviaciones con los valores del libro, sin leer MS
    Project (por ejemplo, después de cargar a mano el AC del mes)."""
    book_path = Path(book_path)
    progress.stage("libro", "Lectura del libro")
    wb = load_workbook(book_path)
    book = read_book(book_path, wb, allow_fix=True)
    check_sheets(wb.sheetnames, book.meta, charts=True)
    if book.status_date is None:
        raise InputError("ACT-008",
                         "El libro no tiene fecha de corte (hoja Resumen, celda B5): se debe ejecutar 'actualizar' "
                         "antes de 'analizar'.")
    idx = book.horizon.index(book.status_date)
    if not 0 <= idx < book.horizon.months:
        raise InputError("ACT-006",
                         f"La fecha de corte {book.status_date:%d/%m/%Y} está fuera del horizonte del libro.")
    warnings: list[Issue] = []
    if book.fixable:
        cells = ", ".join(f"{sh}!{col(c)}{r}" for sh, r, c, _v, _why in book.fixable[:8])
        more = f" y {len(book.fixable) - 8} más" if len(book.fixable) > 8 else ""
        warnings.append(Issue("ACT-025", message=f"Se reemplazaron por su valor numérico celdas de ingreso con texto "
                                                 f"o fórmulas: {cells}{more}."))
        _write_fixable(wb, book)
    decimals = int(book.meta.get("decimales") or 0)
    reference = changes.read_reference(wb)
    changed = changes.detect(book, reference, idx)
    warnings += changes.issues(changed, decimals)
    if reference is None:
        warnings.append(changes.first_reference())
    warnings += consistency.ev_ac_checks(book, idx)
    warnings += consistency.month_gaps(book, idx, latest_only=False)
    warnings += consistency.eac_realism(book, idx)
    warnings += consistency.loe_share(book, cfg.max_loe_share)
    progress.stage("analizar", "Análisis de desviaciones")
    th = analysis.thresholds(book.meta, cfg)
    found = analysis.deviations(book, idx, th)
    entries, notes = _analysis_entries(wb, book, idx, found, book.status_date)
    warnings += notes
    write_summary_trend(wb[SHEET_SUMMARY], book.meta["fila_total"], book.horizon.months, th)
    _write_analysis(wb, book, idx, entries, found, th)
    _write_changes(wb, book, idx, book.status_date, changed, decimals)
    _refresh_charts(wb, book, idx, book.meta)
    progress.stage("guardar", "Guardado del libro")
    backup_path = _save(wb, book_path, backup, cfg.fast_save)
    return UpdateResult(book.status_date, idx, [], warnings, backup_path, book, deviations=found, thresholds=th)


def _refresh_charts(wb, book: BookData, idx: int | None, meta: dict) -> None:
    """Vuelve a crear los gráficos (openpyxl no los conserva) y fija el mes de reporte en el mes de corte (idx None:
    libro sin fecha de corte)."""
    from evm import dashboard
    currency = str(meta.get("moneda") or "")
    if dashboard.SHEET_DASHBOARD in wb.sheetnames:
        bac = book_totals(book)[3]
        packages = dashboard.chart_packages((s.row, s.code, s.level, s.kind) for s in book.structure)
        try:
            layout = int(meta.get("disposicion_cuadro_mando") or 1)
        except (TypeError, ValueError):
            layout = 1
        if idx is None:
            dashboard.add_dashboard_charts(wb, packages, book.horizon.months, plan_end(book), None, bac, currency,
                                           layout=layout)
        else:
            wb[dashboard.SHEET_DASHBOARD][dashboard.REPORT_CELL].value = book.horizon.month_start(idx)
            loaded = [i for i in range(idx + 1)
                      if any(series[r][i] is not None for series in (book.ev, book.ac) for r in series)]
            dashboard.add_dashboard_charts(wb, packages, book.horizon.months, plan_end(book), idx + 1,
                                           bac, currency, idx + 1 - loaded[0] if loaded else 0, layout)
    elif LEGACY_SHEET_CURVE in wb.sheetnames:
        add_s_curve_chart(wb[LEGACY_SHEET_CURVE], book.horizon.months, str(meta.get("proyecto") or ""), currency)
    from evm import cost_review
    if cost_review.SHEET_CHART in wb.sheetnames and cost_review.SHEET_PASTE in wb.sheetnames:
        months = book.horizon.months
        last = max(plan_end(book), cost_review.last_paste_month(wb[cost_review.SHEET_PASTE], months))
        cost_review.add_chart(wb, months, last, currency)


def plan_end(book: BookData) -> int:
    """Último mes con % planificado (N° de mes), aunque falte ingresar algún BAC."""
    last = 0
    for series in book.pct_pv.values():
        for vals in series.values():
            for i, v in enumerate(vals):
                if abs(v) > 1e-9 and i + 1 > last:
                    last = i + 1
    return last


def book_totals(book: BookData) -> tuple[list[float], list[float], list[float], float]:
    """Totales mensuales de PV, EV y AC (sumando las filas de ingreso) y BAC total."""
    n = book.horizon.months
    pv = [0.0] * n
    ev = [0.0] * n
    ac = [0.0] * n
    for s in book.input_rows:
        for i in range(n):
            pv[i] += book.pv[s.row][i]
            ev[i] += book.ev[s.row][i] or 0.0
            ac[i] += book.ac[s.row][i] or 0.0
    bac = sum(book.bac[s.row] for s in book.input_rows)
    return pv, ev, ac, bac


def indicators_for(book: BookData, status: date):
    """Indicadores del proyecto a la fecha de corte (mismos cálculos que la hoja Resumen)."""
    from evm.indicators import compute
    pv, ev, ac, bac = book_totals(book)
    idx = book.horizon.index(status)
    if not 0 <= idx < book.horizon.months:
        raise InputError("ACT-006", f"La fecha de corte {status:%d/%m/%Y} está fuera del horizonte del libro.")
    loaded = [any(book.ev[s.row][i] is not None for s in book.input_rows) for i in range(book.horizon.months)]
    return compute(pv, ev, ac, bac, idx, loaded)
