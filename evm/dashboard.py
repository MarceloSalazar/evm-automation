"""Hoja "Cuadro de Mando": indicadores al mes de reporte y gráficos nativos de Excel.

El mes de reporte se elige en un desplegable (meses hasta la fecha de corte de la
hoja Resumen). Los valores y los gráficos se calculan con fórmulas de Excel a
partir de las hojas PV, EV y AC; los datos de los gráficos están en la hoja
oculta "_graficos". La biblioteca que escribe el libro no conserva gráficos al
reabrirlo, por eso "actualizar" los vuelve a crear con add_dashboard_charts().
"""
from __future__ import annotations

from dataclasses import dataclass
from datetime import date

from openpyxl.chart import AreaChart, BarChart, LineChart, Reference, ScatterChart, Series
from openpyxl.chart.axis import ChartLines, DateAxis
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.marker import Marker
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText, Text
from openpyxl.chart.title import Title
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
from openpyxl.drawing.text import CharacterProperties, Paragraph, ParagraphProperties, RegularTextRun
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.properties import PageSetupProperties

from evm.config import Config
from evm.wbs import KIND_RESIDUAL, Report
from evm.workbook import (C_AC, C_ACN, C_DATE, C_EV, C_EVN, C_PVC, CURVE_FIRST_ROW, CUT_MONTH_CELL, DARK,
                          FILL_GREEN, FILL_HEADER, FILL_INPUT, FILL_RED, FILL_YELLOW, FMT_DATE, FMT_RATIO,
                          MONTHNUM_ROW, NAME_START, NAME_STATUS, SHEET_AC, SHEET_CURVE, SHEET_EV, SHEET_PV,
                          SHEET_SUMMARY, _Layout, _traffic_light, col, money_format, put_text)

SHEET_DASHBOARD = "Cuadro de Mando"
SHEET_CHARTDATA = "_graficos"

NAME_REPORT = "MesReporte"              # desplegable: mes de reporte (fecha)
NAME_REPORT_N = "MesReporteN"           # mes de reporte efectivo: N° de mes del horizonte
NAME_REPORT_LIST = "ListaMesesReporte"  # meses del desplegable (hasta la fecha de corte)
NAME_EAC_METHOD = "MetodoEAC"
REPORT_CELL = "E5"
EAC_CELL = "R5"
EAC_METHODS = ["BAC / CPI", "AC + (BAC - EV)", "AC + (BAC - EV) / (CPI × SPI)"]

# Escala de los gráficos de CPI y SPI (los valores fuera de escala se dibujan en el borde)
INDEX_MIN, INDEX_MAX = 0.5, 1.5
SCATTER_PAD = 0.1       # margen del gráfico CPI-SPI: deja lugar a las etiquetas y a los rótulos de cuadrante
MAX_SCATTER = 200                       # Excel admite hasta 255 series por gráfico

# Cuadrícula de la hoja: columnas B a U de 50 píxeles
GRID_FIRST, GRID_LAST = 2, 21
GRID_WIDTH = 7.140625                   # ancho que Excel muestra como 50 píxeles (Calibri 11)
ROW_CHARTS = 25                         # gráficos 1 y 2
ROW_CHARTS_END = 44
ROW_SCATTER = 47                        # gráfico 3 y tabla de elementos EDT
ROW_SCATTER_END = 72

# Hoja _graficos: parámetros (A1:B7), cuadrantes (D1:F5), serie mensual y elementos EDT (desde la fila 11)
D_HEADER_ROW, D_FIRST_ROW = 10, 11
(M_NUM, M_DATE, M_PV, M_EV, M_AC, M_BAC, M_MARK, M_CPI, M_SPI, M_RED, M_YELLOW, M_GREEN, M_TARGET,
 M_CPI_RAW, M_SPI_RAW) = range(1, 16)
(P_CODE, P_LEVEL, P_BAC, P_PV, P_EV, P_AC, P_CPI_RAW, P_SPI_RAW, P_CPI, P_SPI) = range(17, 27)

ACCENT_BAC, ACCENT_PV, ACCENT_EV, ACCENT_AC = "7F7F7F", "1F4E79", "548235", "C00000"
COLOR_CPI, COLOR_SPI = "1F4E79", "ED7D31"
GRAY_TEXT = "595959"
FILL_LABEL = PatternFill("solid", fgColor="EEF2F7")
FILL_SUBTITLE = PatternFill("solid", fgColor="D6DCE4")
FILL_LEVEL1 = PatternFill("solid", fgColor="EEF2F7")
FILL_WHITE = PatternFill("solid", fgColor="FFFFFF")
CARD_SIDE = Side(style="thin", color="D9D9D9")
EMU_PX = 9525


@dataclass
class ChartPackage:
    """Elemento EDT del gráfico CPI-SPI (fila de las hojas PV, EV y AC)."""
    row: int
    code: str
    level: int


def chart_packages(rows) -> list[ChartPackage]:
    """Elementos EDT de nivel 1 y 2 (sin las filas de otros costos), en el orden del libro.

    rows: tuplas (fila, código, nivel, tipo). Si son demasiados para un gráfico, solo el nivel 1.
    """
    out = [ChartPackage(r, code, level) for r, code, level, kind in rows if level <= 2 and kind != KIND_RESIDUAL]
    if len(out) > MAX_SCATTER:
        out = [p for p in out if p.level == 1][:MAX_SCATTER]
    return out


def chart_months(months: int, plan_end: int, cutoff: int | None) -> tuple[int, int]:
    """Meses de los gráficos: evolución EVM (hasta el fin del plan o la fecha de corte, lo último)
    y evolución de CPI y SPI (hasta la fecha de corte)."""
    evm = min(months, max(1, plan_end, cutoff or 0))
    idx = min(months, max(3, cutoff or min(12, plan_end or 12)))
    return evm, idx


def plan_end_month(total_pv: list[float]) -> int:
    """Último mes con PV planificado (N° de mes, 1 en adelante)."""
    last = 0
    for i, v in enumerate(total_pv):
        if abs(v) > 0.005:
            last = i + 1
    return last


# ----------------------------------------------------------------------------
# Hoja _graficos
# ----------------------------------------------------------------------------
def _month_cum(sheet: str, row: int, lay: _Layout) -> str:
    """Acumulado de una fila de PV, EV o AC hasta el mes de reporte."""
    months = f"'{sheet}'!${col(lay.first_col)}{row}:${col(lay.last_col)}{row}"
    nums = f"'{SHEET_PV}'!${col(lay.first_col)}${MONTHNUM_ROW}:${col(lay.last_col)}${MONTHNUM_ROW}"
    return f'SUMIFS({months},{nums},"<="&{NAME_REPORT_N})'


def _fmt_threshold(v: float) -> str:
    return f"{v:.2f}".replace(".", ",")


def write_chart_data(ws, cfg: Config, lay: _Layout, packages: list[ChartPackage]) -> None:
    """Hoja oculta con los datos de los gráficos (fórmulas que dependen del mes de reporte)."""
    tr = lay.total_row
    sel = NAME_REPORT
    cut = f"'{SHEET_SUMMARY}'!${CUT_MONTH_CELL[0]}${CUT_MONTH_CELL[1:]}"
    si = f"((YEAR({sel})-YEAR({NAME_START}))*12+MONTH({sel})-MONTH({NAME_START})+1)"
    params = [
        ("Mes de reporte (N°)", f'=IFERROR(IF(AND({sel}="",{cut}=""),"",MAX(1,MIN({lay.months},'
                                f'IF({sel}="",{cut},IF({cut}="",{si},MIN({si},{cut})))))),"")'),
        ("Escala mínima de CPI y SPI", INDEX_MIN),
        ("Escala máxima de CPI y SPI", INDEX_MAX),
        ("Umbral rojo", cfg.threshold_red),
        ("Umbral amarillo", cfg.threshold_yellow),
        ("BAC", f"='{SHEET_PV}'!$D${tr}"),
        ("Altura de la marca del mes de reporte",
         f'=IF(B1="",0,MAX(B6,{_month_cum(SHEET_PV, tr, lay)},{_month_cum(SHEET_EV, tr, lay)},'
         f'{_month_cum(SHEET_AC, tr, lay)}))'),
    ]
    for i, (label, value) in enumerate(params, start=1):
        ws.cell(row=i, column=1, value=label)
        ws.cell(row=i, column=2, value=value)

    # Rótulos de los cuadrantes del gráfico CPI-SPI (series de un punto, sin marcador)
    ws.cell(row=1, column=4, value="Cuadrante")
    ws.cell(row=1, column=5, value="CPI")
    ws.cell(row=1, column=6, value="SPI")
    for i, (label, x, y) in enumerate(_quadrants(), start=2):
        ws.cell(row=i, column=4, value=label)
        ws.cell(row=i, column=5, value=x)
        ws.cell(row=i, column=6, value=y)

    # Líneas de referencia del gráfico CPI-SPI: CPI = 1 (vertical) y SPI = 1 (horizontal)
    ws.cell(row=1, column=8, value="CPI")
    ws.cell(row=1, column=9, value="SPI")
    lo, hi = INDEX_MIN - SCATTER_PAD, INDEX_MAX + SCATTER_PAD
    for i, (x, y) in enumerate(((1, lo), (1, hi), (lo, 1), (hi, 1)), start=2):
        ws.cell(row=i, column=8, value=x)
        ws.cell(row=i, column=9, value=y)

    red, yellow = _fmt_threshold(cfg.threshold_red), _fmt_threshold(cfg.threshold_yellow)
    heads = {M_NUM: "Mes N°", M_DATE: "Mes", M_PV: "PV acumulado", M_EV: "EV acumulado", M_AC: "AC acumulado",
             M_BAC: "BAC", M_MARK: "Mes de reporte", M_CPI: "CPI acumulado", M_SPI: "SPI acumulado",
             M_RED: f"Zona roja (< {red})", M_YELLOW: f"Zona amarilla ({red} a {yellow})",
             M_GREEN: f"Zona verde (≥ {yellow})", M_TARGET: "Objetivo (1,00)", M_CPI_RAW: "CPI (valor)",
             M_SPI_RAW: "SPI (valor)"}
    for c, text in heads.items():
        ws.cell(row=D_HEADER_ROW, column=c, value=text)
    curve = f"'{SHEET_CURVE}'"
    f0 = CURVE_FIRST_ROW
    for i in range(lay.months):
        r = D_FIRST_ROW + i
        cr = f0 + i

        def L(c: int) -> str:
            return col(c)

        cells = {
            M_NUM: i + 1,
            M_DATE: lay.month_dates[i],
            M_PV: f"={curve}!{L(C_PVC)}{cr}",
            M_EV: (f'=IF(OR($B$1="",A{r}>$B$1,{curve}!{L(C_EVN)}{cr}=0),NA(),'
                   f'SUM({curve}!${L(C_EV)}${f0}:{L(C_EV)}{cr}))'),
            M_AC: (f'=IF(OR($B$1="",A{r}>$B$1,{curve}!{L(C_ACN)}{cr}=0),NA(),'
                   f'SUM({curve}!${L(C_AC)}${f0}:{L(C_AC)}{cr}))'),
            M_BAC: "=$B$6",
            M_MARK: f"=IF(A{r}=$B$1,$B$7,NA())",
            M_CPI_RAW: f"=IFERROR({L(M_EV)}{r}/{L(M_AC)}{r},NA())",
            M_SPI_RAW: f"=IFERROR({L(M_EV)}{r}/{L(M_PV)}{r},NA())",
            M_CPI: f"=IF(ISNUMBER({L(M_CPI_RAW)}{r}),MAX($B$2,MIN($B$3,{L(M_CPI_RAW)}{r})),NA())",
            M_SPI: f"=IF(ISNUMBER({L(M_SPI_RAW)}{r}),MAX($B$2,MIN($B$3,{L(M_SPI_RAW)}{r})),NA())",
            M_RED: "=$B$4",
            M_YELLOW: "=$B$5-$B$4",
            M_GREEN: "=$B$3-$B$5",
            M_TARGET: 1,
        }
        for c, v in cells.items():
            cell = ws.cell(row=r, column=c, value=v)
            if c == M_DATE:
                cell.number_format = "mmm-yy"

    p_heads = {P_CODE: "Código EDT", P_LEVEL: "Nivel", P_BAC: "BAC", P_PV: "PV acum.", P_EV: "EV acum.",
               P_AC: "AC acum.", P_CPI_RAW: "CPI", P_SPI_RAW: "SPI", P_CPI: "CPI (gráfico)",
               P_SPI: "SPI (gráfico)"}
    for c, text in p_heads.items():
        ws.cell(row=D_HEADER_ROW, column=c, value=text)
    for k, p in enumerate(packages):
        r = D_FIRST_ROW + k
        pr = p.row
        cells = {
            P_CODE: f"='{SHEET_PV}'!$A${pr}",
            P_LEVEL: p.level,
            P_BAC: f"='{SHEET_PV}'!$D${pr}",
            P_PV: f'=IF($B$1="",NA(),{_month_cum(SHEET_PV, pr, lay)})',
            P_EV: f'=IF($B$1="",NA(),{_month_cum(SHEET_EV, pr, lay)})',
            P_AC: f'=IF($B$1="",NA(),{_month_cum(SHEET_AC, pr, lay)})',
            P_CPI_RAW: f"=IFERROR({col(P_EV)}{r}/{col(P_AC)}{r},NA())",
            P_SPI_RAW: f"=IFERROR({col(P_EV)}{r}/{col(P_PV)}{r},NA())",
            P_CPI: f"=IF(ISNUMBER({col(P_CPI_RAW)}{r}),MAX($B$2,MIN($B$3,{col(P_CPI_RAW)}{r})),NA())",
            P_SPI: f"=IF(ISNUMBER({col(P_SPI_RAW)}{r}),MAX($B$2,MIN($B$3,{col(P_SPI_RAW)}{r})),NA())",
        }
        for c, v in cells.items():
            ws.cell(row=r, column=c, value=v)
    ws.column_dimensions["A"].width = 36


def _quadrants() -> list[tuple[str, float, float]]:
    """Rótulos de los cuadrantes: (texto, CPI, SPI), en el margen superior o inferior del gráfico."""
    left, right = (INDEX_MIN + 1) / 2, (INDEX_MAX + 1) / 2
    top, bottom = INDEX_MAX + SCATTER_PAD / 2, INDEX_MIN - SCATTER_PAD / 2
    return [("Adelantado / sobrecosto", left, top), ("Adelantado / ahorro", right, top),
            ("Atrasado / sobrecosto", left, bottom), ("Atrasado / ahorro", right, bottom)]


# ----------------------------------------------------------------------------
# Hoja Cuadro de Mando
# ----------------------------------------------------------------------------
def _merge(ws, row: int, c1: int, c2: int, value=None, row2: int | None = None):
    ws.merge_cells(start_row=row, start_column=c1, end_row=row2 or row, end_column=c2)
    cell = ws.cell(row=row, column=c1)
    if value is not None:
        cell.value = value
    return cell


def _section(ws, row: int, text: str) -> None:
    cell = ws.cell(row=row, column=GRID_FIRST, value=text)
    cell.font = Font(bold=True, size=12, color=DARK)
    cell.alignment = Alignment(vertical="bottom")
    for c in range(GRID_FIRST, GRID_LAST + 1):
        ws.cell(row=row, column=c).border = Border(bottom=Side(style="medium", color=DARK))
    ws.row_dimensions[row].height = 21


def _card(ws, row: int, c1: int, c2: int, label: str, value, fmt: str, caption=None,
          caption_fmt: str | None = None, accent: str = DARK) -> str:
    """Tarjeta de indicador (rótulo, valor y nota). Devuelve la dirección absoluta del valor."""
    _merge(ws, row, c1, c2, label)
    _merge(ws, row + 1, c1, c2, value)
    _merge(ws, row + 2, c1, c2, caption)
    for r in range(row, row + 3):
        for c in range(c1, c2 + 1):
            cell = ws.cell(row=r, column=c)
            cell.fill = FILL_LABEL if r == row else FILL_WHITE
            cell.border = Border(
                top=Side(style="thick", color=accent) if r == row else None,
                bottom=CARD_SIDE if r == row + 2 else None,
                left=CARD_SIDE if c == c1 else None,
                right=CARD_SIDE if c == c2 else None)
    head = ws.cell(row=row, column=c1)
    head.font = Font(bold=True, size=9, color=GRAY_TEXT)
    head.alignment = Alignment(horizontal="center", vertical="center")
    val = ws.cell(row=row + 1, column=c1)
    val.font = Font(bold=True, size=18, color=DARK)
    val.alignment = Alignment(horizontal="center", vertical="center")
    val.number_format = fmt
    cap = ws.cell(row=row + 2, column=c1)
    cap.font = Font(size=9, italic=True, color=GRAY_TEXT)
    cap.alignment = Alignment(horizontal="center", vertical="top")
    if caption_fmt:
        cap.number_format = caption_fmt
    return f"${col(c1)}${row + 1}"


def _tcpi_light(ws, ref: str, cfg: Config) -> None:
    """Semáforo del TCPI: rojo si supera 1/umbral rojo (o es negativo), amarillo si supera 1/umbral amarillo."""
    red, yellow = 1 / cfg.threshold_red, 1 / cfg.threshold_yellow
    ws.conditional_formatting.add(ref, FormulaRule(
        formula=[f"AND(ISNUMBER({ref}),OR({ref}<0,{ref}>{red:.6f}))"], fill=FILL_RED, stopIfTrue=True))
    ws.conditional_formatting.add(ref, FormulaRule(
        formula=[f"AND(ISNUMBER({ref}),{ref}>{yellow:.6f})"], fill=FILL_YELLOW, stopIfTrue=True))
    ws.conditional_formatting.add(ref, FormulaRule(formula=[f"ISNUMBER({ref})"], fill=FILL_GREEN, stopIfTrue=True))


def write_dashboard(ws, report: Report, cfg: Config, lay: _Layout, packages: list[ChartPackage],
                    status: date | None) -> None:
    p = report.project
    tr = lay.total_row
    money = money_format(cfg.decimals)
    money_signed = f"{money};[Red]-{money}"
    currency = cfg.currency or p.currency_symbol or ""

    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 100
    ws.sheet_properties.tabColor = DARK
    ws.column_dimensions["A"].width = 2
    for c in range(GRID_FIRST, GRID_LAST + 1):
        ws.column_dimensions[col(c)].width = GRID_WIDTH
    ws.column_dimensions[col(GRID_LAST + 1)].width = 2

    # Encabezado
    for r, value, font, fill, height in (
            (1, "Cuadro de Mando EVM", Font(bold=True, size=20, color="FFFFFF"), FILL_HEADER, 36),
            (2, p.name, Font(bold=True, size=12, color="FFFFFF"), FILL_HEADER, 22)):
        cell = _merge(ws, r, GRID_FIRST, GRID_LAST)
        put_text(ws, r, GRID_FIRST, value)
        cell.font = font
        cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        for c in range(GRID_FIRST, GRID_LAST + 1):
            ws.cell(row=r, column=c).fill = fill
        ws.row_dimensions[r].height = height
    info = [f"Planificación: {report.plan}" if report.plan else "Planificación: todas",
            p.baseline_label, f"Moneda: {currency or '-'}"]
    if report.multi:
        info.append(f"Empresas: {report.company_a} y {report.company_b} (valores totales)")
    cell = _merge(ws, 3, GRID_FIRST, GRID_LAST)
    put_text(ws, 3, GRID_FIRST, "   |   ".join(info))
    cell.font = Font(size=9, color=DARK)
    cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    for c in range(GRID_FIRST, GRID_LAST + 1):
        ws.cell(row=3, column=c).fill = FILL_SUBTITLE
    ws.row_dimensions[3].height = 18
    ws.row_dimensions[4].height = 8

    # Selección del mes de reporte y del método de EAC
    ws.row_dimensions[5].height = 26
    label_font = Font(bold=True, size=10, color=DARK)
    right = Alignment(horizontal="right", vertical="center")
    center = Alignment(horizontal="center", vertical="center")
    _merge(ws, 5, 2, 4, "Mes de reporte").font = label_font
    ws.cell(row=5, column=2).alignment = right
    sel = _merge(ws, 5, 5, 7, date(status.year, status.month, 1) if status else None)
    sel.number_format = "mmmm yyyy"
    sel.font = Font(bold=True, size=12, color=DARK)
    sel.alignment = center
    month_n = _merge(ws, 5, 8, 9, f"={NAME_REPORT_N}")
    month_n.number_format = '"Mes "0'
    month_n.font = Font(size=9, italic=True, color=GRAY_TEXT)
    month_n.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    _merge(ws, 5, 10, 12, "Corte de datos").font = label_font
    ws.cell(row=5, column=10).alignment = right
    cut = _merge(ws, 5, 13, 14, f'=IF({NAME_STATUS}="","",{NAME_STATUS})')
    cut.number_format = FMT_DATE
    cut.font = Font(size=10, color=DARK)
    cut.alignment = center
    _merge(ws, 5, 15, 17, "Método EAC").font = label_font
    ws.cell(row=5, column=15).alignment = right
    eac = _merge(ws, 5, 18, 21, EAC_METHODS[0])
    eac.font = Font(bold=True, size=10, color=DARK)
    eac.alignment = center
    for c1, c2 in ((5, 7), (18, 21)):
        for c in range(c1, c2 + 1):
            cell = ws.cell(row=5, column=c)
            cell.fill = FILL_INPUT
            cell.border = Border(top=Side(style="thin", color="BF9000"), bottom=Side(style="thin", color="BF9000"),
                                 left=Side(style="thin", color="BF9000") if c == c1 else None,
                                 right=Side(style="thin", color="BF9000") if c == c2 else None)
    note = _merge(ws, 6, GRID_FIRST, GRID_LAST,
                  "El mes de reporte se elige en la lista (meses hasta la fecha de corte de la hoja Resumen); "
                  "'actualizar' lo fija en el mes de corte. Valores acumulados hasta el mes elegido, inclusive.")
    note.font = Font(size=8, italic=True, color=GRAY_TEXT)
    note.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[7].height = 8

    dv = DataValidation(type="list", formula1=NAME_REPORT_LIST, allow_blank=True, showErrorMessage=True,
                        errorTitle="Mes de reporte",
                        error="Se debe elegir un mes de la lista (meses hasta la fecha de corte).")
    dv.add(REPORT_CELL)
    ws.add_data_validation(dv)
    dv2 = DataValidation(type="list", formula1='"' + ",".join(EAC_METHODS) + '"', allow_blank=False,
                         showErrorMessage=True, errorTitle="Método EAC",
                         error="Se debe elegir un método de la lista.")
    dv2.add(EAC_CELL)
    ws.add_data_validation(dv2)

    # Tarjetas de indicadores
    def cum(sheet: str) -> str:
        return f'=IF({NAME_REPORT_N}="","",{_month_cum(sheet, tr, lay)})'

    _section(ws, 8, "Valores acumulados al mes de reporte")
    for r in (9, 14, 19):
        ws.row_dimensions[r].height = 17
        ws.row_dimensions[r + 1].height = 32
        ws.row_dimensions[r + 2].height = 16
    bac = _card(ws, 9, 2, 6, "BAC  ·  presupuesto", f"='{SHEET_PV}'!$D${tr}", money,
                "Presupuesto a la conclusión", accent=ACCENT_BAC)
    pv = _card(ws, 9, 7, 11, "PV  ·  valor planificado", cum(SHEET_PV), money, None, '0.0%" planificado"',
               ACCENT_PV)
    ev = _card(ws, 9, 12, 16, "EV  ·  valor ganado", cum(SHEET_EV), money, None, '0.0%" completado"', ACCENT_EV)
    ac = _card(ws, 9, 17, 21, "AC  ·  costo real", cum(SHEET_AC), money, None, '0.0%" del BAC gastado"',
               ACCENT_AC)
    for ref, c in ((pv, 7), (ev, 12), (ac, 17)):
        ws.cell(row=11, column=c, value=f'=IFERROR({ref}/{bac},"")')

    ws.row_dimensions[12].height = 10
    _section(ws, 13, "Desviaciones e índices de desempeño")
    # Los indicadores de costo quedan vacíos mientras no haya AC cargado (p. ej. AC a mano en multi-empresa)
    cv = _card(ws, 14, 2, 6, "CV  ·  variación de costo", f'=IF(OR({ev}="",{ac}=0),"",{ev}-{ac})', money_signed,
               None, '0.0%" del EV"')
    sv = _card(ws, 14, 7, 11, "SV  ·  variación de cronograma", f'=IF({ev}="","",{ev}-{pv})', money_signed,
               None, '0.0%" del PV"')
    cpi = _card(ws, 14, 12, 16, "CPI  ·  índice de costo", f'=IFERROR({ev}/{ac},"")', FMT_RATIO,
                "EV / AC  ·  < 1: sobrecosto")
    spi = _card(ws, 14, 17, 21, "SPI  ·  índice de cronograma", f'=IFERROR({ev}/{pv},"")', FMT_RATIO,
                "EV / PV  ·  < 1: atraso")
    ws.cell(row=16, column=2, value=f'=IFERROR({cv}/{ev},"")')
    ws.cell(row=16, column=7, value=f'=IFERROR({sv}/{pv},"")')
    for ref in (cpi, spi):
        _traffic_light(ws, ref.replace("$", ""), cfg)

    ws.row_dimensions[17].height = 10
    _section(ws, 18, "Pronósticos")
    m = NAME_EAC_METHOD
    eac = _card(
        ws, 19, 6, 9, "EAC  ·  costo estimado final",
        f'=IFERROR(IF({ac}=0,"",IF({m}="{EAC_METHODS[1]}",{ac}+{bac}-{ev},IF({m}="{EAC_METHODS[2]}",'
        f'{ac}+({bac}-{ev})/({cpi}*{spi}),{bac}/{cpi}))),"")', money, f'="Método: "&{m}')
    _card(ws, 19, 2, 5, "ETC  ·  costo para terminar", f'=IFERROR({eac}-{ac},"")', money, "EAC - AC")
    vac = _card(ws, 19, 10, 13, "VAC  ·  variación al final", f'=IFERROR({bac}-{eac},"")', money_signed,
                None, '0.0%" del BAC"')
    ws.cell(row=21, column=10, value=f'=IFERROR({vac}/{bac},"")')
    tcpi_bac = _card(ws, 19, 14, 17, "TCPI (BAC)", f'=IFERROR(IF({ac}=0,"",({bac}-{ev})/({bac}-{ac})),"")',
                     FMT_RATIO,
                     "(BAC - EV) / (BAC - AC)")
    tcpi_eac = _card(ws, 19, 18, 21, "TCPI (EAC)", f'=IFERROR(({bac}-{ev})/({eac}-{ac}),"")', FMT_RATIO,
                     "(BAC - EV) / (EAC - AC)")
    for ref in (tcpi_bac, tcpi_eac):
        _tcpi_light(ws, ref.replace("$", ""), cfg)

    # Gráficos
    ws.row_dimensions[22].height = 10
    _section(ws, 23, "Evolución en el tiempo")
    ws.row_dimensions[24].height = 6
    ws.row_dimensions[ROW_CHARTS_END].height = 10
    _section(ws, ROW_CHARTS_END + 1, "Desempeño por elemento EDT (niveles 1 y 2) al mes de reporte")
    ws.row_dimensions[ROW_CHARTS_END + 2].height = 6

    # Tabla de elementos EDT (junto al gráfico CPI-SPI)
    hr = ROW_SCATTER
    for c1, c2, text in ((12, 13, "Código"), (14, 18, "Elemento EDT"), (19, 19, "CPI"), (20, 20, "SPI"),
                         (21, 21, "% Compl.")):
        _merge(ws, hr, c1, c2, text)
        for c in range(c1, c2 + 1):
            ws.cell(row=hr, column=c).fill = FILL_HEADER
        head = ws.cell(row=hr, column=c1)
        head.font = Font(bold=True, size=9, color="FFFFFF")
        head.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[hr].height = 18
    data = f"'{SHEET_CHARTDATA}'"
    bottom = Border(bottom=Side(style="thin", color="E7E6E6"))
    for k, pk in enumerate(packages):
        r = hr + 1 + k
        dr = D_FIRST_ROW + k
        code = _merge(ws, r, 12, 13, f"={data}!${col(P_CODE)}${dr}")
        name = _merge(ws, r, 14, 18, f"='{SHEET_PV}'!$B${pk.row}")
        code.alignment = Alignment(horizontal="left", vertical="center", indent=0 if pk.level == 1 else 1)
        name.alignment = Alignment(horizontal="left", vertical="center", indent=0 if pk.level == 1 else 1,
                                   shrink_to_fit=True)
        values = {19: (f'=IFERROR({data}!${col(P_CPI_RAW)}${dr},"")', FMT_RATIO),
                  20: (f'=IFERROR({data}!${col(P_SPI_RAW)}${dr},"")', FMT_RATIO),
                  21: (f'=IFERROR({data}!${col(P_EV)}${dr}/{data}!${col(P_BAC)}${dr},"")', "0.0%")}
        for c, (v, fmt) in values.items():
            cell = ws.cell(row=r, column=c, value=v)
            cell.number_format = fmt
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for c in range(12, 22):
            cell = ws.cell(row=r, column=c)
            cell.border = bottom
            cell.font = Font(size=9, bold=pk.level == 1, color=DARK if pk.level == 1 else "000000")
            if pk.level == 1:
                cell.fill = FILL_LEVEL1
    if packages:
        last = hr + len(packages)
        _traffic_light(ws, f"S{hr + 1}:T{last}", cfg)
    note_row = max(ROW_SCATTER_END, hr + len(packages)) + 1
    note = _merge(ws, note_row, GRID_FIRST, GRID_LAST,
                  f"Gráficos de índices: escala de {_fmt_threshold(INDEX_MIN)} a {_fmt_threshold(INDEX_MAX)}; los "
                  "valores fuera de escala se dibujan en el límite de la escala (los valores exactos están en la "
                  "tabla). Gráfico CPI-SPI: el centro es (1, 1); a la derecha, costo menor que el presupuestado; "
                  "arriba, avance mayor que el planificado.")
    note.font = Font(size=8, italic=True, color=GRAY_TEXT)
    note.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True, indent=1)
    ws.row_dimensions[note_row].height = 24

    # Impresión: vertical, una página de ancho
    ws.print_area = f"B1:U{note_row}"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5


def define_names(wb, lay: _Layout) -> None:
    dash = f"'{SHEET_DASHBOARD}'"
    cut = f"'{SHEET_SUMMARY}'!${CUT_MONTH_CELL[0]}${CUT_MONTH_CELL[1:]}"
    wb.defined_names[NAME_REPORT] = DefinedName(NAME_REPORT, attr_text=f"{dash}!${REPORT_CELL[0]}${REPORT_CELL[1:]}")
    wb.defined_names[NAME_EAC_METHOD] = DefinedName(NAME_EAC_METHOD,
                                                    attr_text=f"{dash}!${EAC_CELL[0]}${EAC_CELL[1:]}")
    wb.defined_names[NAME_REPORT_N] = DefinedName(NAME_REPORT_N, attr_text=f"'{SHEET_CHARTDATA}'!$B$1")
    first = f"'{SHEET_CURVE}'!${col(C_DATE)}${CURVE_FIRST_ROW}"
    wb.defined_names[NAME_REPORT_LIST] = DefinedName(
        NAME_REPORT_LIST,
        attr_text=f"OFFSET({first},0,0,IF(ISNUMBER({cut}),MAX(1,MIN({lay.months},{cut})),{lay.months}),1)")


# ----------------------------------------------------------------------------
# Gráficos
# ----------------------------------------------------------------------------
def _rich(size: int = 800, color: str = GRAY_TEXT, bold: bool = False, italic: bool = False) -> RichText:
    cp = CharacterProperties(sz=size, b=bold, i=italic, solidFill=color)
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp)])


def _title(text: str, size: int = 1200, bold: bool = True, color: str = DARK) -> Title:
    cp = CharacterProperties(sz=size, b=bold, solidFill=color)
    para = Paragraph(pPr=ParagraphProperties(defRPr=cp), r=[RegularTextRun(rPr=cp, t=text)])
    return Title(tx=Text(rich=RichText(p=[para])), overlay=False)


def _line(color: str, width: float = 0.75, dash: str | None = None) -> LineProperties:
    ln = LineProperties(solidFill=color, w=int(width * 12700))
    if dash:
        ln.prstDash = dash
    return ln


def _style_axis(axis, title: str | None = None, fmt: str | None = None, grid: bool = False) -> None:
    axis.delete = False
    if fmt:
        axis.number_format = fmt
    axis.txPr = _rich(800)
    axis.spPr = GraphicalProperties(ln=_line("BFBFBF"))
    axis.majorGridlines = ChartLines(spPr=GraphicalProperties(ln=_line("E7E6E6", 0.5))) if grid else None
    axis.title = _title(title, 900, bold=False, color=GRAY_TEXT) if title else None


def _frame(chart, title: str, width_px: int, height_px: int) -> None:
    chart.title = _title(title)
    chart.roundedCorners = False
    chart.graphical_properties = GraphicalProperties(solidFill="FFFFFF", ln=_line("D9D9D9"))
    chart.width = width_px / 37.795
    chart.height = height_px / 37.795


def _legend(chart) -> None:
    chart.legend.position = "b"
    chart.legend.overlay = False
    chart.legend.txPr = _rich(800, color="404040")


def _date_axis(months: int) -> DateAxis:
    step = next(s for s in (1, 2, 3, 6, 12, 24) if months / s <= 14)
    axis = DateAxis(crossAx=100)
    axis.baseTimeUnit = "months"
    axis.majorUnit = step
    axis.majorTimeUnit = "months"
    axis.lblOffset = 100
    axis.tickLblPos = "low"
    _style_axis(axis, fmt="mmm-yy")
    return axis


def _money_axis(bac: float, currency: str) -> tuple[str, str]:
    """Formato del eje de montos según la magnitud del BAC."""
    unit = f" de {currency}" if currency else ""
    if bac >= 1e7:
        return '#,##0,,', f"Millones{unit}"
    if bac >= 1e6:
        return '#,##0.0,,', f"Millones{unit}"
    if bac >= 1e4:
        return '#,##0,', f"Miles{unit}"
    return "#,##0", currency or "Monto"


def _anchor(c1: int, r1: int, c2: int, r2: int, off1: int = 0, off2: int = 0) -> TwoCellAnchor:
    """Ancla entre celdas (columnas y filas base 0; desplazamientos en píxeles)."""
    anchor = TwoCellAnchor()
    anchor._from = AnchorMarker(col=c1, colOff=off1 * EMU_PX, row=r1, rowOff=0)
    anchor.to = AnchorMarker(col=c2, colOff=off2 * EMU_PX, row=r2, rowOff=0)
    return anchor


def _marker(series, symbol: str | None, color: str, show: bool) -> None:
    """Marcadores pequeños en una serie medida (EV, AC, CPI o SPI)."""
    if symbol and show:
        series.marker = Marker(symbol=symbol, size=5)
        series.marker.graphicalProperties = GraphicalProperties(solidFill=color, ln=_line(color))
    else:
        series.marker = Marker(symbol="none")


def _evm_chart(data, months: int, bac: float, currency: str, data_months: int):
    last = D_FIRST_ROW + months - 1
    fmt, unit = _money_axis(bac, currency)
    cats = Reference(data, min_col=M_DATE, min_row=D_FIRST_ROW, max_row=last)

    line = LineChart()
    for c, color, width, dash, marker in ((M_PV, ACCENT_PV, 2.25, None, None),
                                          (M_EV, ACCENT_EV, 2.25, None, "circle"),
                                          (M_AC, ACCENT_AC, 2.25, None, "square"),
                                          (M_BAC, "7F7F7F", 1.25, "dash", None)):
        line.add_data(Reference(data, min_col=c, min_row=D_HEADER_ROW, max_row=last), titles_from_data=True)
        s = line.series[-1]
        s.smooth = False
        s.graphicalProperties = GraphicalProperties(ln=_line(color, width, dash))
        _marker(s, marker, color, data_months <= 2)   # con uno o dos meses de datos, la línea no se ve

    # Marca vertical del mes de reporte (columna fina)
    bar = BarChart()
    bar.add_data(Reference(data, min_col=M_MARK, min_row=D_HEADER_ROW, max_row=last), titles_from_data=True)
    bar.gapWidth = 500
    bar.series[0].graphicalProperties = GraphicalProperties(solidFill="A6A6A6", ln=LineProperties(noFill=True))

    x = _date_axis(months)
    line.x_axis = bar.x_axis = x
    line.y_axis.crossAx = x.axId
    _style_axis(line.y_axis, title=unit, fmt=fmt, grid=True)
    line.y_axis.scaling.min = 0
    bar.y_axis = line.y_axis
    line.set_categories(cats)
    bar.set_categories(cats)
    line += bar
    _legend(line)
    return line


def _index_chart(data, months: int):
    last = D_FIRST_ROW + months - 1
    cats = Reference(data, min_col=M_DATE, min_row=D_FIRST_ROW, max_row=last)
    line = LineChart()
    for c, color, width, dash, marker in ((M_CPI, COLOR_CPI, 2.25, None, "circle"),
                                          (M_SPI, COLOR_SPI, 2.25, None, "diamond"),
                                          (M_TARGET, "404040", 1.0, "dash", None)):
        line.add_data(Reference(data, min_col=c, min_row=D_HEADER_ROW, max_row=last), titles_from_data=True)
        s = line.series[-1]
        s.smooth = False
        s.graphicalProperties = GraphicalProperties(ln=_line(color, width, dash))
        _marker(s, marker, color, months <= 36)
    # Bandas de alarma (área apilada detrás de las líneas)
    area = AreaChart(grouping="stacked")
    for c, color in ((M_RED, "F4CCCC"), (M_YELLOW, "FFF2CC"), (M_GREEN, "E2EFDA")):
        area.add_data(Reference(data, min_col=c, min_row=D_HEADER_ROW, max_row=last), titles_from_data=True)
        area.series[-1].graphicalProperties = GraphicalProperties(solidFill=color, ln=LineProperties(noFill=True))
    x = _date_axis(months)
    line.x_axis = area.x_axis = x
    line.y_axis.crossAx = x.axId
    _style_axis(line.y_axis, title="Índice acumulado", fmt="0.00", grid=True)
    line.y_axis.scaling.min = INDEX_MIN
    line.y_axis.scaling.max = INDEX_MAX
    line.y_axis.majorUnit = 0.1
    line.y_axis.crossBetween = "midCat"
    area.y_axis = line.y_axis
    line.set_categories(cats)
    area.set_categories(cats)
    line += area
    _legend(line)
    return line


def _scatter_chart(data, packages: list[ChartPackage]):
    chart = ScatterChart()
    chart.scatterStyle = "lineMarker"
    chart.varyColors = False
    for first, title in ((2, "CPI = 1"), (4, "SPI = 1")):
        s = Series(Reference(data, min_col=9, min_row=first, max_row=first + 1),
                   Reference(data, min_col=8, min_row=first, max_row=first + 1), title=title)
        s.marker = Marker(symbol="none")
        s.graphicalProperties = GraphicalProperties(ln=_line("404040", 1.25))
        chart.series.append(s)
    # Etiquetas alternadas (nivel 1: arriba/abajo; nivel 2: derecha/izquierda) para que no se superpongan
    # cuando un elemento y sus subelementos tienen índices parecidos
    positions, count = [], {1: 0, 2: 0}
    for p in packages:
        positions.append(("t", "b") if p.level == 1 else ("r", "l"))
        positions[-1] = positions[-1][count[p.level] % 2]
        count[p.level] += 1
    for k, p in enumerate(packages):
        r = D_FIRST_ROW + k
        s = Series(Reference(data, min_col=P_SPI, min_row=r), Reference(data, min_col=P_CPI, min_row=r),
                   title=p.code)
        color = DARK if p.level == 1 else "2E75B6"
        s.marker = Marker(symbol="diamond" if p.level == 1 else "circle", size=10 if p.level == 1 else 8)
        s.marker.graphicalProperties = GraphicalProperties(solidFill=color, ln=_line("FFFFFF", 0.75))
        s.graphicalProperties = GraphicalProperties(ln=LineProperties(noFill=True))
        s.dLbls = DataLabelList(showSerName=True, showVal=False, showCatName=False, showLegendKey=False,
                                showPercent=False, showBubbleSize=False, dLblPos=positions[k])
        s.dLbls.txPr = _rich(800, color=DARK if p.level == 1 else "404040", bold=p.level == 1)
        chart.series.append(s)
    for k, (label, _x, _y) in enumerate(_quadrants()):
        r = 2 + k
        s = Series(Reference(data, min_col=6, min_row=r), Reference(data, min_col=5, min_row=r), title=label)
        s.marker = Marker(symbol="none")
        s.graphicalProperties = GraphicalProperties(ln=LineProperties(noFill=True))
        s.dLbls = DataLabelList(showSerName=True, showVal=False, showCatName=False, showLegendKey=False,
                                showPercent=False, showBubbleSize=False, dLblPos="ctr")
        s.dLbls.txPr = _rich(800, color="808080", italic=True)
        chart.series.append(s)
    # Margen de SCATTER_PAD alrededor de la escala: sus marcas no muestran número
    fmt = f'[<{INDEX_MIN - SCATTER_PAD / 2}]"";[>{INDEX_MAX + SCATTER_PAD / 2}]"";0.0'
    for axis, title in ((chart.x_axis, "CPI - índice de costo"), (chart.y_axis, "SPI - índice de cronograma")):
        _style_axis(axis, title=title, fmt=fmt, grid=True)
        axis.scaling.min = INDEX_MIN - SCATTER_PAD
        axis.scaling.max = INDEX_MAX + SCATTER_PAD
        axis.majorUnit = 0.1
        axis.crosses = "min"
    chart.legend = None
    return chart


def add_dashboard_charts(wb, packages: list[ChartPackage], months: int, plan_end: int, cutoff: int | None,
                         bac: float, currency: str, data_months: int = 0) -> None:
    """Agrega (o vuelve a agregar) los tres gráficos de la hoja Cuadro de Mando.

    data_months: meses con EV o AC cargados hasta la fecha de corte (con 1 o 2, EV y AC llevan marcadores).
    """
    ws = wb[SHEET_DASHBOARD]
    data = wb[SHEET_CHARTDATA]
    ws._charts = []  # openpyxl no conserva gráficos al reabrir: se recrean siempre
    evm_months, index_months = chart_months(months, plan_end, cutoff)
    r1, r2 = ROW_CHARTS - 1, ROW_CHARTS_END - 1

    chart = _evm_chart(data, evm_months, bac, currency, data_months)
    _frame(chart, "Evolución EVM - valores acumulados", 496, 380)
    chart.anchor = _anchor(GRID_FIRST - 1, r1, 10, r2, 0, 46)
    ws.add_chart(chart)

    chart = _index_chart(data, index_months)
    _frame(chart, "Evolución de CPI y SPI acumulados", 496, 380)
    chart.anchor = _anchor(11, r1, GRID_LAST, r2, 4, 0)
    ws.add_chart(chart)

    chart = _scatter_chart(data, packages)
    _frame(chart, "CPI vs SPI por elemento EDT", 496, 500)
    chart.anchor = _anchor(GRID_FIRST - 1, ROW_SCATTER - 1, 10, ROW_SCATTER_END - 1, 0, 46)
    ws.add_chart(chart)

