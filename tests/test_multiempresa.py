"""Multi-empresa: % planificado y % ganado de cada empresa, con su BAC y su AC en el libro."""
import shutil
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from evm.cli import main
from evm.companies import resolve_shares
from evm.config import Config, ConfigError, load_config
from evm.model import Task
from evm.reader_xml import read_xml
from evm.update import read_book, run_update
from evm.wbs import SEV_INFO, build_report
from evm.workbook import FILL_INPUT, FIRST_MONTH_COL, FIRST_ROW, read_meta, read_structure


ROOT = Path(__file__).resolve().parents[1]
SHEETS = ["Cuadro de Mando", "Resumen", "PV", "PV-EmpresaA", "PV-EmpresaB", "% PV-EmpresaA", "% PV-EmpresaB",
          "EV", "EV-EmpresaA", "EV-EmpresaB", "% EV-EmpresaA", "% EV-EmpresaB", "AC", "AC-EmpresaA", "AC-EmpresaB"]


# ----------------------------------------------------------------------------
# Utilidades
# ----------------------------------------------------------------------------
def multi_config(tmp_path, example_files, extra: str = "", name: str = "multi.toml") -> Path:
    path = tmp_path / name
    path.write_text(example_files["config"].read_text(encoding="utf-8") +
                    "\n[multiempresa]\nactivar = true\n" + extra, encoding="utf-8")
    return path


def report_for(xml: Path, cfg: Config):
    return build_report(read_xml(xml, cfg), cfg)


def interno(plans):
    """Hijos de la planificación INTERNO: arranque, gestión, ingeniería, construcción y puesta en marcha."""
    return plans[0].children


def write_variant(generator, path: Path, mutate, status: date | None = None) -> Path:
    plans = generator.build_plans()
    mutate(interno(plans))
    generator.write(plans, path, status)
    return path


def rows_by_code(report):
    return {r.code: r for r in report.rows}


def book_row(book, code):
    return next(s.row for s in book.structure if s.code == code)


def task(uid, name, company="", split=None, parent=None, **kw):
    t = Task(uid=uid, id=uid, name=name, outline_level=1 if parent is None else parent.outline_level + 1,
             company=company, split_a=split, **kw)
    if parent is not None:
        t.parent = parent
        parent.children.append(t)
    return t


@pytest.fixture(scope="module")
def multi_template(tmp_path_factory, example_files, generator):
    """Libro multi-empresa del ejemplo con el BAC de cada empresa, generado una vez por módulo, y su configuración
    (las pruebas no la modifican)."""
    folder = tmp_path_factory.mktemp("multi")
    cfg_path = multi_config(folder, example_files)
    out = folder / "EVM_multi.xlsx"
    assert main(["generar", "--origen", str(example_files["base"]), "--config", str(cfg_path),
                 "--salida", str(out)]) == 0
    generator.fill_book(out)                        # BAC de cada empresa, que se ingresa a mano
    return out, cfg_path


@pytest.fixture()
def multi_book(tmp_path, multi_template):
    template, cfg_path = multi_template
    out = tmp_path / "EVM_multi.xlsx"
    shutil.copy(template, out)
    return out, cfg_path


def set_bac(path, values: dict[str, tuple[float, float]]) -> None:
    """BAC de cada empresa en las filas indicadas: {código: (BAC de A, BAC de B)}."""
    wb = load_workbook(path)
    rows = {s.code: s.row for s in read_structure(wb)}
    for code, (a, b) in values.items():
        wb["PV-EmpresaA"].cell(rows[code], 4, a)
        wb["PV-EmpresaB"].cell(rows[code], 4, b)
    wb.save(path)


# ----------------------------------------------------------------------------
# Configuración
# ----------------------------------------------------------------------------
def test_example_config_defaults():
    cfg = load_config(ROOT / "config.ejemplo.toml")
    assert not cfg.multi_enabled
    assert (cfg.company_field, cfg.company_a, cfg.company_b, cfg.company_both) == \
        ("Empresa", "EmpresaA", "EmpresaB", "Ambas")
    assert (cfg.split_field, cfg.split_default) == ("% Reparto", 50.0)
    assert cfg.wbs_field == "Item EDT" and cfg.plan == "INTERNO" and cfg.currency == "EUR"


def test_example_config_reads_item_edt(example_files):
    """El campo personalizado "Item EDT" del ejemplo tiene los mismos códigos que el campo EDT."""
    cfg = load_config(ROOT / "config.ejemplo.toml")
    by_item = read_xml(example_files["base"], cfg)
    by_wbs = read_xml(example_files["base"], Config())
    assert [t.wbs for t in by_item.tasks] == [t.wbs for t in by_wbs.tasks]
    assert len([r for r in report_for(example_files["base"], cfg).rows if r.is_input]) == 24


def test_config_from_toml(tmp_path):
    path = tmp_path / "c.toml"
    path.write_text('[multiempresa]\nactivar = true\nempresa_a = "Acme"\nempresa_b = "Beta SL"\n'
                    'ambas = "Compartida"\ncampo_empresa = "Text3"\ncampo_reparto = "Number2"\n'
                    'reparto_por_defecto = 70\n', encoding="utf-8")
    cfg = load_config(path)
    assert cfg.multi_enabled
    assert (cfg.company_a, cfg.company_b, cfg.company_both) == ("Acme", "Beta SL", "Compartida")
    assert (cfg.company_field, cfg.split_field, cfg.split_default) == ("Text3", "Number2", 70.0)


@pytest.mark.parametrize("kw, message", [
    ({"company_a": "NombreMuyLargo16"}, "15 caracteres"),
    ({"company_a": ""}, "vacío"),
    ({"company_b": "A/B"}, "no admitidos"),
    ({"company_b": "O'Neil"}, "no admitidos"),
    ({"company_a": "Acme", "company_b": "ACME"}, "distintas"),
    ({"company_both": "empresaa"}, "ambas"),
    ({"split_default": 120.0}, "reparto_por_defecto"),
])
def test_config_validation(kw, message):
    with pytest.raises(ConfigError, match=message):
        Config(multi_enabled=True, **kw).validate()
    Config(multi_enabled=False, **kw).validate()          # desactivado: no se valida


# ----------------------------------------------------------------------------
# Reparto por tarea
# ----------------------------------------------------------------------------
def test_resolve_shares_inheritance_and_notes():
    cfg = Config(multi_enabled=True)
    root = task(1, "Programa")
    both = task(2, "Ambas 80", "Ambas", 80, root)
    inherit = task(3, "Hereda", "", None, both)
    b = task(4, "B", "empresab", 30, root)               # % Reparto ignorado fuera de "Ambas"
    b_both = task(5, "Ambas sin %", "Ambas", None, b)
    zero = task(6, "Ambas con 0", "Ambas", 0, root)       # campo numérico vacío en MS Project
    bad_split = task(7, "Ambas 150", "Ambas", 150, root)
    unknown = task(8, "Otra", "Otra empresa", None, root)
    a = task(9, "A", "EmpresaA", None, root)

    shares, notes = resolve_shares([root], cfg)
    assert shares[both.uid] == pytest.approx(0.8)
    assert shares[inherit.uid] == pytest.approx(0.8)
    assert shares[b.uid] == 0.0
    assert shares[b_both.uid] == pytest.approx(0.5)
    assert shares[zero.uid] == pytest.approx(0.5)
    assert shares[bad_split.uid] == pytest.approx(0.5)
    assert shares[unknown.uid] == pytest.approx(0.5)
    assert shares[a.uid] == 1.0
    assert [t.uid for t in notes["sin_empresa"]] == [root.uid]
    assert [t.uid for t in notes["reparto"]] == [bad_split.uid]
    assert [t.uid for t in notes["desconocida"]] == [unknown.uid]

    # Una rama sin empresa hereda de sus tareas superiores aunque no se recorran
    shares2, notes2 = resolve_shares([inherit], cfg)
    assert shares2[inherit.uid] == pytest.approx(0.8) and not notes2["sin_empresa"]

    # Reparto por defecto configurable
    shares3, _ = resolve_shares([root], Config(multi_enabled=True, split_default=25))
    assert shares3[root.uid] == pytest.approx(0.25) and shares3[b_both.uid] == pytest.approx(0.25)


# ----------------------------------------------------------------------------
# Lectores
# ----------------------------------------------------------------------------
def test_xml_reads_company_fields(tmp_path, example_files):
    cfg = load_config(multi_config(tmp_path, example_files))
    p = read_xml(example_files["base"], cfg)
    by_name = {}
    for t in p.tasks:
        by_name.setdefault(t.name, t)
    assert (by_name["Gestión del programa"].company, by_name["Gestión del programa"].split_a) == ("Ambas", 60)
    assert by_name["Ingeniería"].company == "EmpresaA"
    assert by_name["Estudios de terreno"].company == "EmpresaB"
    assert by_name["Detalle civil"].company == "" and by_name["Detalle civil"].split_a is None
    assert not any("empresa" in w.message for w in p.warnings)

    missing = load_config(multi_config(tmp_path, example_files, 'campo_empresa = "No existe"\n', "m2.toml"))
    p2 = read_xml(example_files["base"], missing)
    assert any("No se encontró el campo de empresa 'No existe'" in w.message for w in p2.warnings)


def test_com_reads_company_fields():
    from test_reader_com_fake import FakeApp, make_session, program

    class App(FakeApp):
        def CustomFieldGetName(self, fid):
            return {188743737: "Empresa", 188743768: "% Reparto"}.get(fid) or super().CustomFieldGetName(fid)

    tasks = program()
    values = {1: ("Ambas", 0.0), 2: ("", 80.0), 4: ("EmpresaB", 0.0)}   # uid -> (Texto3, Número2)
    for t in tasks:
        if t is not None:
            t.Text3, t.Number2 = values.get(t.UniqueID, ("", 0.0))
    s = make_session(tasks)
    s.app = App()
    cfg = Config(multi_enabled=True)
    p = s.read_project(cfg)
    by_uid = {t.uid: t for t in p.tasks}
    assert (by_uid[1].company, by_uid[1].split_a) == ("Ambas", 0.0)
    assert by_uid[4].company == "EmpresaB"
    rep = build_report(p, cfg)
    rows = rows_by_code(rep)
    # 1.1.1 hereda "Ambas" de la fase: las dos empresas tienen % planificado; 1.1.2 es de la empresa B (también su
    # subtarea sin empresa): la empresa A no participa
    assert sum(rows["1.1.1"].pct_vec["A"]) == pytest.approx(1.0) and sum(rows["1.1.1"].pct_vec["B"]) == 1.0
    assert any("EmpresaA no participa en las tareas de 1 fila (1.1.2)" in i.message for i in rep.issues)


# ----------------------------------------------------------------------------
# Informe (% planificado por empresa)
# ----------------------------------------------------------------------------
def test_report_planned_percent_per_company(tmp_path, example_files):
    cfg = load_config(multi_config(tmp_path, example_files))
    rep = report_for(example_files["base"], cfg)
    assert rep.multi and (rep.company_a, rep.company_b) == ("EmpresaA", "EmpresaB")
    for r in rep.rows:
        if r.is_input:
            assert set(r.pct_vec) == {"A", "B"}
            assert sum(r.pct_vec["A"]) == pytest.approx(1.0) and sum(r.pct_vec["B"]) == pytest.approx(1.0)
    messages = [i.message for i in rep.issues]
    # 2.1.1 es de EmpresaB dentro de Ingeniería (EmpresaA); la construcción es de EmpresaB
    assert any(m.startswith("EmpresaA no participa en las tareas de 7 filas (2.1.1, 3.1.1, 3.1.2, 3.1.3, 3.2.1, "
                            "3.2.2, 3.2.3)") for m in messages)
    assert any(m.startswith("EmpresaB no participa en las tareas de 7 filas (2.1.2, 2.1.3, 2.1.4, 2.2.1, 2.2.2, "
                            "2.2.3, 2.2.4)") for m in messages)
    assert not [i for i in rep.issues if "empresa" in i.message.lower() and i.severity != SEV_INFO]


def test_report_company_notes_and_residual(tmp_path, example_files, generator):
    def mutate(nodes):
        nodes[1].split = 150                                     # Gestión: % fuera de rango -> 50 %
        nodes[3].company = "Otra"                                # Construcción: empresa no reconocida
        nodes[4].company = ""                                    # Puesta en marcha: sin empresa
        nodes[1].children[0].children.append(generator.T(
            "Gasto sin código", 3, date(2026, 3, 2), date(2026, 3, 31), 10_000, code="", company="EmpresaB"))

    xml = write_variant(generator, tmp_path / "variante.xml", mutate)
    cfg = load_config(multi_config(tmp_path, example_files))
    rep = report_for(xml, cfg)
    rows = rows_by_code(rep)
    issues = [(i.code, i.message) for i in rep.issues]
    assert any(code == "4.0.0" and "Sin empresa" in msg for code, msg in issues)
    assert any(msg.startswith("7 tareas sin empresa") for _, msg in issues)
    assert any(code == "3.0.0" and "'Otra' no reconocida" in msg for code, msg in issues)
    assert any(code == "1.0.0" and "fuera de rango" in msg for code, msg in issues)
    # Fila de otras tareas: la tarea sin código es de EmpresaB
    residual = rows["1.1.0-D"]
    assert [t.name for t in residual.leaves] == ["Gasto sin código"]


# ----------------------------------------------------------------------------
# Libro
# ----------------------------------------------------------------------------
def test_workbook_sheets_and_formulas(multi_book):
    path, _ = multi_book
    wb = load_workbook(path)
    assert wb.sheetnames[:len(SHEETS)] == SHEETS
    r = FIRST_ROW
    month = FIRST_MONTH_COL
    pv, pv_a, ev_a = wb["PV"], wb["PV-EmpresaA"], wb["EV-EmpresaA"]
    assert pv.cell(r, 4).value == f"='PV-EmpresaA'!D{r}+'PV-EmpresaB'!D{r}"
    letter = pv.cell(r, month).column_letter
    assert pv.cell(r, month).value == f"='PV-EmpresaA'!{letter}{r}+'PV-EmpresaB'!{letter}{r}"
    assert wb["EV"].cell(r, month).value == f"='EV-EmpresaA'!{letter}{r}+'EV-EmpresaB'!{letter}{r}"
    leaf = next(x for x in range(FIRST_ROW, pv.max_row) if pv.cell(x, 1).value == "1.1.1")
    assert pv_a.cell(leaf, 4).value == pytest.approx(1_080_000)            # BAC del ejemplo: Ambas 60 %
    assert wb["PV-EmpresaB"].cell(leaf, 4).value == pytest.approx(720_000)
    assert pv_a.cell(leaf, 4).fill.fgColor.rgb.endswith(FILL_INPUT.fgColor.rgb[-6:])
    assert pv_a.cell(leaf, month).value == f"=N($D{leaf})*'% PV-EmpresaA'!{letter}{leaf}"
    assert ev_a.cell(leaf, 4).value == f"=IF('PV-EmpresaA'!D{leaf}=\"\",\"\",'PV-EmpresaA'!D{leaf})"
    assert ev_a.cell(leaf, month).value == f"=N($D{leaf})*'% EV-EmpresaA'!{letter}{leaf}"
    ac_a = wb["AC-EmpresaA"]
    assert ac_a.cell(leaf, month).fill.fgColor.rgb.endswith(FILL_INPUT.fgColor.rgb[-6:])
    assert not wb["AC"].cell(leaf, month).fill.fgColor.rgb.endswith(FILL_INPUT.fgColor.rgb[-6:])
    assert pv.data_validations.dataValidation and not pv_a.data_validations.dataValidation
    meta = read_meta(wb)
    assert meta["multiempresa"] is True and meta["empresa_a"] == "EmpresaA" and meta["ambas"] == "Ambas"


def test_single_company_book_unchanged(tmp_path, example_files):
    out = tmp_path / "EVM.xlsx"
    assert main(["generar", "--origen", str(example_files["base"]), "--config", str(example_files["config"]),
                 "--salida", str(out)]) == 0
    wb = load_workbook(out)
    assert not [s for s in wb.sheetnames if "-" in s]
    assert read_meta(wb)["multiempresa"] is False


# ----------------------------------------------------------------------------
# Actualización
# ----------------------------------------------------------------------------
def test_update_splits_ev_and_leaves_ac(multi_book, example_files):
    path, cfg_path = multi_book
    cfg = load_config(cfg_path)
    res = run_update(read_xml(example_files["jun"], cfg), cfg, path, backup=False)
    b = read_book(path)
    assert b.multi
    m = 17
    loe = book_row(b, "1.1.1")
    pv_a = b.bac_role["A"][loe] * sum(b.pct_pv["A"][loe][:m + 1])
    assert b.ev_a[loe][m] == pytest.approx(pv_a, abs=0.01)                       # LOE: PV de A
    assert b.ev_b[loe][m] == pytest.approx(sum(b.pv[loe][:m + 1]) - pv_a, abs=0.01)
    assert b.ev_a[loe][m] == pytest.approx(0.6 * b.ev[loe][m], abs=1)
    terreno = book_row(b, "2.1.1")                                                # EmpresaB
    assert b.ev_a[terreno][m] == 0 and b.ev_b[terreno][m] > 0
    civil = book_row(b, "2.1.3")                                                  # EmpresaA
    assert b.ev_b[civil][m] == 0 and b.ev_a[civil][m] > 0
    for s in b.input_rows:
        assert (b.ev[s.row][m] or 0) == pytest.approx((b.ev_a[s.row][m] or 0) + (b.ev_b[s.row][m] or 0))
        assert not any(b.ac_a[s.row]) and not any(b.ac_b[s.row])                # AC: a mano
    assert {k for r in res.rows for k in r.by_company} == {"EV-EmpresaA", "EV-EmpresaB"}


def test_update_ev_by_company_progress(tmp_path, example_files, generator):
    """En una fila con tareas de las dos empresas, el EV de cada una usa el avance de sus tareas y su BAC."""
    def mutate(nodes):
        conceptual = nodes[2].children[0].children[1]           # 2.1.2 Diseño conceptual
        conceptual.method = "% Completado"
        conceptual.children[2].company = "EmpresaB"            # Alternativas y selección
        nodes[1].children[0].children.append(generator.T(
            "Gasto sin código", 3, date(2026, 3, 2), date(2026, 3, 31), 10_000, code="", company="EmpresaB"))

    base = write_variant(generator, tmp_path / "v_base.xml", mutate)
    status = write_variant(generator, tmp_path / "v_2026-06.xml", mutate, date(2026, 6, 30))
    cfg_path = multi_config(tmp_path, example_files)
    book = tmp_path / "EVM.xlsx"
    assert main(["generar", "--origen", str(base), "--config", str(cfg_path), "--salida", str(book)]) == 0
    set_bac(book, {"2.1.2": (340_000, 260_000), "1.1.0-D": (0, 10_000)})
    cfg = load_config(cfg_path)
    project = read_xml(status, cfg)
    run_update(project, cfg, book, backup=False)

    b = read_book(book)
    m = 5
    row = book_row(b, "2.1.2")
    alt = next(t for t in project.tasks if t.name == "Alternativas y selección")
    assert 0 < alt.percent_complete < 100                                  # avance distinto por empresa
    # EmpresaB solo tiene esa tarea en la fila: su % ganado es el de la tarea
    assert b.ev_b[row][m] == pytest.approx(260_000 * alt.percent_complete / 100, abs=0.01)
    assert 0 < b.ev_a[row][m] < 340_000
    assert b.ev_a[row][m] / 340_000 != pytest.approx(alt.percent_complete / 100)

    residual = book_row(b, "1.1.0-D")
    gasto = next(t for t in project.tasks if t.name == "Gasto sin código")
    assert not b.ev_a[residual][m]
    assert b.ev_b[residual][m] == pytest.approx(10_000 * gasto.percent_complete / 100, abs=0.01)


def test_cli_dry_run_and_book_settings(multi_book, example_files, capsys):
    path, cfg_path = multi_book
    base = ["actualizar", "--origen", str(example_files["jun"]), "--libro", str(path)]
    assert main(base + ["--config", str(cfg_path), "--simular"]) == 0
    out = capsys.readouterr().out
    assert "EV-EmpresaA" in out and "EV-EmpresaB" in out and "se carga a mano" in out
    assert not any(any(v) for v in read_book(path).ev_a.values())               # simulación: sin cambios

    # Con una configuración sin [multiempresa] se usan los datos del libro
    assert main(base + ["--config", str(example_files["config"])]) == 0
    b = read_book(path)
    assert any(b.ev_a[s.row][17] for s in b.input_rows)


# ----------------------------------------------------------------------------
# Migración
# ----------------------------------------------------------------------------
def test_migration_multi_to_multi_and_single_rejected(tmp_path, multi_book, example_files, capsys):
    path, cfg_path = multi_book
    cfg = load_config(cfg_path)
    run_update(read_xml(example_files["jun"], cfg), cfg, path, backup=False)
    wb = load_workbook(path)
    leaf = book_row(read_book(path), "2.1.3")
    wb["AC-EmpresaA"].cell(leaf, FIRST_MONTH_COL + 17, 12_345)
    wb.save(path)
    old = read_book(path)

    new = tmp_path / "EVM_v2.xlsx"
    assert main(["generar", "--origen", str(example_files["base"]), "--config", str(cfg_path),
                 "--salida", str(new), "--migrar-desde", str(path)]) == 0
    b = read_book(new)
    for attr in ("ev_a", "ev_b", "ac_a"):
        before = sum(v or 0 for vals in getattr(old, attr).values() for v in vals)
        after = sum(v or 0 for vals in getattr(b, attr).values() for v in vals)
        assert after == pytest.approx(before)
    assert b.ac_a[book_row(b, "2.1.3")][17] == 12_345

    single = tmp_path / "EVM_single.xlsx"
    assert main(["generar", "--origen", str(example_files["base"]), "--config", str(example_files["config"]),
                 "--salida", str(single)]) == 0
    capsys.readouterr()
    assert main(["generar", "--origen", str(example_files["base"]), "--config", str(cfg_path),
                 "--salida", str(tmp_path / "x.xlsx"), "--migrar-desde", str(single)]) == 2
    assert "multi-empresa" in capsys.readouterr().err


# ----------------------------------------------------------------------------
# Fórmulas (LibreOffice)
# ----------------------------------------------------------------------------
def test_multi_formulas_recalculate(multi_book, example_files, recalc):
    path, cfg_path = multi_book
    cfg = load_config(cfg_path)
    run_update(read_xml(example_files["jun"], cfg), cfg, path, backup=False)
    wb = load_workbook(path)
    rows = {wb["PV"].cell(r, 1).value: r for r in range(FIRST_ROW, wb["PV"].max_row + 1)}
    wb["AC-EmpresaA"].cell(rows["2.1.3"], FIRST_MONTH_COL + 17, 1_000)
    wb["AC-EmpresaB"].cell(rows["3.1.1"], FIRST_MONTH_COL + 17, 500)
    wb.save(path)

    out = load_workbook(recalc(path), data_only=True)
    errors = [(ws.title, c.coordinate, c.value) for ws in out.worksheets for row in ws.iter_rows() for c in row
              if isinstance(c.value, str) and c.value.startswith("#") and c.value != "#N/A"]
    assert not errors
    total_row = rows["TOTAL"]
    for kind in ("PV", "EV", "AC"):
        t, a, b = out[kind], out[f"{kind}-EmpresaA"], out[f"{kind}-EmpresaB"]
        for c in [4] + list(range(FIRST_MONTH_COL, FIRST_MONTH_COL + 121)):
            for r in (rows["1.1.1"], rows["2.0.0"], total_row):
                assert (t.cell(r, c).value or 0) == pytest.approx((a.cell(r, c).value or 0) +
                                                                 (b.cell(r, c).value or 0), abs=0.01)
    assert out["PV"].cell(total_row, 4).value == pytest.approx(44_210_000)
    assert out["PV-EmpresaA"].cell(total_row, 4).value == pytest.approx(9_947_000)
    assert out["AC"].cell(total_row, FIRST_MONTH_COL + 17).value == pytest.approx(1_500)
    summary = {out["Resumen"].cell(r, 1).value: out["Resumen"].cell(r, 2).value for r in range(1, 60)}
    assert summary["AC - Costo real acumulado"] == pytest.approx(1_500)
    assert out["Cuadro de Mando"]["Q10"].value == pytest.approx(1_500)                 # AC total en el tablero
    ev_total = sum(out["EV"].cell(total_row, c).value or 0 for c in range(FIRST_MONTH_COL, FIRST_MONTH_COL + 18))
    assert summary["EV - Valor ganado acumulado"] == pytest.approx(ev_total)
