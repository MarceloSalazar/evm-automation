"""Pruebas de punta a punta con proyectos aleatorios (semilla fija).

Cada proyecto se arma con el generador de ejemplos (XML de MS Project): EDT de
hasta 4 niveles, tareas sin código EDT, hitos con peso, todos los métodos de
medición, empresas y % de reparto. Se genera el libro, se ingresa el BAC, se
actualiza en varias fechas de corte (con el AC de cada mes) y se verifican las
reglas contables. La cantidad de proyectos se puede aumentar con la variable de
entorno EVM_SEMILLAS (p. ej. 50).
"""
import os
import random
from datetime import date, timedelta

import pytest
from openpyxl import load_workbook

from evm.cli import main
from evm.config import load_config
from evm.reader_xml import read_xml
from evm.update import book_totals, indicators_for, read_book
from evm.wbs import build_report
from evm.workbook import FIRST_MONTH_COL, read_meta, read_structure

from builders import assert_report_invariants, set_bac


SEEDS = int(os.environ.get("EVM_SEMILLAS", "1"))
STATUSES = [date(2026, 3, 31), date(2027, 6, 30), date(2028, 12, 31)]
METHODS = ["0/100", "50/50", "20/80", "Hitos ponderados", "% Completado", "% Físico completado",
           "Unidades completadas", "Avance lineal (LOE)", "Por definir", ""]
COMPANIES = ["EmpresaA", "EmpresaB", "Ambas", "", "", "Otra"]
START = date(2026, 1, 5)


def random_tree(gen, rng: random.Random, depth: int = 1, max_depth: int = 4) -> list:
    """Ramas aleatorias de tareas del generador de ejemplos."""
    nodes = []
    for _ in range(rng.randint(1, 4 if depth == 1 else 3)):
        t = gen.T(f"Tarea {depth}-{rng.randint(1, 9999)}", depth)
        if rng.random() < 0.25:
            t.company = rng.choice(COMPANIES)
            if t.company == "Ambas" and rng.random() < 0.6:
                t.split = rng.choice([0, 25, 60, 100, 150])
        if depth > 1 and rng.random() < 0.12:
            t.code = ""                                           # sin código EDT
        if depth < max_depth and rng.random() < 0.55:
            t.children = random_tree(gen, rng, depth + 1, max_depth)
            if rng.random() < 0.5:
                t.method = rng.choice(METHODS)
            if t.method == "Hitos ponderados" or rng.random() < 0.15:
                s = gen.add_wd(START + timedelta(days=rng.randint(0, 700)), 1)
                t.children.append(gen.T("Hito", depth + 1, s, s, 0.0, milestone=True,
                                        weight=rng.choice([None, 0, 10, 30]), code=""))
        else:
            s = gen.add_wd(START + timedelta(days=rng.randint(0, 900)), 1)
            if rng.random() < 0.08:
                t.start = t.finish = s
                t.milestone = True
                t.cost = rng.choice([0.0, 5000.0])
            else:
                t.start, t.finish = s, gen.add_wd(s, rng.choice([1, 3, 20, 45, 130, 400]))
                t.cost = rng.choice([0.0, 999.99, 12_345.67, 250_000.0, 3_000_000.0])
            t.method = rng.choice(METHODS)
        nodes.append(t)
    return nodes


def write_project(gen, seed: int, path, status: date | None = None):
    """Escribe el proyecto de la semilla (siempre el mismo) con avance a la fecha indicada."""
    plan = gen.T("Planificación interna", 1, code="INTERNO", children=random_tree(gen, random.Random(seed)))
    gen.write([plan], path, status)
    return path


@pytest.fixture(scope="module")
def configs(tmp_path_factory):
    d = tmp_path_factory.mktemp("config_aleatorio")
    base = ('[campos]\nmetodo = "Método EVM"\npeso_hito = "Peso hito"\n'
            '[horizonte]\nmes_inicio = "2026-01"\n[exportacion]\nplanificacion = "INTERNO"\n')
    single, multi = d / "una_empresa.toml", d / "multiempresa.toml"
    single.write_text(base, encoding="utf-8")
    multi.write_text(base + "[multiempresa]\nactivar = true\n", encoding="utf-8")
    return {False: single, True: multi}


def cum(series: dict, row: int, idx: int) -> float:
    return sum(v or 0.0 for v in series[row][: idx + 1])


def load_ac(path, idx: int, value: float = 1_000.0) -> None:
    """AC del mes `idx` en todas las filas de ingreso (hoja AC o las de cada empresa), como se ingresa a mano."""
    wb = load_workbook(path)
    meta = read_meta(wb)
    sheets = [f"AC-{meta['empresa_a']}", f"AC-{meta['empresa_b']}"] if meta.get("multiempresa") else ["AC"]
    for st in read_structure(wb):
        if st.is_input:
            for sheet in sheets:
                wb[sheet].cell(st.row, FIRST_MONTH_COL + idx, value)
    wb.save(path)


@pytest.mark.parametrize("multi", [False, True], ids=["una_empresa", "multiempresa"])
@pytest.mark.parametrize("seed", range(SEEDS))
def test_random_project_invariants(tmp_path, generator, configs, seed, multi):
    cfg_path = configs[multi]
    cfg = load_config(cfg_path)
    base = write_project(generator, seed, tmp_path / "base.xml")

    # Informe: el % planificado de cada fila (y empresa) está entre 0 y 100 % y los subtotales suman a sus hijos
    rep = build_report(read_xml(base, cfg), cfg)
    assert_report_invariants(rep)

    book = tmp_path / "EVM.xlsx"
    assert main(["generar", "--origen", str(base), "--config", str(cfg_path), "--salida", str(book)]) == 0
    set_bac(book)
    for status in STATUSES:
        xml = write_project(generator, seed, tmp_path / f"{status:%Y-%m}.xml", status)
        assert main(["actualizar", "--origen", str(xml), "--config", str(cfg_path), "--libro", str(book),
                     "--sin-respaldo"]) == 0, (seed, status)
        b = read_book(book)
        idx = b.horizon.index(status)
        load_ac(book, idx)
        for s in b.input_rows:
            ev = cum(b.ev, s.row, idx)
            assert -0.05 <= ev <= b.bac[s.row] + 0.05, (seed, status, s.code, ev, b.bac[s.row])
            if multi:
                for i in range(idx + 1):
                    parts = (b.ev_a[s.row][i] or 0.0) + (b.ev_b[s.row][i] or 0.0)
                    assert (b.ev[s.row][i] or 0.0) == pytest.approx(parts, abs=0.011), (seed, s.code, i)

    # Regenerar con --migrar-desde conserva el BAC y el historial de EV y AC
    new = tmp_path / "EVM_v2.xlsx"
    assert main(["generar", "--origen", str(base), "--config", str(cfg_path), "--salida", str(new),
                 "--migrar-desde", str(book)]) == 0
    before, after = book_totals(read_book(book)), book_totals(read_book(new))
    assert after[3] == pytest.approx(before[3]), seed
    for k in (1, 2):
        assert sum(after[k]) == pytest.approx(sum(before[k]), abs=0.05), (seed, k)


def test_random_project_recalculates_like_python(tmp_path, generator, configs, recalc):
    """Proyecto con fila de otras tareas, multi-empresa: Excel (LibreOffice) y Python calculan lo mismo."""
    cfg_path = configs[True]
    base = write_project(generator, 0, tmp_path / "base.xml")
    book = tmp_path / "EVM.xlsx"
    assert main(["generar", "--origen", str(base), "--config", str(cfg_path), "--salida", str(book)]) == 0
    assert any(s.kind == "residual" for s in read_book(book).structure)
    set_bac(book, 250_000.0)
    for status in STATUSES[:2]:
        xml = write_project(generator, 0, tmp_path / f"{status:%Y-%m}.xml", status)
        assert main(["actualizar", "--origen", str(xml), "--config", str(cfg_path), "--libro", str(book),
                     "--sin-respaldo"]) == 0
        load_ac(book, read_book(book).horizon.index(status))
    out = load_workbook(recalc(book), data_only=True)
    errors = [(ws.title, c.coordinate, c.value) for ws in out.worksheets for row in ws.iter_rows() for c in row
              if isinstance(c.value, str) and c.value.startswith("#") and c.value != "#N/A"]
    assert not errors, errors[:10]
    ind = indicators_for(read_book(book), STATUSES[1])
    summary = {out["Resumen"].cell(r, 1).value: out["Resumen"].cell(r, 2).value for r in range(1, 45)}
    assert summary["BAC - Presupuesto a la conclusión"] == pytest.approx(ind.bac, abs=0.01)
    assert summary["PV - Valor planificado acumulado"] == pytest.approx(ind.pv, abs=0.01)
    assert summary["EV - Valor ganado acumulado"] == pytest.approx(ind.ev, abs=0.01)
    assert summary["AC - Costo real acumulado"] == pytest.approx(ind.ac, abs=0.01)
    dash = out["Cuadro de Mando"]
    assert (dash["L10"].value, dash["Q10"].value) == (pytest.approx(ind.ev, abs=0.01), pytest.approx(ind.ac, abs=0.01))
    pv = out["PV"]
    total = next(r for r in range(5, pv.max_row + 1) if pv.cell(r, 1).value == "TOTAL")
    for kind in ("PV", "EV", "AC"):
        for c in range(4, 7 + 121):
            a, b = (out[f"{kind}-{name}"].cell(total, c).value or 0 for name in ("EmpresaA", "EmpresaB"))
            assert (out[kind].cell(total, c).value or 0) == pytest.approx(a + b, abs=0.01), (kind, c)
