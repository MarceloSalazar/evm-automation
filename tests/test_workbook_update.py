from datetime import date

import pytest
from openpyxl import load_workbook

from evm.cli import main
from evm.config import load_config
from evm.reader_xml import read_xml
from evm.update import book_totals, read_book, run_update
from evm.workbook import FIRST_MONTH_COL, FIRST_ROW, HEADER_ROW


@pytest.fixture()
def book(tmp_path, example_files):
    out = tmp_path / "EVM.xlsx"
    assert main(["generar", "--origen", str(example_files["base"]), "--config", str(example_files["config"]),
                 "--salida", str(out)]) == 0
    return out


def test_layout_matches_requested_format(book):
    wb = load_workbook(book)
    assert wb.sheetnames[:4] == ["Resumen", "PV", "EV", "AC"]
    ws = wb["PV"]
    heads = [ws.cell(HEADER_ROW, c).value for c in range(1, 7)]
    assert heads == ["Código EDT", "Actividad", "Método de medición", "BAC (Presupuesto)",
                     "Fecha inicio Plan (PV)", "Fecha fin Plan (PV)"]
    assert ws.cell(3, FIRST_MONTH_COL).value == 1                       # G3 = mes 1
    assert ws.cell(3, FIRST_MONTH_COL + 119).value == 120               # DV3 = mes 120
    assert ws.cell(HEADER_ROW, FIRST_MONTH_COL).value.date() == date(2026, 1, 1)
    assert ws.cell(FIRST_ROW, 1).value == "1.0.0"
    assert str(ws.cell(FIRST_ROW, 4).value).startswith("=SUM(")         # fila resumen = fórmula
    leaf = next(r for r in range(FIRST_ROW, ws.max_row) if ws.cell(r, 1).value == "1.1.1")
    assert ws.cell(leaf, 3).value == "Avance lineal (LOE)"
    assert isinstance(ws.cell(leaf, 4).value, (int, float))               # fila de ingreso = valor
    assert ws.data_validations.dataValidation
    assert wb["_meta"].sheet_state == "hidden"
    assert "FechaCorte" in wb.defined_names


def test_cli_refuses_to_overwrite(book, example_files, capsys):
    rc = main(["generar", "--origen", str(example_files["base"]), "--salida", str(book)])
    assert rc == 2
    assert "ya existe" in capsys.readouterr().err


def test_update_cycle(book, example_files):
    cfg = load_config(example_files["config"])
    jun = read_xml(example_files["jun"], cfg)
    res = run_update(jun, cfg, book, backup=False)
    assert res.status_date == date(2027, 6, 30) and res.month_index == 17
    b = read_book(book)
    pv, ev, ac, bac = book_totals(b)
    assert bac == pytest.approx(44_210_000)
    ev_jun, ac_jun = sum(ev), sum(ac)
    assert ev_jun > 0 and ac_jun > 0
    # AC = ACWP de MS Project del proyecto completo (suma de las tareas de nivel 1)
    assert ac_jun == pytest.approx(sum(t.acwp for t in jun.tasks if t.outline_level == 1), abs=0.05)
    # Repetir el mismo mes no duplica valores
    run_update(jun, cfg, book, backup=False)
    b2 = read_book(book)
    assert sum(book_totals(b2)[1]) == pytest.approx(ev_jun)
    # Mes siguiente: el valor del mes es la diferencia con el acumulado anterior
    jul = read_xml(example_files["jul"], cfg)
    run_update(jul, cfg, book, backup=False)
    b3 = read_book(book)
    _, ev3, ac3, _ = book_totals(b3)
    assert ev3[17] == pytest.approx(ev_jun)
    assert ev3[18] > 0
    assert b3.status_date == date(2027, 7, 31)


def test_update_refuses_when_later_months_loaded(book, example_files):
    cfg = load_config(example_files["config"])
    run_update(read_xml(example_files["jul"], cfg), cfg, book, backup=False)
    with pytest.raises(ValueError, match="posteriores"):
        run_update(read_xml(example_files["jun"], cfg), cfg, book, backup=False)


def test_update_only_ac_and_backup(book, example_files):
    cfg = load_config(example_files["config"])
    res = run_update(read_xml(example_files["jun"], cfg), cfg, book, only="ac")
    assert res.backup is not None and res.backup.exists()
    b = read_book(book)
    _, ev, ac, _ = book_totals(b)
    assert sum(ev) == 0 and sum(ac) > 0


def test_ev_matches_methods(book, example_files):
    cfg = load_config(example_files["config"])
    res = run_update(read_xml(example_files["jun"], cfg), cfg, book, backup=False)
    rows = {r.code: r for r in res.rows}
    # 0/100 con 85 % de avance no gana valor
    assert rows["1.2.1"].ev_cum == 0
    # LOE: EV = PV acumulado de la fila
    b = read_book(book)
    loe_row = next(s for s in b.input_rows if s.code == "1.1.1")
    assert rows["1.1.1"].ev_cum == pytest.approx(sum(b.pv[loe_row.row][:18]), abs=0.01)


def test_migrate_keeps_ev_ac(book, tmp_path, example_files):
    cfg = load_config(example_files["config"])
    run_update(read_xml(example_files["jun"], cfg), cfg, book, backup=False)
    new = tmp_path / "EVM_v2.xlsx"
    assert main(["generar", "--origen", str(example_files["base"]), "--config", str(example_files["config"]),
                 "--salida", str(new), "--migrar-desde", str(book)]) == 0
    a, b = read_book(book), read_book(new)
    assert sum(book_totals(a)[1]) == pytest.approx(sum(book_totals(b)[1]))
    assert sum(book_totals(a)[2]) == pytest.approx(sum(book_totals(b)[2]))
    assert b.status_date == date(2027, 6, 30)


def test_curve_png(book, tmp_path, example_files):
    cfg = load_config(example_files["config"])
    run_update(read_xml(example_files["jun"], cfg), cfg, book, backup=False)
    png = tmp_path / "curva.png"
    assert main(["curva", "--libro", str(book), "--salida", str(png)]) == 0
    assert png.stat().st_size > 10_000


def test_formulas_recalculate_without_errors(book, example_files, recalc):
    """Recalcula con LibreOffice y compara con el cálculo en Python."""
    from evm.chart import indicators_for
    cfg = load_config(example_files["config"])
    run_update(read_xml(example_files["jun"], cfg), cfg, book, backup=False)
    out = recalc(book)
    wb = load_workbook(out, data_only=True)
    errors = [(ws.title, c.coordinate, c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row
              if isinstance(c.value, str) and c.value.startswith("#") and c.value != "#N/A"]
    assert not errors, errors[:10]
    ind = indicators_for(read_book(book), date(2027, 6, 30))
    ws = wb["Resumen"]
    assert ws["B16"].value == pytest.approx(ind.bac, abs=0.01)
    assert ws["B17"].value == pytest.approx(ind.pv, abs=0.01)
    assert ws["B18"].value == pytest.approx(ind.ev, abs=0.01)
    assert ws["B19"].value == pytest.approx(ind.ac, abs=0.01)
    assert ws["B22"].value == pytest.approx(ind.spi, abs=1e-6)
    assert ws["B23"].value == pytest.approx(ind.cpi, abs=1e-6)
    assert ws["B35"].value == pytest.approx(ind.es, abs=1e-6)
    pv = wb["PV"]
    total_row = next(r for r in range(FIRST_ROW, pv.max_row + 1) if pv.cell(r, 1).value == "TOTAL")
    assert pv.cell(total_row, 4).value == pytest.approx(44_210_000, abs=0.01)
    curve = wb["Curva S"]
    assert curve.cell(5 + 17, 6).value == pytest.approx(ind.ev, abs=0.01)   # EV acumulado en el mes 18
    assert curve.cell(5 + 18, 6).value == "#N/A"                            # después del corte
    assert curve.cell(5, 6).value == "#N/A"                                 # antes del primer dato
