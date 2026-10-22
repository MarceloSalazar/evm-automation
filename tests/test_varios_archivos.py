"""Varios archivos de MS Project en un mismo libro (secciones [[proyectos]]).

El programa de ejemplo (códigos EDT 1.0.0 a 4.0.0, planificación INTERNO) se combina
con el segundo proyecto de ejemplo (código EDT 9.0.0, sin planificaciones). Los dos
archivos numeran sus tareas desde el UniqueID 1, de modo que los UniqueID se repiten.
"""
from datetime import date

import pytest
from openpyxl import load_workbook

from evm import cli
from evm.cli import main
from evm.config import ConfigError, load_config
from evm.update import indicators_for, read_book
from evm.workbook import split_key

pytestmark = pytest.mark.modo("costo")

COMMON = ('[campos]\nmetodo = "Método EVM"\npeso_hito = "Peso hito"\n\n[horizonte]\nmes_inicio = "2026-01"\n\n'
          '[exportacion]\nplanificacion = "INTERNO"\n')
MULTI_COMPANY = '\n[multiempresa]\nactivar = true\nac_desde_project = true\n'


@pytest.fixture(scope="module")
def xml(tmp_path_factory, generator):
    """XML de los dos proyectos: línea base y avance a junio y julio de 2027."""
    d = tmp_path_factory.mktemp("varios")
    out = {"dir": d}
    for project, build in ((1, generator.build_plans), (2, generator.build_second)):
        for label, status in (("base", None), ("jun", date(2027, 6, 30)), ("jul", date(2027, 7, 31))):
            path = d / f"p{project}_{label}.xml"
            generator.write(build(), path, status, project)
            out[(project, label)] = path
    return out


def projects_toml(entries) -> str:
    """Secciones [[proyectos]]: (nombre, archivo, texto adicional)."""
    return "".join(f'\n[[proyectos]]\nnombre = "{n}"\norigen = "{f}"\n{extra}' for n, f, extra in entries)


def write_config(xml, name: str, entries=None, extra: str = "", common: str = COMMON):
    """Archivo de configuración en la carpeta de los XML (las rutas de [[proyectos]] son relativas a ella)."""
    if entries is None:
        entries = [("P1", "p1_base.xml", ""), ("P2", "p2_base.xml", 'planificacion = ""\n')]
    path = xml["dir"] / name
    path.write_text(common + extra + projects_toml(entries), encoding="utf-8")
    return path


def generate(cfg, out, *extra) -> None:
    assert main(["generar", "--config", str(cfg), "--salida", str(out), "--sobrescribir", *extra]) == 0


def update(cfg, book, *extra) -> None:
    assert main(["actualizar", "--config", str(cfg), "--libro", str(book), "--sin-respaldo", *extra]) == 0


def by_code(path) -> dict:
    """Valores de cada fila de ingreso por código EDT: BAC, PV, EV, AC y (multi-empresa) los de cada empresa."""
    b = read_book(path)
    out = {}
    for s in b.input_rows:
        values = [b.bac[s.row], b.pv[s.row], b.ev[s.row], b.ac[s.row]]
        if b.multi:
            values += [b.bac_a[s.row], b.pv_a[s.row], b.ev_a[s.row], b.ev_b[s.row], b.ac_a[s.row], b.ac_b[s.row]]
        out[s.code] = values
    return out


def single_books(xml, tmp_path, extra: str = "", months=("jun", "jul")):
    """Libros de cada proyecto por separado (un solo archivo), con los mismos meses cargados."""
    cfg1 = write_config(xml, "solo1.toml", [], extra)
    cfg2 = write_config(xml, "solo2.toml", [], extra, COMMON.replace('"INTERNO"', '""'))
    books = []
    for project, cfg in ((1, cfg1), (2, cfg2)):
        book = tmp_path / f"solo{project}.xlsx"
        generate(cfg, book, "--origen", str(xml[(project, "base")]))
        for m in months:
            update(cfg, book, "--origen", str(xml[(project, m)]))
        books.append(book)
    return books


def combined_book(xml, tmp_path, extra: str = "", months=("jun", "jul"), name="EVM.xlsx"):
    cfg = write_config(xml, "varios.toml", extra=extra)
    book = tmp_path / name
    generate(cfg, book)
    for m in months:
        update(cfg, book, "--origen", f"P1={xml[(1, m)]}", "--origen", f"P2={xml[(2, m)]}")
    return cfg, book


# ----------------------------------------------------------------------------
# Configuración
# ----------------------------------------------------------------------------
def test_config_projects(tmp_path):
    (tmp_path / "datos").mkdir()
    path = tmp_path / "c.toml"
    path.write_text(COMMON + projects_toml([
        ("P1", "datos/p1.mpp", ""),
        ("P2", str(tmp_path / "p2.xml").replace("\\", "/"), 'planificacion = ""\nlinea_base = 1\n')]),
        encoding="utf-8")
    cfg = load_config(path)
    assert [p.name for p in cfg.projects] == ["P1", "P2"]
    assert cfg.projects[0].origin == str(tmp_path / "datos" / "p1.mpp")      # relativa a la configuración
    assert cfg.project_config("P1").plan == "INTERNO" and cfg.project_config("P1").baseline is None
    assert cfg.project_config("P2").plan == "" and cfg.project_config("P2").baseline == 1
    assert cfg.project_config("otro") is cfg


@pytest.mark.parametrize("text, message", [
    ('[proyectos]\nnombre = "P1"\norigen = "a.xml"\n', "dos corchetes"),
    (projects_toml([("P1", "a.xml", ""), ("p1", "b.xml", "")]), "repetido"),
    (projects_toml([("", "a.xml", "")]), "nombre"),
    (projects_toml([("P1", "", "")]), "falta 'origen'"),
    (projects_toml([("P1", "activo", "")]), "activo"),
    (projects_toml([("P:1", "a.xml", "")]), "sin ':'"),
    (projects_toml([("P1", "a.xml", "linea_base = 11\n")]), "linea_base"),
])
def test_config_project_errors(tmp_path, text, message):
    path = tmp_path / "c.toml"
    path.write_text(text, encoding="utf-8")
    with pytest.raises(ConfigError, match=message):
        load_config(path)


# ----------------------------------------------------------------------------
# generar
# ----------------------------------------------------------------------------
def test_generate_combines_files(xml, tmp_path, capsys):
    cfg = write_config(xml, "varios.toml")
    book = tmp_path / "EVM.xlsx"
    generate(cfg, book)
    out = capsys.readouterr().out
    assert "Leyendo P1:" in out and "Leyendo P2:" in out and "BAC total: 49,010,000.00" in out

    b = read_book(book)
    assert sum(b.bac[s.row] for s in b.input_rows) == pytest.approx(44_210_000 + 4_800_000)
    codes = [s.code for s in b.structure]
    assert codes[0] == "1.0.0" and "9.0.0" in codes and codes.index("9.0.0") > codes.index("4.0.0")
    # Claves con el nombre del archivo: el UniqueID se repite entre archivos
    keys = [split_key(s.key) for s in b.structure]
    assert {k[0] for k in keys} == {"P1", "P2"}
    assert {k[1] for k in keys if k[0] == "P1"} & {k[1] for k in keys if k[0] == "P2"}
    assert all(s.key.startswith("P2:") for s in b.structure if s.code.startswith("9."))

    meta = b.meta
    assert meta["proyectos"] == "P1;P2"
    assert meta["proyecto.P1"] == "Programa de ejemplo EVM" and meta["planificacion.P1"] == "INTERNO"
    assert meta["proyecto.P2"].startswith("Accesos y servicios") and not meta["planificacion.P2"]
    wb = load_workbook(book)
    assert wb["Resumen"]["A2"].value == "Programa de ejemplo EVM + Accesos y servicios (proyecto 2 de ejemplo)"
    assert wb["Resumen"]["B14"].value == "P1: INTERNO  |  P2: todas"
    assert wb["Cuadro de Mando"]["B2"].value == wb["Resumen"]["A2"].value
    details = [r[3] for r in wb["Validación"].iter_rows(min_row=4, values_only=True) if r[3]]
    assert any(d.startswith("[P1] Planificación 'INTERNO'") for d in details)
    assert any(d.startswith("[P2] ") for d in details)


def test_combined_equals_separate_books(xml, tmp_path):
    """El libro combinado tiene, fila por fila, los mismos BAC, PV, EV y AC que el libro de cada archivo."""
    p1, p2 = (by_code(b) for b in single_books(xml, tmp_path))
    assert not set(p1) & set(p2)
    _cfg, book = combined_book(xml, tmp_path)
    assert by_code(book) == {**p1, **p2}
    assert read_book(book).status_date == date(2027, 7, 31)


def test_combined_multicompany_equals_separate_books(xml, tmp_path):
    """Multi-empresa: la parte de cada empresa se calcula con las tareas de cada archivo (UniqueID repetidos)."""
    p1, p2 = (by_code(b) for b in single_books(xml, tmp_path, MULTI_COMPANY, months=("jun",)))
    _cfg, book = combined_book(xml, tmp_path, MULTI_COMPANY, months=("jun",))
    combined = by_code(book)
    assert combined == {**p1, **p2}
    assert any(v[4] != v[0] for v in combined.values())    # hay filas repartidas entre empresas


def test_title_order_and_start_month(xml, tmp_path):
    entries = [("P2", "p2_base.xml", 'planificacion = ""\n'), ("P1", "p1_base.xml", "")]
    common = ('[campos]\nmetodo = "Método EVM"\n\n'                 # sin mes de inicio
              '[exportacion]\nplanificacion = "INTERNO"\ntitulo = "Programa"\n')
    cfg = write_config(xml, "orden.toml", entries, common=common)
    book = tmp_path / "orden.xlsx"
    generate(cfg, book)
    b = read_book(book)
    assert [s.code for s in b.structure if s.level == 1] == ["1.0.0", "2.0.0", "3.0.0", "4.0.0", "9.0.0"]
    assert b.horizon.start == date(2026, 1, 1)              # el primer mes con PV de los dos archivos
    assert b.meta["proyectos"] == "P2;P1" and b.meta["titulo"] == "Programa"


def test_duplicate_code_in_two_files_stops(xml, tmp_path, capsys):
    cfg = write_config(xml, "dup.toml", [("P1", "p1_base.xml", ""), ("P1B", "p1_base.xml", "")])
    out = tmp_path / "dup.xlsx"
    assert main(["generar", "--config", str(cfg), "--salida", str(out)]) == 2
    assert "está en los archivos 'P1' y 'P1B'" in capsys.readouterr().err
    assert not out.exists()


def test_each_file_is_read_with_its_own_settings(xml, tmp_path, monkeypatch):
    entries = [("P1", "p1_base.xml", ""), ("P2", "p2_base.xml", 'planificacion = ""\nlinea_base = 0\n')]
    cfg = write_config(xml, "propia.toml", entries)
    seen = []
    original = cli.open_project

    def spy(origen, cfg, motor, visible=False, status=None):
        seen.append((origen, cfg.plan, cfg.baseline))
        return original(origen, cfg, motor, visible, status)

    monkeypatch.setattr(cli, "open_project", spy)
    generate(cfg, tmp_path / "propia.xlsx")
    assert seen == [(str(xml[(1, "base")]), "INTERNO", None), (str(xml[(2, "base")]), "", 0)]


@pytest.mark.parametrize("args, message", [
    (["--origen", "x.xml"], "NOMBRE=RUTA"),
    (["--origen", "P9=x.xml"], "NOMBRE=RUTA"),
    (["--origen", "P1=activo"], "no se admite 'activo'"),
    (["--planificacion", "CLIENTE"], "--planificacion"),
    (["--linea-base", "1"], "--linea-base"),
])
def test_generate_option_errors(xml, tmp_path, capsys, args, message):
    cfg = write_config(xml, "varios.toml")
    assert main(["generar", "--config", str(cfg), "--salida", str(tmp_path / "e.xlsx"), *args]) == 2
    assert message in capsys.readouterr().err


def test_generate_requires_output_and_origin(xml, tmp_path, capsys):
    cfg = write_config(xml, "varios.toml")
    assert main(["generar", "--config", str(cfg)]) == 2
    assert "--salida" in capsys.readouterr().err
    single = write_config(xml, "solo.toml", [])
    assert main(["generar", "--config", str(single), "--salida", str(tmp_path / "e.xlsx")]) == 2
    assert "--origen" in capsys.readouterr().err
    assert main(["generar", "--config", str(single), "--salida", str(tmp_path / "e.xlsx"),
                 "--origen", str(xml[(1, "base")]), "--origen", str(xml[(2, "base")])]) == 2
    assert "[[proyectos]]" in capsys.readouterr().err


# ----------------------------------------------------------------------------
# actualizar
# ----------------------------------------------------------------------------
def test_update_requires_the_files_of_the_book(xml, tmp_path, capsys):
    cfg = write_config(xml, "varios.toml")
    book = tmp_path / "EVM.xlsx"
    generate(cfg, book)
    single = write_config(xml, "solo.toml", [])
    other = write_config(xml, "otro.toml", [("P1", "p1_jun.xml", ""), ("P3", "p2_jun.xml", 'planificacion = ""\n')])
    cases = [(single, ["--origen", str(xml[(1, "jun")])], "varios archivos de MS Project (P1, P2)"),
             (other, [], "no coinciden con los del libro")]
    for config, args, message in cases:
        assert main(["actualizar", "--config", str(config), "--libro", str(book), *args]) == 2
        assert message in capsys.readouterr().err
    # Libro de un solo archivo y configuración con [[proyectos]]
    single_book = tmp_path / "solo.xlsx"
    generate(single, single_book, "--origen", str(xml[(1, "base")]))
    assert main(["actualizar", "--config", str(cfg), "--libro", str(single_book)]) == 2
    assert "un solo archivo" in capsys.readouterr().err
    assert read_book(book).status_date is None and read_book(single_book).status_date is None


def test_update_status_dates(xml, tmp_path, capsys):
    cfg = write_config(xml, "varios.toml")
    book = tmp_path / "EVM.xlsx"
    generate(cfg, book)
    mixed = ["--origen", f"P1={xml[(1, 'jun')]}", "--origen", f"P2={xml[(2, 'jul')]}"]
    assert main(["actualizar", "--config", str(cfg), "--libro", str(book), *mixed]) == 2
    err = capsys.readouterr().err
    assert "fechas de estado distintas (P1: 30/06/2027, P2: 31/07/2027)" in err
    assert main(["actualizar", "--config", str(cfg), "--libro", str(book), "--sin-respaldo",
                 "--fecha-corte", "2027-06-30", *mixed]) == 0
    out = capsys.readouterr().out
    assert "[P2] La fecha de estado de MS Project (31/07/2027) difiere de la fecha de corte (30/06/2027)" in out
    assert "[P1] La fecha de estado" not in out
    assert read_book(book).status_date == date(2027, 6, 30)


def test_update_checks_ac_per_file(xml, tmp_path, capsys):
    """El AC de cada archivo coincide con el CRTR de su planificación (P1) o de su proyecto (P2)."""
    cfg, book = combined_book(xml, tmp_path, months=("jun",))
    out = capsys.readouterr().out
    assert "difiere del costo real" not in out
    b = read_book(book)
    ac = {src: 0.0 for src in ("P1", "P2")}
    for s in b.input_rows:
        ac[split_key(s.key)[0]] += sum(v or 0.0 for v in b.ac[s.row])
    from evm.reader_xml import read_xml
    config = load_config(cfg)
    p1 = read_xml(xml[(1, "jun")], config.project_config("P1"))
    p2 = read_xml(xml[(2, "jun")], config.project_config("P2"))
    assert ac["P1"] == pytest.approx(next(t.acwp for t in p1.tasks if t.wbs == "INTERNO"), abs=0.05)
    assert ac["P2"] == pytest.approx(p2.summary_acwp, abs=0.05)


def test_update_analysis_of_deviations(xml, tmp_path):
    """La hoja Análisis de desviaciones asocia cada paquete con su archivo (UniqueID repetidos)."""
    from evm.analysis import A_BAC, SHEET_ANALYSIS, read_entries
    _cfg, book = combined_book(xml, tmp_path, months=("jun",))
    b = read_book(book)
    rows = {s.key: (s.code, s.row) for s in b.input_rows}
    ws = load_workbook(book)[SHEET_ANALYSIS]
    found = read_entries(ws)
    assert {split_key(e.key)[0] for e in found} == {"P1", "P2"}
    for i, e in enumerate(found, start=5):
        code, data_row = rows[e.key]
        assert e.code == code and ws.cell(i, A_BAC).value == f"='PV'!$D${data_row}"


def test_update_indicators_recalculate(xml, tmp_path, recalc):
    _cfg, book = combined_book(xml, tmp_path, months=("jun",))
    out = load_workbook(recalc(book), data_only=True)
    errors = [(ws.title, c.coordinate, c.value) for ws in out.worksheets for row in ws.iter_rows() for c in row
              if isinstance(c.value, str) and c.value.startswith("#") and c.value != "#N/A"]
    assert not errors, errors[:10]
    ind = indicators_for(read_book(book), date(2027, 6, 30))
    ws = out["Resumen"]
    for cell, value in (("B16", ind.bac), ("B17", ind.pv), ("B18", ind.ev), ("B19", ind.ac)):
        assert ws[cell].value == pytest.approx(value, abs=0.01), cell


# ----------------------------------------------------------------------------
# --migrar-desde
# ----------------------------------------------------------------------------
def validation(book) -> list[tuple]:
    return [r for r in load_workbook(book)["Validación"].iter_rows(min_row=4, values_only=True) if r[0]]


def test_migrate_combined_book_keeps_history(xml, tmp_path):
    cfg, book = combined_book(xml, tmp_path)
    new = tmp_path / "EVM_LB2.xlsx"
    generate(cfg, new, "--migrar-desde", str(book))
    assert by_code(new) == by_code(book)
    assert not [r for r in validation(new) if r[0] in ("Error", "Advertencia")]
    assert read_book(new).status_date == date(2027, 7, 31)


def test_migrate_without_one_file(xml, tmp_path):
    """Si se quita un archivo, su historial no se traslada y se informa como error."""
    _cfg, book = combined_book(xml, tmp_path, months=("jun",))
    only_p1 = write_config(xml, "solo_p1.toml", [("P1", "p1_base.xml", "")])
    new = tmp_path / "EVM_P1.xlsx"
    generate(only_p1, new, "--migrar-desde", str(book))
    before = by_code(book)
    after = by_code(new)
    assert set(after) == {c for c in before if not c.startswith("9.")}
    assert all(after[c] == before[c] for c in after)
    rows = validation(new)
    assert any("El archivo 'P2' del libro anterior no forma parte del libro nuevo" in r[3] for r in rows)
    assert {r[1] for r in rows if r[0] == "Error"} == {c for c, v in before.items()
                                                      if c.startswith("9.") and any(v[2] + v[3])}


def test_migrate_single_file_book_by_code(xml, tmp_path):
    single = write_config(xml, "solo.toml", [])
    old = tmp_path / "solo.xlsx"
    generate(single, old, "--origen", str(xml[(1, "base")]))
    update(single, old, "--origen", str(xml[(1, "jun")]))
    cfg = write_config(xml, "varios.toml")
    new = tmp_path / "EVM.xlsx"
    generate(cfg, new, "--migrar-desde", str(old))
    before, after = by_code(old), by_code(new)
    assert all(after[c] == before[c] for c in before)
    assert all(not any(v[2] + v[3]) for c, v in after.items() if c.startswith("9."))
    assert any("es de un solo archivo de MS Project" in r[3] for r in validation(new))
