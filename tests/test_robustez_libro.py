"""Robustez del libro: actualizar, migrar, escritura, Cuadro de Mando, CLI y configuración.

Incluye los casos de los problemas detectados en la revisión del código, ya corregidos, para que no
vuelvan a aparecer.
"""
import re
import shutil
import sys
import zipfile
from datetime import date, datetime

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.writer.excel import ExcelWriter

from evm import dashboard
from evm.cli import main
from evm.config import Config, ConfigError, load_config
from evm.migrate import plan_migration, prior_from_book
from evm.model import ProjectData, Task
from evm.update import book_totals, read_book, run_update
from evm.wbs import build_report
from evm.workbook import COL_BAC, FIRST_MONTH_COL, FIRST_ROW, WorkbookFormatError, read_structure, write_workbook


S, F = datetime(2026, 1, 5, 8), datetime(2026, 3, 27, 17)
ERROR = "error"     # resultado de una lectura rechazada con WorkbookFormatError


# ----------------------------------------------------------------------------
# Utilidades
# ----------------------------------------------------------------------------
def T(uid, level, wbs, name, **kw):
    return Task(uid=uid, id=kw.pop("id", uid), name=name, outline_level=level, wbs=wbs,
                baseline_start=kw.pop("bs", S), baseline_finish=kw.pop("bf", F), has_baseline=True,
                start=S, finish=F, **kw)


def P(tasks, name="Programa", status=None, **kw):
    return ProjectData(name=name, source="p.xml", backend="xml", tasks=tasks, status_date=status, **kw)


def make_book(path, project, cfg=None, prior=None, bac=None):
    """Genera el libro e ingresa el BAC de las filas indicadas ({código: BAC}; en multi-empresa, el BAC es de la
    empresa A)."""
    cfg = cfg or Config()
    rep = build_report(project, cfg)
    if prior is not None:
        rep.issues.extend(plan_migration(rep, prior, cfg.max_level))
    write_workbook(rep, cfg, path, prior)
    if bac:
        wb = load_workbook(path)
        rows = {s.code: s.row for s in read_structure(wb)}
        sheet = f"PV-{cfg.company_a}" if cfg.multi_enabled else "PV"
        for code, value in bac.items():
            wb[sheet].cell(rows[code], COL_BAC, value)
        wb.save(path)
    return rep


def base_tasks():
    return [T(1, 1, "1", "Fase"), T(2, 2, "1.1", "B"), T(3, 2, "1.2", "C")]


BAC = {"1.1.0": 100, "1.2.0": 200}


def progress_tasks():
    """Avance: B 50 % (EV 50), C 25 % (EV 50)."""
    tasks = base_tasks()
    tasks[1].percent_complete = 50
    tasks[2].percent_complete = 25
    return tasks


def row_of(ws, code):
    """Fila de la hoja cuyo código EDT (columna A) es `code`."""
    return next(r for r in range(FIRST_ROW, ws.max_row + 1) if ws.cell(r, 1).value == code)


def book_row(book, code):
    return next(s.row for s in book.structure if s.code == code)


def tamper(path, fn):
    """Modifica el libro como lo haría el usuario en Excel: abre, aplica `fn(wb)` y guarda."""
    wb = load_workbook(path)
    fn(wb)
    wb.save(path)
    return path


def run_cli(capsys, *args):
    """Ejecuta la línea de comandos y devuelve (código de salida, salida estándar, salida de error)."""
    capsys.readouterr()
    rc = main([str(a) for a in args])
    captured = capsys.readouterr()
    return rc, captured.out, captured.err


def backups(folder):
    return sorted(p.name for p in folder.glob("*.respaldo-*"))


SMALL = Config(months=24)      # horizonte corto: libros más rápidos de leer y guardar


@pytest.fixture(scope="session")
def small_template(tmp_path_factory):
    """Libro de base_tasks() generado una sola vez (se copia en cada prueba)."""
    path = tmp_path_factory.mktemp("plantilla_chica") / "b.xlsx"
    make_book(path, P(base_tasks()), SMALL, bac=BAC)
    return path


@pytest.fixture()
def small_book(tmp_path, small_template):
    path = tmp_path / "b.xlsx"
    shutil.copy(small_template, path)
    return path


@pytest.fixture(scope="session")
def example_template(tmp_path_factory, example_files):
    """Libro del proyecto de ejemplo generado una sola vez (se copia en cada prueba)."""
    out = tmp_path_factory.mktemp("plantilla") / "EVM.xlsx"
    assert main(["generar", "--origen", str(example_files["base"]), "--config", str(example_files["config"]),
                 "--salida", str(out)]) == 0
    return out


@pytest.fixture()
def example_book(tmp_path, example_template):
    path = tmp_path / "EVM.xlsx"
    shutil.copy(example_template, path)
    return path


# ----------------------------------------------------------------------------
# actualizar: cálculo y asociación de filas
# ----------------------------------------------------------------------------
def residual_tasks():
    """Fase con un gasto sin código EDT (fila 1.0.0-D) terminado y los paquetes con fila propia sin avance."""
    return [T(1, 1, "1", "Fase", company="EmpresaA"), T(2, 2, "1.1", "B"), T(3, 2, "1.2", "C"),
            T(4, 2, "", "Gasto sin código", company="EmpresaA", percent_complete=100)]


@pytest.mark.parametrize("multi", [False, True], ids=["una_empresa", "multiempresa"])
def test_residual_row_ev_uses_its_own_tasks(tmp_path, multi):
    cfg = Config(multi_enabled=multi, months=24)
    book = tmp_path / "b.xlsx"
    make_book(book, P(residual_tasks()), cfg, bac={"1.0.0-D": 100})
    res = run_update(P(residual_tasks()), cfg, book, date(2026, 1, 31), backup=False)
    residual = next(r for r in res.rows if r.code == "1.0.0-D")
    assert residual.bac == pytest.approx(100)
    assert residual.ev_cum == pytest.approx(100)          # su única tarea está terminada


def test_update_rejects_shifted_rows(small_book):
    def insert_row(wb):
        ws = wb["EV"]
        ws.insert_rows(6)                                   # como "Insertar fila" en Excel
        ws.cell(6, 1, "Nota")
    tamper(small_book, insert_row)
    before = small_book.read_bytes()
    with pytest.raises(WorkbookFormatError, match=r"hoja 'EV'.*fila 6.*'1\.1\.0'.*'Nota'"):
        run_update(P(progress_tasks()), Config(), small_book, date(2026, 1, 31), backup=False)
    assert small_book.read_bytes() == before
    with pytest.raises(WorkbookFormatError, match="filas insertadas, eliminadas u ordenadas"):
        prior_from_book(small_book, "Programa")            # tampoco se migra un historial desplazado


def test_update_rejects_deleted_rows_in_company_sheet(tmp_path):
    cfg = Config(multi_enabled=True, months=24)
    book = tmp_path / "multi.xlsx"
    make_book(book, P(base_tasks()), cfg)
    tamper(book, lambda wb: wb["AC-EmpresaA"].delete_rows(6))
    with pytest.raises(WorkbookFormatError, match="hoja 'AC-EmpresaA'"):
        run_update(P(progress_tasks()), cfg, book, date(2026, 1, 31), backup=False)


@pytest.mark.parametrize("with_residual", [False, True], ids=["sin_fila_D", "con_fila_D"])
def test_new_coded_package_without_row_warns(tmp_path, with_residual):
    tasks = base_tasks()
    if with_residual:
        tasks.append(T(4, 2, "", "Gasto sin código"))
    book = tmp_path / "b.xlsx"
    make_book(book, P(tasks), SMALL, bac={**BAC, "1.0.0-D": 100} if with_residual else BAC)
    now = [T(t.uid, t.outline_level, t.wbs, t.name) for t in tasks]
    now.append(T(9, 2, "1.3", "Paquete nuevo", percent_complete=50))
    res = run_update(P(now), Config(), book, date(2026, 1, 31), dry_run=True)
    [warning] = [w for w in res.warnings if "1.3 'Paquete nuevo'" in w]
    assert "no tiene fila en el libro" in warning and "--migrar-desde" in warning
    assert all(not r.ev_cum for r in res.rows)               # el avance del paquete nuevo no va a la fila -D


def test_overwritten_total_formula_is_reported(small_book):
    total = row_of(load_workbook(small_book)["AC"], "TOTAL")
    tamper(small_book, lambda wb: wb["AC"].cell(total, FIRST_MONTH_COL, 0))
    res = run_update(P(progress_tasks()), Config(), small_book, date(2026, 1, 31), backup=False)
    gaps = [w for w in res.warnings if "sin la fórmula del libro" in w]
    assert len(gaps) == 1 and f"AC!G{total}" in gaps[0]
    assert load_workbook(small_book)["AC"].cell(total, FIRST_MONTH_COL).value == 0   # no se modifica


def test_overwritten_company_sum_formula_is_reported(tmp_path):
    cfg = Config(multi_enabled=True, months=24)
    book = tmp_path / "multi.xlsx"
    make_book(book, P(base_tasks()), cfg)
    row = row_of(load_workbook(book)["EV"], "1.1.0")
    tamper(book, lambda wb: wb["EV"].cell(row, FIRST_MONTH_COL + 1, 5))      # suma de empresas pisada
    res = run_update(P(progress_tasks()), cfg, book, date(2026, 1, 31), dry_run=True)
    assert any(f"EV!H{row}" in w for w in res.warnings if "sin la fórmula del libro" in w)


def test_untouched_book_reports_no_formula_gaps(small_book):
    assert read_book(small_book).formula_gaps == []


def test_task_made_inactive_after_generation_is_reported(small_book):
    tasks = progress_tasks()
    tasks[2].active = False                                  # C se desactiva en MS Project
    res = run_update(P(tasks), Config(), small_book, date(2026, 1, 31), dry_run=True)
    assert {r.code for r in res.rows} == {"1.1.0"}
    assert any("1.2.0" in w and "inactiva" in w for w in res.warnings)


def renamed_tasks(shift=0):
    """Fase con dos paquetes y una tarea sin código (fila 1.0.0-D), con avance; UID desplazados en `shift`."""
    tasks = [T(1 + shift, 1, "1", "Fase"), T(2 + shift, 2, "1.1", "B"), T(3 + shift, 2, "1.2", "C"),
             T(4 + shift, 2, "", "Sin código")]
    tasks[1].percent_complete = 50
    tasks[2].percent_complete = 25
    tasks[3].percent_complete = 100
    return tasks


@pytest.mark.parametrize("cfg", [Config(), Config(max_level=4)], ids=["mismo_nivel", "nivel_4"])
def test_renamed_project_update_matches_by_code(tmp_path, cfg):
    book = tmp_path / "b.xlsx"
    make_book(book, P(renamed_tasks()), SMALL, bac={**BAC, "1.0.0-D": 100})
    status = date(2026, 1, 31)
    same = run_update(P(renamed_tasks()), Config(), book, status, dry_run=True)
    renamed = run_update(P(renamed_tasks(100), name="Programa (copia)"), cfg, book, status, dry_run=True)

    def values(res):
        return {r.code: r.ev_cum for r in res.rows}

    assert values(renamed) == values(same) == {"1.1.0": 50, "1.2.0": 50, "1.0.0-D": 100}
    assert set(values(same)) == {"1.1.0", "1.2.0", "1.0.0-D"}
    assert any("se asocian por código EDT" in w for w in renamed.warnings)


# ----------------------------------------------------------------------------
# actualizar: fecha de corte, simulación y --forzar
# ----------------------------------------------------------------------------
@pytest.mark.parametrize("status", [date(2025, 12, 31), date(2027, 1, 31)], ids=["antes", "despues"])
def test_cutoff_outside_horizon_is_rejected(tmp_path, status):
    book = tmp_path / "b.xlsx"
    make_book(book, P(base_tasks()), Config(months=12))
    before = book.read_bytes()
    with pytest.raises(ValueError, match="fuera del horizonte"):
        run_update(P(progress_tasks()), Config(), book, status, backup=False)
    assert book.read_bytes() == before


def test_cutoff_warnings_and_negative_corrections(tmp_path):
    book = tmp_path / "b.xlsx"
    make_book(book, P(base_tasks()), Config(months=12), bac=BAC)
    res = run_update(P(progress_tasks()), Config(), book, date(2026, 1, 15), dry_run=True)
    assert any("último día del mes" in w for w in res.warnings)

    # Menos avance que en el mes anterior: valor del mes negativo, informado como corrección
    run_update(P(progress_tasks()), Config(), book, date(2026, 2, 28), backup=False)
    less = progress_tasks()
    less[1].percent_complete = 10
    res = run_update(P(less), Config(), book, date(2026, 3, 31), backup=False)
    row = next(r for r in res.rows if r.code == "1.1.0")
    assert row.ev_month == pytest.approx(-40)
    assert "% EV del mes negativo (corrección)" in row.note

    # Fecha de corte anterior a la última cargada (sin valores posteriores): se advierte
    other = tmp_path / "b2.xlsx"
    make_book(other, P(base_tasks()), Config(months=12))
    tamper(other, lambda wb: wb["Resumen"].__setitem__("B5", date(2026, 2, 28)))
    res = run_update(P(progress_tasks()), Config(), other, date(2026, 1, 31), backup=False)
    assert any("anterior a la última cargada" in w for w in res.warnings)


def test_force_allows_loading_an_earlier_month(small_book):
    run_update(P(progress_tasks()), Config(), small_book, date(2026, 2, 28), backup=False)
    before = small_book.read_bytes()
    with pytest.raises(ValueError, match="posteriores"):
        run_update(P(progress_tasks()), Config(), small_book, date(2026, 1, 31), backup=False)
    assert small_book.read_bytes() == before                 # sin --forzar no se modifica
    res = run_update(P(progress_tasks()), Config(), small_book, date(2026, 1, 31), backup=False, force=True)
    assert any("anterior a la última cargada" in w for w in res.warnings)
    forced = [w for w in res.warnings if "--forzar" in w]
    assert len(forced) == 1
    assert "02/2026" in forced[0] and "% EV 1.1.0" in forced[0] and "% EV 1.2.0" in forced[0]
    assert "ya no coinciden con MS Project" in forced[0] and "'actualizar'" in forced[0]
    b = read_book(small_book)
    assert b.status_date == date(2026, 1, 31)
    assert b.ev[book_row(b, "1.1.0")][0] == pytest.approx(50)   # EV acumulado a enero (sin meses previos)


def test_dry_run_is_side_effect_free(tmp_path, small_template):
    folder = tmp_path / "simulacion"
    folder.mkdir()
    book = folder / "b.xlsx"
    shutil.copy(small_template, book)
    tamper(book, lambda wb: wb["AC"].cell(row_of(wb["AC"], "1.1.0"), FIRST_MONTH_COL, "60"))
    data, mtime = book.read_bytes(), book.stat().st_mtime_ns
    res = run_update(P(progress_tasks()), Config(), book, date(2026, 2, 28), dry_run=True)   # respaldo activado
    assert res.dry_run and res.backup is None
    assert book.read_bytes() == data and book.stat().st_mtime_ns == mtime
    assert [p.name for p in folder.iterdir()] == ["b.xlsx"]
    assert res.book.ac[book_row(res.book, "1.1.0")][0] == 60                # valor convertido solo en memoria
    assert any("AC!G" in w for w in res.warnings)


# ----------------------------------------------------------------------------
# actualizar: celdas de ingreso con texto
# ----------------------------------------------------------------------------
def load_text_in_ac(path, text):
    """Valor de AC!G (fila 1.1.0) después de 'actualizar' con `text` escrito como texto, o ERROR."""
    tamper(path, lambda wb: wb["AC"].cell(row_of(wb["AC"], "1.1.0"), FIRST_MONTH_COL, text))
    try:
        run_update(P(progress_tasks()), Config(), path, date(2026, 2, 28), backup=False)
    except WorkbookFormatError:
        return ERROR
    ws = load_workbook(path)["AC"]
    return ws.cell(row_of(ws, "1.1.0"), FIRST_MONTH_COL).value


@pytest.mark.parametrize("text, accepted", [
    ("60", {60}),
    ("12,5", {12.5}),
    ("1 234", {1234}),
    ("1.234,5", {1234.5, ERROR}),
    ("1,234.50", {1234.5, ERROR}),
])
def test_text_numbers_in_input_cells(small_book, text, accepted):
    assert load_text_in_ac(small_book, text) in accepted


# ----------------------------------------------------------------------------
# Guardado del libro
# ----------------------------------------------------------------------------
def test_failed_save_keeps_original_book(small_book, monkeypatch):
    before = small_book.read_bytes()

    def fail(self, ws):
        raise OSError("disco lleno (simulado)")

    monkeypatch.setattr(ExcelWriter, "write_worksheet", fail)
    with pytest.raises(OSError, match="simulado"):
        run_update(P(progress_tasks()), Config(), small_book, date(2026, 1, 31), backup=False)
    monkeypatch.undo()
    assert small_book.read_bytes() == before
    assert [p.name for p in small_book.parent.iterdir()] == ["b.xlsx"]     # sin archivos temporales


def test_failed_generation_keeps_existing_book(small_book, monkeypatch):
    before = small_book.read_bytes()

    def fail(self, ws):
        raise OSError("disco lleno (simulado)")

    monkeypatch.setattr(ExcelWriter, "write_worksheet", fail)
    with pytest.raises(OSError, match="simulado"):
        make_book(small_book, P(progress_tasks()), SMALL)       # como generar --sobrescribir
    monkeypatch.undo()
    assert small_book.read_bytes() == before
    assert [p.name for p in small_book.parent.iterdir()] == ["b.xlsx"]


def lock_on_save(monkeypatch):
    """Simula el libro abierto en Excel: guardar falla con PermissionError."""
    def locked(self, filename):
        raise PermissionError(13, "El proceso no tiene acceso al archivo", str(filename))
    monkeypatch.setattr(Workbook, "save", locked)


def test_locked_book_reports_clear_error(small_book, monkeypatch):
    before = small_book.read_bytes()
    lock_on_save(monkeypatch)
    with pytest.raises(PermissionError, match="se debe cerrar"):
        run_update(P(progress_tasks()), Config(), small_book, date(2026, 1, 31))
    assert small_book.read_bytes() == before


def test_locked_book_leaves_no_orphan_backup(small_book, monkeypatch):
    lock_on_save(monkeypatch)
    with pytest.raises(PermissionError):
        run_update(P(progress_tasks()), Config(), small_book, date(2026, 1, 31))
    assert backups(small_book.parent) == []


@pytest.mark.skipif(sys.platform != "win32", reason="bloqueo de archivos de Windows")
def test_real_windows_lock_reports_clear_error(small_book):
    win32file = pytest.importorskip("win32file")
    before = small_book.read_bytes()
    handle = win32file.CreateFile(str(small_book), win32file.GENERIC_READ, win32file.FILE_SHARE_READ, None,
                                  win32file.OPEN_EXISTING, 0, None)        # como Excel: solo lectura compartida
    try:
        with pytest.raises(PermissionError, match="se debe cerrar"):
            run_update(P(progress_tasks()), Config(), small_book, date(2026, 1, 31), backup=False)
    finally:
        handle.Close()
    assert small_book.read_bytes() == before
    assert [p.name for p in small_book.parent.iterdir()] == ["b.xlsx"]


@pytest.mark.skipif(sys.platform != "win32", reason="bloqueo de archivos de Windows")
def test_real_windows_lock_is_reported_before_reading_the_project(example_book, example_files, capsys):
    win32file = pytest.importorskip("win32file")
    handle = win32file.CreateFile(str(example_book), win32file.GENERIC_READ, win32file.FILE_SHARE_READ, None,
                                  win32file.OPEN_EXISTING, 0, None)
    try:
        rc, out, err = run_cli(capsys, "actualizar", "--origen", example_files["jun"], "--config",
                               example_files["config"], "--libro", example_book)
    finally:
        handle.Close()
    assert rc == 2 and "se debe cerrar" in err and "Leyendo" not in out
    assert backups(example_book.parent) == []


# ----------------------------------------------------------------------------
# Libro modificado o dañado
# ----------------------------------------------------------------------------
def _rename_ev(wb):
    wb["EV"].title = "EV mensual"


def _delete(name):
    def fn(wb):
        del wb[name]
    return fn


def _meta(key, value=None, delete=False):
    def fn(wb):
        ws = wb["_meta"]
        r = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == key)
        if delete:
            ws.delete_rows(r)
        else:
            ws.cell(r, 2).value = value
    return fn


@pytest.mark.parametrize("mutate", [
    pytest.param(_rename_ev, id="hoja_EV_renombrada"),
    pytest.param(_delete("Resumen"), id="sin_Resumen"),
    pytest.param(_delete("_graficos"), id="sin_graficos"),
    pytest.param(_meta("horizonte_inicio", "2026-01-01"), id="inicio_como_texto"),
    pytest.param(_meta("meses", "abc"), id="meses_no_numerico"),
    pytest.param(_meta("fila_total", delete=True), id="falta_fila_total"),
    pytest.param(_delete("_meta"), id="sin_meta"),
])
def test_tampered_book_fails_with_format_error(example_book, example_files, capsys, mutate):
    tamper(example_book, mutate)
    before = example_book.read_bytes()
    rc, out, err = run_cli(capsys, "actualizar", "--origen", example_files["jun"], "--config",
                            example_files["config"], "--libro", example_book, "--sin-respaldo")
    assert example_book.read_bytes() == before
    assert "inesperado" not in err and "ERROR:" in err
    assert rc == 2
    assert "Leyendo" not in out                             # se informa antes de leer MS Project


def test_legacy_book_without_dashboard_still_updates(small_book):
    """Libros anteriores a la versión 1.3.0: sin Cuadro de Mando ni _graficos, con hoja 'Curva S'."""
    def legacy(wb):
        del wb["Cuadro de Mando"]
        del wb["_graficos"]
        wb["Evolución mensual"].title = "Curva S"
    tamper(small_book, legacy)
    res = run_update(P(progress_tasks()), Config(), small_book, date(2026, 1, 31), backup=False)
    assert {r.code for r in res.rows} == {"1.1.0", "1.2.0"}
    assert chart_count(small_book) == 1


def test_missing_company_sheet_in_multi_book(tmp_path):
    cfg = Config(multi_enabled=True, months=24)
    book = tmp_path / "multi.xlsx"
    make_book(book, P(base_tasks()), cfg)
    tamper(book, _delete("EV-EmpresaB"))
    with pytest.raises(WorkbookFormatError, match="Falta la hoja 'EV-EmpresaB'"):
        run_update(P(progress_tasks()), cfg, book, date(2026, 1, 31), backup=False)


# ----------------------------------------------------------------------------
# Escritura del libro y Cuadro de Mando
# ----------------------------------------------------------------------------
def feature_snapshot(path):
    """Elementos del libro que openpyxl debe conservar al reabrirlo y guardarlo."""
    wb = load_workbook(path)
    sheets = {ws.title: (len(ws.data_validations.dataValidation), len(list(ws.conditional_formatting)),
                         ws.sheet_state, str(ws.freeze_panes), len(ws.merged_cells.ranges)) for ws in wb.worksheets}
    return sheets, sorted(wb.defined_names.keys()), wb["Cuadro de Mando"].print_area


def chart_count(path):
    with zipfile.ZipFile(path) as z:
        return sum(1 for n in z.namelist() if re.fullmatch(r"xl/charts/chart\d+\.xml", n))


def test_update_roundtrip_preserves_workbook_features(small_book):
    before = feature_snapshot(small_book)
    assert chart_count(small_book) == 3
    run_update(P(progress_tasks()), Config(), small_book, date(2026, 1, 31), backup=False)
    run_update(P(progress_tasks()), Config(), small_book, date(2026, 2, 28), backup=False)
    assert feature_snapshot(small_book) == before
    assert chart_count(small_book) == 3


@pytest.mark.parametrize("months", [1, 14, 336, 337, 600, 1000])
def test_date_axis_supports_every_horizon(months):
    step = dashboard._date_axis(months).majorUnit
    assert step >= 1 and (months / step <= 14 or step == 60)


def test_generate_and_update_book_beyond_336_months(tmp_path):
    tasks = base_tasks()
    tasks[1].baseline_finish = datetime(2055, 12, 28, 17)       # 30 años de plan
    book = tmp_path / "largo.xlsx"
    make_book(book, P(tasks), Config(months=400))
    res = run_update(P(progress_tasks()), Config(), book, date(2054, 12, 31), backup=False)
    assert res.month_index == 347
    assert chart_count(book) == 3


@pytest.mark.parametrize("where", ["tarea", "proyecto"])
def test_control_characters_in_names(tmp_path, where):
    tasks, name = base_tasks(), "Programa"
    if where == "tarea":
        tasks[1].name = "B\x07campana"
    else:
        name = "Prog\x0bVT"
    book = tmp_path / "b.xlsx"
    make_book(book, P(tasks, name=name), SMALL)
    ws = load_workbook(book)["PV"]
    assert ws.cell(row_of(ws, "1.1.0"), 2).value == ("Bcampana" if where == "tarea" else "B")
    meta = read_book(book).meta
    assert meta["proyecto"] == ("Programa" if where == "tarea" else "ProgVT")
    res = run_update(P(progress_tasks(), name=name), Config(), book, date(2026, 1, 31), dry_run=True)
    assert not any("fue generado para el proyecto" in w for w in res.warnings)
    assert {r.code for r in res.rows} == {"1.1.0", "1.2.0"}


# ----------------------------------------------------------------------------
# --migrar-desde
# ----------------------------------------------------------------------------
def old_book_with_ev(path):
    """Libro anterior con EV en enero (10), febrero (20) y mayo (5) de 2026 y fecha de corte en mayo."""
    make_book(path, P(base_tasks()), SMALL, bac=BAC)

    def load(wb):
        ws = wb["% EV"]
        r = row_of(ws, "1.1.0")
        for i, v in ((0, 0.1), (1, 0.2), (4, 0.05)):                # BAC 100
            ws.cell(r, FIRST_MONTH_COL + i, v)
        wb["Resumen"]["B5"].value = date(2026, 5, 31)
    tamper(path, load)
    return path


def test_migration_renamed_project_by_code(tmp_path):
    old = tmp_path / "old.xlsx"
    make_book(old, P(base_tasks()), SMALL, bac=BAC)

    def load(wb):
        for code in ("1.1.0", "1.2.0"):
            wb["% EV"].cell(row_of(wb["% EV"], code), FIRST_MONTH_COL, 0.1)     # EV 10 y 20
    tamper(old, load)
    new_tasks = [T(11, 1, "1", "Fase"), T(12, 2, "1.1", "B"), T(13, 2, "1.2", "C renombrada")]
    new = tmp_path / "new.xlsx"
    rep = make_book(new, P(new_tasks, name="Otro nombre"), SMALL, prior=prior_from_book(old, "Otro nombre"))
    b = read_book(new)
    assert {s.code: sum(v or 0 for v in b.ev[s.row]) for s in b.input_rows} == \
        {"1.1.0": pytest.approx(10), "1.2.0": pytest.approx(20)}
    messages = [(i.severity, i.code, i.message) for i in rep.issues]
    assert any(sev == "Advertencia" and "corresponde al proyecto 'Programa'" in m for sev, _c, m in messages)
    assert any(code == "1.2.0" and "Recibió el historial" in m for _s, code, m in messages)


def test_migration_values_after_new_horizon_end_are_accumulated(tmp_path):
    old = old_book_with_ev(tmp_path / "old.xlsx")
    new = tmp_path / "new.xlsx"
    rep = make_book(new, P(base_tasks()), Config(months=2), prior=prior_from_book(old, "Programa"))
    b = read_book(new)
    assert b.ev[book_row(b, "1.1.0")] == pytest.approx([10, 25])
    assert any("después del último mes" in i.message for i in rep.issues)


def test_migration_values_before_new_horizon_start_warn(tmp_path):
    old = old_book_with_ev(tmp_path / "old.xlsx")
    new = tmp_path / "new.xlsx"
    rep = make_book(new, P(base_tasks()), Config(start_month=date(2026, 4, 1), months=12),
                    prior=prior_from_book(old, "Programa"))
    b = read_book(new)
    assert sum(book_totals(b)[1]) == pytest.approx(35)        # el total se conserva
    assert b.ev[book_row(b, "1.1.0")][:2] == pytest.approx([30, 5])    # enero y febrero, en el mes 1 (abril)
    assert any("EV/AC por 30.00" in i.message and "antes del primer mes" in i.message for i in rep.issues)


def test_migration_keeps_explicit_zero_months(tmp_path):
    old = tmp_path / "old.xlsx"
    make_book(old, P(base_tasks()), SMALL, bac=BAC)
    run_update(P(base_tasks()), Config(), old, date(2026, 1, 31), backup=False)         # enero cargado con 0
    run_update(P(progress_tasks()), Config(), old, date(2026, 2, 28), backup=False)
    new = tmp_path / "new.xlsx"
    make_book(new, P(base_tasks()), SMALL, prior=prior_from_book(old, "Programa"))
    a, b = read_book(old), read_book(new)
    row = book_row(a, "1.1.0")
    assert a.ev[row][:2] == [0.0, 50.0]
    assert b.ev[book_row(b, "1.1.0")][:2] == [0.0, 50.0]
    assert b.ev[book_row(b, "1.1.0")][2] is None                # los meses sin cargar siguen vacíos


def test_migration_zero_only_rows_are_not_reported_as_lost(tmp_path):
    old = tmp_path / "old.xlsx"
    make_book(old, P(base_tasks()), SMALL)                                            # sin BAC
    run_update(P(base_tasks()), Config(), old, date(2026, 1, 31), backup=False)         # solo ceros
    prior = prior_from_book(old, "Programa")
    rep = make_book(tmp_path / "new.xlsx", P(base_tasks()[:2]), SMALL, prior=prior)  # C eliminada
    assert prior.unmatched == [] and prior.matched == 0
    assert not any(i.severity == "Error" for i in rep.issues)


# ----------------------------------------------------------------------------
# Línea de comandos
# ----------------------------------------------------------------------------
@pytest.mark.parametrize("command", [[], ["generar"], ["actualizar"], ["analizar"]],
                         ids=["programa", "generar", "actualizar", "analizar"])
def test_cli_help_of_every_command(capsys, command):
    """La ayuda se muestra sin errores (un % sin escapar en un texto de ayuda la rompe)."""
    with pytest.raises(SystemExit) as exc:
        main(command + ["--help"])
    assert exc.value.code == 0 and "usage" in capsys.readouterr().out.lower()


@pytest.mark.parametrize("args, message", [
    (["generar", "--origen", "{tmp}/no_existe.xml", "--salida", "{tmp}/x.xlsx"], "No existe el archivo de origen"),
    (["generar", "--origen", "{tmp}/p.txt", "--salida", "{tmp}/x.xlsx"], "No se reconoce la extensión"),
    (["generar", "--origen", "{tmp}/p.xml", "--salida", "{tmp}/x.xls"], "debe ser un archivo .xlsx"),
    (["generar", "--origen", "{tmp}/p.xml", "--salida", "{tmp}/x.xlsx", "--migrar-desde", "{tmp}/x.xlsx"],
     "debe ser un archivo distinto"),
    (["actualizar", "--origen", "{tmp}/p.xml", "--libro", "{tmp}/no_existe.xlsx"], "No existe el libro"),
    (["actualizar", "--origen", "{tmp}/p.xml", "--libro", "{tmp}/x.xlsx", "--fecha-corte", "2026-02-30"],
     "Fecha inválida"),
], ids=["origen_inexistente", "extension", "salida_xls", "migrar_igual_salida", "libro_inexistente",
        "fecha_invalida"])
def test_cli_errors_exit_with_code_2(tmp_path, capsys, args, message):
    rc, _out, err = run_cli(capsys, *[a.replace("{tmp}", str(tmp_path)) for a in args])
    assert rc == 2
    assert err.startswith("ERROR: ") and message in err


@pytest.mark.parametrize("error, code, message", [
    (KeyError("x"), 1, "ERROR inesperado: KeyError"),
    (KeyboardInterrupt(), 130, "Cancelado."),
    (RuntimeError("falla de COM"), 2, "ERROR: falla de COM"),
], ids=["inesperado", "cancelado", "runtime"])
def test_cli_unexpected_errors(tmp_path, capsys, monkeypatch, error, code, message):
    def boom(args):
        raise error
    monkeypatch.setattr("evm.cli.cmd_generar", boom)
    rc, _out, err = run_cli(capsys, "generar", "--origen", tmp_path / "p.xml")
    assert rc == code and message in err
    if code == 1:
        assert "-v" in err                                   # indica cómo ver el detalle técnico


@pytest.mark.parametrize("option, message", [
    pytest.param(["--meses", "0"], "meses debe estar entre", id="meses_0"),
    pytest.param(["--nivel-max", "0"], "nivel_maximo debe estar entre", id="nivel_max_0"),
    pytest.param(["--mes-inicio", ""], "Mes inválido", id="mes_inicio_vacio"),
])
def test_cli_rejects_zero_overrides(tmp_path, capsys, example_files, option, message):
    rc, out, err = run_cli(capsys, "generar", "--origen", example_files["base"], "--config",
                           example_files["config"], "--salida", tmp_path / "x.xlsx", *option)
    assert rc == 2 and message in err
    assert "Leyendo" not in out and not (tmp_path / "x.xlsx").exists()


# ----------------------------------------------------------------------------
# Configuración
# ----------------------------------------------------------------------------
@pytest.mark.parametrize("content, message", [
    pytest.param(b'[exportacion]\nmoneda = 5\n', "debe ser texto", id="moneda_numerica"),
    pytest.param(b'[metodos]\nlista = "0/100"\n', "lista de textos", id="lista_texto"),
    pytest.param(b'[exportacion]\nferiados = "2026-01-01"\n', "lista de fechas", id="feriados_texto"),
    pytest.param(b'[horizonte]\nmeses = 0\n', "meses debe estar entre", id="meses_0"),
    pytest.param(b'[exportacion]\ndecimales = 5\n', "decimales", id="decimales_5"),
    pytest.param(b'[semaforo]\nrojo = 0.97\namarillo = 0.95\n', "umbrales", id="umbrales_invertidos"),
    pytest.param(b'[campos\n', "Error de sintaxis", id="sintaxis"),
    pytest.param('[campos]\nmetodo = "Método EVM"\n'.encode("cp1252"), "UTF-8", id="ansi_cp1252"),
    # Sección que no es tabla y enteros con decimales: ver test_robustez_calculo.py
])
def test_config_rejects_invalid_files(tmp_path, content, message):
    path = tmp_path / "config.toml"
    path.write_bytes(content)
    with pytest.raises(ConfigError, match=message):
        load_config(path)


def test_config_missing_file(tmp_path):
    with pytest.raises(ConfigError, match="No se encontró"):
        load_config(tmp_path / "no_existe.toml")


def test_config_accepts_utf8_with_bom(tmp_path):
    path = tmp_path / "config.toml"
    path.write_bytes('[campos]\nmetodo = "Método EVM"\n'.encode("utf-8-sig"))
    assert load_config(path).method_field == "Método EVM"
