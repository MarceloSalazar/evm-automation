"""Comando validar y controles del archivo de MS Project: códigos EDT, fechas y duración, pesos, método,
empresas, avance a la fecha de estado y lógica del cronograma."""
from datetime import date, datetime

import pytest
from openpyxl import load_workbook

from evm.cli import main
from evm.config import Config, ConfigError, load_config
from evm.model import CONSTRAINT_MSO, LINK_FF, LINK_FS, LINK_SF, LINK_SS, Link
from evm.reader_xml import duration_minutes, read_xml
from evm.validation import logic_checks, progress_checks, validate_part
from evm.wbs import (SEV_ERROR, SEV_INFO, SEV_WARN, TOPIC_DATES, TOPIC_LOGIC, TOPIC_METHOD, TOPIC_PROGRESS,
                     TOPIC_WBS, TOPIC_WEIGHTS, build_part, build_report)

from builders import P, T, by_name, tree
from mspdi import NUMBER1, TEXT10, project_xml, task_xml, write_xml

S, F = datetime(2026, 1, 5, 8), datetime(2026, 3, 27, 17)


def report(tasks, cfg=None, **project_kw):
    cfg = cfg or Config()
    cfg.validate()
    return build_report(P(tasks, **project_kw), cfg)


def found(rep_or_issues, topic=None, severity=None, text=""):
    issues = getattr(rep_or_issues, "issues", rep_or_issues)
    return [i for i in issues if (topic is None or i.topic == topic) and (severity is None or i.severity == severity)
            and text in i.message]


# ----------------------------------------------------------------------------
# Códigos EDT
# ----------------------------------------------------------------------------
def test_child_code_must_extend_parent_code():
    tasks = tree('''
        1 Fase
          1.2 Grupo
            1.2.1 Bien
            1.14.2 Mal
    ''')
    rep = report(tasks)
    errors = found(rep, TOPIC_WBS, SEV_ERROR)
    assert [(i.code, i.name) for i in errors] == [("1.14.2", "Mal")]
    assert "no depende del de su tarea superior 1.2.0 «Grupo»: se esperaba 1.2.x" in errors[0].message


def test_valid_codes_have_no_wbs_observations():
    tasks = tree('''
        1 Fase
          1.1 A
            1.1.1 A1
            1.1.2 A2
              - Actividad sin código
          1.2 B
        2 Fase 2
    ''')
    rep = report(tasks)
    assert not found(rep, TOPIC_WBS, SEV_ERROR) and not found(rep, TOPIC_WBS, SEV_WARN)


def test_code_format_errors():
    tasks = tree('''
        1 Fase
          1..3 Vacio
          1.A Letra
    ''')
    rep = report(tasks)
    by = {i.name: i for i in found(rep, TOPIC_WBS) if i.name in ("Vacio", "Letra")}
    assert by["Vacio"].severity == SEV_ERROR and "formato no válido" in by["Vacio"].message
    assert by["Letra"].severity == SEV_WARN and "no numérico" in by["Letra"].message


def test_skipped_level_order_gaps_and_duplicates():
    tasks = tree('''
        1 Fase
          1.2 B
            1.2.3.1 Salta
          1.1 A
        2 Fase 2
          2.1 X
          2.5 Y
          2.5 Z
    ''')
    rep = report(tasks)
    wbs = found(rep, TOPIC_WBS)
    salta = next(i for i in wbs if i.name == "Salta")
    assert salta.severity == SEV_WARN and "salta un nivel" in salta.message
    order = next(i for i in wbs if i.name == "A")
    assert order.severity == SEV_WARN and "fuera de orden" in order.message and "1.2.0 «B»" in order.message
    gaps = [i for i in wbs if "saltos" in i.message]
    assert len(gaps) == 1 and gaps[0].severity == SEV_INFO and "2.1.0 → 2.5.0" in gaps[0].message
    dup = [i for i in wbs if "duplicado" in i.message]
    assert len(dup) == 1 and dup[0].severity == SEV_ERROR and dup[0].code == "2.5.0" and "2 tareas" in dup[0].message


def test_code_without_upper_code():
    tasks = tree('''
        - Contenedor
          1.2 Suelto
    ''')
    rep = report(tasks)
    issue = next(i for i in found(rep, TOPIC_WBS) if i.name == "Suelto")
    assert issue.severity == SEV_WARN and "debería depender de la tarea 1.0.0" in issue.message


def test_plan_code_is_not_compared():
    """Sin planificación elegida, INTERNO y CLIENTE son contenedores: sus subtareas no se comparan con ellos (pero
    los códigos repetidos entre planificaciones son duplicados)."""
    tasks = tree('''
        INTERNO Interna
          1 Fase
            1.1 A
        CLIENTE Cliente
          1 Fase
    ''')
    rep = report(tasks)
    wbs = found(rep, TOPIC_WBS)
    assert not [i for i in wbs if "no depende" in i.message or "no numérico" in i.message]
    assert [i.code for i in wbs if "duplicado" in i.message] == ["1.0.0"]
    assert "planificacion" in next(i for i in wbs if "duplicado" in i.message).message
    rep = report(tasks, Config(plan="INTERNO"))
    assert not found(rep, TOPIC_WBS, SEV_ERROR) and not found(rep, TOPIC_WBS, SEV_WARN)


def test_project_code_zero_is_a_container():
    tasks = tree('''
        0 Proyecto
          1 Fase
            1.1 A
          2 Fase 2
    ''')
    rep = report(tasks)
    assert not found(rep, TOPIC_WBS, SEV_ERROR) and not found(rep, TOPIC_WBS, SEV_WARN)


def test_visible_no_branch_codes_are_checked():
    tasks = tree('''
        1 Fase
          1.1 A
          2.7 Oculta
    ''')
    by_name(tasks)["Oculta"].visible = "No"
    rep = report(tasks)
    assert [i.name for i in found(rep, TOPIC_WBS, SEV_ERROR)] == ["Oculta"]


# ----------------------------------------------------------------------------
# Fechas y duración
# ----------------------------------------------------------------------------
def test_dates_and_duration():
    tasks = tree('''
        1 Fase
          1.1 A
            - Al reves
            - Hito largo
            - Cero
            - Estimada
            - Manual
    ''')
    t = by_name(tasks)
    t["Al reves"].start, t["Al reves"].finish = F, S
    t["Al reves"].baseline_start, t["Al reves"].baseline_finish = F, S
    t["Hito largo"].milestone, t["Hito largo"].duration = True, 480.0
    t["Cero"].duration = 0.0
    t["Estimada"].estimated, t["Estimada"].duration = True, 2400.0
    t["Manual"].manual = True
    rep = report(tasks)
    dates = found(rep, TOPIC_DATES)
    reverse = [i for i in dates if i.name.startswith("Al reves")]
    assert {i.severity for i in reverse} == {SEV_ERROR} and len(reverse) == 2
    assert reverse[0].name == "Al reves (Id 3)"                 # sin código EDT: con su Id
    assert any(i.name.startswith("Hito largo") and "marcada como hito y tiene duración" in i.message for i in dates)
    assert any(i.name.startswith("Cero") and "Duración 0" in i.message for i in dates)
    assert any("1 tarea tiene duración estimada" in i.message for i in dates)
    assert any(i.name.startswith("Estimada") and i.message == "Duración estimada («?»)." for i in dates)
    manual = [i for i in dates if "manualmente" in i.message]
    assert len(manual) == 1 and manual[0].severity == SEV_WARN and "Id 7" in manual[0].message


def test_duration_unknown_is_not_checked():
    """Sin el dato de la duración (XML sin <Duration>), no se revisan los hitos con duración ni la duración 0."""
    tasks = tree('''
        1 Fase
          1.1 A
            - Hito
    ''')
    by_name(tasks)["Hito"].milestone = True
    rep = report(tasks)
    assert not found(rep, TOPIC_DATES)


# ----------------------------------------------------------------------------
# Hitos y pesos, método de medición, empresas
# ----------------------------------------------------------------------------
def test_milestone_weight_checks():
    tasks = tree('''
        1 Fase
          1.1 A
            - Tarea con peso
          1.2 B
            - Tarea
    ''')
    t = by_name(tasks)
    t["Tarea con peso"].milestone_weight = 30
    t["B"].method = "Hitos ponderados"
    rep = report(tasks, Config(milestone_weight_field="Peso hito", method_field="Método EVM"))
    weights = found(rep, TOPIC_WEIGHTS)
    assert any(i.name.startswith("Tarea con peso") and "no es un hito" in i.message for i in weights)
    assert any(i.code == "1.2.0" and "no tiene hitos" in i.message for i in weights)


def test_weight_field_without_values():
    tasks = tree('''
        1 Fase
          1.1 A
            - Tarea
            - Hito
    ''')
    by_name(tasks)["Hito"].milestone = True
    rep = report(tasks, Config(milestone_weight_field="Peso hito"), fields_found={"peso_hito": "Peso hito"})
    assert len(found(rep, text="ningún hito tiene peso")) == 1


def test_method_checks():
    long_s, long_f = datetime(2026, 1, 5, 8), datetime(2026, 6, 30, 17)
    tasks = tree('''
        1 Fase
          1.1 Corto
          1.2 Largo
          1.3 Medio
          1.4 Raro
          1.5 Sin metodo
            - Sub
    ''')
    t = by_name(tasks)
    t["Corto"].method = "0/100"
    t["Largo"].method = "0/100"
    t["Medio"].method = "50/50"
    t["Raro"].method = "60/30"
    t["Sub"].method = "50/50"
    for name in ("Largo", "Medio"):
        t[name].baseline_start, t[name].baseline_finish = long_s, long_f
    t["Medio"].baseline_finish = datetime(2026, 3, 20, 17)          # 2,5 meses
    t["Corto"].baseline_start, t["Corto"].baseline_finish = datetime(2026, 1, 5, 8), datetime(2026, 1, 30, 17)
    rep = report(tasks, Config(method_field="Método EVM"))
    method = {i.code: i for i in found(rep, TOPIC_METHOD) if i.code}
    assert "1.1.0" not in method
    assert method["1.2.0"].severity == SEV_WARN and "0/100 en un paquete de 5,8 meses" in method["1.2.0"].message
    assert "el EV queda en el 50 %" in method["1.3.0"].message and "meses_formula_fija" in method["1.3.0"].message
    assert "no reconocido" in method["1.4.0"].message
    assert method["1.5.0"].severity == SEV_WARN and "La fila no tiene método" in method["1.5.0"].message

    # Umbrales configurables; 0 = no se controla
    rep = report(tasks, Config(method_field="Método EVM", max_months_0_100=6, max_months_fixed=0))
    codes = {i.code for i in found(rep, TOPIC_METHOD) if "paquete de" in i.message}
    assert codes == set()


def test_validation_thresholds_in_config(tmp_path):
    path = tmp_path / "config.toml"
    path.write_text("[validacion]\nmeses_0_100 = 2\nmeses_formula_fija = 3.5\n", encoding="utf-8")
    cfg = load_config(path)
    assert (cfg.max_months_0_100, cfg.max_months_fixed) == (2.0, 3.5)
    path.write_text("[validacion]\nmeses_0_100 = -1\n", encoding="utf-8")
    with pytest.raises(ConfigError, match="meses_0_100"):
        load_config(path)
    assert cfg.max_loe_share == 15.0
    path.write_text("[validacion]\nloe_maximo = 25\n", encoding="utf-8")
    assert load_config(path).max_loe_share == 25.0
    path.write_text("[validacion]\nloe_maximo = 120\n", encoding="utf-8")
    with pytest.raises(ConfigError, match="loe_maximo") as exc:
        load_config(path)
    assert exc.value.id == "CFG-050"


def test_split_on_task_that_is_not_both():
    tasks = tree('''
        1 Fase
          1.1 A
          1.2 B
    ''')
    t = by_name(tasks)
    t["A"].company, t["A"].split_a = "EmpresaA", 70
    t["B"].company, t["B"].split_a = "Ambas", 70
    rep = report(tasks, Config(multi_enabled=True))
    ignored = found(rep, text="el reparto solo se usa")
    assert [i.code for i in ignored] == ["1.1.0"] and ignored[0].severity == SEV_INFO


# ----------------------------------------------------------------------------
# Avance a la fecha de estado (solo validar)
# ----------------------------------------------------------------------------
def progress(tasks, status, cfg=None, **project_kw):
    project = P(tasks, **project_kw)
    part = build_part(project, cfg or Config())
    return progress_checks(project, part.scope, part.code_of, status)


def test_actual_dates_consistency():
    tasks = tree('''
        1 Fase
          1.1 Fin sin comienzo
          1.2 Al reves
          1.3 Avance sin comienzo
          1.4 Completa sin fin
          1.5 Fin con avance parcial
    ''')
    t = by_name(tasks)
    t["Fin sin comienzo"].actual_finish = datetime(2026, 2, 1)
    t["Al reves"].actual_start, t["Al reves"].actual_finish = datetime(2026, 2, 10), datetime(2026, 2, 1)
    t["Al reves"].percent_complete = 100
    t["Avance sin comienzo"].percent_complete = 40
    t["Completa sin fin"].percent_complete, t["Completa sin fin"].actual_start = 100, S
    t["Fin con avance parcial"].actual_start, t["Fin con avance parcial"].actual_finish = S, F
    t["Fin con avance parcial"].percent_complete = 80
    issues = progress(tasks, None)
    by = {}
    for i in issues:
        by.setdefault(i.name, []).append(i)
    assert by["Fin sin comienzo"][0].severity == SEV_ERROR
    assert by["Al reves"][0].severity == SEV_ERROR and "anterior al Comienzo real" in by["Al reves"][0].message
    assert "40 % y no tiene Comienzo real" in by["Avance sin comienzo"][0].message
    assert "100 % completado y no tiene Fin real" in by["Completa sin fin"][0].message
    assert "80 %" in by["Fin con avance parcial"][0].message
    assert any("no tiene fecha de estado" in i.message and i.severity == SEV_WARN for i in issues)


def test_progress_against_status_date():
    tasks = tree('''
        1 Fase
          1.1 Futuro
          1.2 Atrasada
          1.3 Sin comenzar
          1.4 Terminada
          1.5 Pendiente
    ''')
    t = by_name(tasks)
    status = date(2026, 2, 28)
    t["Futuro"].actual_start, t["Futuro"].percent_complete = datetime(2026, 3, 2, 8), 10
    t["Atrasada"].actual_start, t["Atrasada"].percent_complete = S, 50
    t["Atrasada"].finish = datetime(2026, 2, 20, 17)
    t["Terminada"].actual_start, t["Terminada"].percent_complete = S, 100
    t["Terminada"].actual_finish = datetime(2026, 2, 13, 17)
    for name in ("Atrasada", "Terminada", "Sin comenzar"):
        t[name].baseline_finish = datetime(2026, 2, 13, 17)
    t["Pendiente"].start, t["Pendiente"].finish = datetime(2026, 3, 2, 8), datetime(2026, 3, 27, 17)
    t["Pendiente"].baseline_start = datetime(2026, 3, 2, 8)
    issues = progress(tasks, status)
    future = [i for i in issues if "posterior a la fecha de estado" in i.message]
    assert [(i.name, i.severity) for i in future] == [("Futuro", SEV_ERROR)]
    late = {i.name: i.message for i in issues if i.code and "anterior a la fecha de estado" in i.message}
    assert set(late) == {"Atrasada", "Sin comenzar"}
    assert "no está terminada (50 %)" in late["Atrasada"] and "no tiene avance" in late["Sin comenzar"]
    assert any("2 tareas tienen trabajo pendiente" in i.message for i in issues)
    bei = next(i for i in issues if "BEI" in i.message)
    assert "= 0,33" in bei.message and bei.severity == SEV_WARN and "1.2.0" in bei.message
    started = next(i for i in issues if "debían comenzar" in i.message or "debía comenzar" in i.message)
    assert "1.3.0" in started.message and "1.5.0" not in started.message


def missed_and_cpli_tasks(forecast_finish: datetime) -> list:
    """Tres tareas que debían terminar el 13/02/2026 (a tiempo, tarde y sin terminar) y una que termina, según la
    línea base, el 27/03/2026 y, según el programa actual, en `forecast_finish`."""
    tasks = tree('''
        1 Fase
          1.1 A tiempo
          1.2 Tarde
          1.3 Sin terminar
          1.4 Futura
    ''')
    t = by_name(tasks)
    for name, finish in (("A tiempo", datetime(2026, 2, 12, 17)), ("Tarde", datetime(2026, 2, 20, 17)),
                         ("Sin terminar", datetime(2026, 3, 6, 17))):
        t[name].baseline_finish, t[name].finish, t[name].actual_start = datetime(2026, 2, 13, 17), finish, S
        if name != "Sin terminar":
            t[name].actual_finish, t[name].percent_complete = finish, 100
    t["Sin terminar"].percent_complete = 50
    t["Futura"].start, t["Futura"].finish = datetime(2026, 3, 2, 8), forecast_finish
    return tasks


def test_missed_tasks_and_cpli():
    """DCMA 11 y 13: tareas incumplidas respecto de la línea base y CPLI con el fin previsto 10 días atrasado."""
    issues = progress(missed_and_cpli_tasks(datetime(2026, 4, 10, 17)), date(2026, 2, 27))
    missed = next(i for i in issues if i.id == "AVA-014")
    assert missed.severity == SEV_WARN
    assert missed.message == ("De 3 tareas que debían terminar a la fecha de estado según la línea base, 2 no "
                              "terminaron a tiempo (66,7 %; DCMA: hasta 5 %): 1 sin terminar y 1 terminada después "
                              "de su fin de línea base (1.2.0, 1.3.0).")
    cpli = next(i for i in issues if i.id == "AVA-015")
    assert cpli.severity == SEV_WARN
    assert cpli.message == ("Índice de longitud del camino crítico (CPLI) = 0,67: quedan 30 días laborables hasta el "
                            "fin previsto (10/04/2026) y el fin de la línea base es el 27/03/2026 (10 días laborables "
                            "de atraso) (DCMA: al menos 0,95).")


def test_cpli_ahead_of_baseline_and_not_computed():
    issues = progress(missed_and_cpli_tasks(datetime(2026, 3, 20, 17)), date(2026, 2, 27))
    cpli = next(i for i in issues if i.id == "AVA-015")
    assert cpli.severity == SEV_INFO and "= 1,33" in cpli.message and "(5 días laborables de adelanto)" in cpli.message
    # Sin fecha de estado, sin línea base o sin trabajo pendiente después de la fecha de estado: no se calcula
    assert not [i for i in progress(missed_and_cpli_tasks(F), None) if i.id in ("AVA-014", "AVA-015")]
    no_baseline = progress(missed_and_cpli_tasks(F), date(2026, 2, 27), baseline_found=False)
    assert not [i for i in no_baseline if i.id in ("AVA-014", "AVA-015")]
    assert not [i for i in progress(missed_and_cpli_tasks(F), date(2026, 4, 30)) if i.id == "AVA-015"]


def test_missed_tasks_within_threshold_is_info():
    lines = ["1 Fase"] + [f"  1.{n} T{n}" for n in range(1, 22)]
    tasks = tree("\n".join(lines), bf=datetime(2026, 2, 13, 17))
    for task in tasks[1:]:
        task.actual_start, task.actual_finish, task.percent_complete = S, datetime(2026, 2, 13, 17), 100
    tasks[1].actual_finish = datetime(2026, 2, 16, 17)                 # 1 de 21 terminó tarde: 4,8 %
    missed = next(i for i in progress(tasks, date(2026, 2, 27)) if i.id == "AVA-014")
    assert missed.severity == SEV_INFO and "(4,8 %; DCMA: hasta 5 %)" in missed.message


def test_out_of_sequence_progress():
    tasks = tree('''
        1 Fase
          1.1 Comenzada
          1.2 Sin comenzar
          1.3 FC
          1.4 CC bien
          1.5 CC
          1.6 FF
          1.7 Sin avance
    ''')
    t = by_name(tasks)
    uid = {n: t[n].uid for n in t}
    t["Comenzada"].actual_start, t["Comenzada"].percent_complete = S, 40
    for name, pred, kind in (("FC", "Comenzada", LINK_FS), ("CC bien", "Comenzada", LINK_SS),
                             ("CC", "Sin comenzar", LINK_SS), ("FF", "Comenzada", LINK_FF),
                             ("Sin avance", "Comenzada", LINK_FS)):
        t[name].predecessors = [Link(uid[pred], kind)]
        if name != "Sin avance":
            t[name].actual_start, t[name].percent_complete = S, 10
    t["FF"].actual_finish, t["FF"].percent_complete = F, 100
    issues = progress(tasks, None)
    assert [(i.name, i.message) for i in issues if i.id == "AVA-017"] == [
        ("FC", "Tiene avance y su predecesora 1.1.0 «Comenzada» no terminó (vínculo FC)."),
        ("CC", "Tiene avance y su predecesora 1.2.0 «Sin comenzar» no comenzó (vínculo CC)."),
        ("FF", "Está terminada y su predecesora 1.1.0 «Comenzada» no terminó (vínculo FF)."),
    ]
    summary = next(i for i in issues if i.id == "AVA-016")
    assert summary.severity == SEV_WARN and summary.topic == TOPIC_PROGRESS
    assert summary.message.startswith("3 vínculos tienen avance fuera de secuencia, en 3 tareas:")


# ----------------------------------------------------------------------------
# Lógica del cronograma (solo validar)
# ----------------------------------------------------------------------------
def logic(tasks, cfg=None, **project_kw):
    project = P(tasks, **project_kw)
    part = build_part(project, cfg or Config())
    return logic_checks(project, part.scope, part.code_of, part.rows)


def test_schedule_without_links():
    issues = logic(tree('''
        1 Fase
          1.1 A
    '''))
    assert len(issues) == 1 and "no tiene vínculos" in issues[0].message and issues[0].topic == TOPIC_LOGIC


def test_links_and_constraints():
    tasks = tree('''
        1 Fase
          1.1 Inicio
          1.2 Suelta
          1.3 Adelanto
          1.4 CF
          1.5 Fin
        2 Resumen
          2.1 Detalle
    ''')
    t = by_name(tasks)
    days = {"Inicio": 0, "Suelta": 10, "Adelanto": 20, "CF": 30, "Fin": 40, "Detalle": 25}
    for name, d in days.items():
        t[name].start = datetime(2026, 1, 5 + d % 20, 8) if d < 20 else datetime(2026, 2, d - 15, 8)
        t[name].finish = t[name].start.replace(hour=17)
    uid = {n: t[n].uid for n in t}
    t["Adelanto"].predecessors = [Link(uid["Inicio"], LINK_FS, -480)]
    t["CF"].predecessors = [Link(uid["Adelanto"], LINK_SF, 0)]
    t["Fin"].predecessors = [Link(uid["CF"], LINK_FS, 960), Link(uid["Detalle"], LINK_SS, 0)]
    t["Resumen"].predecessors = [Link(uid["Inicio"], LINK_FS, 0)]
    t["Detalle"].constraint_type, t["Detalle"].constraint_date = CONSTRAINT_MSO, datetime(2026, 2, 10, 8)
    issues = logic(tasks)
    by = {}
    for i in issues:
        by.setdefault(i.name, []).append(i.message)
    assert by["Suelta"] == ["Sin predecesoras ni sucesoras."]
    assert "Inicio" not in by and "Fin" not in by               # comienzo y fin del proyecto
    assert any("adelanto" in m for m in by["Adelanto"])
    assert any("comienzo a fin (CF)" in m for m in by["CF"])
    assert by["Resumen"] == ["Tarea resumen con vínculos: se recomienda vincular sus tareas de detalle."]
    assert by["Detalle"] == ["Delimitación «Debe comenzar el 10/02/2026»."]   # su predecesora es la de la resumen
    general = [i.message for i in issues if not i.code]
    assert any("1 tarea sin predecesoras y 1 tarea sin sucesoras" in m for m in general)
    assert any("1 vínculo tiene posición positiva" in m for m in general)
    assert any("son fin a comienzo (FC)" in m for m in general)
    assert any("1 tarea tiene delimitaciones" in m for m in general)


def test_float_and_duration():
    """DCMA 6, 7 y 8: holgura alta, holgura negativa y duración alta de las tareas de detalle sin terminar. Las filas
    LOE no cuentan para la holgura alta ni para la duración alta; los hitos, tampoco para la duración."""
    tasks = tree('''
        1 Fase
          1.1 Holgada
          1.2 Negativa
          1.3 Larga
          1.4 Esfuerzo
          1.5 Hito
          1.6 Terminada
    ''')
    t = by_name(tasks)
    day = 480
    for task in tasks[1:]:
        task.total_slack, task.baseline_duration = 0.0, 10.0 * day
    t["Negativa"].predecessors = [Link(t["Holgada"].uid)]
    t["Holgada"].total_slack = 60.0 * day
    t["Negativa"].total_slack = -5.5 * day
    t["Larga"].baseline_duration = 100.0 * day
    t["Esfuerzo"].method = "LOE"
    t["Esfuerzo"].total_slack, t["Esfuerzo"].baseline_duration = 90.0 * day, 200.0 * day
    t["Hito"].milestone, t["Hito"].baseline_duration = True, 0.0
    t["Terminada"].percent_complete, t["Terminada"].actual_start, t["Terminada"].actual_finish = 100, S, F
    t["Terminada"].total_slack, t["Terminada"].baseline_duration = 100.0 * day, 100.0 * day
    issues = logic(tasks)
    high = next(i for i in issues if i.id == "LOG-013")
    assert high.severity == SEV_WARN and high.topic == TOPIC_LOGIC
    assert high.message == ("1 tarea sin terminar tiene más de 44 días laborables de holgura total (25 % de las "
                            "tareas sin terminar; DCMA: hasta 5 %): 1.1.0 (60 días). Una holgura alta suele indicar "
                            "que faltan sucesoras.")
    negative = next(i for i in issues if i.id == "LOG-014")
    assert "1 tarea sin terminar tiene holgura negativa (20 % de las tareas sin terminar" in negative.message
    assert [(i.name, i.message) for i in issues if i.id == "LOG-015"] == [
        ("Negativa", "Holgura total de -5,5 días laborables.")]
    long = next(i for i in issues if i.id == "LOG-016")
    assert long.severity == SEV_WARN
    assert long.message.startswith("1 tarea sin terminar dura más de 44 días laborables en la línea base (33,3 % de "
                                   "las tareas sin terminar; DCMA: hasta 5 %): 1.3.0 (100 días).")
    current = next(i for i in logic(tasks, baseline_found=False) if i.id == "LOG-016")
    assert "en el programa actual" in current.message


def test_high_float_threshold_and_schedule_without_links():
    lines = ["1 Fase"] + [f"  1.{n} T{n}" for n in range(1, 22)]
    tasks = tree("\n".join(lines))
    for task in tasks[1:]:
        task.total_slack = 0.0
    tasks[1].total_slack = 45.0 * 480                               # 1 de 21: 4,8 %
    tasks[2].predecessors = [Link(tasks[1].uid)]
    high = next(i for i in logic(tasks) if i.id == "LOG-013")
    assert high.severity == SEV_INFO and "(4,8 % de las tareas sin terminar" in high.message
    tasks[2].predecessors = []                                      # sin vínculos: toda la holgura llega al fin
    assert not [i for i in logic(tasks) if i.id == "LOG-013"]


def test_validate_part_runs_dcma_checks():
    """validar: el CPLI usa el calendario del proyecto y la holgura alta, las filas del libro (LOE)."""
    tasks = missed_and_cpli_tasks(datetime(2026, 4, 10, 17))
    t = by_name(tasks)
    t["Futura"].predecessors = [Link(t["Sin terminar"].uid)]
    t["Futura"].total_slack = 50.0 * 480
    part = validate_part(P(tasks, status_date=datetime(2026, 2, 27, 17)), Config())
    ids = {i.id for i in part.checks}
    assert {"AVA-014", "AVA-015", "LOG-013"} <= ids
    t["Futura"].method = "Avance lineal (LOE)"
    part = validate_part(P(tasks, status_date=datetime(2026, 2, 27, 17)), Config())
    assert "LOG-013" not in {i.id for i in part.checks}


# ----------------------------------------------------------------------------
# Lectura de los campos nuevos
# ----------------------------------------------------------------------------
def test_duration_minutes():
    assert duration_minutes("PT40H0M0S") == 2400
    assert duration_minutes("PT7H30M0S") == 450
    assert duration_minutes("P1DT2H") == 1560
    assert duration_minutes("") is None and duration_minutes("x") is None


def test_xml_reads_schedule_fields(tmp_path):
    path = write_xml(tmp_path / "p.xml", project_xml([
        task_xml(1, 1, "A", wbs="1", start=S, finish=F, baselines=[(0, S, F, 0)],
                 extra="<Duration>PT480H0M0S</Duration><Estimated>1</Estimated><Manual>1</Manual>"
                       "<ConstraintType>2</ConstraintType><ConstraintDate>2026-01-05T08:00:00</ConstraintDate>"),
        task_xml(2, 1, "B", wbs="2", start=S, finish=F, baselines=[(0, S, F, 0)],
                 extra="<Duration>PT0H0M0S</Duration><PredecessorLink><PredecessorUID>1</PredecessorUID>"
                       "<Type>3</Type><LinkLag>-4800</LinkLag></PredecessorLink>"),
    ], definitions=[(TEXT10, "Text10", "Método EVM"), (NUMBER1, "Number1", "Peso hito")]))
    p = read_xml(path, Config(method_field="Método EVM", milestone_weight_field="Peso hito"))
    a, b = p.tasks
    assert (a.duration, a.estimated, a.manual, a.constraint_type) == (28800, True, True, CONSTRAINT_MSO)
    assert a.constraint_date == datetime(2026, 1, 5, 8) and a.predecessors == []
    assert b.duration == 0 and b.predecessors == [Link(1, LINK_SS, -4800)]
    assert p.fields_found == {"metodo": "Método EVM", "peso_hito": "Peso hito"}


def test_com_reads_schedule_fields():
    from test_reader_com_fake import FakeTask, make_session

    class Dep:
        def __init__(self, a, b, kind=LINK_FS, lag=0):
            self.From, self.To, self.Type, self.Lag = a, b, kind, lag

    a = FakeTask(1, 1, "A", "1")
    b = FakeTask(2, 1, "B", "2")
    a.Duration, a.Estimated, a.Manual, a.ConstraintType = 4800, True, False, 2
    a.ConstraintDate = datetime(2026, 1, 5, 8)
    b.Duration = "1 día?"                                       # no numérico: duración desconocida
    dep = Dep(a, b, LINK_SF, -60)
    a.TaskDependencies = [dep]
    b.TaskDependencies = [dep]
    p = make_session([a, b]).read_project(Config(method_field="Método EVM"))
    ta, tb = p.tasks
    assert (ta.duration, ta.estimated, ta.constraint_type, ta.predecessors) == (4800, True, 2, [])
    assert ta.constraint_date == datetime(2026, 1, 5, 8)
    assert tb.duration is None and tb.predecessors == [Link(1, LINK_SF, -60)]
    assert p.fields_found == {"metodo": "Método EVM"}


# ----------------------------------------------------------------------------
# Comando validar y hoja Validación
# ----------------------------------------------------------------------------
def problem_xml(tmp_path):
    return write_xml(tmp_path / "proyecto.xml", project_xml([
        task_xml(1, 1, "Fase", wbs="1", start=S, finish=F, baselines=[(0, S, F, 0)]),
        task_xml(2, 2, "Grupo", wbs="1.2", start=S, finish=F, baselines=[(0, S, F, 0)]),
        task_xml(3, 3, "Mal", wbs="1.14.2", start=S, finish=F, baselines=[(0, S, F, 0)]),
    ], header="<StatusDate>2026-02-27T17:00:00</StatusDate>"))


def test_cli_validar_reports_errors_and_creates_no_file(tmp_path, capsys):
    path = problem_xml(tmp_path)
    before = sorted(tmp_path.iterdir())
    assert main(["validar", "--origen", str(path)]) == 1
    out = capsys.readouterr().out
    assert "Estructura y códigos EDT: 1 error" in out
    assert "[X] EDT-005 1.14.2 Mal: El código EDT no depende" in out
    assert "Fecha de estado: 27/02/2026" in out
    assert "Lógica del cronograma: 1 advertencia" in out
    assert "No se creó ni se modificó ningún archivo." in out
    assert "El PV y el EV de cada fila" not in out                # recordatorio del libro
    assert sorted(tmp_path.iterdir()) == before


def test_cli_validar_without_errors_and_limit(tmp_path, capsys):
    tasks = [task_xml(1, 1, "Fase", wbs="1", start=S, finish=F, baselines=[(0, S, F, 0)])]
    tasks += [task_xml(i, 2, f"T{i}", wbs=f"1.{i}", start=S, finish=F, baselines=[(0, S, F, 0)],
                       extra="<Estimated>1</Estimated>") for i in range(2, 30)]
    path = write_xml(tmp_path / "p.xml", project_xml(tasks))
    assert main(["validar", "--origen", str(path), "--fecha-corte", "2026-01-31"]) == 0
    out = capsys.readouterr().out
    assert "Fecha de estado: 31/01/2026 (--fecha-corte)" in out
    assert "Fechas y duración: 29 advertencias" in out and "... y 9 más (--todas para verlas todas)." in out
    main(["validar", "--origen", str(path), "--todas"])
    assert "--todas para verlas todas" not in capsys.readouterr().out


def test_cli_validar_missing_plan_reviews_whole_project(tmp_path, capsys):
    path = problem_xml(tmp_path)
    cfg = tmp_path / "config.toml"
    cfg.write_text('[exportacion]\nplanificacion = "INTERNO"\n', encoding="utf-8")
    assert main(["validar", "--origen", str(path), "--config", str(cfg)]) == 1
    out = capsys.readouterr().out
    assert "No se encontró la planificación 'INTERNO'" in out and "Se revisa el proyecto completo." in out
    assert "1.14.2 Mal" in out


def test_cli_validar_several_files_with_repeated_codes(tmp_path, capsys):
    one = problem_xml(tmp_path)
    two = write_xml(tmp_path / "dos.xml", project_xml([
        task_xml(1, 1, "Otra fase", wbs="1", start=S, finish=F, baselines=[(0, S, F, 0)])]))
    cfg = tmp_path / "config.toml"
    cfg.write_text(f"[[proyectos]]\nnombre = \"P1\"\norigen = '{one}'\n\n"
                   f"[[proyectos]]\nnombre = \"P2\"\norigen = '{two}'\n", encoding="utf-8")
    assert main(["validar", "--config", str(cfg)]) == 1
    out = capsys.readouterr().out
    assert "El código EDT '1.0.0' está en los archivos 'P1' y 'P2'" in out
    assert "[P1] El código EDT no depende" in out


def test_part_keeps_scope_for_schedule_checks():
    tasks = tree('''
        1 Fase
          1.1 A
    ''')
    part = validate_part(P(tasks), Config(), "", date(2026, 2, 1))
    assert part.scope is not None and part.code_of(tasks[1]) == "1.1.0"
    assert any(i.topic == TOPIC_PROGRESS for i in part.checks) and any(i.topic == TOPIC_LOGIC for i in part.checks)


def test_generar_writes_topic_and_new_checks(tmp_path):
    path = problem_xml(tmp_path)
    out = tmp_path / "libro.xlsx"
    assert main(["generar", "--origen", str(path), "--salida", str(out)]) == 0
    ws = load_workbook(out)["Validación"]
    assert [ws.cell(row=3, column=c).value for c in (2, 6)] == ["Código", "Tema"]
    rows = [r for r in ws.iter_rows(min_row=4, values_only=True) if r[0]]
    assert ("Error", "EDT-005", "1.14.2", "Mal") == rows[[r[2] for r in rows].index("1.14.2")][:4]
    assert {r[5] for r in rows} >= {TOPIC_WBS}
    assert not [r for r in rows if r[5] in (TOPIC_PROGRESS, TOPIC_LOGIC)]     # solo en validar


def test_lead_allows_progress_before_the_predecessor_finishes():
    tasks = tree('''
        1 Fase
          1.1 Comenzada
          1.2 Con adelanto
    ''')
    t = by_name(tasks)
    t["Comenzada"].actual_start, t["Comenzada"].percent_complete = S, 40
    t["Con adelanto"].actual_start, t["Con adelanto"].percent_complete = S, 10
    t["Con adelanto"].predecessors = [Link(t["Comenzada"].uid, LINK_FS, -4800)]
    assert not [i for i in progress(tasks, None) if i.id in ("AVA-016", "AVA-017")]
    assert [i.id for i in logic(tasks) if i.id == "LOG-006"] == ["LOG-006"]       # el adelanto se informa aparte


def test_dcma_metrics_count_detail_tasks_only():
    tasks = tree('''
        1 Fase
          1.1 A
          1.2 B
    ''')
    t = by_name(tasks)
    t["Fase"].constraint_type = CONSTRAINT_MSO                  # delimitación en la tarea resumen
    t["B"].predecessors = [Link(t["A"].uid, LINK_FS)]
    t["Fase"].predecessors = [Link(t["A"].uid, LINK_SS)]        # vínculo de la resumen: LOG-010
    issues = logic(tasks)
    ids = [i.id for i in issues]
    assert "LOG-011" not in ids and "LOG-012" not in ids
    assert "LOG-009" not in ids                                 # el vínculo CC de la resumen no cuenta
    assert [i.name for i in issues if i.id == "LOG-010"] == ["Fase"]


def test_wbs_code_issues_are_capped_per_code(monkeypatch):
    import evm.validation as validation
    monkeypatch.setattr(validation, "LIMIT", 3)
    tasks = [T(1, 1, "1", "Fase")] + [T(i, 2, f"2.{i}", f"Mala {i}") for i in range(2, 9)]     # 7 × EDT-005
    tasks += [T(20, 1, "5", "Otra"), T(21, 2, "5.1", "A"), T(22, 2, "5.1", "B")]                # 1 × EDT-009
    issues = build_report(P(tasks), Config()).issues
    edt5 = [i for i in issues if i.id == "EDT-005"]
    assert [i.name for i in edt5] == ["Mala 2", "Mala 3", "Mala 4", ""]
    assert edt5[-1].message.startswith("Y 4 observaciones más con este código: se muestran las primeras 3.")
    assert edt5[-1].severity == SEV_ERROR
    assert len([i for i in issues if i.id == "EDT-009"]) == 1                # sin exceso: sin línea de resumen
