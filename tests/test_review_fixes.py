"""Pruebas de regresión de los problemas detectados en la revisión del código."""
from datetime import date, datetime

import pytest
from openpyxl import load_workbook

from evm.config import Config
from evm.indicators import earned_schedule, planned_duration
from evm.methods import parse_method
from evm.migrate import plan_migration, prior_from_book
from evm.model import ProjectData, Task
from evm.percent import RowWeights, TaskTime, earned_fraction
from evm.update import book_totals, read_book, run_update
from evm.wbs import KIND_RESIDUAL, build_report
from evm.workbook import (COL_BAC, FIRST_MONTH_COL, FIRST_ROW, WorkbookFormatError, _sum_formula, read_structure,
                          write_workbook)


S, F = datetime(2026, 1, 5, 8), datetime(2026, 3, 27, 17)


def T(uid, level, wbs, name, **kw):
    return Task(uid=uid, id=kw.pop("id", uid), name=name, outline_level=level, wbs=wbs,
                baseline_start=kw.pop("bs", S), baseline_finish=kw.pop("bf", F), has_baseline=True,
                start=S, finish=F, **kw)


def P(tasks, name="Programa", status=None, **kw):
    return ProjectData(name=name, source="p.xml", backend="xml", tasks=tasks, status_date=status, **kw)


def make_book(path, project, cfg=None, prior=None, bac=None):
    """Genera el libro e ingresa el BAC de las filas indicadas ({código: BAC})."""
    cfg = cfg or Config()
    rep = build_report(project, cfg)
    if prior is not None:
        rep.issues.extend(plan_migration(rep, prior, cfg.max_level))
    write_workbook(rep, cfg, path, prior)
    if bac:
        wb = load_workbook(path)
        rows = {s.code: s.row for s in read_structure(wb)}
        for code, value in bac.items():
            wb["PV"].cell(rows[code], COL_BAC, value)
        wb.save(path)
    return rep


def set_month(path, sheet, code, month_index, value):
    wb = load_workbook(path)
    ws = wb[sheet]
    row = next(r for r in range(FIRST_ROW, ws.max_row + 1) if ws.cell(r, 1).value == code)
    ws.cell(row=row, column=FIRST_MONTH_COL + month_index, value=value)
    wb.save(path)


def row_values(path, sheet, code):
    book = read_book(path)
    s = next(s for s in book.input_rows if s.code == code)
    return (book.ev if sheet == "EV" else book.ac)[s.row]


# ------------------------------------------------------------------ migración
def base_tasks():
    return [T(1, 1, "1", "Fase"), T(2, 2, "1.1", "B"), T(3, 2, "1.2", "C")]


BAC = {"1.1.0": 100, "1.2.0": 200}


def test_migration_insert_does_not_duplicate_history(tmp_path):
    old = tmp_path / "old.xlsx"
    make_book(old, P(base_tasks()), bac=BAC)
    set_month(old, "% EV", "1.1.0", 0, 0.1)                       # EV 10
    set_month(old, "% EV", "1.2.0", 0, 0.1)                       # EV 20
    # Se inserta una tarea nueva en 1.2: la tarea C (UID 3) pasa a 1.3
    new_tasks = [T(1, 1, "1", "Fase"), T(2, 2, "1.1", "B"), T(4, 2, "1.2", "Nueva"), T(3, 2, "1.3", "C")]
    new = tmp_path / "new.xlsx"
    prior = prior_from_book(old, "Programa")
    make_book(new, P(new_tasks), prior=prior)
    assert sum(v or 0 for v in row_values(new, "EV", "1.3.0")) == pytest.approx(20)
    assert sum(v or 0 for v in row_values(new, "EV", "1.2.0")) == 0
    assert sum(book_totals(read_book(new))[1]) == pytest.approx(30)


def test_migration_reports_deleted_task_and_keeps_methods(tmp_path):
    old = tmp_path / "old.xlsx"
    make_book(old, P(base_tasks()), bac=BAC)
    set_month(old, "% EV", "1.2.0", 0, 0.1)
    wb = load_workbook(old)
    ws = wb["PV"]
    r = next(r for r in range(FIRST_ROW, ws.max_row + 1) if ws.cell(r, 1).value == "1.1.0")
    ws.cell(r, 3).value = "0/100"
    wb.save(old)
    prior = prior_from_book(old, "Programa")
    new = tmp_path / "new.xlsx"
    rep = make_book(new, P([T(1, 1, "1", "Fase"), T(2, 2, "1.1", "B")]), prior=prior)
    assert [u.code for u in prior.unmatched] == ["1.2.0"]
    assert any(i.severity == "Error" and i.code == "1.2.0" for i in rep.issues)
    ws = load_workbook(new)["PV"]
    r = next(r for r in range(FIRST_ROW, ws.max_row + 1) if ws.cell(r, 1).value == "1.1.0")
    assert ws.cell(r, 3).value == "0/100"


def test_migration_lower_level_aggregates_history(tmp_path):
    old = tmp_path / "old.xlsx"
    make_book(old, P(base_tasks()), Config(max_level=2), bac={"1.1": 100, "1.2": 200})
    set_month(old, "% EV", "1.1", 0, 0.1)                         # EV 10
    set_month(old, "AC", "1.2", 1, 5)
    prior = prior_from_book(old, "Programa")
    new = tmp_path / "new.xlsx"
    make_book(new, P(base_tasks()), Config(max_level=1), prior=prior)
    book = read_book(new)
    _, ev, ac, bac = book_totals(book)
    assert sum(ev) == pytest.approx(10) and sum(ac) == 5 and bac == 300


# ------------------------------------------------------------------ actualizar
def progress_tasks():
    tasks = base_tasks()
    tasks[1].percent_complete = 50
    tasks[2].percent_complete = 25
    return tasks


def test_update_requires_status_date(tmp_path):
    book = tmp_path / "b.xlsx"
    make_book(book, P(base_tasks()))
    with pytest.raises(ValueError, match="fecha de estado"):
        run_update(P(progress_tasks()), Config(), book, backup=False)


def test_update_skips_deleted_uid(tmp_path):
    book = tmp_path / "b.xlsx"
    make_book(book, P(base_tasks()), bac=BAC)
    tasks = progress_tasks()
    tasks.pop(2)                                                     # C eliminada
    tasks.append(T(9, 2, "1.2", "Otra", percent_complete=80))        # nueva, ocupa el código de C
    res = run_update(P(tasks), Config(), book, date(2026, 1, 31), backup=False)
    assert "ya no existe" in " ".join(w.message for w in res.warnings)
    _, ev, _, _ = book_totals(res.book)
    assert sum(ev) == pytest.approx(50)                              # la tarea nueva no gana el EV de C


def test_dry_run_reports_simulated_values_without_saving(tmp_path):
    book = tmp_path / "b.xlsx"
    make_book(book, P(base_tasks()), bac=BAC)
    before = book.read_bytes()
    res = run_update(P(progress_tasks()), Config(), book, date(2026, 1, 31), dry_run=True)
    assert book.read_bytes() == before
    _, ev, _, _ = book_totals(res.book)
    assert sum(ev) == pytest.approx(100 * 0.5 + 200 * 0.25)


def test_text_numbers_and_formulas_in_input_cells(tmp_path):
    book = tmp_path / "b.xlsx"
    make_book(book, P(base_tasks()))
    set_month(book, "AC", "1.1.0", 0, "60")
    with pytest.raises(WorkbookFormatError, match="texto"):
        read_book(book)
    res = run_update(P(progress_tasks()), Config(), book, date(2026, 2, 28), backup=False)
    assert any("texto o fórmulas" in w.message for w in res.warnings)
    wb = load_workbook(book)
    assert any(isinstance(c.value, (int, float)) and c.value == 60
               for c in wb["AC"].iter_rows(min_row=FIRST_ROW, min_col=FIRST_MONTH_COL, max_col=FIRST_MONTH_COL)
               for c in c)
    read_book(book)                                                  # ya no hay texto


def test_backups_never_overwrite(tmp_path):
    book = tmp_path / "b.xlsx"
    make_book(book, P(base_tasks()))
    status = date(2026, 1, 31)
    r1 = run_update(P(progress_tasks()), Config(), book, status)
    r2 = run_update(P(progress_tasks()), Config(), book, status)
    assert r1.backup != r2.backup and r1.backup.exists() and r2.backup.exists()


# ------------------------------------------------------------------ estructura
def test_summary_tasks_do_not_form_other_tasks_rows():
    tasks = [T(1, 1, "1", "S1"), T(2, 2, "1.1", "S2"), T(3, 3, "1.1.1", "X"), T(4, 3, "1.1.2", "Y"),
             T(5, 2, "1.2", "L")]
    rep = build_report(P(tasks), Config())
    assert [r.code for r in rep.rows] == ["1.0.0", "1.1.0", "1.1.1", "1.1.2", "1.2.0"]
    assert not [r for r in rep.rows if r.kind == KIND_RESIDUAL]


def test_inactive_children_do_not_form_other_tasks_rows():
    tasks = [T(1, 1, "1", "Fase"), T(2, 2, "1.1", "A"), T(3, 2, "1.2", "B", active=False)]
    rep = build_report(P(tasks), Config())
    assert [r.code for r in rep.rows] == ["1.0.0", "1.1.0"]
    assert not [r for r in rep.rows if r.kind == KIND_RESIDUAL]


def test_long_sum_formulas_are_compressed_or_use_sumifs():
    assert _sum_formula("G", [5, 6, 7, 9]) == "=SUM(G5:G7,G9)"
    f = _sum_formula("DV", list(range(5, 3005, 2)), (5, 3004))
    assert f.startswith("=SUMIFS(DV5:DV3004,'_estructura'!$E$2:$E$3001") and len(f) < 200


# ------------------------------------------------------------------ métodos
def test_percent_fixed_formulas_and_single_leaf_physical():
    assert parse_method("0%/100%").start_fraction == 0.0
    assert parse_method("50 % / 50 %").start_fraction == 0.5
    leaf = Task(uid=2, id=2, name="A", outline_level=4, physical_percent_complete=60)
    rw = RowWeights({2: 1.0}, [TaskTime(leaf, 480, {(2026, 1): 1.0}, True)])
    assert earned_fraction(parse_method("% Físico completado"), rw, 0)[0] == pytest.approx(0.6)
    frac, note = earned_fraction(parse_method("Hitos ponderados"), rw, 0)
    assert frac == pytest.approx(0.6) and "físico" in note.message


# ------------------------------------------------------------------ cronograma ganado
def test_earned_schedule_plateaus_and_horizon():
    cum = [104.76, 200, 200, 200, 300]
    assert earned_schedule(cum, 200, 300, at=2) == pytest.approx(2.0)      # en el plan
    assert earned_schedule(cum, 200, 300, at=5) == pytest.approx(4.0)      # tramo plano: máximo 4
    assert earned_schedule([0, 0, 100], 0, 100, at=1) == pytest.approx(1.0)
    assert earned_schedule([100, 200, 300], 350, 500, at=3) is None        # la curva no llega
    assert planned_duration([100, 200], 500) is None


def test_earned_schedule_formula_matches_python_in_spreadsheet(tmp_path, recalc):
    """Proyecto con un mes sin PV: el ES de la hoja Resumen debe coincidir con Python."""
    from evm.update import indicators_for
    tasks = [T(1, 1, "1", "A", bs=datetime(2026, 1, 5, 8), bf=datetime(2026, 1, 30, 17)),
             T(2, 1, "2", "B", bs=datetime(2026, 4, 6, 8), bf=datetime(2026, 4, 30, 17))]
    tasks[0].percent_complete = 100
    book = tmp_path / "es.xlsx"
    make_book(book, P(tasks), Config(months=12), bac={"1.0.0": 300, "2.0.0": 300})
    run_update(P(tasks), Config(months=12), book, date(2026, 2, 28), backup=False)
    ind = indicators_for(read_book(book), date(2026, 2, 28))
    assert ind.es == pytest.approx(2.0)                                     # EV en el tramo plano, AT = 2
    ws = load_workbook(recalc(book), data_only=True)["Resumen"]
    assert ws["B35"].value == pytest.approx(ind.es)
    assert ws["B38"].value == pytest.approx(ind.spi_t)


# ------------------------------------------------------------------ COM
def test_com_open_passes_all_arguments_and_activates_open_project():
    pythoncom = pytest.importorskip("pythoncom")
    from evm.reader_com import PJ_POOL_READ_ONLY, ComSession

    class Proj:
        def __init__(self, name):
            self.FullName = name
            self.activated = 0

        def Activate(self):
            self.activated += 1

    class Projects:
        def __init__(self, items):
            self.items = items

        @property
        def Count(self):
            return len(self.items)

        def Item(self, i):
            return self.items[i - 1]

    class App:
        def __init__(self, items):
            self.Projects = Projects(items)
            self.ActiveProject = items[0] if items else None
            self.calls = []
            self.DisplayAlerts = True
            self.AutomationSecurity = 1

        def FileOpenEx(self, *args):
            self.calls.append(args)
            p = Proj(args[0])
            self.Projects.items.append(p)
            self.ActiveProject = p
            return True

    other = Proj(r"C:\otro.mpp")
    s = ComSession(r"C:\programa.mpp")
    s.app, s._pythoncom, s.started_by_us = App([other]), pythoncom, False
    s._open()
    args = s.app.calls[0]
    assert len(args) == 17 and args[1] is True and args[6] is True
    assert args[11] == PJ_POOL_READ_ONLY and args[14] is True and args[16] is True
    assert s.app.DisplayAlerts is False
    s.opened_by_us = False                     # evita FileCloseEx en el objeto simulado
    s._pythoncom = None
    s._cleanup()
    assert other.activated == 1                # se reactivó el proyecto que estaba activo

    # Proyecto ya abierto (no activo): se activa antes de leer
    target = Proj(str(__import__("pathlib").Path(r"C:\programa.mpp").resolve()))
    s2 = ComSession(r"C:\programa.mpp")
    s2.app, s2._pythoncom, s2.started_by_us = App([other, target]), pythoncom, False
    s2._open()
    assert s2.project is target and target.activated == 1 and not s2.app.calls


# ------------------------------------------------------------------ segunda revisión
def test_names_starting_with_equals_are_text(tmp_path):
    import zipfile
    tasks = [T(1, 1, "1", "=== Fase 1 ==="), T(2, 2, "1.1", "=PMO (oficina)")]
    book = tmp_path / "eq.xlsx"
    make_book(book, P(tasks, name="=Programa"))
    wb = load_workbook(book)
    assert wb["PV"].cell(FIRST_ROW, 2).value == "=== Fase 1 ==="
    assert wb["PV"].cell(FIRST_ROW, 2).data_type == "s"
    with zipfile.ZipFile(book) as z:
        sheets = [n for n in z.namelist() if n.startswith("xl/worksheets/sheet")]
        assert not any("<f>== Fase" in z.read(n).decode() or "<f>PMO" in z.read(n).decode() for n in sheets)
    read_book(book)


def test_migration_uses_new_method_changed_in_ms_project(tmp_path):
    old = tmp_path / "old.xlsx"
    tasks = base_tasks()
    tasks[1].method = "Unidades completadas"
    make_book(old, P(tasks))
    new_tasks = base_tasks()
    new_tasks[1].method = "50/50"
    new = tmp_path / "new.xlsx"
    make_book(new, P(new_tasks), prior=prior_from_book(old, "Programa"))
    ws = load_workbook(new)["PV"]
    r = next(r for r in range(FIRST_ROW, ws.max_row + 1) if ws.cell(r, 1).value == "1.1.0")
    assert ws.cell(r, 3).value == "50/50"


def test_listas_note_of_old_books_is_not_a_method(tmp_path):
    old = tmp_path / "old.xlsx"
    make_book(old, P(base_tasks()))
    wb = load_workbook(old)
    ws = wb["Listas"]
    n = ws.max_row
    ws.cell(n + 2, 1).value = "Puede agregar o modificar métodos en esta lista (columna A). Nota antigua larga."
    wb.save(old)
    assert not any("Nota antigua" in m for m in prior_from_book(old, "Programa").methods_list)


def test_earned_schedule_tolerates_rounding_noise():
    cum = [458051.04, 469364.71, 487729.80999999994, 487729.80999999994, 487729.80999999994,
           487729.80999999994, 600000.0]
    assert earned_schedule(cum, 487729.81, 600000.0, at=4) == pytest.approx(4.0)
    assert earned_schedule([87709.24, 143533.52, 200300.81000000003, 200300.81000000003, 300000.0],
                           200300.81, 300000.0, at=4) == pytest.approx(4.0)


def test_formula_in_bac_cell_is_replaced_by_its_value(tmp_path, recalc):
    book = tmp_path / "f.xlsx"
    make_book(book, P(base_tasks()))
    wb = load_workbook(book)
    ws = wb["PV"]
    r = next(r for r in range(FIRST_ROW, ws.max_row + 1) if ws.cell(r, 1).value == "1.1.0")
    ws.cell(r, 4).value = "=50+50"
    wb.save(book)
    calc = recalc(book)                                     # guarda valores calculados, como Excel
    res = run_update(P(progress_tasks()), Config(), calc, date(2026, 1, 31), backup=False)
    assert any("texto o fórmulas" in w.message for w in res.warnings)
    assert load_workbook(calc)["PV"].cell(r, 4).value == 100
    read_book(calc)                                          # el siguiente uso no falla
