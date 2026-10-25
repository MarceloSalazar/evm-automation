"""Plazo según el cronograma de MS Project (fin previsto, desvío, holgura y camino crítico) y tendencia de los
índices (últimos 3 meses y mes anterior con datos)."""
from datetime import date, datetime

import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from evm import analysis, dashboard
from evm.config import Config
from evm.indicators import compute
from evm.schedule import FINISHED, RowSchedule, ScheduleSummary, combine, row_schedule, working_days
from evm.update import indicators_for, read_book, run_analysis, run_update
from evm.workbook import (I_BL_FINISH, I_CRITICAL, I_SLACK, I_VARIANCE, ROW_BL_FINISH, ROW_CRITICAL, ROW_FINISH,
                          ROW_LATE, ROW_PLAN_STATUS, ROW_VARIANCE, SHEET_IND, SHEET_META, SHEET_SUMMARY, read_meta,
                          write_meta, write_workbook)
from evm.workcalendar import WorkCalendar
from evm.wbs import build_report

from builders import P, T, set_bac

CAL = WorkCalendar()
DAY = 480.0     # minutos de un día laborable


def task(uid, *, bf, f, slack=None, critical=False, done=False, **kw):
    t = T(uid, 2, f"1.{uid}", f"T{uid}", bf=bf, finish=f, **kw)
    t.total_slack = None if slack is None else slack * DAY
    t.critical = critical
    if done:
        t.percent_complete, t.actual_finish = 100.0, f
    return t


# ----------------------------------------------------------------------------
# Cálculo
# ----------------------------------------------------------------------------
def test_working_days_signed_and_without_weekends_or_holidays():
    mon, wed, next_mon = date(2026, 3, 2), date(2026, 3, 4), date(2026, 3, 9)
    assert working_days(CAL, mon, wed) == 2
    assert working_days(CAL, wed, mon) == -2
    assert working_days(CAL, mon, next_mon) == 5          # sábado y domingo no cuentan
    assert working_days(CAL, datetime(2026, 3, 2, 17), datetime(2026, 3, 2, 8)) == 0
    cal = WorkCalendar()
    cal.add_holidays([date(2026, 3, 3)])
    assert working_days(cal, mon, wed) == 1


def test_row_schedule_uses_latest_finish_lowest_open_slack_and_open_critical_tasks():
    tasks = [task(1, bf=datetime(2026, 3, 2, 17), f=datetime(2026, 3, 6, 17), slack=12, critical=False),
             task(2, bf=datetime(2026, 3, 9, 17), f=datetime(2026, 3, 11, 17), slack=-2, critical=True),
             task(3, bf=datetime(2026, 1, 9, 17), f=datetime(2026, 1, 9, 17), slack=0, critical=True, done=True)]
    s = row_schedule(tasks, CAL, DAY)
    assert s.baseline_finish == datetime(2026, 3, 9, 17) and s.finish == datetime(2026, 3, 11, 17)
    assert s.variance == 2
    assert s.slack == -2 and s.critical and s.status == "Sí" and s.late_critical
    # Las tareas terminadas no cuentan para la holgura ni el camino crítico
    done = row_schedule(tasks[2:], CAL, DAY)
    assert done.status == FINISHED and done.slack is None and not done.critical
    # Sin línea base: sin desvío
    t = task(4, bf=None, f=datetime(2026, 3, 11, 17), bs=None)
    assert row_schedule([t], CAL, DAY).variance is None


def test_summary_rows_and_total_combine_their_rows():
    a = RowSchedule(datetime(2026, 3, 2), datetime(2026, 3, 4), 2, 5.0, False, True)
    b = RowSchedule(datetime(2026, 4, 1), datetime(2026, 4, 8), 5, 0.0, True, True)
    c = RowSchedule(datetime(2026, 2, 2), datetime(2026, 2, 2), 0, None, False, False)
    assert combine([a, b], CAL).variance == 5 and combine([a, b], CAL).slack == 0.0
    plan = ScheduleSummary.build(date(2026, 3, 31), {6: a, 7: b, 9: c}, [6, 7, 9], {5: [6, 7], 8: [9]},
                                 {}, CAL)
    assert plan.rows[5].finish == datetime(2026, 4, 8) and plan.rows[5].critical
    assert plan.total.finish == datetime(2026, 4, 8) and plan.total.variance == 5
    assert (plan.critical, plan.late_critical) == (1, 1)


# ----------------------------------------------------------------------------
# Libro
# ----------------------------------------------------------------------------
def schedule_tasks():
    """1 Fase (resumen): 1.1 crítica atrasada 3 días; 1.2 con holgura y a tiempo; 2 terminada."""
    t = [T(1, 1, "1", "Fase"),
         T(2, 2, "1.1", "A", bf=datetime(2026, 3, 27, 17), finish=datetime(2026, 4, 1, 17)),
         T(3, 2, "1.2", "B"),
         T(4, 1, "2", "Fase 2", bs=datetime(2026, 1, 5, 8), bf=datetime(2026, 2, 27, 17),
           finish=datetime(2026, 2, 27, 17), percent_complete=100.0, actual_finish=datetime(2026, 2, 27, 17))]
    t[1].total_slack, t[1].critical = 0.0, True
    t[2].total_slack = 15 * DAY
    t[3].total_slack, t[3].critical = 0.0, True
    return t


def test_generar_writes_schedule_columns_summary_and_new_dashboard(tmp_path):
    cfg = Config(max_level=2)
    book = tmp_path / "libro.xlsx"
    project = P(schedule_tasks(), status_date=datetime(2026, 3, 31))
    write_workbook(build_report(project, cfg), cfg, book)
    wb = load_workbook(book)
    ind = wb[SHEET_IND]
    rows = {ind.cell(r, 1).value: r for r in range(5, 9)}
    assert [ind.cell(4, c).value for c in range(I_BL_FINISH, I_CRITICAL + 1)] == [
        "Fin línea base", "Fin previsto", "Desvío fin (días lab.)", "Holgura (días lab.)", "Crítica"]
    a, b, phase = rows["='PV'!A6"], rows["='PV'!A7"], rows["='PV'!A5"]
    assert ind.cell(a, I_VARIANCE).value == 3 and ind.cell(a, I_SLACK).value == 0 and ind.cell(a, I_CRITICAL).value == "Sí"
    assert ind.cell(b, I_VARIANCE).value == 0 and ind.cell(b, I_SLACK).value == 15 and ind.cell(b, I_CRITICAL).value == "No"
    assert ind.cell(phase, I_VARIANCE).value == 3 and ind.cell(phase, I_CRITICAL).value == "Sí"
    assert ind.cell(rows["='PV'!A8"], I_CRITICAL).value == "Terminado"
    res = wb[SHEET_SUMMARY]
    assert res.cell(ROW_PLAN_STATUS, 2).value == datetime(2026, 3, 31)
    assert res.cell(ROW_BL_FINISH, 2).value == datetime(2026, 3, 27, 17)
    assert res.cell(ROW_FINISH, 2).value == datetime(2026, 4, 1, 17)
    assert res.cell(ROW_VARIANCE, 2).value == 3
    assert (res.cell(ROW_CRITICAL, 2).value, res.cell(ROW_LATE, 2).value) == (1, 1)
    meta = read_meta(wb)
    assert meta["disposicion_cuadro_mando"] == dashboard.LAYOUT
    dash = wb[dashboard.SHEET_DASHBOARD]
    assert dash.cell(dashboard.ROW_TREND, 2).value.startswith("Tendencia")
    assert dash.cell(dashboard.ROW_PLAN, 2).value.startswith("Plazo")
    charts, _end, scatter, _send = dashboard.chart_rows()
    assert [c.anchor._from.row for c in dash._charts] == [charts - 1, charts - 1, scatter - 1]


def test_actualizar_refreshes_schedule_and_analysis_and_upgrades_old_books(tmp_path):
    cfg = Config(max_level=2)
    tasks = schedule_tasks()
    book = tmp_path / "libro.xlsx"
    write_workbook(build_report(P(tasks), cfg), cfg, book)
    set_bac(book)
    # Libro de una versión anterior: sin disposición del Cuadro de Mando ni columnas de plazo
    wb = load_workbook(book)
    meta = read_meta(wb)
    meta.pop("disposicion_cuadro_mando")
    write_meta(wb[SHEET_META], meta)
    wb[SHEET_IND].delete_cols(I_BL_FINISH, 5)
    wb.save(book)

    tasks[1].finish = datetime(2026, 4, 3, 17)          # el atraso crece a 5 días
    res = run_update(P(tasks, status_date=datetime(2026, 3, 31)), cfg, book, backup=False)
    assert res.schedule.total.variance == 5 and res.schedule.late_critical == 1
    wb = load_workbook(book)
    ind = wb[SHEET_IND]
    a = next(r for r in range(5, 9) if ind.cell(r, 1).value == "='PV'!A6")
    assert ind.cell(4, I_BL_FINISH).value == "Fin línea base" and ind.cell(a, I_VARIANCE).value == 5
    assert "actualizar" in ind.cell(3, I_BL_FINISH).value
    # Análisis de desviaciones: el paquete atrasado tiene el plazo del mes
    ws = wb[analysis.SHEET_ANALYSIS]
    entries = {e.code: e for e in analysis.read_entries(ws)}
    assert entries["1.1"].plan == (5, 0, "Sí")
    assert ws.cell(analysis.HEAD_ROW, analysis.A_TREND).value == "Tendencia"
    # Libro anterior: los gráficos quedan en las filas de su disposición
    charts, _end, scatter, _send = dashboard.chart_rows(1)
    assert [c.anchor._from.row for c in wb[dashboard.SHEET_DASHBOARD]._charts] == [charts - 1, charts - 1,
                                                                                   scatter - 1]
    # 'analizar' conserva el plazo de la última lectura
    run_analysis(cfg, book, backup=False)
    entries = {e.code: e for e in analysis.read_entries(load_workbook(book)[analysis.SHEET_ANALYSIS])}
    assert entries["1.1"].plan == (5, 0, "Sí")


def test_analysis_sheet_of_previous_versions_is_read_by_position(tmp_path):
    """Hoja de la versión 2.2: 19 columnas, con la clave en la S."""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = analysis.SHEET_ANALYSIS
    ws.column_dimensions["S"].hidden = True
    old = ["Mes", "Código EDT", "Actividad", "BAC", "PV acum.", "EV acum.", "AC acum.", "SV", "CV", "SPI", "CPI",
           "Semáforo", "Antecedente", "Causa", "Impacto en plazo y costo", "Acción correctiva", "Responsable",
           "Fecha compromiso", "clave"]
    for c, text in enumerate(old, start=1):
        ws.cell(analysis.HEAD_ROW, c, text)
    ws.append([datetime(2027, 6, 1), "1.1.1", "A", "=x", None, None, None, None, None, None, None, None,
               "Nuevo", "Lluvias", "2 semanas", "Turno extra", "JP", datetime(2027, 7, 15), "12"])
    ws.append([datetime(2027, 6, 1), "9.9.9", "Agregada a mano", 5, None, None, None, None, None, None, None, None,
               None, "Nota"])
    entries = analysis.read_entries(ws)
    assert entries[0].key == "12" and entries[0].texts[0] == "Lluvias" and entries[0].origin == "Nuevo"
    assert entries[0].plan == (None, None, None)
    raw = entries[1].raw
    assert raw[analysis.A_BAC - 1] == 5 and raw[analysis.A_CAUSE - 1] == "Nota"
    # Con los encabezados de la versión actual, por nombre
    analysis.write_sheet(wb, entries[:1], {"12": 6}, 12, analysis.thresholds({}, Config()), 0, date(2027, 6, 1), 1)
    again = analysis.read_entries(wb[analysis.SHEET_ANALYSIS])
    assert again[0].texts[0] == "Lluvias" and again[0].key == "12"
    sheet = wb[analysis.SHEET_ANALYSIS]
    assert not sheet.column_dimensions["S"].hidden                      # ahora es la columna Antecedente
    assert sheet.column_dimensions[get_column_letter(analysis.A_KEY)].hidden


# ----------------------------------------------------------------------------
# Tendencia
# ----------------------------------------------------------------------------
def test_trend_uses_last_loaded_months():
    pv = [10.0] * 8
    ev = [9.0] * 8
    ac = [10.0] * 8
    # Todos los meses con datos: mes anterior = mes 5 (índice 4); inicio de los últimos 3 meses = índice 2
    ind = compute(pv, ev, ac, 100.0, 5)
    assert ind.cpi_prev == pytest.approx(0.9) and ind.spi_prev == pytest.approx(0.9)
    assert ind.cpi_3m == pytest.approx(27 / 30) and ind.eac_prev == pytest.approx(100 / 0.9)
    # Informes bimestrales: el mes anterior con datos es el del informe anterior
    loaded = [True, False, True, False, True, False, True, False]
    ev2 = [9.0, 0.0, 18.0, 0.0, 14.0, 0.0, 20.0, 0.0]
    ind = compute(pv, ev2, ac, 100.0, 6, loaded)
    assert ind.spi_prev == pytest.approx(41 / 50)              # acumulado del mes 5 (índice 4)
    assert ind.spi_3m == pytest.approx((61 - 27) / (70 - 30))   # desde el índice 2 (último con datos hasta 3 antes)
    # Primer mes cargado: sin mes anterior ni últimos 3 meses
    first = [False] * 5 + [True, False, False]
    ind = compute(pv, ev, ac, 100.0, 5, first)
    assert ind.cpi_prev is None and ind.spi_3m is None and ind.eac_prev is None


def test_dashboard_and_summary_trend_recalculate(tmp_path, example_files, generator, recalc):
    """Las fórmulas de tendencia del Resumen, el Cuadro de Mando y el análisis dan lo mismo que 'actualizar'."""
    from evm.cli import main
    from evm.config import load_config
    from evm.reader_xml import read_xml
    from evm.workbook import ROW_CPI1, ROW_CPI3, ROW_EAC1, ROW_SPI1, ROW_SPI3
    book = tmp_path / "EVM.xlsx"
    assert main(["generar", "--origen", str(example_files["base"]), "--config", str(example_files["config"]),
                 "--salida", str(book)]) == 0
    generator.fill_book(book)
    cfg = load_config(example_files["config"])
    for month, status in ((4, date(2027, 4, 30)), (5, date(2027, 5, 31)), (6, date(2027, 6, 30)),
                          (7, date(2027, 7, 31))):
        xml = tmp_path / f"av{month}.xml"
        generator.write(generator.build_plans(), xml, status)
        generator.fill_book(book, status)
        run_update(read_xml(xml, cfg), cfg, book, status, backup=False)
    ind = indicators_for(read_book(book), date(2027, 7, 31))
    assert ind.cpi_3m is not None and ind.spi_prev is not None
    out = load_workbook(recalc(book), data_only=True)
    res = out[SHEET_SUMMARY]
    for row, value in ((ROW_CPI3, ind.cpi_3m), (ROW_SPI3, ind.spi_3m), (ROW_CPI1, ind.cpi_prev),
                       (ROW_SPI1, ind.spi_prev), (ROW_EAC1, ind.eac_prev)):
        assert res.cell(row, 2).value == pytest.approx(value, rel=1e-9), row
    dash = out[dashboard.SHEET_DASHBOARD]
    r = dashboard.ROW_TREND + 1
    assert dash.cell(r + 1, 2).value == pytest.approx(ind.cpi_3m, rel=1e-9)
    assert dash.cell(r + 1, 6).value == pytest.approx(ind.spi_3m, rel=1e-9)
    assert dash.cell(r + 1, 10).value == pytest.approx(round(ind.cpi - ind.cpi_prev, 2), abs=1e-9)
    assert dash.cell(r + 1, 14).value == pytest.approx(round(ind.spi - ind.spi_prev, 2), abs=1e-9)
    assert dash.cell(r + 1, 18).value == pytest.approx(round(ind.eac_cpi - ind.eac_prev), abs=1e-6)
    plan = dashboard.ROW_PLAN + 1
    assert dash.cell(plan + 1, 10).value == res.cell(ROW_VARIANCE, 2).value
    assert dash.cell(plan + 1, 14).value == res.cell(ROW_CRITICAL, 2).value
    # Análisis: SPI y CPI del mes anterior de cada paquete, con los valores de la hoja PV, EV y AC
    an = out[analysis.SHEET_ANALYSIS]
    trends = {an.cell(i, analysis.A_TREND).value for i in range(analysis.FIRST, an.max_row + 1)}
    assert trends <= {"Mejora", "Estable", "Empeora", None}
    b = read_book(book)
    row_of = {s.key: s.row for s in b.input_rows}
    june = b.horizon.index(date(2027, 6, 30))                  # mes anterior con datos del análisis de julio
    checked = 0
    for i in range(analysis.FIRST, an.max_row + 1):
        r = row_of.get(an.cell(i, analysis.A_KEY).value)
        pv = sum(b.pv[r][: june + 1]) if r else 0.0
        if pv:
            ev = sum(v or 0.0 for v in b.ev[r][: june + 1])
            assert an.cell(i, analysis.A_SPI_PREV).value == pytest.approx(ev / pv, abs=1e-9)
            checked += 1
    assert checked
    errors = [(ws.title, c.coordinate, c.value) for ws in out.worksheets for row in ws.iter_rows() for c in row
              if isinstance(c.value, str) and c.value.startswith("#") and c.value != "#N/A"]
    assert not errors, errors[:10]
