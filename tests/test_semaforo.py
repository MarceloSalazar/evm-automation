"""Semáforo de SPI y CPI: umbrales inferiores y superiores (configuración, hojas, gráficos y análisis)."""
import re
import zipfile

import pytest
from openpyxl import load_workbook

from evm import dashboard
from evm.cli import main
from evm.config import Config, ConfigError, load_config
from evm.workbook import Thresholds

CUSTOM = "[semaforo]\nrojo_inferior = 0.85\namarillo_inferior = 0.9\namarillo_superior = 1.1\nrojo_superior = 1.2\n"


def config(tmp_path, text: str, name: str = "c.toml"):
    path = tmp_path / name
    path.write_text(text, encoding="utf-8")
    return path


@pytest.fixture()
def book(tmp_path, example_files):
    """Libro del ejemplo con umbrales propios (0,85 / 0,90 / 1,10 / 1,20)."""
    text = example_files["config"].read_text(encoding="utf-8") + "\n" + CUSTOM
    cfg = config(tmp_path, text, "propio.toml")
    out = tmp_path / "EVM.xlsx"
    assert main(["generar", "--origen", str(example_files["base"]), "--config", str(cfg), "--salida", str(out)]) == 0
    return out


def test_config_defaults_and_keys(tmp_path):
    th = Thresholds.from_config(Config())
    assert (th.red_low, th.yellow_low, th.yellow_high, th.red_high) == (0.9, 0.95, 1.05, 1.1)
    cfg = load_config(config(tmp_path, CUSTOM))
    assert Thresholds.from_config(cfg) == Thresholds(0.85, 0.9, 1.1, 1.2)
    # Nombres anteriores de los umbrales inferiores
    cfg = load_config(config(tmp_path, "[semaforo]\nrojo = 0.8\namarillo = 0.9\n"))
    assert Thresholds.from_config(cfg) == Thresholds(0.8, 0.9, 1.05, 1.1)


@pytest.mark.parametrize("text, message", [
    ("[semaforo]\nrojo = 0.9\nrojo_inferior = 0.9\n", "misma clave"),
    ("[semaforo]\namarillo_superior = 0.95\n", "amarillo_inferior < amarillo_superior"),
    ("[semaforo]\nrojo_superior = 1.0\n", "amarillo_superior <= rojo_superior"),
    ("[semaforo]\nrojo_inferior = 0\n", "0 < rojo_inferior"),
    ("[semaforo]\namarillo_superior = true\n", "numérico"),
])
def test_config_errors(tmp_path, text, message):
    with pytest.raises(ConfigError, match=re.escape(message)):
        load_config(config(tmp_path, text))


@pytest.mark.parametrize("value, light", [
    (None, ""), (0.84, "Rojo"), (0.85, "Amarillo"), (0.89, "Amarillo"), (0.9, "Verde"), (1.0, "Verde"),
    (1.1, "Verde"), (1.11, "Amarillo"), (1.2, "Amarillo"), (1.21, "Rojo"),
])
def test_light_boundaries(value, light):
    """En el umbral exacto se aplica el color más favorable (los rangos son cerrados)."""
    assert Thresholds(0.85, 0.9, 1.1, 1.2).light(value) == light


def test_dashboard_shows_the_thresholds(book):
    wb = load_workbook(book)
    ws = wb[dashboard.SHEET_DASHBOARD]
    assert ws["B17"].value == "Semáforo CPI y SPI"
    assert [ws.cell(17, c).value for c in (5, 8, 11, 15, 18)] == [
        "< 0,85", "0,85 a 0,90", "0,90 a 1,10", "1,10 a 1,20", "> 1,20"]
    assert [ws.cell(17, c).fill.fgColor.rgb[-6:] for c in (5, 8, 11, 15, 18)] == [
        "F8CBAD", "FFE699", "C6E0B4", "FFE699", "F8CBAD"]
    meta = {r[0]: r[1] for r in wb["_meta"].iter_rows(min_row=2, values_only=True)}
    assert [meta[k] for k in Thresholds.META_KEYS] == [0.85, 0.9, 1.1, 1.2]
    data = wb[dashboard.SHEET_CHARTDATA]
    assert [data[c].value for c in ("B4", "B5", "B8", "B9")] == [0.85, 0.9, 1.1, 1.2]
    assert data.cell(dashboard.D_HEADER_ROW, dashboard.M_GREEN).value == "Zona verde (0,90 a 1,10)"


def test_conditional_formatting_uses_both_limits(book):
    """Indicadores EDT, Resumen y Cuadro de Mando: rojo y amarillo por debajo y por encima del rango verde."""
    wb = load_workbook(book)
    for sheet in ("Indicadores EDT", "Resumen", dashboard.SHEET_DASHBOARD, "Análisis de desviaciones"):
        formulas = [f for cf in wb[sheet].conditional_formatting for rule in cf.rules for f in rule.formula]
        assert any("<0.85" in f and ">1.2" in f for f in formulas), sheet
        assert any("<0.9" in f and ">1.1" in f for f in formulas), sheet


def test_index_chart_bands(book, recalc):
    """Bandas del gráfico de CPI y SPI: roja, amarilla, verde, amarilla y roja hasta la escala máxima."""
    out = load_workbook(recalc(book), data_only=True)
    data = out[dashboard.SHEET_CHARTDATA]
    r = dashboard.D_FIRST_ROW
    bands = [data.cell(r, c).value for c in (dashboard.M_RED, dashboard.M_YELLOW, dashboard.M_GREEN,
                                             dashboard.M_YELLOW_HIGH, dashboard.M_RED_HIGH)]
    assert bands == pytest.approx([0.85, 0.05, 0.2, 0.1, dashboard.INDEX_MAX - 1.2])
    # Las bandas superiores no repiten la leyenda (la de las inferiores indica los dos rangos)
    with zipfile.ZipFile(book) as z:
        charts = [z.read(n).decode() for n in z.namelist() if n.startswith("xl/charts/chart")]
    xml = next(x for x in charts if "'_graficos'!AB10" in x)
    series = dict(re.findall(r"<idx val=\"(\d+)\" /><order val=\"\d+\" /><tx><strRef><f>'_graficos'!([A-Z]+)10</f>",
                             xml))
    hidden = re.findall(r"<legendEntry><idx val=\"(\d+)\" /><delete val=\"1\" />", xml)
    assert sorted(series[i] for i in hidden) == ["AA", "AB"]
    # Excel solo respeta las entradas eliminadas si todas las series (valores y categorías) tienen caché
    assert xml.count("<numCache>") == 2 * len(series)


def test_analysis_lists_packages_above_the_upper_limit(book, example_files):
    """Con el amarillo superior en 1,10, los paquetes LOE (CPI 1,09 a 1,12) solo se listan si superan 1,10."""
    from evm.analysis import SHEET_ANALYSIS, read_entries
    cfg = book.parent / "propio.toml"
    assert main(["actualizar", "--origen", str(example_files["jun"]), "--config", str(cfg), "--libro", str(book),
                 "--sin-respaldo"]) == 0
    codes = {e.code for e in read_entries(load_workbook(book)[SHEET_ANALYSIS])}
    assert {"1.1.1", "1.1.2", "1.2.3"} <= codes             # CPI 1,11 y 1,12
    assert not {"1.1.3", "2.1.2"} & codes                   # CPI 1,09: dentro del rango verde (0,90 a 1,10)
    assert {"2.2.1", "1.2.1"} <= codes                      # por debajo del amarillo inferior
