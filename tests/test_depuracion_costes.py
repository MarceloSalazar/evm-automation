"""Depuración de costes: filas «(coste)» con la previsión pegada en la hoja PV-depuración."""
import zipfile
from datetime import date, datetime

import pytest
from openpyxl import load_workbook

from evm import cost_review as cr
from evm.config import MODE_COST, MODE_PERCENT, Config, ConfigError, load_config
from evm.migrate import plan_migration, prior_from_book
from evm.update import indicators_for, read_book, run_analysis, run_update
from evm.wbs import SEV_INFO, SEV_WARN, LinearPVSource, build_report
from evm.workbook import FILL_COST, KIND_COST, SHEET_STRUCT, read_structure, write_workbook

from builders import P, T, issues_with, rows_by_code

pytestmark = pytest.mark.modo(MODE_PERCENT)

JAN = (datetime(2027, 1, 4, 8), datetime(2027, 1, 29, 17))
FEB = (datetime(2027, 2, 1, 8), datetime(2027, 2, 26, 17))
MAR = (datetime(2027, 3, 1, 8), datetime(2027, 3, 26, 17))
START = date(2027, 1, 1)


def tasks() -> list:
    """EDT con un resumen de nivel 2 (1.2), una tarea suelta bajo 1 (fila «Otras tareas» 1.0.0-D) y tareas de las
    dos empresas."""
    return [
        T(1, 1, "1", "Gestión", bs=JAN[0], bf=MAR[1]),
        T(2, 2, "1.1", "Dirección", bs=JAN[0], bf=FEB[1], company="Ambas", method="% Completado"),
        T(3, 2, "1.2", "Contratos", bs=JAN[0], bf=MAR[1]),
        T(4, 3, "1.2.1", "Licitación", bs=JAN[0], bf=JAN[1], company="EmpresaA", method="0/100"),
        T(5, 3, "1.2.2", "Obras", bs=FEB[0], bf=MAR[1], company="Ambas", method="% Completado"),
        T(6, 2, "", "Asesoría", bs=MAR[0], bf=MAR[1], company="Ambas"),
        T(7, 1, "2", "Ingeniería", bs=JAN[0], bf=MAR[1]),
        T(8, 2, "2.1", "Básica", bs=JAN[0], bf=MAR[1], company="Ambas", method="% Completado"),
    ]


def config(review: bool = True, multi: bool = True, **kw) -> Config:
    kw.setdefault("start_month", START)
    cfg = Config(months=24, multi_enabled=multi, cost_review=review, **kw)
    cfg.validate()
    return cfg


def make_book(path, cfg, prior=None, task_list=None):
    rep = build_report(P(task_list or tasks()), cfg, LinearPVSource())
    if prior is not None:
        rep.issues.extend(plan_migration(rep, prior, cfg.max_level))
    write_workbook(rep, cfg, path, prior)
    return rep


def rows(wb) -> dict[str, int]:
    return {s.code: s.row for s in read_structure(wb)}


def set_bac(path, sheet: str, values: dict[str, float]) -> None:
    wb = load_workbook(path)
    r = rows(wb)
    for code, v in values.items():
        wb[sheet].cell(row=r[code], column=4, value=v)
    wb.save(path)


def paste(path, lines: list[list]) -> None:
    """Pega líneas en la hoja PV-depuración: [código, descripción, BAC, mes 1, mes 2...]."""
    wb = load_workbook(path)
    ws = wb[cr.SHEET_PASTE]
    for k, line in enumerate(lines):
        for c, v in enumerate(line, start=1):
            if v is not None:
                ws.cell(row=cr.P_FIRST + k, column=c, value=v)
    wb.save(path)


def months(*pairs) -> list:
    """Importes por mes: months((1, 100), (3, 50)) -> [100, None, 50]."""
    out = [None] * max(m for m, _v in pairs)
    for m, v in pairs:
        out[m - 1] = v
    return out


# ----------------------------------------------------------------------------
# Configuración
# ----------------------------------------------------------------------------
def test_config_keys_and_defaults(tmp_path):
    assert Config().cost_review is False and Config().cost_review_threshold == 10
    path = tmp_path / "c.toml"
    path.write_text('[exportacion]\ndepuracion_costes = true\ndepuracion_empresa = "empresab"\n'
                    'depuracion_alineamiento = true\ndepuracion_umbral = 15\n\n[multiempresa]\nactivar = true\n',
                    encoding="utf-8")
    cfg = load_config(path)
    assert cfg.cost_review and cfg.cost_review_company == "EmpresaB" and cfg.cost_review_role == "B"
    assert cfg.cost_review_alignment and cfg.cost_review_threshold == 15
    assert config().cost_review_role == "A"                      # por defecto, la empresa A
    assert config(multi=False).cost_review_role == ""            # una sola empresa: la hoja PV
    assert config(review=False).cost_review_role is None
    with pytest.raises(ConfigError, match="depuracion_empresa"):
        config(cost_review_company="Otra")
    with pytest.raises(ConfigError, match="depuracion_umbral"):
        config(cost_review_threshold=0)


def test_cost_mode_ignores_the_review(tmp_path):
    cfg = Config(months=24, mode=MODE_COST, cost_review=True)
    cfg.validate()
    rep = make_book(tmp_path / "libro.xlsx", cfg)
    assert rep.cost_review is None and issues_with(rep, "solo se aplica en el modo porcentaje", SEV_WARN)
    assert cr.SHEET_PASTE not in load_workbook(tmp_path / "libro.xlsx", read_only=True).sheetnames


def test_codes_without_trailing_zeros():
    assert cr.norm_code("1.2.0") == "1.2" and cr.norm_code("1.0.0") == "1" and cr.norm_code(" 10.0 ") == "10"
    assert cr.norm_code("1.10") == "1.10" and cr.norm_code("1,2") == "1.2" and cr.norm_code(3) == "3"
    for text in ("XRF01", "", None, "1..2", ".1", "1.", "1.2-D", "1 2", "1e5"):
        assert cr.norm_code(text) is None, text


# ----------------------------------------------------------------------------
# Libro
# ----------------------------------------------------------------------------
def test_cost_rows_below_every_row_and_hidden_rows_elsewhere(tmp_path):
    path = tmp_path / "libro.xlsx"
    rep = make_book(path, config())
    wb = load_workbook(path)
    names = wb.sheetnames
    assert names.index(cr.SHEET_CHART) == names.index("% PV-EmpresaB") + 1
    assert names.index(cr.SHEET_PASTE) == names.index(cr.SHEET_CHART) + 1
    assert wb[cr.SHEET_HELPER].sheet_state == "hidden"
    r = rows(wb)
    assert sorted(r.values()) == list(range(5, 5 + 2 * len(rep.rows), 2))      # una fila «(coste)» por medio
    pv_a = wb["PV-EmpresaA"]
    total = max(r.values()) + 2
    for code, row in list(r.items()) + [("TOTAL", total)]:
        cost = pv_a.cell(row=row + 1, column=2)
        assert pv_a.cell(row=row + 1, column=1).value == code and cost.value.endswith(" (coste)"), code
        assert cost.fill.fgColor.rgb.endswith(FILL_COST.fgColor.rgb[-6:])
        assert pv_a.cell(row=row + 1, column=4).value.startswith("=SUM(")
    assert pv_a.cell(row=r["1.2.1"] + 1, column=7).value == cr.paste_sum(0, r["1.2.1"])
    assert pv_a.cell(row=r["1.2.0"] + 1, column=7).value == f"=SUM(G{r['1.2.1'] + 1},G{r['1.2.2'] + 1})"
    assert pv_a.row_dimensions[r["1.2.1"] + 1].outline_level == 3            # un nivel más que la EDT
    for sheet in ("PV", "PV-EmpresaB", "% PV-EmpresaA", "EV-EmpresaA", "AC", "Indicadores EDT"):
        ws = wb[sheet]
        assert ws.row_dimensions[r["1.2.1"] + 1].hidden, sheet
        assert all(ws.cell(row=r["1.2.1"] + 1, column=c).value is None for c in range(1, 12)), sheet
    struct = wb[SHEET_STRUCT]
    assert struct.cell(row=r["1.2.1"] - 3, column=3).value == "1.2.1"
    assert struct.cell(row=r["1.2.1"] - 2, column=5).value == KIND_COST
    meta = {k: v for k, v in wb["_meta"].iter_rows(min_row=2, values_only=True)}
    assert meta["depuracion_costes"] is True and meta["depuracion_hoja"] == "PV-EmpresaA"
    assert issues_with(rep, "Depuración de costes de EmpresaA", SEV_INFO)


def test_single_company_book_reviews_the_pv_sheet(tmp_path):
    path = tmp_path / "libro.xlsx"
    make_book(path, config(multi=False))
    wb = load_workbook(path)
    r = rows(wb)
    pv = wb["PV"]
    assert pv.cell(row=r["2.1.0"] + 1, column=2).value == "Básica (coste)"
    validations = " ".join(str(dv.sqref) for dv in pv.data_validations.dataValidation)
    assert f"C{r['2.1.0']}" in validations and f"C{r['2.1.0'] + 1}" not in validations
    assert wb["% PV"].row_dimensions[r["2.1.0"] + 1].hidden


def test_update_and_analysis_match_a_book_without_review(tmp_path):
    """La depuración no cambia el PV, el EV ni los indicadores: las filas «(coste)» no forman parte del libro."""
    results = []
    for review in (True, False):
        path = tmp_path / f"libro_{review}.xlsx"
        make_book(path, config(review=review))
        for sheet, bac in (("PV-EmpresaA", 1000.0), ("PV-EmpresaB", 500.0)):
            set_bac(path, sheet, {c: bac for c in ("1.1.0", "1.2.1", "1.2.2", "1.0.0-D", "2.1.0")})
        progress = tasks()
        progress[1].percent_complete = 50
        progress[3].percent_complete = 100
        run_update(P(progress), config(review=review), path, date(2027, 2, 28), backup=False)
        run_analysis(config(review=review), path, backup=False)
        book = read_book(path)
        code = {s.row: s.code for s in book.structure}
        results.append(({code[r]: v for r, v in book.pv.items()}, {code[r]: v for r, v in book.ev.items()},
                        indicators_for(book, date(2027, 2, 28))))
        with zipfile.ZipFile(path) as z:            # los gráficos se vuelven a crear al guardar
            charts = [n for n in z.namelist() if n.startswith("xl/charts/chart")]
        assert len(charts) == (4 if review else 3)
    assert results[0][0] == results[1][0] and results[0][1] == results[1][1]
    assert results[0][2] == results[1][2] and results[0][2].ev > 0


def test_lookup_of_book_codes(tmp_path):
    rep = build_report(P(tasks()), config(), LinearPVSource())
    entries, nonnumeric, repeated = cr.lookup_rows(rep)
    by = rows_by_code(rep)
    assert entries["1.2.1"] == (by["1.2.1"], by["1.2.1"])
    assert entries["1"] == (cr.NO_ROW, by["1.0.0-D"])          # resumen: lo que depende va a «Otras tareas»
    assert entries["1.2"] == (cr.NO_ROW, cr.NO_ROW)            # resumen sin fila propia
    assert "1.0.0-D" not in entries and not nonnumeric and not repeated


# ----------------------------------------------------------------------------
# --migrar-desde
# ----------------------------------------------------------------------------
LINES = [["1.1", "Dirección", 200, 100, 100], ["1.2.1", "Licitación", None, 50]]


def test_migration_keeps_the_pasted_forecast_and_the_bac(tmp_path):
    old = tmp_path / "old.xlsx"
    make_book(old, config())
    set_bac(old, "PV-EmpresaA", {"1.1.0": 1234.0})
    paste(old, LINES)
    new = tmp_path / "new.xlsx"
    rep = make_book(new, config(), prior_from_book(old, "Programa"))
    ws = load_workbook(new)[cr.SHEET_PASTE]
    assert [ws.cell(row=cr.P_FIRST, column=c).value for c in range(1, 6)] == LINES[0]
    assert issues_with(rep, f"Se conservó el contenido de la hoja {cr.SHEET_PASTE} (2 líneas)")
    assert read_book(new).bac_role["A"][rows(load_workbook(new))["1.1.0"]] == 1234.0

    # Libro de entrega: sin depuración, conserva el BAC y no traslada la previsión
    final = tmp_path / "final.xlsx"
    rep = make_book(final, config(review=False), prior_from_book(new, "Programa"))
    assert cr.SHEET_PASTE not in load_workbook(final, read_only=True).sheetnames
    assert issues_with(rep, "no se trasladó: el libro nuevo no tiene depuración", SEV_INFO)
    book = read_book(final)
    assert book.bac_role["A"][rows(load_workbook(final))["1.1.0"]] == 1234.0
    assert sorted(s.row for s in book.structure) == list(range(5, 5 + len(book.structure)))


def test_migration_shifts_months_and_checks_the_company(tmp_path):
    old = tmp_path / "old.xlsx"
    make_book(old, config())
    paste(old, [["1.1", "Dirección", 300, 100, 100, 100]])
    later = tmp_path / "later.xlsx"                 # el horizonte empieza un mes después
    rep = make_book(later, config(start_month=date(2027, 2, 1)), prior_from_book(old, "Programa"))
    ws = load_workbook(later)[cr.SHEET_PASTE]
    assert [ws.cell(row=cr.P_FIRST, column=c).value for c in range(4, 7)] == [100, 100, None]
    assert issues_with(rep, "se desplazaron 1 columnas", SEV_WARN) and issues_with(rep, "quedaron fuera 100.00")
    other = tmp_path / "other.xlsx"
    rep = make_book(other, config(cost_review_company="EmpresaB"), prior_from_book(old, "Programa"))
    assert issues_with(rep, "tiene la previsión de EmpresaA y el libro nuevo depura EmpresaB", SEV_WARN)
    assert load_workbook(other)[cr.SHEET_PASTE].cell(row=cr.P_FIRST, column=1).value is None


# ----------------------------------------------------------------------------
# Fórmulas de Excel (LibreOffice)
# ----------------------------------------------------------------------------
@pytest.mark.libreoffice
def test_pasted_lines_are_assigned_to_book_rows(tmp_path, recalc):
    path = tmp_path / "libro.xlsx"
    make_book(path, config(cost_review_alignment=True))
    set_bac(path, "PV-EmpresaA", {"1.1.0": 200.0, "1.2.1": 50.0, "1.2.2": 60.0, "1.0.0-D": 10.0, "2.1.0": 400.0})
    lines = [
        ["1", "Gestión (cabecera, sin importes)", None],
        ["XRF01", "Subtarea suelta de Gestión", None, *months((3, 10))],      # -> 1.0.0-D («Otras tareas»)
        ["1.1", "Dirección", 200, 100, 100, "50"],                           # 1.1 = 1.1.0; "50": texto
        [1.2, "Contratos (subtotal con importes propios)", None, *months((1, 999))],
        ["1.2.1.5", "Licitación, tarea de detalle", None, *months((1, 50))],  # -> 1.2.1
        ["1.2.2", "Obras", 70, *months((2, 30), (3, 30))],                    # BAC distinto de los meses
        [None, None, None],
        [None, "Nota sin código ni importes", None],
        ["2.1", "Básica", None, *months((2, 200), (30, 200))],               # mes 30: fuera del horizonte
        ["1.2.7", "Tarea de un resumen sin fila propia", None, *months((1, 5))],
        ["9.9", "Código inexistente", None, *months((1, 7))],
    ]
    paste(path, lines)
    wb = load_workbook(path)
    wb[cr.SHEET_CHART][cr.G_SELECT] = "1.2.0"
    wb.save(path)
    wv = load_workbook(recalc(path), data_only=True)
    ws = wv[cr.SHEET_PASTE]
    status_col = cr.P_FIRST_MONTH + cr.paste_months(24) + 1
    status = [ws.cell(row=cr.P_FIRST + k, column=status_col).value or None for k in range(len(lines))]
    assert status == [cr.STATUS_SUBTOTAL, "1.0.0-D", "1.1.0", cr.STATUS_SUBTOTAL, "1.2.1", "1.2.2", None, None,
                      "2.1.0", cr.STATUS_SUMMARY, cr.STATUS_NONE]
    obs = [ws.cell(row=cr.P_FIRST + k, column=status_col + 1).value for k in range(len(lines))]
    assert "Subtotal distinto" in obs[3] and "BAC distinto" in obs[5] and "fuera del horizonte" in obs[8]
    assert "como texto" in obs[2] and not obs[0] and not obs[4]
    pasted = 10 + 200 + 50 + 60 + 400 + 5 + 7
    assert ws.cell(row=cr.P_TOTAL_ROW, column=3).value == pytest.approx(pasted)
    assert ws.cell(row=cr.P_MISSING_ROW, column=3).value == pytest.approx(200 + 5 + 7)

    pv_a = wv["PV-EmpresaA"]
    r = rows(wb)

    def cost(code: str) -> list:
        return [pv_a.cell(row=r[code] + 1, column=7 + i).value or 0 for i in range(3)]

    assert cost("1.1.0") == [100, 100, 0] and cost("1.2.1") == [50, 0, 0] and cost("1.2.2") == [0, 30, 30]
    assert cost("1.0.0-D") == [0, 0, 10] and cost("2.1.0") == [0, 200, 0]
    assert cost("1.2.0") == [50, 30, 30] and cost("1.0.0") == [150, 130, 40]
    total = max(r.values()) + 2
    assert pv_a.cell(row=total + 1, column=4).value == pytest.approx(pasted - 200 - 5 - 7)
    assert pv_a.cell(row=r["1.2.2"] + 1, column=4).value == 60
    # Indicador: diferencia de BAC y diferencia máxima acumulada en % del BAC
    last = 7 + 24
    assert pv_a.cell(row=r["1.1.0"] + 1, column=last + 1).value == pytest.approx(0)
    assert pv_a.cell(row=r["2.1.0"] + 1, column=last + 1).value == pytest.approx(-0.5)
    assert 0 < pv_a.cell(row=r["2.1.0"] + 1, column=last + 2).value <= 1
    assert pv_a.cell(row=r["1.1.0"], column=last + 1).value is None

    chart = wv[cr.SHEET_CHART]
    assert chart["B5"].value == pytest.approx(110) and chart["B6"].value == pytest.approx(110)
    assert chart.cell(row=cr.G_FIRST + 2, column=cr.G_COSTC).value == pytest.approx(110)
    assert chart["C4"].value == "Contratos"


@pytest.mark.libreoffice
def test_example_forecast_is_fully_assigned(tmp_path, generator, recalc):
    """El ayudante del ejemplo pega una previsión ficticia (hitos de pago y subtareas) que se asigna completa."""
    xml = tmp_path / "base.xml"
    generator.write(generator.build_plans(), xml, None)
    cfg = config(milestone_weight_field="Peso hito", method_field="Método EVM", plan="INTERNO")
    cfg.start_month = None
    from evm.reader_xml import read_xml
    from evm.wbs import build_part, assemble_report
    project = read_xml(xml, cfg)
    rep = assemble_report([build_part(project, cfg, LinearPVSource())], cfg)
    path = tmp_path / "ejemplo.xlsx"
    write_workbook(rep, cfg, path)
    generator.fill_book(path)
    ws = load_workbook(path)[cr.SHEET_PASTE]
    codes = [ws.cell(row=cr.P_FIRST + k, column=1).value for k in range(40)]
    assert "HP-01" in codes and "XRF01" in codes
    wv = load_workbook(recalc(path), data_only=True)
    pa = wv[cr.SHEET_PASTE]
    assert pa.cell(row=cr.P_TOTAL_ROW, column=3).value > 0
    assert pa.cell(row=cr.P_MISSING_ROW, column=3).value == pytest.approx(0)
    total = max(rows(load_workbook(path)).values()) + 2
    assert wv["PV-EmpresaA"].cell(row=total + 1, column=4).value == pytest.approx(
        pa.cell(row=cr.P_TOTAL_ROW, column=3).value)
