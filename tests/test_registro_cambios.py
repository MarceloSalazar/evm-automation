"""Controles del cierre mensual de la norma EIA-748-E: registro de cambios del BAC y de los meses ya informados
(#24, #31), coherencia del EV y el AC (#32), análisis de desviaciones incompleto y acciones vencidas (#33) y meses
sin informe (#34)."""
from datetime import date, datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

from evm import analysis, changes, consistency
from evm.analysis import A_ACTION, A_CAUSE, A_DATE, SHEET_ANALYSIS, TEXT_COLS
from evm.cli import main
from evm.timeline import Horizon
from evm.update import BookData, read_book
from evm.wbs import KIND_LEAF
from evm.workbook import COL_BAC, FIRST_MONTH_COL, StructRow

START = date(2027, 1, 1)


# ----------------------------------------------------------------------------
# Libro en memoria
# ----------------------------------------------------------------------------
def fake_book(rows: dict[str, dict], months: int = 6) -> BookData:
    """Libro de una sola empresa con una fila de ingreso por código: {código: {bac, ev, ac, pv, method}} (ev y ac:
    valores de cada mes desde enero de 2027; None = vacío)."""
    structure = [StructRow(10 + i, str(100 + i), code, 3, KIND_LEAF, None, f"Paquete {code}")
                 for i, code in enumerate(rows)]

    def series(values, fill=None):
        return list(values) + [fill] * (months - len(values))

    data = {s.row: rows[s.code] for s in structure}
    bac = {r: float(d.get("bac", 1000.0)) for r, d in data.items()}
    pct_pv = {r: series(d.get("pv", []), 0.0) for r, d in data.items()}
    pct_ev = {r: series(d.get("ev", [])) for r, d in data.items()}
    ac = {r: series(d.get("ac", [])) for r, d in data.items()}
    ev = {r: [None if v is None else bac[r] * v for v in vals] for r, vals in pct_ev.items()}
    pv = {r: [bac[r] * v for v in vals] for r, vals in pct_pv.items()}
    return BookData(Path("libro.xlsx"), {"decimales": 0}, structure, Horizon(START, months), pv, ev, ac, bac,
                    {r: d.get("method", "% Completado") for r, d in data.items()}, None,
                    bac_role={"": bac}, pct_pv={"": pct_pv}, pct_ev={"": pct_ev})


def ids(issues) -> list[tuple[str, str]]:
    return [(i.id, i.code) for i in issues]


# ----------------------------------------------------------------------------
# Coherencia del EV y el AC (#32)
# ----------------------------------------------------------------------------
def test_ev_and_ac_consistency():
    book = fake_book({
        "1.1": {"ev": [0.2, 0.1], "ac": [0, None]},                     # avance sin costo
        "1.2": {"ev": [0, 0], "ac": [300, 100]},                        # costo sin avance
        "1.3": {"ev": [0, 0], "ac": [300, 100], "method": "0/100"},     # 0/100: es lo esperado
        "1.6": {"ev": [0, 0], "ac": [300, 100], "method": "Hitos ponderados"},   # antes del primer hito
        "1.4": {"ev": [0.6, 0.4, 0], "ac": [500, 400, 50]},             # terminó en febrero, AC en marzo
        "1.5": {"ev": [0.5, 0.2, 0.1], "ac": [400, 300, 100]},          # coherente
    })
    issues = consistency.ev_ac_checks(book, 2)
    assert ids(issues) == [("ACT-039", "1.1"), ("ACT-040", "1.2"), ("ACT-041", "1.4")]
    assert issues[0].message == "EV acumulado 300 (30 % ganado) y sin AC a 03/2027: el CPI no se puede calcular."
    assert issues[1].message == "AC acumulado 400 y sin % ganado a 03/2027."
    assert issues[2].message == "Llegó al 100 % en 02/2027 y tiene AC en 03/2027 (50)."


def test_ev_and_ac_until_the_last_month_with_ac():
    """Después de 'actualizar', el AC del mes de corte todavía no se cargó: se revisa hasta el mes anterior."""
    book = fake_book({"1.1": {"ev": [0.2, 0.1, 0.1], "ac": [200, 100]}, "1.2": {"ev": [0, 0, 0.3], "ac": [0, 0]}})
    issues = consistency.ev_ac_checks(book, 2)
    assert ids(issues) == [("ACT-042", "")]
    assert issues[0].message.startswith("El AC de 03/2027 no está cargado: la coherencia del EV y el AC se revisó "
                                        "hasta 02/2027.")
    empty = fake_book({"1.1": {"ev": [0.2]}})
    assert [i.message.startswith("Ningún mes tiene AC") for i in consistency.ev_ac_checks(empty, 0)] == [True]


# ----------------------------------------------------------------------------
# Meses sin informe (#34)
# ----------------------------------------------------------------------------
def test_month_gaps():
    book = fake_book({"1.1": {"ev": [0.1, None, 0.1, None, None, 0.2]}, "1.2": {"ev": [0.0]}})
    latest = consistency.month_gaps(book, 5, latest_only=True)
    assert ids(latest) == [("ACT-046", "")]
    assert latest[0].message.startswith("04/2027 y 05/2027 no tienen % ganado cargado: su avance quedó en 06/2027.")
    every = consistency.month_gaps(book, 5, latest_only=False)
    assert every[0].message.startswith("02/2027, 04/2027 y 05/2027 no tienen % ganado cargado: su avance quedó en el "
                                       "mes cargado siguiente.")
    (single,) = consistency.month_gaps(book, 2, latest_only=True)
    assert single.message.startswith("02/2027 no tiene % ganado cargado: su avance quedó en 03/2027.")
    assert consistency.month_gaps(fake_book({"1.1": {"ev": [None, None, 0.3]}}), 2, latest_only=False) == []


# ----------------------------------------------------------------------------
# Análisis de desviaciones incompleto y acciones vencidas (#33)
# ----------------------------------------------------------------------------
def entry(month: date, code: str, **texts) -> analysis.Entry:
    values = [None] * len(TEXT_COLS)
    for column, value in texts.items():
        values[TEXT_COLS.index({"cause": A_CAUSE, "action": A_ACTION, "date": A_DATE}[column])] = value
    return analysis.Entry(month, f"k{code}", code, f"Paquete {code}", values)


def test_incomplete_analysis_of_the_previous_month():
    jun, jul, may = date(2027, 6, 1), date(2027, 7, 1), date(2027, 5, 1)
    entries = [entry(jun, "1.1", cause="Retraso", action="Reforzar"), entry(jun, "1.2", cause="Lluvias"),
               entry(jun, "1.3"), entry(may, "1.4", cause="Antigua"), entry(jul, "1.5")]
    issues = analysis.incomplete(entries, jul)
    assert ids(issues) == [("ACT-043", "1.2"), ("ACT-044", "")]
    assert issues[0].message == "El análisis de 06/2027 no tiene acción correctiva."
    assert issues[1].message.startswith("1 paquete fuera de umbral en 06/2027 quedó sin analizar, sin texto en la hoja "
                                        "Análisis de desviaciones (1.3): se quitan de la hoja")
    assert analysis.incomplete(entries, may) == []


def test_overdue_actions_of_packages_still_out_of_threshold():
    jul = date(2027, 7, 1)
    entries = [entry(jul, "1.1", action="Reforzar", date=datetime(2027, 7, 15)),
               entry(jul, "1.2", action="Reforzar", date=datetime(2027, 8, 15)),
               entry(jul, "1.3", action="Reforzar", date="15/07/2027"),             # texto: no se interpreta
               entry(jul, "1.4", action="Reforzar", date=date(2027, 7, 1))]          # volvió al rango verde
    found = [analysis.Deviation(f"k{c}", c, "", 0.8, 0.9) for c in ("1.1", "1.2", "1.3")]
    issues = analysis.overdue(entries, found, jul, date(2027, 7, 31))
    assert ids(issues) == [("ACT-045", "1.1")]
    assert "(15/07/2027) es anterior a la fecha de corte" in issues[0].message


# ----------------------------------------------------------------------------
# Detección de cambios (#24, #31)
# ----------------------------------------------------------------------------
def test_detect_bac_and_closed_month_changes():
    before = fake_book({"1.1": {"pv": [0.3, 0.3, 0.4], "ev": [0.2, 0.2], "ac": [200, 250]},
                        "1.2": {"pv": [0, 0, 0.5], "ac": [100]}})
    ref = changes.reference_from_book(before, 1, date(2027, 2, 28))          # informado hasta febrero
    after = fake_book({"1.1": {"bac": 1200, "pv": [0.3, 0.3, 0.4], "ev": [0.25, 0.2, 0.1], "ac": [200, 270, 90]},
                       "1.2": {"bac": 1500, "pv": [0, 0, 0.5], "ac": [100, 80]}})
    found = changes.detect(after, ref, 2)                                     # actualizar a marzo
    assert [(c.code, c.data, c.month, c.old, c.new, c.affected) for c in found] == [
        ("1.1", "BAC", None, 1000.0, 1200.0, 2),          # PV y EV de enero y febrero recalculados
        ("1.1", "% ganado", date(2027, 1, 1), 0.2, 0.25, None),
        ("1.1", "AC", date(2027, 2, 1), 250.0, 270.0, None),
        ("1.2", "BAC", None, 1000.0, 1500.0, 0),          # sin PV ni EV antes de marzo
    ]                                                     # 1.2: AC de febrero cargado en el último mes informado
    late = fake_book({"1.1": {"pv": [0.3, 0.3, 0.4], "ev": [0.2, 0.2], "ac": [200, 250]},
                      "1.2": {"pv": [0, 0, 0.5], "ac": [100]}})
    late.ac[11][0] = None                                 # AC de enero borrado
    assert [(c.code, c.data, c.old, c.new) for c in changes.detect(late, ref, 2)] == [("1.2", "AC", 100.0, 0.0)]
    assert changes.detect(after, None, 2) == []
    issues = changes.issues(found, 0)
    assert [(i.id, i.code, i.severity) for i in issues] == [
        ("ACT-036", "1.1", "Advertencia"), ("ACT-036", "1.2", "Info"), ("ACT-037", "1.1", "Advertencia")]
    assert issues[0].message == ("BAC 1.000 → 1.200 (+200): se recalcularon el PV y el EV de 2 meses ya informados "
                                 "(cambio retroactivo). Se agregó a la hoja Registro de cambios.")
    assert issues[2].message.startswith("Cambiaron valores de meses ya informados (% ganado de 01/2027: 20 % → 25 %; "
                                        "AC de 02/2027: 250 → 270).")


def test_reference_follows_the_calendar_month():
    """--migrar-desde puede cambiar el mes 1 del horizonte: la referencia se compara por mes calendario."""
    old = fake_book({"1.1": {"ev": [0.1, 0.2], "ac": [10, 20]}})
    ref = changes.reference_from_book(old, 1, date(2027, 2, 28))
    new = fake_book({"1.1": {"ev": [None, 0.1, 0.2], "ac": [None, 10, 20]}})
    new.horizon = Horizon(date(2026, 12, 1), 6)
    assert changes.detect(new, ref, 3) == []


# ----------------------------------------------------------------------------
# Comandos con el proyecto de ejemplo
# ----------------------------------------------------------------------------
@pytest.fixture()
def book(tmp_path, example_files, generator):
    out = tmp_path / "EVM.xlsx"
    assert main(["generar", "--origen", str(example_files["base"]), "--config", str(example_files["config"]),
                 "--salida", str(out)]) == 0
    generator.fill_book(out)
    return out


def update(book, example_files, month: str, *extra, ac: bool = True) -> int:
    import generar_ejemplo
    status = {"jun": date(2027, 6, 30), "jul": date(2027, 7, 31)}[month]
    if ac:
        generar_ejemplo.fill_book(book, status)
    return main(["actualizar", "--origen", str(example_files[month]), "--config", str(example_files["config"]),
                 "--libro", str(book), "--sin-respaldo", *extra])


def log_rows(path) -> list[list]:
    wb = load_workbook(path)
    return changes.read_rows(wb[changes.SHEET_LOG])


def set_cell(path, sheet: str, code: str, column: int, value) -> None:
    b = read_book(path)
    row = next(s.row for s in b.input_rows if s.code == code)
    wb = load_workbook(path)
    wb[sheet].cell(row, column, value)
    wb.save(path)


def month_col(path, month: date) -> int:
    return FIRST_MONTH_COL + read_book(path).horizon.index(month)


def test_generate_creates_the_change_log(book):
    wb = load_workbook(book)
    assert wb.sheetnames.index(changes.SHEET_LOG) == wb.sheetnames.index(SHEET_ANALYSIS) + 1
    ws = wb[changes.SHEET_LOG]
    assert [ws.cell(changes.HEAD_ROW, c).value for c in (changes.L_DATA, changes.L_OLD, changes.L_REASON)] == [
        "Dato", "Valor anterior", "Motivo"]
    assert changes.SHEET_REF not in wb.sheetnames and log_rows(book) == []


def test_bac_change_is_logged(book, example_files, capsys):
    assert update(book, example_files, "jun") == 0
    out = capsys.readouterr().out
    assert "ACT-038" in out and "ACT-036" not in out
    wb = load_workbook(book)
    assert wb[changes.SHEET_REF].sheet_state == "hidden"
    bac = read_book(book).bac[next(s.row for s in read_book(book).input_rows if s.code == "2.2.1")]
    set_cell(book, "PV", "2.2.1", COL_BAC, bac + 50_000)

    assert update(book, example_files, "jul", "--simular") == 0
    assert "Se registrará en la hoja Registro de cambios" in capsys.readouterr().out
    assert log_rows(book) == []

    assert update(book, example_files, "jul") == 0
    out = capsys.readouterr().out
    assert "[!] ACT-036 2.2.1" in out and "cambio retroactivo" in out and "ACT-038" not in out
    (row,) = log_rows(book)
    assert row[changes.L_CODE - 1] == "2.2.1" and row[changes.L_DATA - 1] == "BAC"
    assert row[changes.L_NEW - 1] - row[changes.L_OLD - 1] == pytest.approx(50_000)
    assert row[changes.L_AFFECTED - 1] > 0 and row[changes.L_KEY - 1]

    assert main(["analizar", "--libro", str(book), "--sin-respaldo"]) == 0     # sin cambios nuevos
    assert "ACT-036" not in capsys.readouterr().out and len(log_rows(book)) == 1


def test_closed_month_changes_are_logged(book, example_files, capsys):
    """AC del mes cargado después de 'actualizar': no es un cambio. Modificar un mes ya informado, sí."""
    assert update(book, example_files, "jun", ac=False) == 0
    import generar_ejemplo
    generar_ejemplo.fill_book(book, date(2027, 6, 30))           # AC de junio, después de actualizar
    assert update(book, example_files, "jul") == 0
    assert "ACT-037" not in capsys.readouterr().out and log_rows(book) == []

    jun = month_col(book, date(2027, 6, 1))
    set_cell(book, "AC", "1.1.1", jun, 123_456)
    assert main(["analizar", "--libro", str(book), "--sin-respaldo"]) == 0
    out = capsys.readouterr().out
    assert "[!] ACT-037 1.1.1" in out and "AC de 06/2027" in out
    (row,) = log_rows(book)
    assert (row[changes.L_DATA - 1], row[changes.L_MONTH - 1].date(), row[changes.L_NEW - 1]) == (
        "AC", date(2027, 6, 1), 123_456)


def test_migrate_keeps_the_log_and_the_reference(book, example_files, tmp_path, capsys):
    assert update(book, example_files, "jun") == 0
    bac = read_book(book).bac[next(s.row for s in read_book(book).input_rows if s.code == "2.2.1")]
    set_cell(book, "PV", "2.2.1", COL_BAC, bac + 1000)
    assert update(book, example_files, "jul") == 0
    wb = load_workbook(book)
    ws = wb[changes.SHEET_LOG]
    ws.cell(changes.FIRST, changes.L_REASON, "Cambio de alcance aprobado")
    wb.save(book)
    new = tmp_path / "EVM_LB2.xlsx"
    assert main(["generar", "--origen", str(example_files["base"]), "--config", str(example_files["config"]),
                 "--salida", str(new), "--migrar-desde", str(book)]) == 0
    (row,) = log_rows(new)
    assert row[changes.L_REASON - 1] == "Cambio de alcance aprobado" and row[changes.L_KEY - 1]
    details = [r[4] for r in load_workbook(new)["Validación"].iter_rows(min_row=4, values_only=True) if r[4]]
    assert any(d.startswith("Registro de cambios trasladado: 1 fila de la hoja y la referencia de") for d in details)
    capsys.readouterr()
    assert update(new, example_files, "jul", ac=False) == 0
    out = capsys.readouterr().out
    assert not any(code in out for code in ("ACT-036", "ACT-037", "ACT-038"))


def test_incomplete_analysis_and_overdue_action(book, example_files, capsys):
    assert update(book, example_files, "jun") == 0
    wb = load_workbook(book)
    ws = wb[SHEET_ANALYSIS]
    rows = {ws.cell(r, 2).value: r for r in range(analysis.FIRST, ws.max_row + 1)}
    ws.cell(rows["2.2.1"], A_CAUSE, "Planos atrasados")
    ws.cell(rows["2.2.1"], A_DATE, datetime(2027, 7, 15))
    wb.save(book)
    capsys.readouterr()
    assert update(book, example_files, "jul") == 0
    out = capsys.readouterr().out
    assert "[!] ACT-043 2.2.1" in out and "no tiene acción correctiva" in out
    assert any(line.startswith("  [!] ACT-044 12 paquetes fuera de umbral en 06/2027 quedaron sin analizar")
               and "(1.1.1, 1.1.2" in line and "y 2 más)" in line and "2.2.1" not in line for line in out.splitlines())
    assert "[!] ACT-045 2.2.1" in out and "(15/07/2027)" in out


def test_month_gap_when_a_month_is_skipped(book, example_files, capsys):
    assert update(book, example_files, "jun") == 0
    capsys.readouterr()
    assert main(["actualizar", "--origen", str(example_files["jul"]), "--config", str(example_files["config"]),
                 "--libro", str(book), "--sin-respaldo", "--fecha-corte", "2027-08-31"]) == 0
    out = capsys.readouterr().out
    assert "[!] ACT-046 07/2027 no tiene % ganado cargado: su avance quedó en 08/2027." in out
