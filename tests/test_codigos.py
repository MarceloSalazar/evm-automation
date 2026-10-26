"""Códigos de errores y advertencias: catálogo (evm/messages.py), uso en el programa, documento docs/codigos.md y
silenciar ([validacion] silenciar y --silenciar)."""
import ast
import builtins
import importlib.util
from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

from evm.cli import main
from evm.config import ConfigError, load_config
from evm.messages import (AREAS, CATALOG, ID_PATTERN, SEV_ERROR, SEV_FATAL, SEV_INFO, SEV_WARN, TOPIC_DATES, Issue,
                          parse_silence)

from mspdi import project_xml, task_xml, write_xml

ROOT = Path(__file__).resolve().parents[1]
SOURCES = sorted((ROOT / "evm").glob("*.py"))
# Excepciones del programa: llevan su código del catálogo como primer argumento
CODED_ERRORS = {"EvmError", "ConfigError", "PlanNotFound", "WorkbookFormatError", "XmlFormatError", "InputError",
                "MissingFileError", "SaveError", "ProjectAppError"}
BUILTIN_ERRORS = {n for n in dir(builtins) if isinstance(getattr(builtins, n), type)
                  and issubclass(getattr(builtins, n), BaseException)} - {"SystemExit"}
S, F = datetime(2026, 1, 5, 8), datetime(2026, 3, 27, 17)


def trees():
    for path in SOURCES:
        yield path, ast.parse(path.read_text(encoding="utf-8"))


def call_name(node: ast.Call) -> str:
    f = node.func
    return f.id if isinstance(f, ast.Name) else f.attr if isinstance(f, ast.Attribute) else ""


def used_codes() -> dict[str, list[str]]:
    """Códigos que usa el programa: los textos ÁREA-NNN del código fuente, sin las definiciones del catálogo."""
    used: dict[str, list[str]] = {}
    for path, tree in trees():
        definitions = {id(n.args[0]) for n in ast.walk(tree) if isinstance(n, ast.Call) and call_name(n) == "_code"}
        for node in ast.walk(tree):
            if (isinstance(node, ast.Constant) and isinstance(node.value, str) and id(node) not in definitions
                    and ID_PATTERN.fullmatch(node.value)):
                used.setdefault(node.value, []).append(f"{path.name}:{node.lineno}")
    return used


# ----------------------------------------------------------------------------
# Catálogo
# ----------------------------------------------------------------------------
def test_catalog_codes_are_well_formed_and_never_renumbered():
    """Formato ÁREA-NNN, numeración correlativa por área (un código eliminado deja un hueco: se debe retirar, no
    borrar) y textos completos."""
    by_area: dict[str, list[int]] = {}
    for key, code in CATALOG.items():
        assert key == code.id and ID_PATTERN.fullmatch(code.id), code.id
        assert code.id[:3] in AREAS, code.id
        by_area.setdefault(code.id[:3], []).append(int(code.id[4:]))
        assert code.severities and set(code.severities) <= {SEV_FATAL, SEV_ERROR, SEV_WARN, SEV_INFO}, code.id
        assert len(set(code.severities)) == len(code.severities), code.id
        assert code.title.strip() == code.title and code.title and not code.title.endswith("."), code.id
        for text in (code.cause, code.action):
            assert text.strip() == text and text.endswith((".", ")")), (code.id, text)
    for prefix, numbers in by_area.items():
        assert numbers == list(range(1, len(numbers) + 1)), f"{prefix}: numeración con huecos o desordenada"
    assert set(by_area) == set(AREAS)


def test_every_code_of_the_program_is_in_the_catalog_and_every_code_is_used():
    used = used_codes()
    unknown = {c: where for c, where in used.items() if c not in CATALOG}
    assert not unknown, f"códigos sin entrada en el catálogo: {unknown}"
    unused = [c.id for c in CATALOG.values() if not c.retired and c.id not in used]
    assert not unused, f"códigos del catálogo que el programa no usa (se deben retirar): {unused}"


def test_messages_and_errors_always_have_a_code():
    """Toda observación (Issue) y todo error previsto (raise) lleva un código; no quedan avisos de texto suelto."""
    problems = []
    dynamic = {"id", "code_id", "SUMMARY_ID"}       # códigos que se reciben de una variable
    for path, tree in trees():
        for node in ast.walk(tree):
            where = f"{path.name}:{getattr(node, 'lineno', '?')}"
            if isinstance(node, ast.Call) and call_name(node) in ("Issue", "_issue"):
                first = node.args[0] if node.args else None
                ok = ((isinstance(first, ast.Constant) and first.value in CATALOG)
                      or (isinstance(first, ast.Attribute) and first.attr == "id")
                      or (isinstance(first, ast.Name) and first.id in dynamic))
                if not ok:
                    problems.append(f"{where}: Issue sin código del catálogo")
            if isinstance(node, ast.Raise) and isinstance(node.exc, ast.Call):
                name = call_name(node.exc)
                first = node.exc.args[0] if node.exc.args else None
                if name in CODED_ERRORS and not (isinstance(first, ast.Constant) and first.value in CATALOG):
                    problems.append(f"{where}: {name} sin código del catálogo")
                elif name in BUILTIN_ERRORS:
                    problems.append(f"{where}: raise {name}(...): se debe usar un error con código (messages.py)")
            if (isinstance(node, ast.Call) and call_name(node) == "append" and isinstance(node.func.value, ast.Name)
                    and node.func.value.id in ("warnings", "notes") and node.args
                    and isinstance(node.args[0], (ast.Constant, ast.JoinedStr, ast.BinOp))):
                problems.append(f"{where}: aviso de texto suelto, sin código")
            if (isinstance(node, ast.Constant) and isinstance(node.value, str)
                    and node.value.lstrip().startswith(("[!] ", "[i] ", "[X] ", "ERROR:"))):
                problems.append(f"{where}: mensaje de consola sin código")
    assert not problems, "\n".join(problems)


def test_document_is_up_to_date():
    spec = importlib.util.spec_from_file_location("codigos", ROOT / "docs" / "herramientas" / "codigos.py")
    tool = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(tool)
    current = (ROOT / "docs" / "codigos.md").read_text(encoding="utf-8")
    assert current == tool.render(), "docs/codigos.md no está al día: python docs\\herramientas\\codigos.py"
    for code in CATALOG:
        assert f"\n#### {code}\n" in current


def test_issue_takes_severity_and_topic_from_the_catalog():
    i = Issue("FEC-007", "1.2.3", "Montaje", "Duración estimada.")
    assert (i.severity, i.topic) == (SEV_WARN, TOPIC_DATES)
    assert Issue("LOG-008", message="x", severity=SEV_INFO).severity == SEV_INFO       # severidad variable
    assert Issue("EDT-010").severity == SEV_ERROR       # error que detiene el programa, informado por validar


# ----------------------------------------------------------------------------
# Silenciar
# ----------------------------------------------------------------------------
def test_silence_rules_and_invalid_codes():
    s = parse_silence(["fec-007", " MET-004:1.2.1.0 ", "XXX-001", "nada", "CFG-001", "CFG-049", "EDT-001:", ""])
    assert s.rules == {"FEC-007": None, "MET-004": {"1.2.1"}}
    assert [i.id for i in s.issues] == ["CFG-047", "CFG-047", "CFG-048", "CFG-048", "CFG-047"]
    assert all(i.severity == SEV_WARN for i in s.issues)
    assert parse_silence(["MET-004:1.2.1", "MET-004"]).rules == {"MET-004": None}
    assert parse_silence(["MET-004", "MET-004:1.2.1"]).rules == {"MET-004": None}


def test_silence_by_code_and_by_wbs_code():
    s = parse_silence(["FEC-007", "MET-004:1.2.1", "MET-004:2.1-D"])
    issues = [Issue("FEC-007", "1.1.0", "A", "x"), Issue("MET-004", "1.2.1.0", "B", "y"),
              Issue("MET-004", "1.2.10", "C", "z"), Issue("MET-004", "2.1.0-D", "D", "w"),
              Issue("EDT-009", "3", "E", "v")]
    out = s.apply(issues)
    assert [(i.id, i.code) for i in out] == [("MET-004", "1.2.10"), ("EDT-009", "3"), ("CFG-049", "")]
    assert out[-1].severity == SEV_INFO and "3 observaciones silenciadas" in out[-1].message
    assert "FEC-007 (1), MET-004 (2)" in out[-1].message
    assert parse_silence([]).apply(issues) == issues                 # sin silenciar: sin resumen


def test_config_key_silenciar(tmp_path):
    path = tmp_path / "c.toml"
    path.write_text('[validacion]\nsilenciar = ["FEC-007", "LOG-004:1.2"]\n', encoding="utf-8")
    assert load_config(path).silence == ["FEC-007", "LOG-004:1.2"]
    path.write_text('[validacion]\nsilenciar = "FEC-007, LOG-004"\n', encoding="utf-8")
    assert load_config(path).silence == ["FEC-007", "LOG-004"]
    path.write_text("[validacion]\nsilenciar = 3\n", encoding="utf-8")
    with pytest.raises(ConfigError) as exc:
        load_config(path)
    assert exc.value.id == "CFG-004"


def problem_xml(tmp_path):
    """Un error de código EDT (EDT-005: 1.14.2 bajo 1.2) y tareas sin vínculos."""
    return write_xml(tmp_path / "proyecto.xml", project_xml([
        task_xml(1, 1, "Fase", wbs="1", start=S, finish=F, baselines=[(0, S, F, 0)]),
        task_xml(2, 2, "Grupo", wbs="1.2", start=S, finish=F, baselines=[(0, S, F, 0)]),
        task_xml(3, 3, "Mal", wbs="1.14.2", start=S, finish=F, baselines=[(0, S, F, 0)]),
    ], header="<StatusDate>2026-02-27T17:00:00</StatusDate>"))


def test_validar_shows_codes_and_silenced_errors_do_not_count(tmp_path, capsys):
    path = problem_xml(tmp_path)
    assert main(["validar", "--origen", str(path)]) == 1
    out = capsys.readouterr().out
    assert "[X] EDT-005 1.14.2 Mal: El código EDT no depende" in out and "CFG-049" not in out
    assert main(["validar", "--origen", str(path), "--silenciar", "EDT-005,LOG-001", "--silenciar", "ZZZ"]) == 0
    out = capsys.readouterr().out
    assert "EDT-005 1.14.2" not in out and "LOG-001" not in out.replace("LOG-001 (1)", "")
    assert "[!] CFG-047 'ZZZ' no es un código del catálogo" in out
    assert "[i] CFG-049 2 observaciones silenciadas ([validacion] silenciar o --silenciar): EDT-005 (1), LOG-001 (1)." \
        in out


def test_generar_writes_codes_and_hides_silenced_issues_in_the_validation_sheet(tmp_path, capsys):
    path = problem_xml(tmp_path)
    cfg = tmp_path / "c.toml"
    cfg.write_text('[validacion]\nsilenciar = ["EDT-005:1.14.2", "CFG-001"]\n', encoding="utf-8")
    out_path = tmp_path / "libro.xlsx"
    assert main(["generar", "--origen", str(path), "--config", str(cfg), "--salida", str(out_path)]) == 0
    out = capsys.readouterr().out
    assert "[i] LIB-001 El PV y el EV de cada fila" in out and "EDT-005 1.14.2" not in out
    rows = [r for r in load_workbook(out_path)["Validación"].iter_rows(min_row=4, values_only=True) if r[0]]
    ids = [r[1] for r in rows]
    assert "EDT-005" not in ids and "LIB-001" in ids
    assert ("Advertencia", "CFG-048") == rows[ids.index("CFG-048")][:2]
    summary = rows[ids.index("CFG-049")]
    assert summary[0] == "Info" and "1 observación silenciada" in summary[4] and "EDT-005 (1)" in summary[4]


def test_fatal_errors_show_their_code(tmp_path, capsys):
    assert main(["generar", "--origen", str(tmp_path / "no_existe.xml")]) == 2
    assert capsys.readouterr().err.startswith("ERROR LEC-001: No existe el archivo de origen")
    assert main(["validar", "--origen", str(problem_xml(tmp_path)), "--config", str(tmp_path / "x.toml")]) == 2
    assert capsys.readouterr().err.startswith("ERROR CFG-001: No se encontró el archivo de configuración")


def test_actualizar_and_analizar_silence_their_notices(example_files, tmp_path, capsys):
    book = tmp_path / "EVM.xlsx"
    common = ["--config", str(example_files["config"])]
    assert main(["generar", "--origen", str(example_files["base"]), *common, "--salida", str(book)]) == 0
    capsys.readouterr()
    assert main(["actualizar", "--origen", str(example_files["jun"]), *common, "--libro", str(book),
                 "--sin-respaldo", "--simular"]) == 0
    out = capsys.readouterr().out
    assert "[i] ACT-026 El AC se carga a mano" in out and "[!] ACT-022 Ninguna fila tiene BAC" in out
    assert main(["actualizar", "--origen", str(example_files["jun"]), *common, "--libro", str(book),
                 "--sin-respaldo", "--silenciar", "ACT-026,ACT-022"]) == 0
    out = capsys.readouterr().out
    assert "[i] ACT-026" not in out and "[!] ACT-022" not in out
    assert "observaciones silenciadas ([validacion] silenciar o --silenciar): ACT-022 (1), ACT-026 (1)." in out
    assert main(["analizar", "--libro", str(book), "--sin-respaldo", "--silenciar", "FEC-999"]) == 0
    assert "[!] CFG-047 'FEC-999' no es un código del catálogo" in capsys.readouterr().out
