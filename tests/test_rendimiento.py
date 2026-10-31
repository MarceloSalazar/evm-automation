"""Optimizaciones del tiempo de lectura de MS Project y de escritura del libro: dan los mismos resultados.

- Estilos de las celdas (workbook.styled): mismo estilo que la asignación de openpyxl.
- Hojas '% PV' y '% EV': las celdas vacías de las filas de ingreso no llevan formato; 'actualizar' da formato a las
  que carga.
- read_book: el libro con los valores calculados se lee solo si una celda de ingreso tiene una fórmula.
- Calendario: el reparto de la duración entre los meses es el mismo que el día a día.
- Lectura de MS Project (COM simulado): generar y actualizar no leen vínculos ni delimitaciones; una tarea sin
  predecesoras no recorre sus vínculos; los alias de los campos personalizados se leen una vez.
- Guardado rápido (evm/fast_save.py, [exportacion] guardado_rapido): el archivo es el mismo, byte a byte, que con el
  guardado de openpyxl; con la opción desactivada, o con otra versión de openpyxl, se guarda en el modo normal.
"""
import io
import random
import re
import zipfile
from datetime import date, datetime, time, timedelta

import pytest
from openpyxl import Workbook, load_workbook

from evm import fast_save
from evm.cli import main
from evm.config import Config, load_config
from evm.messages import EvmError
from evm.update import read_book, run_update
from evm.wbs import build_report
from evm.workbook import (FILL_INPUT, FILL_L1, FIRST_MONTH_COL, FMT_PCT_MONTH, FONT_BOLD, FONT_ITALIC, indent,
                          read_structure, styled, write_workbook)
from evm.workcalendar import WorkCalendar

from builders import P, T, set_bac
from mspdi import project_xml, task_xml, write_xml

S, F = datetime(2026, 1, 5, 8), datetime(2026, 3, 27, 17)


# ----------------------------------------------------------------------------
# Estilos
# ----------------------------------------------------------------------------
def test_styled_gives_the_same_style_as_openpyxl():
    wb = Workbook()
    ws = wb.active
    a = ws.cell(row=1, column=1)
    a.font, a.fill, a.number_format, a.alignment = FONT_BOLD, FILL_L1, FMT_PCT_MONTH, indent(3)
    attrs = dict(font=FONT_BOLD, fill=FILL_L1, number_format=FMT_PCT_MONTH, alignment=indent(3))
    b = styled(ws.cell(row=1, column=2), **attrs)                # primera vez: lo calcula openpyxl
    c = styled(ws.cell(row=1, column=3), **attrs)                # siguientes: índices guardados
    assert tuple(b._style) == tuple(a._style) == tuple(c._style)
    d = ws.cell(row=1, column=4)
    d.font = FONT_ITALIC
    styled(d, font=None, fill=FILL_INPUT)                        # None no se asigna
    assert d.font.i and d.fill.fgColor.rgb.endswith("FFF2CC")
    other = Workbook()                                           # los índices son de cada libro
    e = styled(other.active.cell(row=1, column=1), fill=FILL_INPUT, font=FONT_BOLD)
    assert e.font.b and e.fill.fgColor.rgb.endswith("FFF2CC")


def test_styles_survive_saving(tmp_path):
    wb = Workbook()
    for c in range(1, 50):
        styled(wb.active.cell(row=1, column=c, value=c / 100), number_format=FMT_PCT_MONTH, font=FONT_BOLD,
               fill=FILL_L1)
    wb.save(tmp_path / "l.xlsx")
    ws = load_workbook(tmp_path / "l.xlsx").active
    assert all(ws.cell(row=1, column=c).number_format == FMT_PCT_MONTH and ws.cell(row=1, column=c).font.b
               and ws.cell(row=1, column=c).fill.fgColor.rgb.endswith("BDD7EE") for c in range(1, 50))


# ----------------------------------------------------------------------------
# Hojas de % y lectura del libro
# ----------------------------------------------------------------------------
def small_book(tmp_path, cfg=None):
    tasks = [T(1, 1, "1", "Fase"), T(2, 2, "1.1", "A", method="% Completado"),
             T(3, 2, "1.2", "B", method="% Completado", bs=datetime(2026, 2, 2, 8))]
    cfg = cfg or Config(months=12)
    path = tmp_path / "EVM.xlsx"
    write_workbook(build_report(P(tasks), cfg), cfg, path)
    set_bac(path, 1000.0)
    return path, tasks, cfg


def test_empty_percent_cells_have_no_format_and_update_formats_what_it_loads(tmp_path):
    path, tasks, cfg = small_book(tmp_path)
    wb = load_workbook(path)
    rows = {s.code: s.row for s in read_structure(wb)}
    pv, ev = wb["% PV"], wb["% EV"]
    r = rows["1.1.0"]
    assert pv.cell(r, FIRST_MONTH_COL).number_format == FMT_PCT_MONTH          # mes con % planificado
    assert pv.cell(r, FIRST_MONTH_COL + 5).number_format == "General"          # mes sin %: celda vacía
    assert ev.cell(r, FIRST_MONTH_COL).number_format == "General"
    assert ev.cell(rows["1.0.0"], FIRST_MONTH_COL + 5).number_format == FMT_PCT_MONTH   # subtotal: fórmula
    tasks[1].percent_complete = 40
    run_update(P(tasks), cfg, path, date(2026, 1, 31), backup=False)
    cell = load_workbook(path)["% EV"].cell(r, FIRST_MONTH_COL)
    assert cell.value == pytest.approx(0.4) and cell.number_format == FMT_PCT_MONTH


def test_read_book_loads_the_calculated_values_only_for_formulas(tmp_path, monkeypatch):
    import evm.update as update
    path, _tasks, _cfg = small_book(tmp_path)
    calls = []
    original = update.load_workbook

    def counting(*args, **kw):
        calls.append(kw.get("data_only", False))
        return original(*args, **kw)

    monkeypatch.setattr(update, "load_workbook", counting)
    read_book(path)
    assert calls == [False]                         # sin fórmulas en las celdas de ingreso: una sola lectura

    # Fórmula escrita a mano en una celda de AC: sin valor calculado (libro guardado sin Excel), LIB-012
    wb = load_workbook(path)
    r = next(s.row for s in read_structure(wb) if s.code == "1.1.0")
    wb["AC"].cell(r, FIRST_MONTH_COL, "=1000+234.5")
    wb.save(path)
    calls.clear()
    with pytest.raises(EvmError, match="fórmula sin valor numérico calculado") as exc:
        read_book(path)
    assert exc.value.id == "LIB-012" and calls == [False, True]

    # Con el valor que calculó Excel: se usa ese valor y la celda se puede convertir en número (ACT-025)
    with zipfile.ZipFile(path) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    ac = next(n for n, d in parts.items() if n.startswith("xl/worksheets/") and b"1000+234.5" in d)
    parts[ac] = re.sub(rb"<f>1000\+234\.5</f><v\s*/?>(</v>)?", b"<f>1000+234.5</f><v>1234.5</v>", parts[ac])
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for n, d in parts.items():
            z.writestr(n, d)
    book = read_book(path, allow_fix=True)
    assert book.ac[r][0] == 1234.5
    assert [(sheet, value) for sheet, _r, _c, value, _why in book.fixable] == [("AC", 1234.5)]


# ----------------------------------------------------------------------------
# Calendario
# ----------------------------------------------------------------------------
def day_by_day(cal: WorkCalendar, start: datetime, finish: datetime) -> dict:
    """Reparto de referencia: cada ventana de cada día, recortada al intervalo."""
    out: dict = {}
    d = start.date()
    while d <= finish.date():
        day0 = datetime.combine(d, time(0, 0))
        for a, b in cal.windows(d):
            lo, hi = max(day0 + timedelta(minutes=a), start), min(day0 + timedelta(minutes=b), finish)
            if hi > lo:
                out[(d.year, d.month)] = out.get((d.year, d.month), 0.0) + (hi - lo).total_seconds() / 60
        d += timedelta(days=1)
    return out


def test_month_split_equals_the_day_by_day_split():
    cal = WorkCalendar()
    cal.add_holidays([date(2026, 1, 1), date(2026, 5, 1), date(2026, 12, 25)])
    cal.work_weeks.append((date(2026, 7, 1), date(2026, 8, 31), {0: [(420, 900)], 4: [], 5: [(480, 720)]}))
    cal.exceptions[date(2026, 3, 3)] = [(600, 660)]
    cal.exceptions[date(2026, 3, 7)] = [(0, 1440)]               # sábado laborable todo el día
    rng = random.Random(7)
    for _ in range(3000):
        s = datetime(2025, 12, 1) + timedelta(minutes=rng.randint(0, 60 * 24 * 500))
        f = s + timedelta(minutes=rng.choice([0, 5, 60, 600, 3000, 60 * 24 * rng.randint(1, 400)]))
        assert cal.minutes_by_month(s, f) == day_by_day(cal, s, f), (s, f)


# ----------------------------------------------------------------------------
# Lectura de MS Project
# ----------------------------------------------------------------------------
def counting_tasks():
    """Tareas COM simuladas que anotan cada propiedad leída (cada lectura es una llamada a MS Project)."""
    from test_reader_com_fake import FakeTask

    reads: list[str] = []

    class Counted:
        def __getattribute__(self, name):
            if name[:1].isupper():
                reads.append(name)
            return object.__getattribute__(self, name)

    class Task(Counted, FakeTask):
        pass

    class Dep(Counted):
        def __init__(self, a, b):
            self.From, self.To, self.Type, self.Lag = a, b, 1, 0

    tasks = [Task(i, 1, f" Tarea {i} ", str(i)) for i in range(1, 6)]
    for t in tasks:
        t.Duration, t.ConstraintType, t.ConstraintDate, t.TaskDependencies = 4800, 0, "NA", []
        t.UniqueIDPredecessors = ""
    for a, b in zip(tasks, tasks[1:]):
        dep = Dep(a, b)
        a.TaskDependencies.append(dep)
        b.TaskDependencies.append(dep)
        b.UniqueIDPredecessors = str(a.UniqueID)
    reads.clear()                                   # solo cuentan las lecturas del programa
    return tasks, reads


def test_com_reads_links_only_when_needed():
    from test_reader_com_fake import make_session
    tasks, reads = counting_tasks()
    p = make_session(tasks).read_project(Config(method_field="Método EVM"), links=False)
    assert not {"TaskDependencies", "UniqueIDPredecessors", "ConstraintType", "ConstraintDate"} & set(reads)
    assert all(t.predecessors == [] for t in p.tasks)
    assert p.tasks[0].name == "Tarea 1"                          # sin espacios, como con XML

    reads.clear()
    p = make_session(tasks).read_project(Config(method_field="Método EVM"))
    assert [len(t.predecessors) for t in p.tasks] == [0, 1, 1, 1, 1]
    assert p.tasks[1].predecessors[0].uid == 1
    # La primera tarea no tiene predecesoras: no recorre sus vínculos (solo los de las demás)
    assert reads.count("TaskDependencies") == len(tasks) - 1


def test_com_reads_custom_field_aliases_once():
    from test_reader_com_fake import FakeApp, FakeTask, make_session
    asked: list[int] = []

    class App(FakeApp):
        def CustomFieldGetName(self, fid):
            asked.append(fid)
            return super().CustomFieldGetName(fid)

    s = make_session([FakeTask(1, 1, "A", "1", method="50/50")])
    s.app = App()
    cfg = Config(method_field="Método EVM", milestone_weight_field="Peso hito", visible_field="Visible",
                 description_field="Sin alias")
    s.read_project(cfg)
    assert len(asked) == len(set(asked)) == 50                   # Texto1..30 y Número1..20, una vez


def test_xml_reads_links_only_when_needed(tmp_path):
    from evm.reader_xml import read_xml
    path = write_xml(tmp_path / "p.xml", project_xml([
        task_xml(1, 1, "A", wbs="1", start=S, finish=F, baselines=[(0, S, F, 0)],
                 extra="<ConstraintType>2</ConstraintType><ConstraintDate>2026-01-05T08:00:00</ConstraintDate>"),
        task_xml(2, 1, "B", wbs="2", start=S, finish=F, baselines=[(0, S, F, 0)],
                 extra="<PredecessorLink><PredecessorUID>1</PredecessorUID><Type>1</Type></PredecessorLink>"),
    ]))
    full, light = read_xml(path, Config()), read_xml(path, Config(), links=False)
    assert full.tasks[1].predecessors and full.tasks[0].constraint_type == 2
    assert light.tasks[1].predecessors == [] and light.tasks[0].constraint_type == 0
    assert light.tasks[0].constraint_date is None


def test_generar_and_actualizar_do_not_read_links(tmp_path, monkeypatch):
    import evm.reader_xml as reader
    seen = []
    original = reader.read_xml

    def spy(path, cfg, links=True):
        seen.append(links)
        return original(path, cfg, links=links)

    monkeypatch.setattr(reader, "read_xml", spy)
    path = write_xml(tmp_path / "p.xml", project_xml([
        task_xml(1, 1, "A", wbs="1", start=S, finish=F, baselines=[(0, S, F, 0)]),
    ], header="<StatusDate>2026-01-30T17:00:00</StatusDate>"))
    book = tmp_path / "EVM.xlsx"
    assert main(["generar", "--origen", str(path), "--salida", str(book), "--meses", "6"]) == 0
    assert main(["actualizar", "--origen", str(path), "--libro", str(book), "--sin-respaldo"]) == 0
    assert main(["validar", "--origen", str(path)]) == 0
    assert main(["generar-fichas", "--origen", str(path), "--salida", str(tmp_path / "f.md")]) == 0
    assert seen == [False, False, True, True]


# ----------------------------------------------------------------------------
# Guardado rápido
# ----------------------------------------------------------------------------
def saved_parts(wb, fast: bool) -> dict[str, bytes]:
    """Partes del .xlsx guardado en memoria (sin docProps/core.xml, que lleva la hora del guardado)."""
    buffer = io.BytesIO()
    with fast_save.fast_cells(fast):
        wb.save(buffer)
    with zipfile.ZipFile(buffer) as z:
        return {n: z.read(n) for n in z.namelist() if n != "docProps/core.xml"}


def assert_same_file(make) -> None:
    """make() devuelve el libro: cada modo guarda uno distinto (openpyxl modifica el libro al guardarlo, por lo que
    un segundo guardado del mismo libro ya no es igual al primero)."""
    normal, fast = saved_parts(make(), False), saved_parts(make(), True)
    assert list(normal) == list(fast)
    assert [n for n in normal if normal[n] != fast[n]] == []


def edge_case_book() -> Workbook:
    """Celdas de todos los tipos: las simples van por el guardado rápido y las demás, por openpyxl."""
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont
    from openpyxl.comments import Comment
    from openpyxl.worksheet.formula import ArrayFormula
    wb = Workbook()
    ws = wb.active
    values = [1, 2.5, -0.1, 10 ** 20, 1e-30, 0, True, False, float("nan"), float("inf"), "texto", " inicio",
              "fin ", "   ", "a & b < c > d", "línea 1\nlínea 2", "ñ € «»", "=SUM(A1:A2)", "=", '="x<y"&"&"',
              datetime(2026, 1, 5, 8, 30), date(2026, 1, 1), time(8, 0), timedelta(hours=5), "#N/A"]
    for i, v in enumerate(values, start=1):
        ws.cell(row=i, column=1, value=v)
        styled(ws.cell(row=i, column=2, value=v), font=FONT_BOLD, number_format=FMT_PCT_MONTH)
    ws.cell(row=1, column=3, value="#N/A").data_type = "s"            # texto que parece un error (put_text)
    styled(ws.cell(row=2, column=3), fill=FILL_INPUT)                  # celda vacía con estilo
    ws["D1"] = ArrayFormula("D1:D2", "=SUM(A1:A2*B1:B2)")
    ws["D3"] = "enlace"
    ws["D3"].hyperlink = "https://example.com/?a=1&b=2"
    ws["D4"] = "con nota"
    ws["D4"].comment = Comment("Nota", "Autor")
    ws["D5"] = CellRichText(["normal ", TextBlock(InlineFont(b=True), "negrita")])
    return wb


def test_fast_save_writes_the_same_file(tmp_path, book_jun):
    assert fast_save.available()
    assert_same_file(edge_case_book)
    assert_same_file(lambda: load_workbook(book_jun))                  # libro del ejemplo actualizado a junio
    for cfg in (Config(months=24, multi_enabled=True), Config(months=24, cost_review=True,
                                                               cost_review_alignment=True)):
        path = tmp_path / "libro.xlsx"
        tasks = [T(1, 1, "1", "Fase", company="Ambas"), T(2, 2, "1.1", "A", company="EmpresaA"),
                 T(3, 2, "1.2", "B", company="EmpresaB"), T(4, 2, "", "Suelta")]
        write_workbook(build_report(P(tasks), cfg), cfg, path)
        assert_same_file(lambda: load_workbook(path))


def test_fast_save_is_limited_to_the_save_and_restored_after_errors():
    from openpyxl.cell._writer import etree_write_cell
    from openpyxl.worksheet import _writer as sheet_writer
    with fast_save.fast_cells(True):
        assert sheet_writer.write_cell is fast_save._write_cell
    assert sheet_writer.write_cell is etree_write_cell
    with fast_save.fast_cells(False):
        assert sheet_writer.write_cell is etree_write_cell
    with pytest.raises(RuntimeError):
        with fast_save.fast_cells(True):
            raise RuntimeError("guardado interrumpido")
    assert sheet_writer.write_cell is etree_write_cell


def test_fast_save_falls_back_with_other_openpyxl_versions_or_writers(monkeypatch):
    import openpyxl
    from openpyxl.cell._writer import etree_write_cell
    from openpyxl.worksheet import _writer as sheet_writer
    monkeypatch.setattr(openpyxl, "__version__", "3.2.0")              # versión no probada
    assert not fast_save.available()
    with fast_save.fast_cells(True):
        assert sheet_writer.write_cell is etree_write_cell
    monkeypatch.undo()

    def other_writer(xf, worksheet, cell, styled=None):                # p. ej. el escritor de lxml
        return etree_write_cell(xf, worksheet, cell, styled)

    monkeypatch.setattr(sheet_writer, "write_cell", other_writer)
    assert not fast_save.available()
    with fast_save.fast_cells(True):
        assert sheet_writer.write_cell is other_writer


def test_fast_save_option(tmp_path, monkeypatch):
    calls = []
    original = fast_save._write_cell

    def counting(*args, **kw):
        calls.append(1)
        return original(*args, **kw)

    monkeypatch.setattr(fast_save, "_write_cell", counting)
    xml = write_xml(tmp_path / "p.xml", project_xml([
        task_xml(1, 1, "A", wbs="1", start=S, finish=F, baselines=[(0, S, F, 0)]),
    ], header="<StatusDate>2026-01-30T17:00:00</StatusDate>"))
    normal = tmp_path / "normal.toml"
    normal.write_text("[exportacion]\nguardado_rapido = false\n", encoding="utf-8")
    assert Config().fast_save and not load_config(normal).fast_save
    book = tmp_path / "EVM.xlsx"
    for argv in (["generar", "--origen", str(xml), "--salida", str(book), "--meses", "6"],
                 ["actualizar", "--origen", str(xml), "--libro", str(book), "--sin-respaldo"],
                 ["analizar", "--libro", str(book), "--sin-respaldo"]):
        assert main([*argv, "--config", str(normal)]) == 0
        assert not calls                                               # modo normal
        assert main([*argv[:1], *argv[1:], *(["--sobrescribir"] if argv[0] == "generar" else [])]) == 0
        assert calls                                                   # guardado rápido (por defecto)
        calls.clear()
