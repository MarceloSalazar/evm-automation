"""Ceros del código EDT. Un 0 intermedio (1.2.0.1): de nivel 4 por defecto (actividad de 1.2) o, con [exportacion]
ignorar_ceros_intermedios = true, de nivel 3 (hijo directo de 1.2, como 1.2.1.0). Los ceros finales son relleno:
1.2, 1.2.0 y 1.2.0.0 son el mismo código (duplicados, asociación por código en actualizar y --migrar-desde)."""
from datetime import date

import pytest
from openpyxl import load_workbook

from evm.cli import main
from evm.config import Config, ConfigError, load_config
from evm.migrate import plan_migration, prior_from_book
from evm.update import book_ignores_interior_zeros, read_book, run_update
from evm.model import ProjectData
from evm.wbs import (KIND_LEAF, KIND_RESIDUAL, KIND_SUMMARY, SEV_ERROR, SEV_WARN, TOPIC_WBS, build_report, code_key,
                     code_level, has_interior_zero)
from evm.workbook import SHEET_VALID, read_meta, write_workbook

from builders import F, P, S, by_name, issues_with, report, rows_by_code, set_bac, tree
from mspdi import project_xml, task_xml, write_xml

TASKS = '''
    1.0.0.0 Fase
      1.0.0.3 Directo de 1
      1.2.0.0 Elemento
        1.2.0.1 Insertado
        1.2.1.0 Paquete
          1.2.1.1 Actividad
'''
IGNORE = Config(ignore_interior_zeros=True)


def wbs_issues(rep, severity=None):
    return [i for i in rep.issues if i.topic == TOPIC_WBS and (severity is None or i.severity == severity)]


def ids(issues, code: str) -> list:
    return [i for i in issues if i.id == code]


# ----------------------------------------------------------------------------
# Nivel del código y configuración
# ----------------------------------------------------------------------------
@pytest.mark.parametrize("code, default, ignoring", [
    ("1.2.0.1", 4, 3),
    ("1.0.0.3", 4, 2),
    ("1.2.0.1.0", 4, 3),            # los ceros finales son relleno en los dos criterios
    ("1.2.0.0.1", 5, 3),
    ("1.2.1.1", 4, 4),
    ("1.2.0", 2, 2),
    ("INTERNO", 1, 1),
    ("1.2.A", 3, 3),
])
def test_code_level(code, default, ignoring):
    assert code_level(code) == default
    assert code_level(code, ignore_interior_zeros=True) == ignoring


def test_has_interior_zero():
    assert [has_interior_zero(c) for c in ("1.2.0.1", "1.0.3", "1.2.0.0", "1.2.1", "INTERNO", "1.0.A", "", "1.2.0-D")] \
        == [True, True, False, False, False, False, False, False]


def test_config_key(tmp_path):
    assert load_config(None).ignore_interior_zeros is False
    path = tmp_path / "config.toml"
    path.write_text("[exportacion]\nignorar_ceros_intermedios = true\n", encoding="utf-8")
    assert load_config(path).ignore_interior_zeros is True
    path.write_text('[exportacion]\nignorar_ceros_intermedios = "si"\n', encoding="utf-8")
    with pytest.raises(ConfigError, match="'ignorar_ceros_intermedios' debe ser true o false"):
        load_config(path)


# ----------------------------------------------------------------------------
# Filas del libro
# ----------------------------------------------------------------------------
def test_interior_zero_is_an_activity_of_the_previous_element_by_default():
    """Por defecto, 1.2.0.1 es de nivel 4: con nivel_maximo = 3 se suma en la fila -D de 1.2.0.0."""
    rep = report(tree(TASKS))
    assert not wbs_issues(rep, SEV_ERROR) and not wbs_issues(rep, SEV_WARN)
    rows = rows_by_code(rep)
    assert list(rows) == ["1.0.0.0", "1.2.0.0", "1.2.1.0", "1.2.0.0-D", "1.0.0.0-D"]
    assert rows["1.2.0.0"].kind == KIND_SUMMARY and rows["1.2.1.0"].kind == KIND_LEAF
    assert rows["1.2.0.0-D"].kind == KIND_RESIDUAL and [t.name for t in rows["1.2.0.0-D"].leaves] == ["Insertado"]
    assert [t.name for t in rows["1.0.0.0-D"].leaves] == ["Directo de 1"]
    assert issues_with(rep, "de nivel mayor a 3: forman la fila 1.2.0.0-D")
    assert rows_by_code(report(tree(TASKS), Config(max_level=4)))["1.2.0.1"].kind == KIND_LEAF


def test_interior_zero_without_other_rows_is_summed_in_its_element():
    rep = report(tree('''
        1 Fase
          1.2 Elemento
            1.2.0.1 Insertado
            1.2.0.2 Otro
    '''))
    rows = rows_by_code(rep)
    assert list(rows) == ["1.0.0", "1.2.0"] and rows["1.2.0"].kind == KIND_LEAF
    assert [t.name for t in rows["1.2.0"].leaves] == ["Insertado", "Otro"]


def test_ignoring_interior_zeros_makes_them_direct_children():
    """1.2.0.1 se inserta debajo de 1.2.0.0 sin renumerar a 1.2.1.0: es su hijo directo, de nivel 3."""
    rep = report(tree(TASKS), Config(max_level=3, ignore_interior_zeros=True))
    assert not wbs_issues(rep, SEV_ERROR) and not wbs_issues(rep, SEV_WARN)
    rows = rows_by_code(rep)
    assert list(rows) == ["1.0.0.0", "1.0.0.3", "1.2.0.0", "1.2.0.1", "1.2.1.0"]
    assert rows["1.2.0.1"].kind == KIND_LEAF and rows["1.2.0.1"].parent.code == "1.2.0.0"
    assert rows["1.2.1.0"].kind == KIND_LEAF and "1.2.1.1" not in rows           # nivel 4: se suma en 1.2.1.0
    assert rows["1.0.0.3"].parent.code == "1.0.0.0"


# ----------------------------------------------------------------------------
# Controles de los códigos EDT
# ----------------------------------------------------------------------------
@pytest.mark.parametrize("ignore", [False, True])
def test_interior_zero_under_another_element_is_an_error(ignore):
    rep = report(tree('''
        1 Fase
          1.0.3 Bien
          1.2 Elemento
            1.2.1 Paquete
              1.2.0.1 Mal ubicado
    '''), Config(ignore_interior_zeros=ignore))
    by = {i.name: i for i in wbs_issues(rep)}
    wrong = by["Mal ubicado"]
    assert wrong.severity == SEV_ERROR and "no depende del de su tarea superior" in wrong.message
    assert "Bien" not in by


@pytest.mark.parametrize("ignore, levels", [(False, (4, 5)), (True, (3, 4))])
def test_reported_levels_follow_the_setting(ignore, levels):
    rep = report(tree('''
        1.2.0.1 Sin superior
        2 Fase
          2.1 Grupo
            2.1.0.3.1 Salta
    '''), Config(ignore_interior_zeros=ignore))
    alone, skip = ids(rep.issues, "EDT-003"), ids(rep.issues, "EDT-006")
    assert len(alone) == 1 and f"Código de nivel {levels[0]} sin una tarea superior" in alone[0].message
    assert "debería depender de la tarea 1.2.0" in alone[0].message
    assert len(skip) == 1 and f"de nivel 2, y su código es de nivel {levels[1]}" in skip[0].message


# ----------------------------------------------------------------------------
# actualizar: el criterio con que se generó el libro
# ----------------------------------------------------------------------------
def book(tmp_path, cfg: Config):
    path = tmp_path / "libro.xlsx"
    write_workbook(build_report(P(tree(TASKS)), cfg), cfg, path)
    set_bac(path, 1000)
    return path


def progress(**done) -> list:
    tasks = tree(TASKS)
    for name, value in done.items():
        by_name(tasks)[name.replace("_", " ")].percent_complete = value
    return tasks


def test_update_keeps_the_rows_of_a_book_that_ignores_interior_zeros(tmp_path):
    path = book(tmp_path, IGNORE)
    assert read_meta(load_workbook(path))["ignorar_ceros_intermedios"] is True
    res = run_update(P(progress(Insertado=50)), Config(), path, date(2026, 3, 31), backup=False, dry_run=True)
    assert {r.code: r.ev_cum for r in res.rows}["1.2.0.1"] == pytest.approx(500)
    warn = ids(res.warnings, "ACT-050")
    assert len(warn) == 1 and "se generó con ignorar_ceros_intermedios = true" in warn[0].message
    assert "(1.0.0.3, 1.2.0.1) tendrían otra fila" in warn[0].message and "--migrar-desde" in warn[0].message
    assert not ids(res.warnings, "ACT-015")
    res = run_update(P(progress(Insertado=50)), IGNORE, path, date(2026, 3, 31), backup=False, dry_run=True)
    assert not ids(res.warnings, "ACT-050")


def test_update_keeps_the_other_tasks_row_of_a_book_that_counts_interior_zeros(tmp_path):
    path = book(tmp_path, Config())
    assert read_meta(load_workbook(path))["ignorar_ceros_intermedios"] is False
    res = run_update(P(progress(Insertado=50)), IGNORE, path, date(2026, 3, 31), backup=False, dry_run=True)
    assert {r.code: r.ev_cum for r in res.rows}["1.2.0.0-D"] == pytest.approx(500)     # sigue en la fila -D
    assert len(ids(res.warnings, "ACT-050")) == 1 and not ids(res.warnings, "ACT-015")


@pytest.mark.parametrize("meta, expected", [
    ({"ignorar_ceros_intermedios": True, "version": "2.4.0"}, True),
    ({"ignorar_ceros_intermedios": False}, False),
    ({"version": "2.4.0"}, False),          # sin la clave: false (no se deduce de la versión)
    ({}, False),
])
def test_criterion_of_the_book(meta, expected):
    assert book_ignores_interior_zeros(meta) is expected


# ----------------------------------------------------------------------------
# generar --migrar-desde
# ----------------------------------------------------------------------------
def test_migration_sums_the_rows_that_move_into_other_tasks_rows(tmp_path):
    old, new = tmp_path / "old.xlsx", tmp_path / "new.xlsx"
    before = Config(months=12, ignore_interior_zeros=True)
    write_workbook(build_report(P(tree(TASKS)), before), before, old)
    set_bac(old, 1000)
    run_update(P(progress(Insertado=50, Directo_de_1=100)), before, old, date(2026, 2, 28), backup=False)
    cfg = Config(months=12)
    rep = build_report(P(tree(TASKS)), cfg)
    prior = prior_from_book(old, "Programa")
    assert prior.ignore_interior_zeros is True
    rep.issues.extend(plan_migration(rep, prior, cfg.max_level, cfg.ignore_interior_zeros))
    write_workbook(rep, cfg, new, prior)
    b = read_book(new)
    rows = {s.code: s.row for s in b.structure}
    assert b.bac[rows["1.2.0.0-D"]] == 1000 and b.bac[rows["1.0.0.0-D"]] == 1000
    assert sum(v or 0 for v in b.ev[rows["1.2.0.0-D"]]) == pytest.approx(500)
    assert sum(v or 0 for v in b.ev[rows["1.0.0.0-D"]]) == pytest.approx(1000)
    assert issues_with(rep, "Se sumó el historial de EV/AC de la fila anterior 1.2.0.1")
    mig = ids(rep.issues, "MIG-020")
    assert len(mig) == 1 and "ignorar_ceros_intermedios = true y el nuevo con false" in mig[0].message
    assert not ids(rep.issues, "MIG-014")


def test_migration_cannot_split_an_other_tasks_row(tmp_path):
    old = tmp_path / "old.xlsx"
    before = Config(months=12)
    write_workbook(build_report(P(tree(TASKS)), before), before, old)
    set_bac(old, 1000)
    run_update(P(progress(Insertado=50)), before, old, date(2026, 2, 28), backup=False)
    rep = build_report(P(tree(TASKS)), IGNORE)
    rep.issues.extend(plan_migration(rep, prior_from_book(old, "Programa"), 3, True))
    assert ids(rep.issues, "MIG-020")
    assert sorted(i.code for i in ids(rep.issues, "MIG-014")) == ["1.0.0.0-D", "1.2.0.0-D"]


def test_migration_without_interior_zeros_does_not_warn(tmp_path):
    old = tmp_path / "old.xlsx"
    text = '''
        1 Fase
          1.1 A
            1.1.1 A1
    '''
    write_workbook(build_report(P(tree(text)), IGNORE), IGNORE, old)
    rep = build_report(P(tree(text)), Config())
    assert not ids(plan_migration(rep, prior_from_book(old, "Programa"), 3), "MIG-020")


# ----------------------------------------------------------------------------
# Ceros finales: 1.2 = 1.2.0 = 1.2.0.0
# ----------------------------------------------------------------------------
@pytest.mark.parametrize("code, key", [
    ("1.2", "1.2"), ("1.2.0", "1.2"), ("1.2.0.0", "1.2"), (" 1.2.0 ", "1.2"), ("1.0.0", "1"),
    ("1.2.0.1", "1.2.0.1"), ("1.2.1", "1.2.1"), ("1.2.0-D", "1.2-D"), ("1.2.0.0-D", "1.2-D"), ("INTERNO", "INTERNO"),
])
def test_code_key(code, key):
    assert code_key(code) == key


@pytest.mark.parametrize("ignore", [False, True])
def test_codes_that_differ_in_trailing_zeros_are_duplicates(ignore):
    rep = report(tree('''
        1 Fase
          1.2 A
          1.2.0.0 B
          1.3 C
          1.3.0 D
          1.4.0.1 E
          1.4.1 F
    '''), Config(ignore_interior_zeros=ignore))
    dup = ids(rep.issues, "EDT-009")
    assert [(i.code, i.severity) for i in dup] == [("1.2.0", SEV_ERROR), ("1.3.0", SEV_ERROR)]
    assert "2 tareas («A» (Id 2, 1.2.0), «B» (Id 3, 1.2.0.0))" in dup[0].message      # códigos distintos en el libro
    assert "2 tareas («C» (Id 4), «D» (Id 5))" in dup[1].message                       # mismo código en el libro


def other_project(tasks: list, name: str = "Otro") -> ProjectData:
    return ProjectData(name=name, source="otro.xml", backend="xml", tasks=tasks)


SHORT = '''
    1 Fase
      1.2 Elemento
        1.2.1 Paquete
        - Suelta
'''
LONG = '''
    1.0.0.0 Fase
      1.2.0.0 Elemento
        1.2.1.0 Paquete
        - Suelta
'''


def test_update_of_another_project_matches_codes_without_trailing_zeros(tmp_path):
    path = tmp_path / "libro.xlsx"
    write_workbook(build_report(P(tree(SHORT)), Config()), Config(), path)
    set_bac(path, 1000)
    tasks = tree(LONG)
    by_name(tasks)["Paquete"].percent_complete = 50
    by_name(tasks)["Suelta"].percent_complete = 100
    res = run_update(other_project(tasks), Config(), path, date(2026, 3, 31), backup=False, dry_run=True)
    assert ids(res.warnings, "ACT-010") and not ids(res.warnings, "ACT-013")
    ev = {r.code: r.ev_cum for r in res.rows}
    assert ev["1.2.1"] == pytest.approx(500) and ev["1.2.0-D"] == pytest.approx(1000)


def test_migration_from_another_project_matches_codes_without_trailing_zeros(tmp_path):
    old = tmp_path / "old.xlsx"
    cfg = Config(months=12)
    write_workbook(build_report(P(tree(SHORT)), cfg), cfg, old)
    set_bac(old, 1000)
    rep = build_report(other_project(tree(LONG)), cfg)
    prior = prior_from_book(old, rep.project.name)
    rep.issues.extend(plan_migration(rep, prior, cfg.max_level))
    assert ids(rep.issues, "MIG-005") and not ids(rep.issues, "MIG-014")              # asociadas por código EDT
    new = tmp_path / "new.xlsx"
    write_workbook(rep, cfg, new, prior)
    b = read_book(new)
    rows = {s.code: s.row for s in b.structure}
    assert b.bac[rows["1.2.1.0"]] == 1000 and b.bac[rows["1.2.0.0-D"]] == 1000


# ----------------------------------------------------------------------------
# Línea de comandos
# ----------------------------------------------------------------------------
def test_cli_generar_saves_the_criterion_and_migrates(tmp_path):
    xml = write_xml(tmp_path / "p.xml", project_xml([
        task_xml(1, 1, "Fase", wbs="1.0.0.0", start=S, finish=F, baselines=[(0, S, F, 0)]),
        task_xml(2, 2, "Elemento", wbs="1.2.0.0", start=S, finish=F, baselines=[(0, S, F, 0)]),
        task_xml(3, 3, "Insertado", wbs="1.2.0.1", start=S, finish=F, baselines=[(0, S, F, 0)]),
        task_xml(4, 3, "Paquete", wbs="1.2.1.0", start=S, finish=F, baselines=[(0, S, F, 0)]),
    ]))
    ignoring, counting = tmp_path / "ignorar.toml", tmp_path / "contar.toml"
    ignoring.write_text('[horizonte]\nmeses = 12\n[exportacion]\nignorar_ceros_intermedios = true\n',
                        encoding="utf-8")
    counting.write_text('[horizonte]\nmeses = 12\n', encoding="utf-8")
    old, new = tmp_path / "old.xlsx", tmp_path / "new.xlsx"
    assert main(["generar", "--origen", str(xml), "--config", str(ignoring), "--salida", str(old)]) == 0
    first = read_book(old)
    assert first.meta["ignorar_ceros_intermedios"] is True
    assert [s.code for s in first.structure if s.is_input] == ["1.2.0.1", "1.2.1.0"]
    assert main(["generar", "--origen", str(xml), "--config", str(counting), "--salida", str(new),
                 "--migrar-desde", str(old)]) == 0
    second = read_book(new)
    assert [s.code for s in second.structure if s.is_input] == ["1.2.1.0", "1.2.0.0-D"]
    codes = [c.value for row in load_workbook(new)[SHEET_VALID].iter_rows() for c in row]
    assert "MIG-020" in codes


def test_cli_shows_the_criterion_only_when_it_is_active(tmp_path, capsys):
    xml = write_xml(tmp_path / "p.xml", project_xml([
        task_xml(1, 1, "Fase", wbs="1.0.0.0", start=S, finish=F, baselines=[(0, S, F, 0)]),
        task_xml(2, 2, "Insertado", wbs="1.0.0.1", start=S, finish=F, baselines=[(0, S, F, 0)]),
    ]))
    cfg = tmp_path / "config.toml"
    cfg.write_text("[exportacion]\nignorar_ceros_intermedios = true\n", encoding="utf-8")
    main(["validar", "--origen", str(xml), "--config", str(cfg)])
    assert "Nivel máximo: 3 (sin los ceros intermedios)" in capsys.readouterr().out
    main(["validar", "--origen", str(xml)])
    out = capsys.readouterr().out
    assert "Nivel máximo: 3" in out and "ceros intermedios" not in out
