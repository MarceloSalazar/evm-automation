"""Hoja Cuadro de Mando: mes de reporte, indicadores y gráficos nativos de Excel."""
import re
import zipfile
from datetime import date, datetime

import pytest
from openpyxl import load_workbook

from evm import dashboard
from evm.cli import main
from evm.config import load_config
from evm.reader_xml import read_xml
from evm.update import indicators_for, plan_end, read_book, run_update
from evm.wbs import KIND_LEAF, KIND_RESIDUAL, KIND_SUMMARY
from evm.workbook import FIRST_ROW

DASH = "Cuadro de Mando"
CARDS = {"bac": "B10", "pv": "G10", "ev": "L10", "ac": "Q10", "cv": "B15", "sv": "G15", "cpi": "L15", "spi": "Q15",
         "etc": "B20", "eac": "F20", "vac": "J20", "tcpi_bac": "N20", "tcpi_eac": "R20"}
TABLE_FIRST = dashboard.chart_rows()[2] + 1


def charts(path):
    """XML de los gráficos del libro, en el orden de la hoja (evolución EVM, índices, CPI-SPI)."""
    with zipfile.ZipFile(path) as z:
        names = sorted((n for n in z.namelist() if re.fullmatch(r"xl/charts/chart\d+\.xml", n)),
                       key=lambda n: int(re.search(r"\d+", n).group()))
        return [z.read(n).decode() for n in names]


def refs(xml):
    return re.findall(r"<(?:c:)?f>([^<]*)</(?:c:)?f>", xml)


def tag(xml, name, attrs=""):
    """Cantidad de elementos <name ...> (con o sin el prefijo c:)."""
    return len(re.findall(rf"<(?:c:)?{name}{attrs}[ />]", xml))


def last_row(xml):
    return max(int(r.rsplit("$", 1)[1]) for r in refs(xml) if "$" in r)


STATUS = {"jun": date(2027, 6, 30), "jul": date(2027, 7, 31)}


def update(book, example_files, which="jun"):
    """'actualizar' del mes; antes, el AC del mes (del ERP), que se carga a mano en la hoja AC."""
    import generar_ejemplo
    generar_ejemplo.fill_book(book, STATUS[which])
    cfg = load_config(example_files["config"])
    return run_update(read_xml(example_files[which], cfg), cfg, book, backup=False)


# ----------------------------------------------------------------------------
# Reglas de los gráficos
# ----------------------------------------------------------------------------
def test_chart_rules():
    assert dashboard.plan_end_month([0, 5, 0, 3, 0, 0]) == 4
    assert dashboard.plan_end_month([0.0, 0.001]) == 0
    assert dashboard.chart_months(120, 55, 18) == (55, 18)      # EVM hasta el fin del plan; índices hasta el corte
    assert dashboard.chart_months(120, 55, 70) == (70, 70)      # proyecto atrasado: hasta el corte
    assert dashboard.chart_months(120, 55, None) == (55, 12)    # sin fecha de corte
    assert dashboard.chart_months(24, 30, 1) == (24, 3)
    rows = [(5, "1.0.0", 1, KIND_SUMMARY), (6, "1.1.0", 2, KIND_SUMMARY), (7, "1.1.1", 3, KIND_LEAF),
            (8, "1.0.0-D", 2, KIND_RESIDUAL), (9, "2.0.0", 1, KIND_LEAF)]
    assert [p.code for p in dashboard.chart_packages(rows)] == ["1.0.0", "1.1.0", "2.0.0"]
    many = [(i, f"1.{i}.0", 2, KIND_SUMMARY) for i in range(dashboard.MAX_SCATTER + 1)] + [(999, "9.0.0", 1, KIND_LEAF)]
    assert [p.code for p in dashboard.chart_packages(many)] == ["9.0.0"]      # demasiados: solo el nivel 1


# ----------------------------------------------------------------------------
# Hoja y gráficos
# ----------------------------------------------------------------------------
def test_dashboard_sheet_and_selectors(book):
    wb = load_workbook(book)
    assert wb.sheetnames[0] == DASH and wb.active.title == DASH
    assert wb["_graficos"].sheet_state == "hidden"
    assert "Evolución mensual" in wb.sheetnames and "Curva S" not in wb.sheetnames
    for name in ("MesReporte", "MesReporteN", "ListaMesesReporte", "MetodoEAC"):
        assert name in wb.defined_names
    ws = wb[DASH]
    assert ws["E5"].value is None                                  # sin fecha de corte todavía
    assert ws["R5"].value == "BAC / CPI"
    lists = {str(dv.sqref): dv.formula1 for dv in ws.data_validations.dataValidation}
    assert lists["E5"] == "ListaMesesReporte"
    assert all(m in lists["R5"] for m in dashboard.EAC_METHODS)
    codes = [ws.cell(r, 12).value for r in range(TABLE_FIRST, TABLE_FIRST + 12)]
    assert all(c and c.startswith("='_graficos'!") for c in codes)
    assert ws.cell(TABLE_FIRST + 12, 12).value is None              # 4 elementos de nivel 1 y 8 de nivel 2
    assert ws.page_setup.orientation == "portrait" and ws.sheet_properties.pageSetUpPr.fitToPage


def test_dashboard_charts(book, example_files):
    xml = charts(book)
    assert len(xml) == 3
    evm, idx, scatter = xml
    assert tag(evm, "lineChart") and tag(evm, "barChart")
    assert tag(idx, "lineChart") and tag(idx, "areaChart") and tag(idx, "grouping", ' val="stacked"')
    assert tag(scatter, "scatterChart") and tag(scatter, "ser") == 12 + 2 + 4
    for x in xml:
        assert refs(x) and all(r.startswith("'_graficos'!") for r in refs(x))
    assert tag(idx, "min", ' val="0.5"') and tag(idx, "max", ' val="1.5"')
    assert tag(scatter, "min", ' val="0.4"') == 2 and tag(scatter, "max", ' val="1.6"') == 2
    b = read_book(book)
    assert last_row(evm) == dashboard.D_FIRST_ROW + plan_end(b) - 1          # hasta el fin del plan

    update(book, example_files)
    wb = load_workbook(book)
    assert wb[DASH]["E5"].value == datetime(2027, 6, 1)                       # mes de corte
    xml = charts(book)
    assert len(xml) == 3
    assert last_row(xml[1]) == dashboard.D_FIRST_ROW + 18 - 1                 # índices hasta el mes 18


def test_migration_keeps_report_month(book, tmp_path, example_files):
    update(book, example_files)
    new = tmp_path / "EVM_v2.xlsx"
    assert main(["generar", "--origen", str(example_files["base"]), "--config", str(example_files["config"]),
                 "--salida", str(new), "--migrar-desde", str(book)]) == 0
    assert load_workbook(new)[DASH]["E5"].value == datetime(2027, 6, 1)
    assert last_row(charts(new)[1]) == dashboard.D_FIRST_ROW + 18 - 1


def test_legacy_book_keeps_s_curve(book, example_files):
    """Libros de versiones anteriores (hoja Curva S, sin Cuadro de Mando): se recrea su curva S."""
    wb = load_workbook(book)
    for name in (DASH, "_graficos"):
        del wb[name]
    wb["Evolución mensual"].title = "Curva S"
    wb.save(book)
    update(book, example_files)
    xml = charts(book)
    assert len(xml) == 1 and all(r.startswith("'Curva S'!") for r in refs(xml[0]))


# ----------------------------------------------------------------------------
# Valores (LibreOffice)
# ----------------------------------------------------------------------------
def card_values(ws):
    return {k: ws[a].value for k, a in CARDS.items()}


def check_cards(values, ind, eac):
    assert values["bac"] == pytest.approx(ind.bac, abs=0.01)
    assert values["pv"] == pytest.approx(ind.pv, abs=0.01)
    assert values["ev"] == pytest.approx(ind.ev, abs=0.01)
    assert values["ac"] == pytest.approx(ind.ac, abs=0.01)
    assert values["cv"] == pytest.approx(ind.cv, abs=0.01)
    assert values["sv"] == pytest.approx(ind.sv, abs=0.01)
    assert values["cpi"] == pytest.approx(ind.cpi, abs=1e-9)
    assert values["spi"] == pytest.approx(ind.spi, abs=1e-9)
    assert values["eac"] == pytest.approx(eac, abs=0.01)
    assert values["etc"] == pytest.approx(eac - ind.ac, abs=0.01)
    assert values["vac"] == pytest.approx(ind.bac - eac, abs=0.01)
    assert values["tcpi_bac"] == pytest.approx(ind.tcpi, abs=1e-9)
    assert values["tcpi_eac"] == pytest.approx((ind.bac - ind.ev) / (eac - ind.ac), abs=1e-9)


def recalc_with(book, recalc, month=None, method=None):
    """Recalcula el libro con LibreOffice, opcionalmente con otro mes de reporte o método de EAC."""
    if month or method:
        wb = load_workbook(book)
        if month:
            wb[DASH]["E5"] = month
        if method:
            wb[DASH]["R5"] = method
        wb.save(book)
    return load_workbook(recalc(book), data_only=True)


def test_dashboard_values_recalculate(book, example_files, recalc):
    update(book, example_files, "jun")
    update(book, example_files, "jul")
    out = recalc_with(book, recalc)
    errors = [(ws.title, c.coordinate, c.value) for ws in out.worksheets for row in ws.iter_rows() for c in row
              if isinstance(c.value, str) and c.value.startswith("#") and c.value != "#N/A"]
    assert not errors, errors[:10]
    b = read_book(book)
    ind = indicators_for(b, date(2027, 7, 31))
    ws = out[DASH]
    check_cards(card_values(ws), ind, ind.eac_cpi)
    assert ws["H5"].value == 19
    # Tabla de elementos EDT: mismos CPI y SPI que la hoja Indicadores EDT a la fecha de corte
    ind_edt = out["Indicadores EDT"]
    assert ws.cell(TABLE_FIRST, 19).value == pytest.approx(ind_edt.cell(FIRST_ROW, 11).value, abs=1e-9)
    assert ws.cell(TABLE_FIRST, 20).value == pytest.approx(ind_edt.cell(FIRST_ROW, 10).value, abs=1e-9)
    data = out["_graficos"]
    first = dashboard.D_FIRST_ROW
    assert data.cell(first + 18, dashboard.M_EV).value == pytest.approx(ind.ev, abs=0.01)
    assert data.cell(first + 18, dashboard.M_MARK).value == pytest.approx(ind.bac, abs=0.01)
    assert data.cell(first + 19, dashboard.M_EV).value == "#N/A"

    # Mes de reporte anterior al corte y otro método de EAC
    out = recalc_with(book, recalc, datetime(2027, 6, 1), dashboard.EAC_METHODS[1])
    ind = indicators_for(b, date(2027, 6, 30))
    check_cards(card_values(out[DASH]), ind, ind.eac_remaining_budget)
    assert out[DASH]["H5"].value == 18
    assert out["_graficos"].cell(first + 18, dashboard.M_EV).value == "#N/A"   # EV solo hasta el mes elegido

    # Un mes posterior al corte se limita al mes de corte; EAC con CPI × SPI
    out = recalc_with(book, recalc, datetime(2028, 1, 1), dashboard.EAC_METHODS[2])
    ind = indicators_for(b, date(2027, 7, 31))
    check_cards(card_values(out[DASH]), ind, ind.eac_cpi_spi)
    assert out[DASH]["H5"].value == 19

    # Sin AC cargado: los indicadores de costo quedan vacíos
    out = recalc_with(book, recalc, datetime(2027, 3, 1))
    values = card_values(out[DASH])
    assert values["pv"] == pytest.approx(indicators_for(b, date(2027, 3, 31)).pv, abs=0.01)
    assert values["ac"] == 0 and values["spi"] == 0
    for key in ("cv", "cpi", "etc", "eac", "vac", "tcpi_bac", "tcpi_eac"):
        assert values[key] is None, key
