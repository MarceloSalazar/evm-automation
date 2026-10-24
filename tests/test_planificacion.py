"""Planificación (INTERNO / CLIENTE), niveles por código EDT y tareas sin código EDT."""
from datetime import date, datetime

import pytest
from openpyxl import load_workbook

from evm.cli import main
from evm.config import Config, ConfigError, load_config
from evm.model import ProjectData, Task
from evm.reader_xml import read_xml
from evm.update import read_book, run_update
from evm.wbs import KIND_LEAF, KIND_RESIDUAL, KIND_SUMMARY, build_report, code_level


S, F = datetime(2026, 1, 5, 8), datetime(2026, 3, 27, 17)


def T(uid, level, wbs, name, **kw):
    return Task(uid=uid, id=uid, name=name, outline_level=level, wbs=wbs, baseline_start=S, baseline_finish=F,
                has_baseline=True, start=S, finish=F, **kw)


def P(tasks, **kw):
    return ProjectData(name="Programa", source="p.xml", backend="xml", tasks=tasks, **kw)


def two_plans():
    """INTERNO y CLIENTE con niveles de esquema que no coinciden con el nivel EDT."""
    out, uid = [], 0
    for plan in ("INTERNO", "CLIENTE"):
        def t(level, wbs, name, **kw):
            nonlocal uid
            uid += 1
            return T(uid, level, wbs, name, **kw)
        out += [
            t(1, plan, f"Planificación {plan}"),
            t(2, "", "Reunión de arranque", milestone=True),             # sin código: se omite
            t(2, "1.0.0", "Ingeniería"),
            t(3, "1.1.0", "Básica"),
            t(4, "1.1.1", "Estudios"),
            t(4, "1.1.2", "Diseño"),
            t(5, "1.1.2.1", "Planos"),                                   # nivel EDT 4: se suma en 1.1.2
            t(5, "", "Revisión interna"),                                # sin código: se suma en 1.1.2
            t(4, "", "Gastos generales"),                                # sin código bajo 1.1.0: fila -D
        ]
    return out


def leaves(row) -> list[str]:
    return [t.name for t in row.leaves]


def test_code_level():
    assert [code_level(c) for c in ("2.0.0", "2.1.0", "2.1.1", "2.1.1.4", "INTERNO", "1")] == [1, 2, 3, 4, 1, 1]


def test_internal_plan_rows_follow_edt_levels():
    rep = build_report(P(two_plans()), Config(plan="INTERNO"))
    rows = {r.code: r for r in rep.rows}
    assert list(rows) == ["1.0.0", "1.1.0", "1.1.1", "1.1.2", "1.1.0-D"]
    assert rows["1.0.0"].kind == KIND_SUMMARY and rows["1.1.0"].kind == KIND_SUMMARY
    assert rows["1.1.2"].kind == KIND_LEAF and leaves(rows["1.1.2"]) == ["Planos", "Revisión interna"]
    assert rows["1.1.0-D"].kind == KIND_RESIDUAL and leaves(rows["1.1.0-D"]) == ["Gastos generales"]
    assert [r.level for r in rep.rows] == [1, 2, 3, 3, 3]
    assert rep.plan == "INTERNO"
    assert any("Planificación 'INTERNO'" in i.message for i in rep.issues)
    assert any("Se omitió 1 tarea sin código EDT" in i.message for i in rep.issues)
    assert not [i for i in rep.issues if "duplicado" in i.message]


def test_client_plan_and_max_level():
    rep = build_report(P(two_plans()), Config(plan="cliente"))
    assert rep.plan == "CLIENTE" and {t.uid for r in rep.rows for t in r.leaves} == {14, 16, 17, 18}
    rep2 = build_report(P(two_plans()), Config(plan="INTERNO", max_level=2))
    assert [r.code for r in rep2.rows] == ["1.0.0", "1.1.0"]
    assert rep2.rows[1].kind == KIND_LEAF
    assert leaves(rep2.rows[1]) == ["Estudios", "Planos", "Revisión interna", "Gastos generales"]
    rep4 = build_report(P(two_plans()), Config(plan="INTERNO", max_level=4))
    assert "1.1.2.1" in [r.code for r in rep4.rows]


def test_without_plan_all_coded_tasks_are_exported():
    rep = build_report(P(two_plans()), Config())
    codes = [r.code for r in rep.rows]
    assert codes[0] == "INTERNO" and "CLIENTE" in codes
    assert any("duplicado" in i.message and "planificacion" in i.message for i in rep.issues)


def test_unknown_plan_is_reported():
    with pytest.raises(ConfigError, match="INTERNO.*CLIENTE"):
        build_report(P(two_plans()), Config(plan="OTRO"))


def test_plan_prefix_is_removed_from_codes():
    tasks = [T(1, 1, "INTERNO", "Plan"), T(2, 2, "INTERNO.1", "Uno"), T(3, 3, "INTERNO.1.1", "A")]
    rep = build_report(P(tasks), Config(plan="INTERNO"))
    assert [r.code for r in rep.rows] == ["1.0.0", "1.1.0"]


def test_uncoded_container_between_edt_levels():
    tasks = [T(1, 1, "INTERNO", "Plan"), T(2, 2, "1.0.0", "Uno"),
             T(3, 3, "", "Agrupador sin código"), T(4, 4, "1.1.0", "A"), T(5, 4, "1.2.0", "B")]
    rep = build_report(P(tasks), Config(plan="INTERNO"))
    rows = {r.code: r for r in rep.rows}
    assert list(rows) == ["1.0.0", "1.1.0", "1.2.0"]
    assert rows["1.1.0"].parent is rows["1.0.0"]


def test_xml_custom_edt_field_without_value_has_no_code(tmp_path):
    xml = tmp_path / "campo.xml"
    xml.write_text("""<?xml version="1.0" encoding="UTF-8"?>
<Project xmlns="http://schemas.microsoft.com/project"><Title>EDT propio</Title>
<ExtendedAttributes><ExtendedAttribute><FieldID>188743731</FieldID><FieldName>Text1</FieldName>
<Alias>EDT</Alias></ExtendedAttribute></ExtendedAttributes><Tasks>
<Task><UID>0</UID><ID>0</ID><Name>EDT propio</Name><OutlineLevel>0</OutlineLevel></Task>
<Task><UID>1</UID><ID>1</ID><Name>Con código</Name><WBS>1</WBS><OutlineLevel>1</OutlineLevel>
<ExtendedAttribute><FieldID>188743731</FieldID><Value>2.0.0</Value></ExtendedAttribute></Task>
<Task><UID>2</UID><ID>2</ID><Name>Sin código</Name><WBS>2</WBS><OutlineLevel>1</OutlineLevel></Task>
</Tasks></Project>""", encoding="utf-8")
    p = read_xml(xml, Config(wbs_field="EDT"))
    assert [t.wbs for t in p.tasks] == ["2.0.0", ""]       # no se usa el EDT automático de MS Project


def test_generate_and_update_with_plan(tmp_path, example_files, generator):
    book = tmp_path / "plan.xlsx"
    assert main(["generar", "--origen", str(example_files["base"]), "--config", str(example_files["config"]),
                 "--salida", str(book)]) == 0
    generator.fill_book(book)
    wb = load_workbook(book)
    assert wb["Resumen"]["B14"].value == "INTERNO"
    codes = [wb["PV"].cell(r, 1).value for r in range(5, wb["PV"].max_row + 1)]
    assert {"1.0.0", "1.1.0", "1.1.1", "2.1.0", "2.1.2"} <= set(codes) and "INTERNO" not in codes
    cfg = load_config(example_files["config"])
    jun = read_xml(example_files["jun"], cfg)
    res = run_update(jun, cfg, book, backup=False)
    assert sum(r.ev_cum or 0 for r in res.rows) > 0 and not any("planificación" in w for w in res.warnings)
    # La planificación del libro manda aunque la configuración indique otra
    res2 = run_update(jun, Config(plan="CLIENTE", method_field="Método EVM"), book, backup=False, dry_run=True)
    assert any("se usa esa planificación" in w for w in res2.warnings)


def test_cli_plan_option_overrides_config(tmp_path, example_files, generator):
    book = tmp_path / "cliente.xlsx"
    assert main(["generar", "--origen", str(example_files["base"]), "--config", str(example_files["config"]),
                 "--planificacion", "CLIENTE", "--salida", str(book)]) == 0
    assert load_workbook(book)["Resumen"]["B14"].value == "CLIENTE"
    generator.fill_book(book)                       # el presupuesto del ejemplo de la planificación CLIENTE
    b = read_book(book)
    assert sum(b.bac.values()) == pytest.approx(44_210_000 * 1.15)
    assert date(2026, 1, 1) == b.horizon.start


def test_edt_name_without_custom_field_uses_ms_project_wbs(tmp_path):
    xml = tmp_path / "wbs.xml"
    xml.write_text("""<?xml version="1.0" encoding="UTF-8"?>
<Project xmlns="http://schemas.microsoft.com/project"><Title>WBS</Title><Tasks>
<Task><UID>0</UID><ID>0</ID><Name>WBS</Name><OutlineLevel>0</OutlineLevel></Task>
<Task><UID>1</UID><ID>1</ID><Name>Uno</Name><WBS>1.2</WBS><OutlineLevel>1</OutlineLevel></Task>
</Tasks></Project>""", encoding="utf-8")
    assert read_xml(xml, Config(wbs_field="EDT")).tasks[0].wbs == "1.2"


def test_com_custom_field_named_edt_has_priority():
    from test_reader_com_fake import FakeApp, FakeTask, make_session
    tasks = [FakeTask(1, 1, "Fase", "1"), FakeTask(2, 1, "Otra", "2")]
    tasks[0].Text1, tasks[1].Text1 = "3.0.0", ""

    class App(FakeApp):
        def CustomFieldGetName(self, fid):
            return "EDT" if fid == 188743731 else ""

    s = make_session(tasks)
    s.app = App()
    p = s.read_project(Config(wbs_field="EDT"))
    assert [t.wbs for t in p.tasks] == ["3.0.0", ""]
