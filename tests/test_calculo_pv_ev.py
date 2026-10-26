"""Cálculo del PV y el EV: % planificado y % ganado de MS Project × BAC ingresado en el libro."""
from datetime import date, datetime

import pytest
from openpyxl import load_workbook

from evm.cli import main
from evm.config import DEFAULT_METHODS, MODE_PERCENT, Config, load_config
from evm.methods import parse_method
from evm.migrate import plan_migration, prior_from_book
from evm.percent import earned_fraction, full_weights, leaf_times, row_weights
from evm.update import indicators_for, read_book, run_update
from evm.wbs import KIND_RESIDUAL, build_report, build_tree
from evm.workbook import FILL_INPUT, WorkbookFormatError, read_structure, write_workbook
from evm.workcalendar import WorkCalendar

from builders import P, T, issues_with, rows_by_code


CAL = WorkCalendar()
FEB = (datetime(2027, 2, 1, 8), datetime(2027, 2, 26, 17))        # 20 días hábiles, todos en febrero
MAR_APR = (datetime(2027, 3, 18, 8), datetime(2027, 4, 28, 17))   # 30 días hábiles: 10 en marzo y 20 en abril
FAB = (datetime(2027, 1, 4, 8), datetime(2027, 10, 22, 17))       # Fabricación: 210 días hábiles
TRA = (datetime(2027, 10, 25, 8), datetime(2027, 12, 17, 17))     # Transporte: 40 días hábiles


def config(**kw) -> Config:
    cfg = Config(months=24, **kw)
    cfg.validate()
    return cfg


def multi_config(**kw) -> Config:
    return config(multi_enabled=True, **kw)


def excavation_tasks() -> list:
    """Ejemplo de la documentación: paquete 3.1.1 con una tarea de A y una tarea 'Ambas' con % Reparto 40."""
    return [
        T(1, 1, "3", "Construcción", bs=FEB[0], bf=MAR_APR[1], company="EmpresaA"),
        T(2, 2, "3.1", "Obra civil", bs=FEB[0], bf=MAR_APR[1]),
        T(3, 3, "3.1.1", "Movimiento de tierras", bs=FEB[0], bf=MAR_APR[1], method="% Completado"),
        T(4, 4, "", "Excavación", bs=FEB[0], bf=FEB[1], company="EmpresaA"),
        T(5, 4, "", "Estructura", bs=MAR_APR[0], bf=MAR_APR[1], company="Ambas", split_a=40),
    ]


def supply_tasks(fat: float | None = 10, site: float | None = 20) -> list:
    """Ejemplo de la documentación: suministro con dos tareas y dos hitos con 'Peso hito'."""
    return [
        T(1, 1, "3", "Construcción", bs=FAB[0], bf=TRA[1]),
        T(2, 2, "3.2", "Equipos", bs=FAB[0], bf=TRA[1]),
        T(3, 3, "3.2.1", "Suministro de equipos", bs=FAB[0], bf=TRA[1], method="% Completado"),
        T(4, 4, "", "Fabricación", bs=FAB[0], bf=FAB[1]),
        T(5, 4, "", "Hito FAT", bs=FAB[1], bf=FAB[1], milestone=True, milestone_weight=fat),
        T(6, 4, "", "Transporte", bs=TRA[0], bf=TRA[1]),
        T(7, 4, "", "Hito equipos en obra", bs=TRA[1], bf=TRA[1], milestone=True, milestone_weight=site),
    ]


def make_book(path, tasks, cfg, prior=None):
    rep = build_report(P(tasks), cfg)
    if prior is not None:
        rep.issues.extend(plan_migration(rep, prior, cfg.max_level))
    write_workbook(rep, cfg, path, prior)
    return rep


def set_bac(path, values: dict[str, dict[str, float]]) -> None:
    """Ingresa el BAC de las filas indicadas: {código: {hoja: BAC}}."""
    wb = load_workbook(path)
    rows = {s.code: s.row for s in read_structure(wb)}
    for code, sheets in values.items():
        for sheet, value in sheets.items():
            wb[sheet].cell(row=rows[code], column=4, value=value)
    wb.save(path)


def row_of(book, code: str) -> int:
    return next(s.row for s in book.structure if s.code == code)


def month(book, d: date) -> int:
    return book.horizon.index(d)


# ----------------------------------------------------------------------------
# Configuración
# ----------------------------------------------------------------------------
def test_old_mode_keys_are_ignored(tmp_path):
    """Las claves de la versión 1.x (modo, ac_desde_project) se ignoran: el cálculo es siempre BAC × %."""
    path = tmp_path / "c.toml"
    path.write_text('[exportacion]\nmodo = "costo"\n[multiempresa]\nactivar = true\nac_desde_project = true\n',
                    encoding="utf-8")
    cfg = load_config(path)
    assert cfg.multi_enabled and not hasattr(cfg, "mode") and not hasattr(cfg, "ac_from_project")


def test_method_list_has_no_bcwp():
    assert Config().methods == DEFAULT_METHODS
    assert not any("CPTR" in m for m in DEFAULT_METHODS)
    custom = Config(methods_list=["0/100", "Otro"])
    assert custom.methods == ["0/100", "Otro"]                # una lista de [metodos] se respeta


# ----------------------------------------------------------------------------
# Pesos y % planificado
# ----------------------------------------------------------------------------
def test_tasks_weigh_by_baseline_duration_and_shares():
    tasks = excavation_tasks()
    build_tree(tasks)
    times = leaf_times(tasks[3:], CAL)
    assert [round(t.minutes / 480) for t in times] == [20, 30]
    a = row_weights(times, lambda t: {4: 1.0, 5: 0.4}[t.uid], "EmpresaA")
    b = row_weights(times, lambda t: {4: 0.0, 5: 0.6}[t.uid], "EmpresaB")
    assert a.weights == pytest.approx({4: 20 / 32, 5: 12 / 32})
    assert b.weights == pytest.approx({5: 1.0})
    assert row_weights(times, lambda t: 0.0) is None              # la empresa no participa


def test_milestone_weight_is_a_share_of_the_bac():
    times = leaf_times(supply_tasks()[3:], CAL)
    rw = full_weights(times)
    fab, tra = (round(t.minutes / 480) for t in (times[0], times[2]))
    assert rw.weights[5] == pytest.approx(0.10) and rw.weights[7] == pytest.approx(0.20)
    assert rw.weights[4] == pytest.approx(0.70 * fab / (fab + tra))
    assert rw.weights[6] == pytest.approx(0.70 * tra / (fab + tra))
    assert not rw.notes


@pytest.mark.parametrize("fat, site, expected, note", [
    (40, 60, {5: 0.4, 7: 0.6}, ""),                                  # suman 100: las tareas no reciben peso
    (80, 70, {5: 80 / 150, 7: 70 / 150}, "suman 150"),               # más de 100: se escalan
    (-10, 20, None, "negativo"),
])
def test_milestone_weight_limits(fat, site, expected, note):
    rw = full_weights(leaf_times(supply_tasks(fat, site)[3:], CAL))
    if expected is not None:
        assert rw.weights == pytest.approx(expected)
    if note:
        assert any(note in n.message for n in rw.notes)
    else:
        assert not rw.notes
    assert sum(rw.weights.values()) == pytest.approx(1.0)


def test_milestones_only_and_no_weights():
    ms = [T(1, 1, "", "H1", bs=FAB[1], bf=FAB[1], milestone=True, milestone_weight=30),
          T(2, 1, "", "H2", bs=TRA[1], bf=TRA[1], milestone=True, milestone_weight=30)]
    rw = full_weights(leaf_times(ms, CAL))
    assert rw.weights == pytest.approx({1: 0.5, 2: 0.5}) and "duración" in rw.notes[0].message
    for m in ms:
        m.milestone_weight = None
    rw = full_weights(leaf_times(ms, CAL))
    assert rw.weights == pytest.approx({1: 0.5, 2: 0.5}) and "pesan igual" in rw.notes[0].message


def test_planned_curve_per_company_and_horizon():
    rep = build_report(P(excavation_tasks()), multi_config())
    row = rows_by_code(rep)["3.1.1"]
    start = rep.horizon.index(date(2027, 2, 1))
    assert rep.horizon.start == date(2027, 2, 1)
    assert row.pct_vec["A"][start:start + 3] == pytest.approx([0.625, 0.125, 0.25])
    assert row.pct_vec["B"][start:start + 3] == pytest.approx([0.0, 1 / 3, 2 / 3])
    assert issues_with(rep, "El PV y el EV de cada fila son el BAC")


def test_tasks_without_baseline_dates_are_warned():
    tasks = excavation_tasks()
    tasks[4].baseline_start = tasks[4].baseline_finish = None
    tasks[4].has_baseline = False
    rep = build_report(P(tasks), config())
    assert issues_with(rep, "1 tarea no tiene fechas de línea base", "Advertencia")
    # Tarea sin código EDT: con su Id, para encontrarla en MS Project
    assert any(i.name == f"Estructura (Id {tasks[4].id})" and "fechas actuales" in i.message for i in rep.issues)
    assert sum(rows_by_code(rep)["3.1.1"].pct_vec[""]) == pytest.approx(1.0)


def test_loose_tasks_form_an_other_tasks_row():
    tasks = [T(1, 1, "1", "Gestión", bs=FEB[0], bf=FEB[1]), T(2, 2, "1.1", "PMO", bs=FEB[0], bf=FEB[1]),
             T(3, 2, "", "Reunión de arranque", bs=FEB[0], bf=FEB[1]),
             T(4, 2, "", "Hito sin peso", bs=FEB[1], bf=FEB[1], milestone=True)]
    rep = build_report(P(tasks), config(max_level=2))
    rows = rows_by_code(rep)
    assert rows["1.0-D"].kind == KIND_RESIDUAL and rows["1.0-D"].name.startswith("Otras tareas")
    tasks.pop(2)                                    # solo un hito sin peso: no forma fila
    assert "1.0-D" not in rows_by_code(build_report(P(tasks), config(max_level=2)))


# ----------------------------------------------------------------------------
# Libro
# ----------------------------------------------------------------------------
def test_book_sheets_formulas_and_input_cells(tmp_path):
    path = tmp_path / "libro.xlsx"
    make_book(path, excavation_tasks(), multi_config())
    wb = load_workbook(path)
    for name in ("% PV-EmpresaA", "% PV-EmpresaB", "% EV-EmpresaA", "% EV-EmpresaB"):
        assert name in wb.sheetnames
    rows = {s.code: s.row for s in read_structure(wb)}
    r = rows["3.1.1"]
    pv_a, ev_a, ac_a, pct_a = wb["PV-EmpresaA"], wb["EV-EmpresaA"], wb["AC-EmpresaA"], wb["% PV-EmpresaA"]
    assert pv_a.cell(row=r, column=4).value is None
    assert pv_a.cell(row=r, column=4).fill.fgColor.rgb.endswith(FILL_INPUT.fgColor.rgb[-6:])
    assert pv_a.cell(row=r, column=7).value == f"=N($D{r})*'% PV-EmpresaA'!G{r}"
    assert ev_a.cell(row=r, column=7).value == f"=N($D{r})*'% EV-EmpresaA'!G{r}"
    assert not ev_a.cell(row=r, column=7).fill.fgColor.rgb.endswith("FFF2CC")
    assert not pct_a.cell(row=r, column=7).fill.fgColor.rgb.endswith("FFF2CC")
    assert ac_a.cell(row=r, column=7).fill.fgColor.rgb.endswith("FFF2CC")
    assert pct_a.cell(row=r, column=7).value == pytest.approx(0.625)
    meta = {k: v for k, v in wb["_meta"].iter_rows(min_row=2, values_only=True)}
    assert meta["modo"] == MODE_PERCENT
    methods = [c.value for c in wb["Listas"]["A"][1:] if c.value]
    assert not any("CPTR" in m for m in methods)


def test_update_loads_earned_percent_per_company(tmp_path):
    path = tmp_path / "libro.xlsx"
    tasks = excavation_tasks()
    make_book(path, tasks, multi_config())
    set_bac(path, {"3.1.1": {"PV-EmpresaA": 300_000, "PV-EmpresaB": 700_000}})
    tasks[3].percent_complete, tasks[3].actual_start, tasks[3].actual_finish = 100, FEB[0], FEB[1]
    tasks[4].percent_complete, tasks[4].actual_start = 30, MAR_APR[0]
    res = run_update(P(tasks), multi_config(), path, date(2027, 3, 31), backup=False)
    book = read_book(path)
    r, i = row_of(book, "3.1.1"), month(book, date(2027, 3, 1))
    assert book.pct_ev["A"][r][i] == pytest.approx(0.7375)
    assert book.pct_ev["B"][r][i] == pytest.approx(0.30)
    assert book.ev_a[r][i] == pytest.approx(221_250) and book.ev_b[r][i] == pytest.approx(210_000)
    ind = indicators_for(book, date(2027, 3, 31))
    assert ind.pv == pytest.approx(458_333.333, abs=0.01) and ind.ev == pytest.approx(431_250)
    assert ind.spi == pytest.approx(0.94, abs=0.005)
    assert all(v is None for v in book.ac[r])                 # el AC se ingresa a mano
    assert all(not r.note for r in res.rows)


def test_supply_example_with_milestones(tmp_path):
    path = tmp_path / "libro.xlsx"
    tasks = supply_tasks()
    make_book(path, tasks, config())
    set_bac(path, {"3.2.1": {"PV": 1_000_000}})
    book = read_book(path)
    r = row_of(book, "3.2.1")
    assert book.pv[r][month(book, date(2027, 10, 1))] >= 100_000          # el hito FAT cae en octubre
    assert sum(book.pv[r]) == pytest.approx(1_000_000)
    tasks[3].percent_complete = tasks[4].percent_complete = 100
    tasks[4].actual_finish = FAB[1]
    tasks[5].percent_complete = 50
    run_update(P(tasks), config(), path, date(2027, 11, 30), backup=False)
    book = read_book(path)
    fab, tra = 210, 40
    expected = 100_000 + 700_000 * fab / (fab + tra) + 0.5 * 700_000 * tra / (fab + tra)
    assert sum(v or 0 for v in book.ev[r]) == pytest.approx(expected)


def test_fixed_and_milestone_methods_are_measured_per_company(tmp_path):
    tasks = excavation_tasks()
    tasks[2].method = "0/100"
    tasks[4].company, tasks[4].split_a = "EmpresaB", None
    tasks[3].percent_complete, tasks[3].actual_finish = 100, FEB[1]
    path = tmp_path / "libro.xlsx"
    make_book(path, tasks, multi_config())
    set_bac(path, {"3.1.1": {"PV-EmpresaA": 300, "PV-EmpresaB": 700}})
    run_update(P(tasks), multi_config(), path, date(2027, 3, 31), backup=False)
    book = read_book(path)
    r, i = row_of(book, "3.1.1"), month(book, date(2027, 3, 1))
    assert book.pct_ev["A"][r][i] == pytest.approx(1.0)       # A terminó su tarea
    assert book.pct_ev["B"][r][i] == pytest.approx(0.0)       # B no comenzó


def test_company_without_tasks_uses_the_whole_row(tmp_path):
    tasks = excavation_tasks()
    tasks[4].company, tasks[4].split_a = "EmpresaA", None
    tasks[3].percent_complete = 50
    path = tmp_path / "libro.xlsx"
    rep = make_book(path, tasks, multi_config())
    assert issues_with(rep, "EmpresaB no participa")
    set_bac(path, {"3.1.1": {"PV-EmpresaA": 300, "PV-EmpresaB": 100}})
    res = run_update(P(tasks), multi_config(), path, date(2027, 3, 31), backup=False, dry_run=True)
    row = next(x for x in res.rows if x.code == "3.1.1")
    assert "EmpresaB no participa" in row.note
    book = read_book(path)
    r = row_of(book, "3.1.1")
    pv_a = book.bac_role["A"][r] * sum(book.pct_pv["A"][r])
    assert pv_a == pytest.approx(300) and sum(book.pv[r]) - pv_a == pytest.approx(100)


def test_bac_change_recalculates_every_month(tmp_path):
    path = tmp_path / "libro.xlsx"
    tasks = excavation_tasks()
    make_book(path, tasks, config())
    set_bac(path, {"3.1.1": {"PV": 1000}})
    tasks[3].percent_complete = 100
    run_update(P(tasks), config(), path, date(2027, 2, 28), backup=False)
    before = read_book(path)
    set_bac(path, {"3.1.1": {"PV": 2000}})
    after = read_book(path)
    r = row_of(after, "3.1.1")
    assert after.pv[r] == pytest.approx([2 * v for v in before.pv[r]])
    assert sum(v or 0 for v in after.ev[r]) == pytest.approx(2 * sum(v or 0 for v in before.ev[r]))


def test_update_writes_the_month_difference_and_warns_without_bac(tmp_path):
    path = tmp_path / "libro.xlsx"
    tasks = excavation_tasks()
    make_book(path, tasks, config())
    tasks[3].percent_complete = 50
    res = run_update(P(tasks), config(), path, date(2027, 2, 28), backup=False)
    assert any("Ninguna fila tiene BAC" in w.message for w in res.warnings)
    tasks[3].percent_complete = 100
    run_update(P(tasks), config(), path, date(2027, 3, 31), backup=False)
    book = read_book(path)
    r = row_of(book, "3.1.1")
    feb, mar = month(book, date(2027, 2, 1)), month(book, date(2027, 3, 1))
    assert book.pct_ev[""][r][feb] == pytest.approx(0.2) and book.pct_ev[""][r][mar] == pytest.approx(0.2)
    with pytest.raises(ValueError, match="posteriores"):
        run_update(P(tasks), config(), path, date(2027, 2, 28), backup=False)


def test_old_bcwp_method_is_measured_as_percent_complete():
    """El método CPTR de la versión 1.x ya no existe: se informa como no reconocido y se usa % completado."""
    tasks = excavation_tasks()
    tasks[3].percent_complete = 100
    text = "CPTR de MS Project"
    frac, note = earned_fraction(parse_method(text), full_weights(leaf_times(tasks[3:], CAL)), 0.0)
    assert frac == pytest.approx(0.4) and "no reconocido" in note.message and "% completado" in note.message


def test_migration_keeps_bac_and_earned_percent(tmp_path):
    old, new = tmp_path / "old.xlsx", tmp_path / "new.xlsx"
    tasks = excavation_tasks()
    make_book(old, tasks, multi_config())
    set_bac(old, {"3.1.1": {"PV-EmpresaA": 300_000, "PV-EmpresaB": 700_000}})
    tasks[3].percent_complete = 100
    run_update(P(tasks), multi_config(), old, date(2027, 2, 28), backup=False)
    rep = make_book(new, excavation_tasks(), multi_config(), prior_from_book(old, "Programa"))
    a, b = read_book(old), read_book(new)
    r = row_of(b, "3.1.1")
    assert b.bac_role["A"][r] == 300_000 and b.bac_role["B"][r] == 700_000
    assert b.pct_ev["A"][r] == a.pct_ev["A"][r] and b.ev == a.ev
    assert issues_with(rep, "Se conservó el BAC")


def test_books_of_version_1_are_rejected(tmp_path):
    """Un libro de la versión 1.x (PV y EV con los costos de MS Project) no se actualiza ni se migra."""
    old = tmp_path / "old.xlsx"
    make_book(old, excavation_tasks(), config())
    wb = load_workbook(old)
    ws = wb["_meta"]
    r = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "modo")
    ws.cell(r, 2).value = "costo"
    wb.save(old)
    with pytest.raises(WorkbookFormatError, match="versión 1.x"):
        read_book(old)
    with pytest.raises(WorkbookFormatError, match="versión 1.x"):
        prior_from_book(old, "Programa")


# ----------------------------------------------------------------------------
# Línea de comandos y fórmulas de Excel
# ----------------------------------------------------------------------------
def test_cli_generates_percent_books_even_with_old_mode_key(tmp_path, example_files, capsys):
    cost = tmp_path / "costo.toml"
    cost.write_text(example_files["config"].read_text(encoding="utf-8").replace(
        "[exportacion]\n", '[exportacion]\nmodo = "costo"\n'), encoding="utf-8")
    out = tmp_path / "libro.xlsx"
    assert main(["generar", "--origen", str(example_files["base"]), "--config", str(cost), "--salida", str(out)]) == 0
    assert "se debe ingresar el BAC de cada fila en la columna D" in capsys.readouterr().out
    assert "% PV" in load_workbook(out, read_only=True).sheetnames


def test_several_ms_project_files(tmp_path, generator):
    """Varios archivos ([[proyectos]]): cada fila calcula su % con las tareas de su archivo."""
    files = {}
    for project, build in ((1, generator.build_plans), (2, generator.build_second)):
        for label, status in (("base", None), ("jun", date(2027, 6, 30))):
            files[(project, label)] = tmp_path / f"p{project}_{label}.xml"
            generator.write(build(), files[(project, label)], status, project)
    cfg = tmp_path / "varios.toml"
    cfg.write_text('[campos]\nmetodo = "Método EVM"\npeso_hito = "Peso hito"\n\n[exportacion]\n'
                   'planificacion = "INTERNO"\n\n[multiempresa]\nactivar = true\n'
                   '\n[[proyectos]]\nnombre = "P1"\norigen = "p1_base.xml"\n'
                   '\n[[proyectos]]\nnombre = "P2"\norigen = "p2_base.xml"\nplanificacion = ""\n', encoding="utf-8")
    out = tmp_path / "varios.xlsx"
    assert main(["generar", "--config", str(cfg), "--salida", str(out)]) == 0
    wb = load_workbook(out)
    for s in read_structure(wb):
        if s.is_input:
            wb["PV-EmpresaA"].cell(row=s.row, column=4, value=1000)
            wb["PV-EmpresaB"].cell(row=s.row, column=4, value=1000)
    wb.save(out)
    assert main(["actualizar", "--config", str(cfg), "--libro", str(out), "--sin-respaldo",
                 "--origen", f"P1={files[(1, 'jun')]}", "--origen", f"P2={files[(2, 'jun')]}"]) == 0
    book = read_book(out)
    earned = {s.key.split(":")[0] for s in book.input_rows
              if any(v for role in ("A", "B") for v in book.pct_ev[role][s.row] if v)}
    assert earned == {"P1", "P2"}
    assert all(sum(book.pct_pv["A"][s.row]) == pytest.approx(1.0) for s in book.input_rows)


def test_excel_formulas_match_the_program(tmp_path, example_files, recalc):
    out = tmp_path / "libro.xlsx"
    cfg = tmp_path / "multi.toml"
    cfg.write_text(example_files["config"].read_text(encoding="utf-8") + "\n[multiempresa]\nactivar = true\n",
                   encoding="utf-8")
    assert main(["generar", "--origen", str(example_files["base"]), "--config", str(cfg), "--salida", str(out)]) == 0
    wb = load_workbook(out)
    for k, s in enumerate(x for x in read_structure(wb) if x.is_input):
        wb["PV-EmpresaA"].cell(row=s.row, column=4, value=100_000 * (k + 1))
        wb["PV-EmpresaB"].cell(row=s.row, column=4, value=50_000 * (k % 3))
        wb["AC-EmpresaA"].cell(row=s.row, column=7 + 17, value=1_000 * k)
    wb.save(out)
    assert main(["actualizar", "--origen", str(example_files["jun"]), "--config", str(cfg), "--libro", str(out),
                 "--sin-respaldo"]) == 0
    book = read_book(out)
    ind = indicators_for(book, date(2027, 6, 30))
    ws = load_workbook(recalc(out), data_only=True)["Resumen"]
    assert ws["B16"].value == pytest.approx(ind.bac)
    assert ws["B17"].value == pytest.approx(ind.pv)
    assert ws["B18"].value == pytest.approx(ind.ev)
    assert ws["B19"].value == pytest.approx(ind.ac)
    assert ind.ev > 0 and ind.pv > 0
