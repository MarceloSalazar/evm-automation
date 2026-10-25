"""Hoja Análisis de desviaciones: paquetes fuera de umbral, texto conservado, 'analizar' y --migrar-desde."""
from datetime import date, datetime

import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from evm import analysis
from evm.analysis import (A_ACTION, A_BAC, A_CAUSE, A_CODE, A_CPI, A_KEY, A_LIGHT, A_MONTH, A_ORIGIN, A_OWNER,
                          A_SPI, FIRST, HEAD_ROW, SHEET_ANALYSIS)
from evm.cli import main
from evm.update import read_book
from evm.workbook import FIRST_MONTH_COL, FIRST_ROW


JUN, JUL = date(2027, 6, 1), date(2027, 7, 1)
JUNE_COUNT = 13         # paquetes fuera de 0,95 a 1,05 en junio de 2027 (8 por debajo y 5 por encima)


STATUS = {"jun": date(2027, 6, 30), "jul": date(2027, 7, 31)}


@pytest.fixture()
def book(tmp_path, example_files, generator):
    out = tmp_path / "EVM.xlsx"
    assert main(["generar", "--origen", str(example_files["base"]), "--config", str(example_files["config"]),
                 "--salida", str(out)]) == 0
    generator.fill_book(out)                        # BAC de cada fila, que se ingresa a mano
    return out


def update(book, example_files, month, *extra, ac: bool = True) -> None:
    """'actualizar' del mes; antes, el AC del mes (del ERP), que se carga a mano en la hoja AC."""
    if ac:
        import generar_ejemplo
        generar_ejemplo.fill_book(book, STATUS[month])
    assert main(["actualizar", "--origen", str(example_files[month]), "--config", str(example_files["config"]),
                 "--libro", str(book), "--sin-respaldo", *extra]) == 0


def entries(book) -> list[analysis.Entry]:
    wb = load_workbook(book)
    return analysis.read_entries(wb[SHEET_ANALYSIS])


def rows_of(book) -> dict[str, int]:
    """Código EDT -> fila de la hoja de análisis."""
    ws = load_workbook(book)[SHEET_ANALYSIS]
    return {ws.cell(r, A_CODE).value: r for r in range(FIRST, ws.max_row + 1) if ws.cell(r, A_KEY).value}


def expected(book, month: date, low: float = 0.95, high: float = 1.05) -> set[str]:
    """Paquetes con SPI o CPI acumulado fuera del rango verde (0,95 a 1,05 por defecto), calculados aparte."""
    b = read_book(book)
    idx = b.horizon.index(month)
    out = set()
    for s in b.input_rows:
        pv = sum(b.pv[s.row][: idx + 1])
        ev = sum(v or 0 for v in b.ev[s.row][: idx + 1])
        ac = sum(v or 0 for v in b.ac[s.row][: idx + 1])
        if any(not low <= x <= high for x in ((ev / pv) if pv else None, (ev / ac) if ac else None)
               if x is not None):
            out.add(s.code)
    return out


def write_cells(book, values: dict[tuple[str, int], object], sheet: str = SHEET_ANALYSIS) -> None:
    """Simula lo que se escribe a mano en Excel: {(código EDT, columna): valor} en la hoja de análisis."""
    wb = load_workbook(book)
    ws = wb[sheet]
    rows = {ws.cell(r, A_CODE).value: r for r in range(FIRST, ws.max_row + 1)}
    for (code, column), value in values.items():
        ws.cell(rows[code], column, value)
    wb.save(book)


def test_generate_creates_empty_sheet(book):
    wb = load_workbook(book)
    assert wb.sheetnames.index(SHEET_ANALYSIS) == wb.sheetnames.index("Indicadores EDT") + 1
    ws = wb[SHEET_ANALYSIS]
    assert [ws.cell(HEAD_ROW, c).value for c in (A_MONTH, A_CODE, A_SPI, A_CPI, A_CAUSE)] == [
        "Mes", "Código EDT", "SPI", "CPI", "Causa"]
    key = get_column_letter(A_KEY)
    assert ws.column_dimensions[key].hidden
    assert ws.auto_filter.ref.startswith(f"A4:{key}")     # al ordenar, la clave acompaña a su fila
    assert entries(book) == []
    meta = read_book(book).meta
    assert [meta[k] for k in ("semaforo_rojo_inferior", "semaforo_amarillo_inferior", "semaforo_amarillo_superior",
                              "semaforo_rojo_superior")] == [0.9, 0.95, 1.05, 1.1]


def test_update_lists_packages_out_of_threshold(book, example_files, capsys):
    update(book, example_files, "jun")
    out = capsys.readouterr().out
    found = entries(book)
    codes = {e.code for e in found}
    n = len(codes)
    assert codes == expected(book, JUN) and n == JUNE_COUNT
    assert {"1.1.1", "2.2.1"} <= codes                    # 1.1.1 (LOE): CPI mayor que 1,05; 2.2.1: atraso
    assert f"Paquetes fuera de umbral (SPI o CPI fuera de 0,95 a 1,05): {n}" in out
    assert all(e.month == JUN and e.origin == analysis.NEW for e in found)
    # La clave y las fórmulas apuntan a la fila del paquete en las hojas PV, EV y AC
    b = read_book(book)
    keys = {s.code: (s.key, s.row) for s in b.input_rows}
    ws = load_workbook(book)[SHEET_ANALYSIS]
    for code, r in rows_of(book).items():
        key, data_row = keys[code]
        assert ws.cell(r, A_KEY).value == key
        assert ws.cell(r, A_BAC).value == f"='PV'!$D${data_row}"
    assert f"Último análisis: 06/2027, {n} paquetes fuera de umbral." == ws.cell(3, 1).value
    # Repetir el mismo mes no duplica filas
    update(book, example_files, "jun")
    assert len(entries(book)) == n


def test_text_is_kept_and_copied_to_the_next_month(book, example_files):
    update(book, example_files, "jun")
    write_cells(book, {("2.2.1", A_CAUSE): "Planos del proveedor atrasados", ("2.2.1", A_OWNER): "Ingeniería",
                       ("2.2.1", A_ACTION): "Reforzar el equipo de detalle"})
    wb = load_workbook(book)                    # fila agregada a mano (sin clave): se conserva
    ws = wb[SHEET_ANALYSIS]
    ws.cell(ws.max_row + 1, A_MONTH, datetime(2027, 6, 1)).value
    manual = ws.max_row
    ws.cell(manual, A_CODE, "2.0.0")
    ws.cell(manual, A_CAUSE, "Comentario del elemento de nivel 1")
    wb.save(book)

    update(book, example_files, "jun")          # mismo mes: se conserva el texto
    june = {e.code: e for e in entries(book) if e.key}
    assert june["2.2.1"].texts[0] == "Planos del proveedor atrasados"
    assert any(not e.key and e.texts[0] == "Comentario del elemento de nivel 1" for e in entries(book))

    update(book, example_files, "jul")
    found = entries(book)
    july = {e.code: e for e in found if e.key and e.month == JUL}
    old = [e for e in found if e.key and e.month == JUN]
    assert set(july) == expected(book, JUL)
    assert [e.code for e in old] == ["2.2.1"]            # junio: solo la fila con texto
    copied = july["2.2.1"]
    assert copied.texts[0] == "Planos del proveedor atrasados" and copied.texts[3] == "Ingeniería"
    assert copied.origin == "Texto copiado de 06/2027: se debe revisar"
    assert july["2.2.2"].origin == "También en 06/2027"
    assert any(not e.key and e.code == "2.0.0" for e in found)
    # Orden: julio primero, en el orden de la EDT
    ws = load_workbook(book)[SHEET_ANALYSIS]
    months = [ws.cell(r, A_MONTH).value.date() for r in range(FIRST, FIRST + len(july) + 1)]
    assert months == [JUL] * len(july) + [JUN]


def test_analizar_after_loading_ac_by_hand(book, example_files, capsys):
    """AC del ERP cargado a mano después de 'actualizar': 'analizar' agrega los paquetes con CPI bajo."""
    update(book, example_files, "jun", ac=False)
    before = {e.code for e in entries(book)}
    assert "1.1.1" not in before                  # LOE: SPI = 1 y, sin AC, el CPI no se calcula
    b = read_book(book)
    row = next(s.row for s in b.input_rows if s.code == "1.1.1")
    ev = sum(v or 0 for v in b.ev[row])
    wb = load_workbook(book)
    wb["AC"].cell(row, FIRST_MONTH_COL + b.horizon.index(JUN), 2 * ev)      # CPI = 0,50
    wb.save(book)
    capsys.readouterr()
    assert main(["analizar", "--libro", str(book), "--sin-respaldo"]) == 0
    out = capsys.readouterr().out
    assert "Fecha de corte: 30/06/2027" in out and "1.1.1" in out
    assert {e.code for e in entries(book)} == before | {"1.1.1"}
    assert {e.code for e in entries(book)} == expected(book, JUN)


def test_text_kept_when_back_within_threshold(book, example_files):
    update(book, example_files, "jun")
    write_cells(book, {("2.1.3", A_CAUSE): "Retraso en la topografía"})
    # 2.1.3 y 2.1.4 vuelven al rango verde (EV y AC acumulados iguales al PV); solo 2.1.3 tiene texto
    b = read_book(book)
    wb = load_workbook(book)
    i = b.horizon.index(JUN)
    for code in ("2.1.3", "2.1.4"):
        row = next(s.row for s in b.input_rows if s.code == code)
        pv = sum(b.pv[row][: i + 1])
        wb["% EV"].cell(row, FIRST_MONTH_COL + i, (pv - sum(v or 0 for v in b.ev[row][:i])) / b.bac[row])
        wb["AC"].cell(row, FIRST_MONTH_COL + i, pv - sum(v or 0 for v in b.ac[row][:i]))
    wb.save(book)
    assert main(["analizar", "--libro", str(book), "--sin-respaldo"]) == 0
    codes = {e.code for e in entries(book)}
    assert "2.1.3" in codes and "2.1.4" not in codes


def test_simulation_does_not_write(book, example_files, capsys):
    update(book, example_files, "jun", "--simular")
    assert f"Paquetes fuera de umbral (SPI o CPI fuera de 0,95 a 1,05): {JUNE_COUNT}" in capsys.readouterr().out
    assert entries(book) == []


def test_analizar_requires_status_date(book, capsys):
    assert main(["analizar", "--libro", str(book)]) == 2
    assert "antes de 'analizar'" in capsys.readouterr().err


def test_old_book_gets_the_sheet(book, example_files):
    """Libro de una versión anterior (sin la hoja ni los umbrales en _meta): 'actualizar' crea la hoja."""
    wb = load_workbook(book)
    del wb[SHEET_ANALYSIS]
    ws = wb["_meta"]
    for r in range(ws.max_row, 1, -1):
        if str(ws.cell(r, 1).value).startswith("semaforo_"):
            ws.delete_rows(r)
    wb.save(book)
    update(book, example_files, "jun")
    wb = load_workbook(book)
    assert wb.sheetnames.index(SHEET_ANALYSIS) == wb.sheetnames.index("Indicadores EDT") + 1
    assert {e.code for e in entries(book)} == expected(book, JUN)


def test_migrate_keeps_the_analysis(book, example_files, tmp_path):
    update(book, example_files, "jun")
    write_cells(book, {("2.2.1", A_CAUSE): "Planos atrasados", ("1.2.2", A_CAUSE): "Licitación desierta"})
    wb = load_workbook(book)                    # 1.2.2: clave inexistente en el libro nuevo
    ws = wb[SHEET_ANALYSIS]
    ws.cell(rows_of(book)["1.2.2"], A_KEY, "99999")
    wb.save(book)
    new = tmp_path / "EVM_LB2.xlsx"
    assert main(["generar", "--origen", str(example_files["base"]), "--config", str(example_files["config"]),
                 "--salida", str(new), "--migrar-desde", str(book)]) == 0
    found = {e.code: e for e in entries(new)}
    assert found["2.2.1"].texts[0] == "Planos atrasados" and found["2.2.1"].key
    assert not found["1.2.2"].key and found["1.2.2"].texts[0] == "Licitación desierta"
    data_row = next(s.row for s in read_book(new).input_rows if s.code == "2.2.1")
    ws = load_workbook(new)[SHEET_ANALYSIS]
    assert ws.cell(rows_of(new)["2.2.1"], A_BAC).value == f"='PV'!$D${data_row}"
    details = [r[3] for r in load_workbook(new)["Validación"].iter_rows(min_row=4, values_only=True) if r[3]]
    assert any(f"trasladaron {JUNE_COUNT - 1} filas del análisis de desviaciones" in d for d in details)
    assert any("ya no tiene fila en el libro nuevo" in d for d in details)


def test_values_recalculate(book, example_files, recalc):
    update(book, example_files, "jun")
    out = load_workbook(recalc(book), data_only=True)
    ws, ind = out[SHEET_ANALYSIS], out["Indicadores EDT"]
    by_code = {ind.cell(r, 1).value: r for r in range(FIRST_ROW, ind.max_row + 1)}
    rows = rows_of(book)
    assert rows
    for code, r in rows.items():
        i = by_code[code]
        for column, ind_column in ((A_SPI, 10), (A_CPI, 11)):
            value, reference = ws.cell(r, column).value, ind.cell(i, ind_column).value
            if reference in ("", None):
                assert value in ("", None), code
            else:
                assert value == pytest.approx(reference, abs=1e-9), code
        values = [x for x in (ws.cell(r, A_SPI).value, ws.cell(r, A_CPI).value) if isinstance(x, (int, float))]
        light = ("Rojo" if min(values) < 0.9 or max(values) > 1.1 else
                 "Amarillo" if min(values) < 0.95 or max(values) > 1.05 else "Verde")
        assert ws.cell(r, A_LIGHT).value == light != "Verde", code
    assert ws.cell(FIRST, A_ORIGIN).value == analysis.NEW
